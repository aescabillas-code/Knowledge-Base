import os
import re
import sqlite3
import mimetypes
import base64
import hashlib
import hmac
from datetime import datetime
from pathlib import Path
from html import escape

import streamlit as st

# Persistent browser cookie support for the one-time access-token gate.
# The dependency is kept optional so the script can still start and display
# a clear setup message if the package has not yet been added to requirements.
try:
    from streamlit_cookies_controller import CookieController
except Exception:
    CookieController = None

try:
    from docx import Document as WordDocument
except Exception:
    WordDocument = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from pptx import Presentation as PowerPointPresentation
except Exception:
    PowerPointPresentation = None

try:
    import fitz
except Exception:
    fitz = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except Exception:
    TfidfVectorizer = None
    cosine_similarity = None


# ============================================================
# PAGE
# ============================================================
# Use the user-supplied HPE logo as the in-app brand mark; Knowledge Base remains live text.
# The logo and favicon are embedded so deployment does not depend on uploaded assets.
HPE_KB_LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAAQgAAABRCAYAAAAuADh0AAAGy0lEQVR4nO3dfUhVZxwH8K9NGGOx3ZepNGoNZnNZoDc3m7RLR2/KYoPFkG2kyZD8Yw2RNll/lSf6ZzGJQXj3R3sJRpNqOVowmpLdWgXLnBPs5eZCNg3zinmxaYxod380N9N7vOeet+fx3u8HDlTnPM95rnm/93dennsycPpQDPbJsLFvs2R43XaOQTaifhfS6WdsuUzRA6C0Ee+NKvMHCIEBQWLNDA2GhYQWiR4A0b9i4OGAdOZWEKVVhjpSFAWhpjqz4xFG2X0AoVAo+XYWvm6jYxDN6/U+tHg8Hni9XixfvhylpaUouN2XTHfTIcGKQgI8xCDTxsbGMDY2prk+OzsbgUAAJSUlKC4uxst3b+jplkEhAQYE2S4SiaC1tRWtra0AAJ/Ph9raWrT7V+HE+HCi5jHYERIGK+WUdfpQ3H/mOQhyXE9PD+rr63H97few9+IQRldXJGrC8xOCMCBImHA4jB07dmDNmjVQz/SLHg7FwYAg4QYHB6GqKl7Z+Rl+da+ab1NWEQ5jQJA0zp07h8LCQjR3z3tegiHhIAYESaexsRGv7v16vk0YEg5hQJCUTp48yfMSEmBAkLRUVcW3Y49qrWYV4QAGBEmtsrISZx9ZJnoYaYsBQdKrq6vD0Ir18VaxirAZ76Qk6YXDYQSDQaB83kug9tG4yzAdMCBSgKIowvYdjUYfWuwSDAZx/d2LeP5m1+xV9tyKTQAYEClDlpm0Fx9fga6uLly6dAkdHR0YGhqypN9oNPqginjjJUv6I314DoIsVTzZj/fzXfiqZgO2d57Avn37UFBQYEnfLS0t6HHlW9IX6cOAINt8ePMKPvDloPfTj7B161bT/d27dw8HDx6Mt4onK23CgCBHfF5ViuPHj5vu59SpUxaMhvRiQKSuDIeXhN544k+0tbWZelF9fX3o9aw21YeFnP4ZO77wJCVZZWZIaJb8b7rvQlVVqKpqeEft7e3Ai08bbg8grS9dJoMVBNlh3qqiuroaixcvNtx5R0eH4baUHAYE2SluSOQO/ozq6mrDnba3txtuS8lhQJDd4oZEWVmZqU5H8jeYak/6MCDICXNC4sd8c+cQIpGIqfakDwOChPhi9HdTt4iPjIzo3VTolZuFzu6rGLyBhaZlYNbvQ05OjuHOIpEIsMTskCgRXuYkYcwHRJbu7WPKZn5YJSEj9E0GwEMMEsjMpc7JyUkLR0JaWEGQMKOjo4bbZmXprx5mWnRkj+F9poO/39r50N8tC4hQKAQswAfPTguJHkDqm1Pimw+IKTPjIR14iEHC9Pcb/9bq7OxsC0dCWhgQ5IQ51UNfVgEuX75suEOjhxiUHAYE2S3u1QOz8ym+vB811Z70YUCQnTQvLR49etRwp263Gx/f4kN1nMCAILtohkPb+GO4cOGC4Y4rKioMt6XkMCDISrEZS1xVTy1Dc3OzqZ2Ul5ebak/6MSDIiJjGktDN+l2mqgeAFYSTeKNU6pLq1uJrS4qwceNGhAYGTPXj9/vxzI2fLBoVJcIKgmz1x3N+7O+7jaKiIgyYDAcACAQCFoyK9GIFQZbrceWjq6sLnZ2deOH7OkxNWXPHY15eHmLvvA4Mh2evSoup1yIwIFKEekbMZb/Zj97r7u6G784dW/bV0NCAbXPDgWzEgEgBoVDowVyYFOb3+7Ft5ZPC9j97ElO64DkIWhAaGhq0VvHwwkaWVRCKokjzAFkjlN0HUv5TeKFqbGxEpfcv0cMA8P8XqaQLVhAktaqqKjS/5tNanVZvVhEYECQtj8eDvobNooeR1hgQJK3e3l70Tka1VrN6cIDdASH84aPzLCQpl8uFq1evYtlvZ7U24f+fQ1hBkFQ2bdqE6HctWHnrF9FDIfA+CJKEy+VCU1MTthcm/Co5Vg8OYkCQUC6XC1u2bEFtbS180SuJNmc4OIwBQULk5uaipqYGNTU1eHbgPPYzHKTEgCDHLF26FOvWrYOiKPihJA+7xoexa+C8nqYMB0EYEGQbt9uNsrIyBAIBrF27FkUT13AYwGEAGB/W0wWDQTAGRApQFMXUk7LNyMzMhMfjgcfjgdfr/e/Pn0wNIxgZwDEAxwBg4loy3TIYJMGASBHq+hWCR3AfQASYiAAThjthMEiGAUEyYDBIigFBIjAQFggGROrim9AGMWWzVF8GbJfpae281ZqINLGCINJh0ZE9oofgiNlfrccKgog0MSCISBMDgog0MSCISBMDgog0MSCISBMDgog0MSCISBMDgog0MSCISBMDgog0MSCISFMGTh+yc/qqzFOOZXjdMowhLaTLNG2rcLo3ESXE6d6UFqY/ESk5rCCISNM/WRyvzGkPNrEAAAAASUVORK5CYII="
HPE_FAVICON_B64 = "iVBORw0KGgoAAAANSUhEUgAAAEAAAABACAYAAACqaXHeAAAGqUlEQVR4nO1aTWhc1xX+zr3vTz94JEuezUzBdrEVPKS0qY2zKTJu0lgmWOAybRbFeGGyCFGXXRRixotACS4txmjZRUspOKilrYNEvRlR6oUw04WCghMhCyVVHMsjSxppRvN+7sli3kxmxvOnQrlQvw8eb+bec8/57jfn3Hfv8IiZ8SJD6CagG5EAugnoRiSAbgK6EQmgm4BuRALoJqAbkQC6CehGJIBuAroRCaCbgG5EAugmoBuRALoJ6EYkgG4CuvHCC2DUfyEiC8BRAGRZFs6fP786OztbbjeYiIiZOZFIjKyvr48CwPDwcGlzc3Ot2nf69On+xcXFb7muywCoAxeemJjAtWvXCpcvX/6y3n9TzASAQQDd/LWMAWCVmd1vWphrF4CXAOwCKNu2XT5z5sx3w3ZRb1dnbzAzksnke0KIMoDy8PDwP8M+k5mRSqXOWZZVBlACUAaw3+oiov2BgYH9I0eObJw8efLelStXfhL6ofAuw/tM6KfYyV+bqwDgpfo5NGRAqKgDQIbi9KQwMxsArPCzXd+nlJLVvi4+sLe3h729PXtjY+O19fX11x48ePA6M79NRALf/Npm6O+/yQAbTWXfLAAABM1GPYDDCwBUfQcRVftYSkmO42zX2daglAIRDQZBYJbLZd7d3fWXl5evTU5OfsnM11OpVJWrD8ADENi2rQzDKDMziLprQUQeEfn1ba0EIBxc2eq4Tn3kOE7p4sWLr46Ojj5xXZcsy6oJ8fTpU4rFYqPLy8sTuVzu/UKh0Oe6brCwsDC1uLj4m6WlpS0AICGGicgEYMbj8T9funTp7bV83hgaGfHbxK7Bsize+eKL7fq2VgLUEASBJCIDgCAi1cJEEhESiURPGcPMXCqV8tPT08/amGwC+HRychJzc3O/dV3Xffbs2dD169dfZebZTCYjXh4b+50keR9CwJHm/du3b+d7id0OHQVIJpNbuVyuk7I+ABw9erTYa0DP80yq5CuhqRTGx8fl/Pw8pqamZmzb/sB1XSNcrFIAZrPZrLi3tPSXOGBuA/g74Lz373sneo3ts2Lb3lnLpNK1p0BbATzPw+zs7DQRbXXwSQDYNM0UMyv0sHYYhiHS6bR48uQJxePxBgFWVlYIgOrv7zfCmmYAHARBAADZbNYfmfnVn8TgwJtqtwgSAmidmS2YEsDsMQffy6TwSY1PO3ulFJRSr/fi2/M8oLJ4dsXu7u723bt329kGADA+Pv7TUqlkAXCllNI0zY+rBpvuPsMlwNtXAAsoRc/nUgsQAWCC12jYsQT+BxDxePz81atX867rUviEgJSSfN/H48ePD/m+/+NcLvcz3/cDAHJoaOirmzdvLhCRYGZFQjCRYBApx3LEoNNX2Sd0W7aJAC+wAhU0WLYVwDAMWJb162Kx+B+0qNfqhACogYGBN4rF4o+6vGyhSqWSMzc399dWjyxmhuu6cF0XzBwAYMdxjBMnTvzi+PHj20inLQBuaExsmzJpDf7t3ZNnP8iXi3JAml0zUDFzaS+/2jDPdsaGYeDChQvTMzMzK90cHzt2zFpbW3uj29smzIxCodA9WYnkoUOHMDY29v78/PzviUjgzp3aBJkZMA3acsufT3379L+68euEtgIwMx49enSYiNYQ/tJtxvuJRGKgl2BCCMRiMW6VAUQEy7KUUmo1Ho/fP3fu3B9v3br1j1rqf/ihbCIIWXlEg5ARjExvi2GLCbSFlDJgZr9KohVpZvaTyWQvwUVfX18pnU7/8PDhwxtBEJCUspYNtm3j7NmzPDExsYbKTg9EJMNyeB7MkIKcdxY/Gn7T/b58p/RR1xIoDnpc/Mwr3Emna7atBKjf1h4EncYwKoca9fDhw0+y2exWJ0dEJNPpNNpOXgjGXknlZfDWHz77+FJlDexCmQj4ij0y+AcAPq02txJA4uBb4frtc8NeIDxQVftkLBbrz2QyO0tLS3Tq1KnnWN+4cYPbThyopF2lhkQJykbT4avDOMAPABU0zLlZAEbl2GiGsXrKhPCA4YafG/4/EEIE1T4A+8wcZDIZ1eqsDwCZTKZzMGYPlfN8QAFLUI/v+SkCiDz2GteyZgFWALyC8Aw+MjKyWon5fP2HCABAKTWtlLoTtpXCuw8AjuMsAPgOAEgp1c7OTj70eaAy47BuWcifI+BfghVDygNkKgNBwDCxWt9KB+Txf4cX/j/BSADdBHQjEkA3Ad2IBNBNQDciAXQT0I1IAN0EdCMSQDcB3YgE0E1ANyIBdBPQjUgA3QR0IxJANwHdeOEF+BoS3gcIsJb5+AAAAABJRU5ErkJggg=="
HPE_LOGO_PATH = Path("knowledge_base_data") / "hpe_kb_logo_v4.png"
HPE_FAVICON_PATH = Path("knowledge_base_data") / "hpe_kb_favicon_v4.png"
HPE_LOGO_PATH.parent.mkdir(exist_ok=True)
try:
    HPE_LOGO_PATH.write_bytes(base64.b64decode(HPE_KB_LOGO_B64))
    HPE_FAVICON_PATH.write_bytes(base64.b64decode(HPE_FAVICON_B64))
except Exception:
    pass

st.set_page_config(
    page_title="HPE Knowledge Base",
    page_icon=str(HPE_FAVICON_PATH) if HPE_FAVICON_PATH.exists() else "◈",
    layout="wide",
    initial_sidebar_state="collapsed",
)

BASE_DIR = Path("knowledge_base_data")
PDF_DIR = BASE_DIR / "pdfs"
IMAGE_DIR = BASE_DIR / "images"
PDF_IMAGE_DIR = IMAGE_DIR / "pdf"
SOP_IMAGE_DIR = IMAGE_DIR / "sop"
SOP_VIDEO_DIR = BASE_DIR / "videos"
DB_PATH = BASE_DIR / "knowledge_base.db"

SUPPORTED_KB_FILES = {
    ".pdf": "PDF",
    ".xlsx": "Excel",
    ".xls": "Excel",
    ".docx": "Word",
    ".doc": "Word",
    ".pptx": "PowerPoint",
    ".ppt": "PowerPoint",
    ".mp4": "Video",
    ".webm": "Video",
    ".mov": "Video",
    ".m4v": "Video",
    ".avi": "Video",
}
RESET_MARKER = BASE_DIR / ".knowledge_base_v2_reset_complete"

BASE_DIR.mkdir(exist_ok=True)
PDF_DIR.mkdir(exist_ok=True)
IMAGE_DIR.mkdir(exist_ok=True)
PDF_IMAGE_DIR.mkdir(exist_ok=True)
SOP_IMAGE_DIR.mkdir(exist_ok=True)
SOP_VIDEO_DIR.mkdir(exist_ok=True)


# ============================================================
# SIX PRIMARY PRODUCT GROUPS
# The home page intentionally has EXACTLY the six tile groups
# visible in the supplied reference image. Each group opens a
# complete product/topic/source directory.
# ============================================================
PRODUCT_GROUPS = {
    "Compute": {
        "icon": "▣",
        "class": "mint",
        "description": "Servers, iLO, Gen10/Gen11, firmware, hardware guides",
        "products": [
            "HPE ProLiant Compute",
            "HPE ProLiant Gen10",
            "HPE ProLiant Gen11",
            "HPE ProLiant Gen12",
            "HPE iLO",
            "HPE Compute Ops Management",
            "HPE OneView",
            "HPE Superdome Flex",
        ],
        "topics": [
            "Server Setup & Configuration",
            "iLO Configuration",
            "BIOS / UEFI",
            "Firmware Updates",
            "Hardware Health",
            "POST & Boot Troubleshooting",
            "Storage Controllers",
            "NIC & Network Adapters",
            "Power & Thermal",
            "Drivers",
            "Lifecycle Management",
            "Server Monitoring",
        ],
        "sources": [
            ("HPE Compute", "https://www.hpe.com/us/en/compute.html"),
            ("HPE ProLiant", "https://www.hpe.com/us/en/compute/proliant.html"),
            ("HPE Support Center", "https://support.hpe.com/"),
            ("HPE Documentation", "https://www.hpe.com/info/docs"),
        ],
    },
    "Networking": {
        "icon": "⌘",
        "class": "blue",
        "description": "Aruba, switches, controllers, access points, routing",
        "products": [
            "HPE Networking",
            "HPE Aruba Networking",
            "Aruba CX Switches",
            "Aruba Access Points",
            "Aruba Gateways",
            "Aruba Central",
            "Aruba AirWave",
            "Juniper Networking",
        ],
        "topics": [
            "Switch Setup",
            "VLAN Configuration",
            "Trunking",
            "Routing",
            "STP",
            "Port Troubleshooting",
            "Wireless Client Connectivity",
            "SSID Configuration",
            "Access Point Troubleshooting",
            "Aruba Central",
            "Gateway & SD-Branch",
            "Network Monitoring",
        ],
        "sources": [
            ("HPE Networking", "https://www.hpe.com/us/en/products/networking.html"),
            ("HPE Aruba Networking", "https://www.hpe.com/us/en/hpe-aruba-networking.html"),
            ("Aruba Networking Documentation", "https://arubanetworking.hpe.com/techdocs/"),
            ("HPE Support Center", "https://support.hpe.com/"),
        ],
    },
    "Storage": {
        "icon": "▤",
        "class": "violet",
        "description": "Alletra, MSA, Primera, StoreOnce, data protection",
        "products": [
            "HPE Alletra Storage",
            "HPE Alletra Storage MP B10000",
            "HPE Alletra Storage MP X10000",
            "HPE Alletra 5000",
            "HPE Alletra 6000",
            "HPE Alletra 9000",
            "HPE MSA",
            "HPE Primera",
            "HPE StoreOnce",
        ],
        "topics": [
            "Storage Setup",
            "Host Connectivity",
            "Volume Mapping",
            "Multipathing",
            "Capacity",
            "Performance",
            "Snapshots",
            "Replication",
            "Data Protection",
            "Storage Networking",
            "Alerts & Health",
            "Backup & Recovery",
        ],
        "sources": [
            ("HPE Alletra Storage", "https://www.hpe.com/us/en/products/storage/alletra.html"),
            ("HPE Storage", "https://www.hpe.com/us/en/storage.html"),
            ("HPE Support Center", "https://support.hpe.com/"),
            ("HPE Documentation", "https://www.hpe.com/info/docs"),
        ],
    },
    "Software & Licensing": {
        "icon": "▥",
        "class": "orange",
        "description": "Licensing, subscriptions, entitlements and software",
        "products": [
            "HPE Software",
            "HPE GreenLake Services",
            "HPE iLO Licensing",
            "HPE OneView",
            "HPE Compute Ops Management",
            "HPE Aruba Networking Subscriptions",
            "HPE Storage Software",
            "HPE Data Services",
        ],
        "topics": [
            "License Activation",
            "License Renewal",
            "Entitlements",
            "Subscriptions",
            "License Assignment",
            "License Consumption",
            "Activation & Registration",
            "Portal Access",
            "Software Downloads",
            "Subscription Management",
            "Licensing Errors",
            "Support Entitlements",
        ],
        "sources": [
            ("HPE Software", "https://buy.hpe.com/us/en/software"),
            ("HPE GreenLake", "https://www.hpe.com/us/en/greenlake.html"),
            ("HPE Support Center", "https://support.hpe.com/"),
            ("HPE Store", "https://buy.hpe.com/"),
        ],
    },
    "Security": {
        "icon": "♢",
        "class": "red",
        "description": "ClearPass, NAC, policy management, zero trust",
        "products": [
            "HPE Aruba Networking ClearPass",
            "ClearPass Policy Manager",
            "ClearPass Guest",
            "ClearPass OnGuard",
            "Aruba Policy Manager",
            "Network Access Control",
            "HPE Aruba Networking SSE",
            "HPE Security",
        ],
        "topics": [
            "ClearPass Troubleshooting",
            "RADIUS",
            "802.1X",
            "MAC Authentication",
            "Role Mapping",
            "Enforcement Policies",
            "Access Tracker",
            "Guest Access",
            "OnGuard",
            "Certificates",
            "Zero Trust",
            "SASE",
        ],
        "sources": [
            ("ClearPass Documentation", "https://arubanetworking.hpe.com/techdocs/ClearPass/"),
            ("HPE Aruba Networking", "https://www.hpe.com/us/en/hpe-aruba-networking.html"),
            ("HPE Security", "https://www.hpe.com/us/en/security.html"),
            ("HPE Support Center", "https://support.hpe.com/"),
        ],
    },
    "Support & Tools": {
        "icon": "⚙",
        "class": "slate",
        "description": "HPE tools, portals, utilities, support and best practices",
        "products": [
            "HPE Support Center",
            "HPE InfoSight",
            "HPE GreenLake",
            "HPE Service Portal",
            "HPE OneView",
            "HPE iLO",
            "HPE Active Health System",
            "HPE Smart Update Manager",
            "HPE Documentation",
        ],
        "topics": [
            "Case Troubleshooting",
            "Support Case Management",
            "Technical Escalation",
            "Documentation Search",
            "Knowledge Articles",
            "Logs & Diagnostics",
            "Health Monitoring",
            "Support Entitlements",
            "Warranty & Care Packs",
            "Best Practices",
            "Customer Communication",
            "Problem Management",
        ],
        "sources": [
            ("HPE Support Center", "https://support.hpe.com/"),
            ("HPE Documentation", "https://www.hpe.com/info/docs"),
            ("HPE Services", "https://www.hpe.com/us/en/services.html"),
            ("HPE GreenLake", "https://www.hpe.com/us/en/greenlake.html"),
        ],
    },
}


# ============================================================
# AI-READY SEED RECORDS
# One question = one self-contained answer.
# ============================================================
SEED_KB = [
    ("CPPM-001","Security","ClearPass Policy Manager",
     "How do I troubleshoot a ClearPass authentication failure?",
     "Check the request, selected ClearPass service, authentication source, identity result, role mapping, and enforcement result. Start with the exact failure reason before changing configuration.",
     "1. Identify the affected request and timestamp.\n2. Check which service processed the request.\n3. Confirm the authentication source is reachable.\n4. Review the exact authentication failure reason.\n5. If authentication succeeds, inspect role mapping and enforcement.\n6. Retest with a controlled request.",
     "ClearPass, authentication failure, service, authentication source, role mapping, enforcement"),
    ("CPPM-002","Security","ClearPass Policy Manager",
     "What is the basic ClearPass authentication flow?",
     "A typical ClearPass flow receives the authentication request, selects the applicable service, validates the identity through the configured authentication source, applies authorization and role mapping logic, and returns the enforcement result to the network device.",
     "1. Identify the network access device.\n2. Identify the authentication method.\n3. Confirm the request matches the intended service.\n4. Confirm the authentication source is available.\n5. Review role mapping and enforcement.\n6. Confirm the network device applies the result.",
     "ClearPass flow, authentication, authorization, service, role mapping, enforcement"),
    ("CPPM-003","Security","ClearPass Role Mapping",
     "What is role mapping used for in ClearPass?",
     "Role mapping assigns one or more roles to an authenticated identity based on configured conditions and attributes. The assigned role can then be used by enforcement logic to determine access.",
     "1. Confirm authentication succeeds.\n2. Review returned identity attributes.\n3. Check role-mapping conditions.\n4. Identify the assigned role.\n5. Confirm the enforcement policy uses that role.",
     "ClearPass role mapping, identity, role, authorization"),
    ("CPPM-004","Security","ClearPass Enforcement",
     "What should I check when ClearPass authenticates a user but gives the wrong VLAN or network access?",
     "Check the selected service, returned identity attributes, role mapping, enforcement policy, and final attributes returned to the network device.",
     "1. Open the affected request.\n2. Confirm the selected service.\n3. Confirm authentication succeeded.\n4. Review assigned roles.\n5. Review enforcement.\n6. Check the attributes returned to the network device.",
     "wrong VLAN, wrong access, ClearPass enforcement, role mapping"),
    ("CPPM-005","Security","ClearPass RADIUS",
     "What should I check when ClearPass does not receive a RADIUS authentication request?",
     "Check reachability between the network access device and ClearPass, RADIUS server configuration, shared secret, configured RADIUS ports, firewall or ACL rules, and whether the network device is sending requests.",
     "1. Verify the ClearPass server address.\n2. Verify the RADIUS shared secret.\n3. Check expected authentication and accounting ports.\n4. Check firewall and ACL rules.\n5. Generate a controlled authentication test.\n6. Confirm whether the request appears in ClearPass request tracking.",
     "RADIUS, ClearPass, no request, shared secret, UDP, firewall"),
    ("LIC-001","Software & Licensing","License Activation",
     "How do I troubleshoot an HPE software license that is not showing?",
     "Verify the product entitlement, account or organization, license identifier, activation or registration state, quantity, and whether the correct portal or workspace is being viewed.",
     "1. Confirm the exact product and license ID.\n2. Confirm the associated account or organization.\n3. Check entitlement status.\n4. Check expected quantity or capacity.\n5. Capture any activation or entitlement error.",
     "license, entitlement, activation, registration, missing license"),
    ("LIC-002","Software & Licensing","License Renewal",
     "What information is needed for an HPE licensing escalation?",
     "Provide the product name, entitlement or license identifier, customer organization, affected account, portal or workspace, exact error, timestamp, expected result, actual result, and supporting evidence.",
     "1. Record the license or entitlement ID.\n2. Capture exact error text.\n3. Record account and organization context.\n4. State expected versus actual behavior.\n5. Attach evidence without exposing secrets.",
     "licensing escalation, entitlement ID, license ID"),
    ("ILO-001","Compute","HPE iLO",
     "What is HPE iLO used for?",
     "HPE Integrated Lights-Out provides remote management capabilities for supported HPE ProLiant servers, including remote hardware monitoring and management functions. Exact capabilities depend on server generation, iLO version, licensing, and configuration.",
     "1. Identify the server generation and iLO version.\n2. Review system health and alerts.\n3. Review management and event logs.\n4. Use supported remote-management functions appropriate to the task.",
     "iLO, remote management, ProLiant, hardware health"),
    ("SRV-001","Compute","HPE ProLiant",
     "What should I check when an HPE ProLiant server does not boot?",
     "Check power state, hardware health indicators, POST behavior, boot-device selection, storage visibility, recent hardware or firmware changes, and system-management logs.",
     "1. Confirm power.\n2. Observe POST or boot messages.\n3. Check hardware health alerts.\n4. Confirm the intended boot device is detected.\n5. Review recent hardware or firmware changes.\n6. Capture the exact error or diagnostic code.",
     "ProLiant boot, POST, boot device, hardware health"),
    ("FW-001","Compute","Firmware Updates",
     "What information should I collect before an HPE server firmware update?",
     "Collect the exact server model and generation, current firmware versions, target firmware version, compatibility requirements, maintenance window, recovery readiness, and rollback or recovery plan.",
     "1. Record current firmware versions.\n2. Confirm the target package applies to the exact model.\n3. Review current applicable HPE documentation.\n4. Confirm the maintenance window.\n5. Confirm recovery planning.",
     "firmware update, compatibility, maintenance window, rollback"),
    ("NET-001","Networking","Aruba CX Switching",
     "What should I check when an Aruba switch port has no connectivity?",
     "Check physical link state, administrative state, port configuration, VLAN assignment, authentication state, interface errors, and the connected endpoint.",
     "1. Check cable and physical link indicators.\n2. Confirm the port is enabled.\n3. Confirm access or trunk configuration.\n4. Confirm the expected VLAN.\n5. Check authentication if applicable.\n6. Review interface error counters.",
     "Aruba switch, CX, port down, VLAN, interface errors"),
    ("NET-002","Networking","VLAN",
     "How should I troubleshoot an Aruba VLAN connectivity problem?",
     "Verify the VLAN exists, the endpoint port assigns or carries the VLAN correctly, trunks carry it where required, the gateway interface exists and is reachable, and routing or ACL policy permits the traffic.",
     "1. Identify source and destination VLANs.\n2. Verify VLAN existence.\n3. Check access-port or trunk VLAN configuration.\n4. Check the gateway or SVI.\n5. Test gateway reachability.\n6. Test the destination.\n7. Check ACL or security policy.",
     "VLAN, Aruba, trunk, SVI, gateway, ACL"),
    ("NET-003","Networking","Wireless",
     "What should I check when an Aruba wireless client cannot connect to an SSID?",
     "Check SSID advertisement, client visibility, authentication method, security credentials, client eligibility, DHCP availability, and whether the issue affects one client or many clients.",
     "1. Confirm the SSID is enabled.\n2. Check whether the client sees it.\n3. Review authentication or security failure details.\n4. Check DHCP and IP assignment.\n5. Compare with a known working client.",
     "Aruba wireless, SSID, client, authentication, DHCP"),
    ("STOR-001","Storage","HPE Alletra",
     "What should I check when a host cannot access an HPE Alletra storage volume?",
     "Check host-to-storage connectivity, initiator and target configuration, zoning or network path, volume presentation or mapping, host multipathing, and storage-side health.",
     "1. Identify host and volume.\n2. Verify host initiator identity.\n3. Verify network or fabric connectivity.\n4. Verify volume mapping or presentation.\n5. Check multipath status.\n6. Check storage alerts and capacity.",
     "Alletra, volume, host connectivity, multipath, mapping"),
    ("STOR-002","Storage","Storage Performance",
     "What should I check when HPE storage performance is slow?",
     "Determine whether the bottleneck is at the host, network or fabric, storage system, volume, workload, or application layer by comparing latency, throughput, I/O rate, path health, and resource utilization.",
     "1. Record the affected workload and time window.\n2. Check host I/O.\n3. Check network or fabric errors.\n4. Check storage latency and throughput.\n5. Compare affected and unaffected workloads.\n6. Correlate with recent changes or unusual load.",
     "storage latency, IOPS, throughput, performance"),
    ("GL-001","Support & Tools","HPE GreenLake",
     "What should I check when a user cannot access an HPE GreenLake service?",
     "Verify the user account, authentication method, organization or workspace membership, assigned permissions, service entitlement, and whether the issue is account-specific or organization-wide.",
     "1. Confirm authentication.\n2. Confirm organization or workspace membership.\n3. Check assigned permissions.\n4. Confirm the service entitlement.\n5. Compare with a known working user when appropriate.",
     "GreenLake, access, workspace, organization, permissions"),
    ("SEC-001","Security","Security",
     "What is the safest approach when troubleshooting a security-sensitive access problem?",
     "Collect the minimum required evidence, avoid exposing credentials or secrets, preserve relevant logs, use least-privilege actions, and follow the approved security and escalation process.",
     "1. Never request passwords or private keys in case notes.\n2. Redact sensitive screenshots.\n3. Preserve timestamps and logs.\n4. Use approved administrative access.\n5. Escalate suspected security incidents.",
     "security, least privilege, credentials, secrets, escalation"),
    ("KB-001","Support & Tools","Knowledge Management",
     "How should an SOP be written so an AI can retrieve an exact answer?",
     "Use one clear question per knowledge record, followed immediately by a direct answer. Add product and topic scope, numbered steps, verification, escalation criteria, and searchable keywords. Keep the record self-contained.",
     "1. Write one question per record.\n2. Put the direct answer first.\n3. Use exact product and feature names.\n4. Keep scope explicit.\n5. Add numbered steps when needed.\n6. Add keywords and metadata.",
     "AI retrieval, RAG, atomic SOP, exact answer, knowledge chunk"),
]


# ============================================================
# DATABASE
# ============================================================
def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    # The old database is deliberately deleted ONE TIME.
    # PDFs are preserved.
    if not RESET_MARKER.exists() and DB_PATH.exists():
        DB_PATH.unlink()

    conn = db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            filename TEXT,
            family TEXT,
            topic TEXT,
            source_url TEXT,
            content TEXT NOT NULL,
            doc_type TEXT DEFAULT 'PDF',
            created_at TEXT NOT NULL
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kb_id TEXT UNIQUE NOT NULL,
            family TEXT NOT NULL,
            topic TEXT NOT NULL,
            question TEXT NOT NULL,
            answer TEXT NOT NULL,
            steps TEXT,
            keywords TEXT,
            source TEXT DEFAULT 'Built-in AI Knowledge',
            created_at TEXT NOT NULL
        )
    """)
\
    conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_images (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kb_id TEXT,
            document_id INTEGER,
            filename TEXT NOT NULL,
            image_path TEXT NOT NULL,
            page_num INTEGER,
            placement TEXT,
            caption TEXT,
            created_at TEXT NOT NULL
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_videos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kb_id TEXT NOT NULL,
            filename TEXT NOT NULL,
            video_path TEXT NOT NULL,
            placement TEXT,
            caption TEXT,
            created_at TEXT NOT NULL
        )
    """)

    if conn.execute("SELECT COUNT(*) FROM kb_records").fetchone()[0] == 0:
        now = datetime.now().isoformat(timespec="seconds")
        conn.executemany("""
            INSERT OR IGNORE INTO kb_records
            (kb_id, family, topic, question, answer, steps, keywords, source, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, [
            (a,b,c,d,e,f,g,"Built-in AI Knowledge",now)
            for a,b,c,d,e,f,g in SEED_KB
        ])

    conn.commit()
    conn.close()

    RESET_MARKER.write_text(
        "New HPE Knowledge Base database initialized.\n",
        encoding="utf-8"
    )


def ensure_hpe_kb_usage_sop():
    """Add the built-in Knowledge Base usage SOP without disturbing existing data."""
    now = datetime.now().isoformat(timespec="seconds")
    kb_id = "KB-USE-001"
    question = "How do I use the HPE Knowledge Base?"
    answer = (
        "Use the HPE Knowledge Base by asking the AI search bar a clear technical question, "
        "choosing the most relevant suggested answer, reviewing the Summary, following the "
        "Troubleshooting Steps, and opening Related Knowledge when more context is needed. "
        "You can also select any of the six HPE & Aruba Product Family tiles to browse structured "
        "knowledge. Admin users can create AI-ready SOPs and attach images to specific procedure steps."
    )
    steps = (
        "1. Start with the AI search bar and enter a clear question.\n"
        "2. Press Enter or the teal arrow to submit the question.\n"
        "3. If several possible answers appear, select the answer that best matches the issue.\n"
        "4. Review Summary for the direct answer and key points.\n"
        "5. Select Troubleshooting Steps and follow the numbered procedure.\n"
        "6. Review Related Knowledge to open connected knowledge records or source documents.\n"
        "7. Use the Compute, Networking, Storage, Software & Licensing, Security, or Support & Tools tiles to browse by product family.\n"
        "8. Admin users can create an SOP, upload screenshots or diagrams, and place each image before or after a selected step.\n"
        "9. Save the SOP so it becomes searchable by the AI Assistant.\n"
        "10. Validate production procedures against the current authoritative HPE documentation for the applicable product and version."
    )
    keywords = "HPE Knowledge Base, AI search, exact answer, Summary, Troubleshooting Steps, Related Knowledge, product families, SOP, images"
    conn = db()
    conn.execute("""
        INSERT OR IGNORE INTO kb_records
        (kb_id, family, topic, question, answer, steps, keywords, source, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (kb_id, "Support & Tools", "Knowledge Base User Guide", question, answer, steps, keywords, "Built-in HPE Knowledge Base SOP", now))
    conn.commit()
    conn.close()
    return kb_id


init_db()
ensure_hpe_kb_usage_sop()


# ============================================================
# DATA
# ============================================================
def load_records():
    conn = db()
    rows = conn.execute("""
        SELECT kb_id, family, topic, question, answer, steps, keywords, source
        FROM kb_records ORDER BY id DESC
    """).fetchall()
    conn.close()
    return [dict(x) for x in rows]


def load_documents():
    conn = db()
    rows = conn.execute("""
        SELECT id, title, filename, family, topic, source_url, content, doc_type, created_at
        FROM documents ORDER BY id DESC
    """).fetchall()
    conn.close()
    return [dict(x) for x in rows]


def index_pdf(uploaded_file, family, topic):
    """Index PDF text and extract embedded PDF images for Related Knowledge."""
    if fitz is None:
        return False, "PyMuPDF is not installed. Add pymupdf to requirements.txt."

    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", uploaded_file.name)
    pdf_path = PDF_DIR / safe
    pdf_path.write_bytes(uploaded_file.getbuffer())

    try:
        pdf = fitz.open(pdf_path)
        page_text = []
        extracted_images = []

        for page_number, page in enumerate(pdf, start=1):
            page_text.append(page.get_text("text") or "")

            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = image_info[0]
                try:
                    image_data = pdf.extract_image(xref)
                    image_bytes = image_data.get("image")
                    ext = image_data.get("ext", "png")
                    if not image_bytes:
                        continue

                    image_name = (
                        f"{Path(safe).stem}_p{page_number}_img{image_index}.{ext}"
                    )
                    image_path = PDF_IMAGE_DIR / image_name
                    image_path.write_bytes(image_bytes)

                    extracted_images.append(
                        (page_number, image_name, str(image_path), f"Page {page_number}")
                    )
                except Exception:
                    # One unsupported/invalid embedded image must not prevent
                    # the rest of the PDF from being indexed.
                    continue

        text = "\n".join(page_text).strip()

        if not text and not extracted_images:
            pdf.close()
            return False, "The PDF contains no extractable text or embedded images."

        conn = db()
        cursor = conn.execute("""
            INSERT INTO documents
            (title, filename, family, topic, source_url, content, doc_type, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            Path(uploaded_file.name).stem,
            uploaded_file.name,
            family,
            topic,
            "",
            text,
            "PDF",
            datetime.now().isoformat(timespec="seconds")
        ))
        document_id = cursor.lastrowid

        now = datetime.now().isoformat(timespec="seconds")
        conn.executemany("""
            INSERT INTO kb_images
            (kb_id, document_id, filename, image_path, page_num, placement, caption, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, [
            (
                None,
                document_id,
                filename,
                image_path,
                page_num,
                placement,
                f"{Path(uploaded_file.name).stem} — {placement}",
                now
            )
            for page_num, filename, image_path, placement in extracted_images
        ])

        conn.commit()
        conn.close()
        pdf.close()

        search.clear()
        return True, (
            f"Indexed {uploaded_file.name}"
            f"{f' and extracted {len(extracted_images)} image(s)' if extracted_images else ''}"
        )
    except Exception as e:
        try:
            pdf.close()
        except Exception:
            pass
        return False, str(e)


def load_kb_images(kb_id=None, document_ids=None, limit=8):
    """Return images attached to an SOP and/or matching PDF documents."""
    conn = db()
    clauses = []
    params = []

    if kb_id:
        clauses.append("kb_id = ?")
        params.append(kb_id)

    if document_ids:
        placeholders = ",".join("?" for _ in document_ids)
        clauses.append(f"document_id IN ({placeholders})")
        params.extend(document_ids)

    if not clauses:
        conn.close()
        return []

    rows = conn.execute(
        f"""
        SELECT id, kb_id, document_id, filename, image_path,
               page_num, placement, caption, created_at
        FROM kb_images
        WHERE {" OR ".join(clauses)}
        ORDER BY
            CASE WHEN kb_id = ? THEN 0 ELSE 1 END,
            page_num ASC,
            id ASC
        LIMIT ?
        """,
        params + [kb_id or "", limit]
    ).fetchall()
    conn.close()
    return [dict(x) for x in rows]


def save_sop_images(kb_id, uploaded_images, placements):
    """Persist admin-uploaded SOP images with an exact step placement."""
    if not uploaded_images:
        return 0

    target_dir = SOP_IMAGE_DIR / kb_id
    target_dir.mkdir(parents=True, exist_ok=True)

    conn = db()
    now = datetime.now().isoformat(timespec="seconds")
    saved = 0

    for index, uploaded_image in enumerate(uploaded_images):
        try:
            safe = re.sub(
                r"[^A-Za-z0-9_.-]+",
                "_",
                uploaded_image.name
            )
            filename = f"{index + 1:02d}_{safe}"
            path = target_dir / filename
            path.write_bytes(uploaded_image.getbuffer())

            placement = placements[index] if index < len(placements) else "End of SOP"

            conn.execute("""
                INSERT INTO kb_images
                (kb_id, document_id, filename, image_path, page_num, placement, caption, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                kb_id,
                None,
                filename,
                str(path),
                None,
                placement,
                f"{uploaded_image.name} — {placement}",
                now
            ))
            saved += 1
        except Exception:
            continue

    conn.commit()
    conn.close()
    return saved


def save_sop_videos(kb_id, uploaded_videos, placements=None):
    """Persist admin-uploaded SOP videos with an exact step placement."""
    if not uploaded_videos:
        return 0

    target_dir = SOP_VIDEO_DIR / kb_id
    target_dir.mkdir(parents=True, exist_ok=True)

    placements = placements or []
    conn = db()
    now = datetime.now().isoformat(timespec="seconds")
    saved = 0

    for index, uploaded_video in enumerate(uploaded_videos):
        try:
            safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", uploaded_video.name)
            filename = f"{index + 1:02d}_{safe}"
            path = target_dir / filename
            path.write_bytes(uploaded_video.getbuffer())

            placement = placements[index] if index < len(placements) else "End of SOP"
            conn.execute("""
                INSERT INTO kb_videos
                (kb_id, filename, video_path, placement, caption, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (
                kb_id,
                filename,
                str(path),
                placement,
                f"{uploaded_video.name} — {placement}",
                now
            ))
            saved += 1
        except Exception:
            continue

    conn.commit()
    conn.close()
    return saved


def load_sop_videos(kb_id, limit=8):
    conn = db()
    rows = conn.execute("""
        SELECT id, kb_id, filename, video_path, placement, caption, created_at
        FROM kb_videos
        WHERE kb_id = ?
        ORDER BY id ASC
        LIMIT ?
    """, (kb_id, limit)).fetchall()
    conn.close()
    return [dict(x) for x in rows]


def delete_sop_media(kb_id):
    """Delete stored SOP images/videos and their database rows."""
    conn = db()
    image_rows = conn.execute(
        "SELECT image_path FROM kb_images WHERE kb_id=?",
        (kb_id,)
    ).fetchall()
    video_rows = conn.execute(
        "SELECT video_path FROM kb_videos WHERE kb_id=?",
        (kb_id,)
    ).fetchall()

    conn.execute("DELETE FROM kb_images WHERE kb_id=?", (kb_id,))
    conn.execute("DELETE FROM kb_videos WHERE kb_id=?", (kb_id,))
    conn.commit()
    conn.close()

    for row in image_rows:
        try:
            Path(row["image_path"]).unlink(missing_ok=True)
        except Exception:
            pass
    for row in video_rows:
        try:
            Path(row["video_path"]).unlink(missing_ok=True)
        except Exception:
            pass


def update_ai_ready_sop(
    kb_id, family, product, topic, question, direct_answer,
    prerequisites, steps, verification, escalation, keywords,
    source_title="", source_url="", uploaded_images=None,
    image_placements=None, uploaded_videos=None, video_placements=None,
    replace_media=False
):
    """Update an existing atomic SOP while preserving its stable KB ID."""
    normalized_steps = steps.strip()
    if prerequisites.strip():
        normalized_steps = (
            "Prerequisites:\n" + prerequisites.strip()
            + "\n\nProcedure:\n" + normalized_steps
        )
    if verification.strip():
        normalized_steps += "\n\nVerification:\n" + verification.strip()
    if escalation.strip():
        normalized_steps += "\n\nEscalation:\n" + escalation.strip()

    source = source_title.strip() or "Admin-created SOP"
    if source_url.strip():
        source += f" | {source_url.strip()}"

    conn = db()
    conn.execute("""
        UPDATE kb_records
        SET family=?, topic=?, question=?, answer=?, steps=?, keywords=?, source=?
        WHERE kb_id=?
    """, (
        family,
        f"{product} — {topic}",
        question.strip(),
        direct_answer.strip(),
        normalized_steps,
        ", ".join([x.strip() for x in keywords.split(",") if x.strip()]),
        source,
        kb_id
    ))
    conn.commit()
    conn.close()

    if replace_media:
        delete_sop_media(kb_id)

    image_count = save_sop_images(kb_id, uploaded_images or [], image_placements or [])
    video_count = save_sop_videos(kb_id, uploaded_videos or [], video_placements or [])
    search.clear()
    return image_count, video_count


def delete_ai_ready_sop(kb_id):
    """Delete an SOP and all media attached to it."""
    delete_sop_media(kb_id)
    conn = db()
    conn.execute("DELETE FROM kb_records WHERE kb_id=?", (kb_id,))
    conn.commit()
    conn.close()
    search.clear()


def extract_office_text(uploaded_file, suffix):
    """Best-effort text extraction for searchable Office files."""
    raw = uploaded_file.getvalue()
    try:
        if suffix in {".docx"} and WordDocument:
            doc = WordDocument(__import__("io").BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        if suffix in {".xlsx", ".xls"} and load_workbook and suffix == ".xlsx":
            wb = load_workbook(__import__("io").BytesIO(raw), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None and str(v).strip()]
                    if vals:
                        parts.append(" | ".join(vals))
            return "\n".join(parts)

        if suffix in {".pptx"} and PowerPointPresentation:
            prs = PowerPointPresentation(__import__("io").BytesIO(raw))
            parts = []
            for slide_no, slide in enumerate(prs.slides, 1):
                parts.append(f"[Slide {slide_no}]")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
            return "\n".join(parts)
    except Exception:
        return ""
    return ""


def index_supported_file(uploaded_file, family, topic):
    """Store PDF/Office/video knowledge sources for Related Files and search."""
    suffix = Path(uploaded_file.name).suffix.lower()
    doc_type = SUPPORTED_KB_FILES.get(suffix)
    if not doc_type:
        return False, "Unsupported file type."

    if suffix == ".pdf":
        return index_pdf(uploaded_file, family, topic)

    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", uploaded_file.name)
    target_dir = BASE_DIR / "files" / doc_type.lower().replace(" ", "_")
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / safe
    path.write_bytes(uploaded_file.getbuffer())

    extracted = extract_office_text(uploaded_file, suffix)
    if doc_type == "Video":
        extracted = f"Video file: {uploaded_file.name}"

    conn = db()
    conn.execute("""
        INSERT INTO documents
        (title, filename, family, topic, source_url, content, doc_type, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        Path(uploaded_file.name).stem,
        uploaded_file.name,
        family,
        topic,
        "",
        extracted.strip() or uploaded_file.name,
        doc_type,
        datetime.now().isoformat(timespec="seconds")
    ))
    conn.commit()
    conn.close()
    search.clear()
    return True, f"Indexed {uploaded_file.name} as {doc_type}."


def file_data_uri(path):
    """Create a browser-openable data URI for a related file."""
    try:
        p = Path(path)
        if not p.exists():
            return None
        mime = mimetypes.guess_type(p.name)[0] or "application/octet-stream"
        encoded = base64.b64encode(p.read_bytes()).decode("ascii")
        return f"data:{mime};base64,{encoded}"
    except Exception:
        return None


def render_clickable_image(path, caption="", max_height=190):
    """Show a compact thumbnail; clicking it opens the actual image in an in-page modal."""
    data_uri = file_data_uri(path)
    if not data_uri:
        return False

    safe_caption = eh(caption)
    # Generate a deterministic, URL-safe fragment ID without adding another
    # browser page or requiring JavaScript.
    token = base64.urlsafe_b64encode(
        f"{path}|{caption}".encode("utf-8")
    ).decode("ascii").rstrip("=")
    modal_id = "kb_img_" + re.sub(r"[^A-Za-z0-9_-]", "_", token)[:80]

    st.markdown(
        f"""
        <div class="kb-image-preview">
          <a href="#{modal_id}" class="kb-image-open" title="View image at actual size">
            <img src="{data_uri}" alt="{safe_caption}"
                 class="kb-clickable-image"
                 style="max-height:{int(max_height)}px;">
          </a>
          <div class="kb-image-caption">{safe_caption}</div>
        </div>

        <div id="{modal_id}" class="kb-image-modal" aria-label="Image preview">
          <div class="kb-image-modal-backdrop">
            <a href="#" class="kb-image-modal-close" aria-label="Close image">×</a>
            <div class="kb-image-modal-content">
              <img src="{data_uri}" alt="{safe_caption}" class="kb-image-full">
              <div class="kb-image-modal-caption">{safe_caption}</div>
            </div>
          </div>
        </div>
        """,
        unsafe_allow_html=True
    )
    return True


# ============================================================
# BUILT-IN VISUALS FOR THE KNOWLEDGE BASE USER-GUIDE SOP
# ============================================================
def ensure_builtin_usage_visuals():
    """Create bundled UI reference images and attach them to KB-USE-001."""
    visuals = [
        ("related_knowledge_reference.png", 'iVBORw0KGgoAAAANSUhEUgAAA3sAAACUCAIAAAClVad+AAAQAElEQVR4AeydB2BVRdbHZ+5L770QQo3SpElVRJEiCGJh7a7iouiqiJ+66ioqq2vZta0F+9oVdde2igVQpPceOoROEtJ7ecl77/vP3OTdZwKSkERT/uPck3PPPXNm5ncfm7Pn3iRGaWUFOwmQAAmQAAmQAAmQAAk0IoENqQfyy+0F5eUF9opCe4Uh2EiABEiABH5/AlwBCZAACbQ2AlIIKaRUUjDjFGwkQAIkQAIkQAIkQAKNTsAlhAtNy5aTcTY6BgYkARIgARIgARIgARJoMgKquimlKZlxNhlmBiYBEiCBVkmAmyIBEiCBOhNwCRQ61cGMs87M6EgCJEACJEACJEACJFAfAlJIuONgxgkOjdwZjgRIgARIgARIgARIAARcaAI1ThczTtBgJwESIAESaIUEuCUSIIHfnYBEE6hzSmacgo0ESIAESIAESIAESKApCKi3OFWNk78dqSnotpiYXCgJkAAJkAAJkAAJNCEBqWNDssapSVCQAAmQAAmQwO9GgBOTQKsloF/j5Hucrfb+cmMkQAIkQAIkQAIk8PsT0K9x8j3O3/9GcAV1I0AvEiABEiABEiCBlkyAT9Vb8t3j2kmABEiABEjgtyTAuUjgZAkw4zxZchxHAiRAAiRAAiRAAiTwqwTMn1WHCzNOQGAngUYjwEAkQAIkQAIkQAJuArJaY8ZZTYJfSYAESIAESIAEWgsB7qO5EWDG2dzuCNdDAiRAAiRAAiRAAq2NwIkzzkJ7+cVff9T57WfRL/nmo8NF+Z4WXMJpalHB0E9ef3bdUjee3XnZA2e/ct4X7+aWlX6zdwfGujtO0T1PMQpjTcvV3/9nb34OLGZ3R16aeqDfh7M+3rkZ02HSif/7IN9ehlHQYXG4nP9Yszj+jX+EvvL3aT9/U1xhhxGXbv7xf2aEVzevXnv0yKDZry5LPYBRWC3snlOYs8MfRjiYp+7FuITr7a3rOr71TNCsR//wzWxsyoxvumEinGIgOwm0HAJcKQmQAAmQAAk0OYHq38dZh7855BSuzNLi0yJjXjr3gk2Z6TNXLDAtcQFBfxs68qY+g/y8vBwuV0ZJUV55mXvhSO/yy8t35WatOXqkT1QcPOGPDgWnyAjzykr/2L2veYpRGGt3OKb3G7otO+OKbz9FSgcjuhl5f0HePYt/6BUZOymppzn7wsP7vk7ZgVFYGyzf7t35zLqlk3ue/tTwcZ/uTH5p48pgH19/L++U/Jzk7KN78rIXH94PCeeEoBCMwmoRGfHRoeD09Oh4LObq7n1ggUONxSw8tO/eJXPHdTzli4lXX96td5ifHyZFNHNHJgQMZCcBEiABEiABEiCB+hBo5b7Vv4+zDhmnm0Skn79NWjVR5FvzDuwprajwNmxuH1NBRfCrlO2D4hKSwiLnHdjdMSRscs/+icGh6FBwCrdKlxNZKXqwjw9O0X1ttktPOe3uAWcdLMzblZcFi7u/u219VmnJE8PG+Nq8TCMiz9q4Eqmheboi/VCH4NC/DBg2pdfpQ+ITfzyYUlJZ0SMiOqesZNHh/YHePijNYq64wOBwX39zSA25tyAX23HHr7GYeQf3xAYEPThkBJZdVll5sCDfHP4rEEwHShIgARIgARIgARJoswTq/bPqyOFGfva2j832wOBzTkgttahwc2Z6sLdvgJf3irRDePx9wiFuB2/DcLpclU6n2wLFSxq55aX/3rIWutkHxSaUOyq3ZmeYpyUVFabiZRi+1Rlwr8iYfHv5hozUc9t3LrCXb8/JRFYa6utnetZFuheD+MhBK5yOVzetum3BNyvSDtZlOH1IgAQagQBDkAAJkAAJtFgC9f5Z9Ylduj925uh8e1l+9aNzVBD/cdZ5Z8QnOlxV2WFRhf1QYT7KiuszUiGzSouRou3Nz92SdbQ2KFQTb+4z6N6Bw21GVd3U4XLhCfgLG1Z0CgnvGRHjOeTPfQffP+jsd7atR2TTjorjn/sMXpdxxDwdGp+4Lz/3s91bv923a3nawX7RcUh22weHlFTY8Vh/bKdTUObE2G7hUaY/5kotLjhSVIC01bSMaN8Z2+kcGo4CLSxw8FzMwNiEg4X5KNyizBnpHwAHs9eGYNopSYAESIAESIAESKBVETipzVS/x+mqyvbqEuSqbn3CfPz+uXaJ3eGAP6qep777r7P+8yYqmjhFf23zalj+/NPXyMySwiLfHjvp1VEXojS45MgBXK3RiyvsV333H/g/s7bq540ySooum/Nxob38pXMvCPf7xbNv1Din9BqQGBTqnh3RJiX1TAqNhIIO/Zoeff+y+Icrvv1kSFzijMEjYOwYHBbhF+Dv5d03Oi4+ICivvAxVT9jRMdeI/77V78NZeNSOU3Rz8ZO+mY1qKE7h4LmYy0897epufR5ZuQCjfAwbdgcf9NoQYGQnARIgARIgARIgARIAger3OOWJM85QH79lV9z0nwlXtgsK2XLd9E/HXxHtHwhL6e1/Q9875e5E/YImFJyiw/OtMZfAAQNRrdx/w1/+OuhsTAk7OhT0K7v1gafZ/37maFggcVo8bSamQMESFrMjOCLjKnLQFVfe7J4dlpiAoHXX3GpOhIrprHMnZt/yQPYtM7668Bo4Y3hcYPDW66YfuvEeVEw/n3g14mNe2DEWOnrmnx8Y1q6jOQVO0c1opoPnYhD/xXMvyLllBobs/tOdKHlid3DGEHSsEEEQmZ0ESKAtE+DeSYAESIAEahAwHx27RH1+cqhGiGZ4irzQr/pHi5pieTZpBHlX/ZxTU8RnTBIgARIgARIgARJoTQSkkNgOjhPXOOHXSJ1hSIAESIAESIAESIAE2goBlDZdaEJAMONsK3ed+yQBEiCBagL8SgIkQAK/BQGUNiWaEBDMOAUbCZAACZAACZAACZBAUxBAdVNVOkV9fla9KdbRTGNyWSRAAiRAAiRAAiRAAg0mgOqmqnSKOvyseoPnYgASIAESIAESOCkCHEQCJNCSCajqJoqcfI+zJd9Erp0ESIAESIAESIAEmjUBVd2U0pR8j7NZ36oTLo4OJEACJEACJEACJNA8CVg1Tr7H2TzvEFdFAiRAAiTQsghwtSRAArUJmNVNJQXf46yNhxYSIAESIAESIAESIIEGE3CpX8TpMiWfqjcYJwPUiQCdSIAESIAESIAE2hYBKfCfxCEka5xt69ZztyRAAiRAAm2cALdPAr8dARemUj+rrqqcrHECBjsJkAAJkAAJkAAJkEAjE5CIJ6UU+I81TrBgJwFPAtRJgARIgARIgAQagwBqm+pNToGv/JtDjQGUMUiABEiABEiABBqZAMO1fAKobQqUOAW+1qfGWVJmz84vTsvKP5KZx04CJEACJEACJEACJNDqCSDxQ/qHJPAkEmDzPU5RrxpnTkFxaXlFWIBfx+iIrnFR7CRAAr87AS6ABEiABEiABJqaABI/pH9IApEK1jfplBggpax7jRNzeNu8EiJCA/18DUNiODsJkAAJkAAJkAAJkIAQonVDQOKH9A9JIFJBJIT12qxLVzf1T6vX4T1O1FFdLhEdElivOehMAiRAAiRAAiRAAiTQagggFURCiLSwPjuSqG8KVeWsw3ucqKOimlqf6PQlARIgAU8C1EmABEiABFoDASSESAvrsxNV5azre5z2ikp/H5/6RKcvCZAACZAACZAACZBAcyPQ0PUgIURaWJ8oqHFKIaSQdahxOl0uPMIXbCRAAiRAAiRAAiRAAm2YABJCpIX1AYAaZ1Xn3xyqDzf6kgAJtG4C3B0JkAAJkEBjEkB1U4g6vscp2EiABEiABEiABEiABEig3gRQ4MQYyDr8rDoc3b22Yl+4oOCOW3InnnfSHcMRpHZkWkiABEiABEiABEiABFoyAdQ4pRBSyDq8xymO35ApFj/7D8felOO7nPgKhiMIQp3YlR4kQAIkQAKaAAUJkAAJtAACLtWEEg2rcZZ9+d/G2m0jhmqsJTEOCZAACZAACZAACZDAyROQaEJASNGgnxxCeVI0UjtmKIfTmZWbX26vQM/NL8QpZGZO/mOvzt6ye3/tmZFDwwHOQqiL7301f/pjr7hPlUmoPNv0cTvD4Wh2LnpxaRl83KeYGjPCYsYpKSvPyS+sqHTAgo5Lm3fuXb5hG/xxWt9+KC0Tu4A0B36/eM15U+5fsHKjeVpDYm3vfD538ZpkTFrjEk9JgARIgARIgARIoLkSQLalcy9XwzLOpt5efmHxpGmPLlqz+bl3Pr/yridSDqZeeOvMdVt3fTl/2ZGjWUgBzQVgN0UlpZDFpWXX//UZ+Jv2UUP73XzlBB9vLyRqcDCNGDVlxnPwcSvQ+198y5g//bX7uClzl641T8fd8MANDzyH3BSjzDgFRSWXTf/79pQDsJSV22968F8TbnrwT/c/M+r6+9IysmGs3ZGMopt2KOhuPS0z54t5S3MLCmFBIvvk659cNWHE0H49sNTSsnIYzY5NFZaUzvl51bL1W6c88OyydVtNOyUJkAAJtBACXCYJkEBbJoDypmiEGqdo4hbo75cYH715575tKQeOHM1O3rVfChkdEeZls017dFaXUde9+vGc1IzswZfdnjTm+ivvfOKNT79btXnHzQ+/sHnnXiztqx+XP/j8u5t37Rs46bYBk2575i3rHYC8guKM7Dy7vQJu6EP69lj28fMTR56xYsN2nPbvkfTFrJnvP31vVHgoTs04peXl0M2+aefeVZt3fvvGY0tn/wuW+cs3QKISiVnmLVs3/bGX0T/8+qdu46YMunTamuRd3y5a3WvC1B7jb5izcNWGbXv6TLz5oltnIoXFKPTZ3yzYtf/wrI++Xrlpe+8Lbuo86rr7nv439nv2NXePuv5e6FP+MPbWay6MiQyLjQqHPzsJkAAJkAAJkAAJtAACqJypn1PHl+Zd4/T18T7tlE5rk3fa7ZVJHeJXb9oRHREaHhpU6XA8P+PWGy4dh4ojkrBXZk6/csKIjTtSTu+VlBAb9fqjd/Tp1sV9GwqLSlA4vO2aC6dePt5tfOiFdydNe2TD9j2mZdWm7cOu+r/vFq0+8/SesMCOq8+9/Tn0Y3YUX7E2TB3g74slFRQV13ZLz8zp2C726XunduvS/t0v5l4zceS1F47+eM7P7345f9SZ/X/49xMhQQFqlBCXnX82lv3k3VPmLFg1Ykjf7958/McVG1IOpaJke+Nl57/88LTC4tJHZ31w39TLu3Vubw6hJAESIAESIAESIIHmTgAlThQLtTSa+VpP7dx+x95DyDLPGth7+cZtKsnz88Oavb1s8dERUFZt2nHH46/075kU4OeL09q9X4+un734EJ5HPzLrQ/fVlx66benH/0Jp07SgqPm/Vx7Z+u2b5w0bAAvsuPrI9OugH7N3ToirrHSs3Lgdz9P3H05vHxftdnM4nOX2SpzeMXnSvVMvR4USU+O0oLhk4GmnTr/uYpsh7fYKe0WlSvhx4Vc7iqy4Tf5+vtP+eNHwgb1/1ZcXSYAESIAEGkCAzO3g2AAAEABJREFUQ0mABBqfgE52IJr5e5zYePu4qKzcgsiwkNN7Ju3efwRFPilhtnpwYEB2bsEjsz5AIRNP4ft17zL5vqcXr0l2eyxZt+XiW/+2avOO7l0S3cYaio+PN55ZY7hpX7FhW9fRk/GI/Gh2rmkxZV5h8XlT7o8bdsXew+m3XjPxlr+9eNbVd50zuM+44QPhEBoU2DOpww0znl27ZRfYPvHa7LueeM3pcnVoF3P9pLGfz10K/3Vbdl88etjcpWuvuPMxLBijPPvkS8YsXLVp/NQZo8/o3zWxnfvSwdSjd//jjY3bG/RbqNzRqJAACZAACZAACZDAb0JASqRtEq3JnqqHvPJm+Dfzavfgx/5Zrx0O6dP98OLZj9/5p5FD+6Uv+/SO6y6JjQxf98XLKEai7Pfig7f1PrXTzrlv7/3x/c1fvz6od7e3nrg7beknZw9S5UA4fDlr5vnDB2377t/7F3xw0xXjMTXSShgx3K1AhwWnuIqOU0yEjlkwFyxmHNQ1YYEdfexZA26+YkLqkk+gP//ALX6+PnCD/OCp+2CE20sPTZs57VosDKvqldRxwjmDDy78EBu59eqJWNuhRbPNBZtP/zELhmDevt27YKmI+c97bkyIjTSNiNy1Q7tNX78GAtB/x86pSYAESIAESIAESKDuBFCAczX1e5yuoqLi556q8YeIyuf/UPdV0pMESIAESIAESKA2AVpIoKUQQG1TSmH2pnqPUwYFBd51b40ap++YcS2FEddJAiRAAiRAAiRAAiTQEALVNU6BSmdTZZwNWR/HkkDDCHA0CZAACZAACZDA709AqoYaJ7402XucfKr++99nroAESIAESIAEfk8CnLutE6iuceJrk2WcfKre1j9l3D8JkAAJkAAJkEDbJiClEFJAojfVU3XWOAUbCZyQAB1IgARIgARIoPUScLmEUD+r3pTvcbLG2Xo/P9wZCZAACZAACbQuAtxN0xCQEhXOJq5xFtw6tcavRjJPCx+8r16bKi4t+/qnpe998X3yrr31Guh2PnI0M+OXv8jdfclUVm7cejD1KHQo+YVF//luwQdfzV24akNRSSn0dz779qv5S8rtFXCA8WBq+k/L15aV23GaW1D4+dyF73/5w5ZjrQ3R9hw44nbGkO8WrsBGVqzfgrFmR4R5S1evTd6B043bdv+weFV2XgE8l67d/Otrhj87CZAACZAACZAACTRnAur9zaaucTbW/lesTx7Up8fkSef3SuqMVAw535K1m+z2in2HUldu2Aod2R6yOlj2Hkpdvn7L/KVr4LZ+y87iktJte/Yj3Vy9afviNRtz8vIXr964etO2SoejxtoqKmBTRihOpysowP/ScSPCQoJ27z8M/coLRl88ZrivjzdGJXVMCAsORvapCAqBWYad3vu6S8alZmTVThARzeGodDtv2rGna4cEbATkERnR0JN3pvTtnuR0OpN37s0vKh7ar9eWnSleXrYAfz+sHw7sJNBMCHAZJEACJEACJFBfAihxYogpm+o9TkzQ8F6K5LGiMiw4CKEMQ67YsKVrh3YxEeHrt+1Ky8zu2D4OWZ1hGNERYfsOp6VnZndOjE+Ii966e196Vo69ojI7Nx+ZaFR46IBe3ZJ37Tulc2JIUOD+w2mIhlQVBUWHwwndXlk5f9nadz777lBaBk4x6YZtu6HHRoXn5BV8/M38VRu3wY5+JD2rpKwMCjrckE2G6rUlxEabCaJnWPh4diwmJjIcltioiIrKSijoqDVjSb4+Ppk5ubiKvLakrLyy0uHv54ur7CRAAiRAAiRAAiTgSaBl6eo9TqFe4sSym3XGicQLBT8kdlgoOpK8oIAAVB9dTpfNsPn5+Hh7efn7+jidLlwyLTgtKy+Hc42eV1CYcuBIcWlpRGgILuUXFOXmFzqcKuP08fIaM2zgny4dnxgfg0s+3t6ndGo/ceSw8JDgiLCQqyaOGdKvJ+w1uufakJh6e6siqGfYGv7ILItKSmAsLC5xOBzbUw6gFyPBtFfYbAaiIYhLCCh+vj5wYycBEiABEiABEiCBFk1AVTdRXZNCyCb77UiikVpSx/bfL16J5+bL1m3u1D5+2frkNZu3x0VH1g5fWl6+cuNWPLzumBCP0uOqjVuRU8ItwN93596DHdrFocpY6XAihYXxgpHDrpo42sfbC3qNjvwPD9NNY35h0ZwFy+b8vNwsYXoasaROCfHzl67+YdHKMrs9Xi+pdlgzApy7dGi3fP0WKAePHMWmenTtiN6pfdzy9ckHU4/2TOqEnHjJmo2RYSGH0zMPHE5D1baoWGWo5qSUJEACdSNALxIgARIggeZCwIVamkvXOF3NPuPsktjuqgtGjzpz4LABfZCWXXDumcjq8PQcdcfw0ODhg/p2TmzXr+cp6P6+vkP79bp4zNnt46KHDeg97pyhF44+C1cHnNYdev+ep1w0ZviQvj1Ra6xxH8wgMEJBzPEjzkChEaeQV00cc9n4kZg0MMAfFkyKZNc0Ykmndk7EVQQfM2wQ8lQ4eHZE69alo9u5XUzUFRNGYRRWhcimZ/euHTEdOlLkCeeeOXb4kN7dumL9iDnyjAFBgQGmGyUJkAAJkAAJkAAJtCgCarGqxomvqHGKZp9xYp117CgiBvr71dGZbiRAAiRAAiRAAiRAAk1KQNU4MQEqna0p44yOCPPRP1GOrbGTAAmQQHMnwPWRAAmQQGsnoGqcUgjdG/STQ7YuXUUjtUYM1UgrYhgSIAESIAESIAESIIGTJ6BqnDhQ43S5GpRx+l1y2cmv4pcja4f65XWekQAJkAAJkAAJkAAJtCgCqG6iwqllgzJOnxEjA+/+awPLkxiOIAjVohBysSRAAiTQVghwnyRAAiRwkgRQ3cRILRuUcSIIMsWQF14N/2beSXcMRxCEYicBEiABEiABEiABEmg9BFR1U6DKid7QjFOwCSIgARIgARIgARIgARKoRcAlqv5r4HuctQLTQAIkQAIkQAK/FwHOSwIk0MwISBQ3q44T1zgNKZ1O/QS+me2CyyEBEiABEiABEiABEvjNCCAhRFpYj+mqKpxqxIkzTh9vr1K7XfnyaOkEuH4SIAESIAESIAESOFkCSAiRFtZjdFV9U404ccbp7+udV1KmfHmQAAmQAAmQAAk0nAAjkEDLJICEEGlhPdZeXePE1xNnnAF+PlKKzILiekxAVxIgARIgARIgARIggVZEAKkgEkKkhfXYU3WNE19PnHEibkRIYIWj8khOfnFZOR7hw8JOAk1KgMFJgARIgARIgASaAwEkfkj/kAQiFURCWL8lobZZ/dPqdco4ER1zoI6KauqBzJyU9Cx2EiABEiABEiCBVk+AGyQBJH5I/5AEIhVEQli/jtpm9U+r1zXjxASoo0aGBsZHhSZEh7GTAAmQAAmQAAmQAAm0egJI/JD+IQlEKljf7tI1TheacNUj46zvNPQngTZAgFskARIgARIgARI4NgGpa5wQ6Mw4j82IVhIgARIgARIggZZDgCttlgTMGqdQ73Iy42yWd4iLIgESIAESIAESIIGWTgC1TSGllsw4W/rN5PpJoG4E6EUCJEACJEACvy0BF5rAod7nZMb527LnbCRAAiRAAiRAAm2ZQFvauypuosapOzPOtnTnuVcSIAESIAESIAES+K0IuNREqsDp4s+qKxQ8SIAEmhMBroUESIAESKB1EDDf4JRoQtajxllSZs/OL07Lyj+SmcdOAiRAAiRAAiRAAiTQigmYW0Pih/QPSeBJJMHqFc76vseZU1BcWl4RFuDXMTqia1wUOwmQAAmQAAmQAAmQQKsngMQP6R+SQKSC9U063e9xijrWODGHt80rISI00M/XMKRgIwESIAESEERAAiRAAq2fABI/pH9IApEKIiGs14ZdwiVQ49TyxE/VUUd1uUR0SGC95qAzCZAACZAACZAACZBAqyGAVBAJIdLCuu9ISpQpJZoQdXiPE3VUVFPFSTQOIQESIAESIAESIAESaC0EkBAiLaz7blxIUVHj1PLENU57RaW/j0/do9OTBEiABEigeRHgakiABEigMQggIURaWPdIUkjtDFmHGqfT5cIjfD2AggRIgARIgARIgARIoI0SQEKItLDum9fvccJdvc154honHFt95wZJgARIgARIgARIgAQal4CUqropVaWzDjXOxp2b0UiABEiABEjgeARoJwESaE0Eqt7jrOPPqv/6zu0LFxTccUvuxPNOumM4gvz6LLxKAiRAAiRAAiRAAiTQ0ghIvWAlG/RUHZli8bP/cOxN0eFOUmA4giDUSY5vW8O4WxIgARIgARIgARJoKQRceqFKNijjLPvyvzpQI4hGDNUIq2EIEiABEiABEvg1ArxGAiRQFwKquika/h4nypOikVojhmqkFTEMCZAACZAACZAACZBAQwjoNzkb5T3OhqyiLmOLS8vyi4rh6Vag17Fv2b3/sVdnm8M9h2D3ufmF5fYK0wgFpzCap78uD6VlIiak6TZ7zgJ0U68tF6zc+MrsrxG/9qUaFiwSG6xhPN4pApoLNhWn04nTo9m5Wbn5DqcTHQpO0eFwvCCmff22PRfdMvPdL+eZpw2R4IC+PeXgJdMeSTmUJhoSS/1JLBc2hfWDDIDjVtaOh1tm+piX3vtq/vTHXsEQ87QuEsGPib3uNw6zpGflXv5/j63ctAP6MTvWuXnn3lc/nrPvcPoxHWgkARIgARIggVZKQEpV5VRHg56q/wZ03vl87kPPv4uJ3Ar0krJyfBeH8isdPvuPHP1y/rKycjvcSsvKkYpBQcelKTOeW7RmM3T0H1dswCmM0NFNTzhDwWmNnltQ+MW8pUjmKioqcWnlxu3oULCeopJSSE99254D85dvKCktq6h0wI4OxRwIHb1EbwSj7n/2bWwQFrPD4o5mWrAYLMnUsXJzwaaSnVeA07OvuXvwZbdffdeTB9Myxt04Y+R1955/44xV1WkQxiKgORw6okFHcvbs25+ddmqnTglxoybfi03BeMyOsRjlvoThOMVesH63ERzQOybE/vWmKxJiImHHVfhAgYQOxewYiwimjjWgm7rnLMWlZdf/9RlsELcPN/HI0Sx3BDcct485fNTQfjdfOcHH28vtYNpNiRkxr6kjFHzQa2CHA9YAnxPeOPi4e2RYyN1TLu2V1NG0uCdC/PzCYsTMKyj68OufFq7aNPm+p3PyC003ShIgARL4jQhwGhL4PQmoNzhRR8ISmnvGiSWWlNmRDCG9gF5QVHL13U92GXXdiGvv2bX/MMppn89beteTr9844znUpS64+SF8j4fbnIWrksZMnvrgv6BXOhz3Pf3vzqOu633BTZt27IXFsyPywy+8t2LDtvuffWv6Yy+fN+X+kZPvXbhyI5wxBAOPHM0eMOm2ecvW4So6xmINE//8cJfR18GIU3RYrvnLP0674KYr7nwcdaxxNz6AIbf//WWHw7luy+7eE28+97p70jKyv120usuoaxNHXINyFz4P7VkAABAASURBVNbp3sjH3/6MLBaVPDNgakY2csekMddfeecT2DUm7X/xLVgMskmcYjp0u70iIzsvr0BVf3GK/tJDt33x0kykm5k5+cGB/rMenjbv7SeHDeiFS7v3Hxk46TYs6Zm3/gsC7q39tHLjTys2vPXZD1fe+fjWPQfGT33w0+8XYfFL123pe+Gfke1hUx99swCQsZhz/vgXbAGLcSMaMOnWnuNvfPmjr5FaYRaz7zlw5JaZL25LOQB0uE3jpz6AaG7PTTtS+lx4c68JU7uOmTx36VrkYd3GTRl06bQVG7fDHwBHXPuXw+mZCPXuF/NWbd5x88MvbE856GWzTXt0FqKBmyecNz79zvRBBRFDvvpx+YPPv3soPRNbGDDpNvA3k3us2Q0Q9cjjYfekhGjHvHGvfDznjsdfefL1T177ZM7YG+5fk7wLmTpmv/VvL63YuK3GRKCH+5gw/KqZL73/z7/cOGxAr44JMQF+vgjOTgIkQAIkQAJtg4Cqbgr1HqdoARnn3CVrUK57/dNvhRBrt+zaue/w4o+eDQsORIrZs2uHRas3HTmahXLmkrXJp3ZqHxociArWe1/Ou2/qFe/98x4MSTmYhirmd28+PmJI3/e+nA+LZ4+NDP/TpPPO6N/zybtvgD0+OuLn95/+dtEaOGMIBqYcSoXdsyOf+/aNxyaOPGPOzytNO1aFBTwy/bq0zBxkKsgmLxt39ow/X2WzGf17dP3fq484HA5kQu9+MfcvN1z21hN3I53C4t0bycotwAIevOXq84YNQMDYqPBXZk6/csKIjTtSUg6q2XH10+dnIJs0K3Dw2bB9z6Rpjzz0gqr+4hQdCdak2x/p3D4uOiK0sLgUKdoNDzyHh864lF+kKm23XXPh1MvHg4B7azERYYiMed9/6t6E2Kjv3nzs9B5JJaXlSOPKKypQk8vOK+jXo+uTd0255aoLDqQeTd69H9HciIb07X7n9ZP+99OK3IIi2D37wdSMH5as/fylh+e/888v5i1zexYUl9oM491/3HPukH7I8NIzczq2i3363qmB/r5IcP827drQINzWTQh12flnY0mvP3pHj64dKh2O52fcesOl47anHPCEc3qvJNOnT7cuGGJ23H03f29vL9OIbX6qAS5bv6U2dhO7JyWM6n+sG/feF/O6dU5M3rVv4/YU3LUla5OjwkNR44S/2d0T4f8OHUjNeOPv/4f133jZ+eu27v7fj8ufuHOKn6+P6UlJAiRAAiRAAm2AgFmVcqHU2QIyzkvGDFv/5St3Xf+H2jfmzNN7LV23FYkdEpFvfl6FMpLbB5kTHua6T+uohAajCHWMnMDhcJbb1WN0xJFSgp/dXuFlq0poYCwpLUMp7t4bLj93aN8vZ81EMXL646/YKysNmxEeEoSg8Kljx6NwFNL690xyj/KqFWRI3x5LP/7XSw/d5o759zuuX/3flz54+j6MCg0K+Pi5+5HpIh+Cw4Bep3z24kPL1m19ZNaHOD1md6I5nPExEUhY//fjikljzvp+8RosHGXCGx/8V5fE+MjQYHOgG1FhUUmnhNj7b74yODDAvFRbVjrU6wRuz0A/P9AL8POJDFPR7pg86d6pl6O6uW7rHjxbL7Pbp15+/oghfWrH8fayIdOFvTYcGD17+9goN39kn+al2gBNu6esQel4Nw4ZZMrBNOTK/bp3mbNwJTJyfz/rA+OeCKl/jy6JD7/w/rUXjerRtWOH+JgZt1yN/zfgOSN1EiABEiABTwLUWyMBfNvHtiRKnU2VcYa88mb4N/Nq9+DH/omZT7oPPO3Ubp3bn33N3XmFxReMGHJqxwRkFfjuftopnbJy87smxiNygJ/vRaPOfOH9L+968jVfH++uHeJHn9F//NQZKNpNvmQMHGp0VMjWJu+86aHnHQ6neQlucMYQDOzfI6lnUocbZjyLQqZ5FQkunswu37ANbkgmFq1Ojo0K65nU8Z6n3njqrf9sTzk08ZaZeMKb1DHBy7CZQyBR3Lp+0lg8177hgWevn3TeyKH93Bu5ZPSZfbt1wVP197/6EZ5I4LJzCx6Z9YHDWbUeGE/Yw0ICkV+igghPwMGz77hhV5iP6Reu3nzxrX9D5bJ7l0Ss2b21pI7t4Ize+5ROgQH+F9/2SGmZHT4ZOXlI+0rKyhEQdVA4PPDc2wXFJVDcHXE279x32yOzvlu02jAk7CYKVHmhd2gXM274wD/c/iju1IQRg92eUnvCAd3pdD3x2uy7nnjN6XKhWolC46Mvf/jAc+/k61cFQoMCkdJNvu/plRt/8RM5nnAC/f1Mn8VrkhHQ7HsOpLr5+3hb/5fAvIoE8XjYPSmZzqasceP6dOuMDxVwjTyj/7Y9B/t0t8qrpr8pQQ9V4byCorc+m/vl/KU/rlg/41/v4NS8SkkCJEACJEACbYOAqm4KoWRTZZyuoqLi556q8YeIyuf/UF++0/540YsPqkqeqaDqNvvZ+9OXfbrow2fwAPmUTgk7vn/r8Tv/hErb/gUfIJ8w4//xwlHw2fnD28s/eT4hJuqf99yI023f/buvzg+QqaAMhhTHdD57UO9Di2Z/9MxfX555uzkX3OCMIRiIZ+gfPHVf6pJP1n3xMq4iPV3z2SxEhgPc8Lh509ev9UrqhFXBZ8lHz501oNfaz2ft++l9PIyeft3FmKhLYvxP7z2FgRPOGYyJEBajQoMDMQQ6NpIYH/O326+Fft3Fo7Gk3qd22jn37b0/vr/569cxCpOiQ0GQ2MhwOGDlCItdmEp0RBhOoeMSOnywVERDN43nDumL1YLPTVeMx5qh4xK2FhYciIEAGx8TuXT2c6v++yIKnIAJpBiI2V+ZeXu7WHUJy8Z6YMRK0DGLGefw4tl4Jm4z1KcImwKKsWcNREH69J6nID5mAf9RyMy++7fpiUfVuIq9IAjgzJx2LXaKifp26wJ/ANz67Zu9u3VGfOR5bz1xd9rSTy4adQa2g6mxTozyhDOodzfTB3cQQ+CA7eAz4ObvZVMZP0ahY1IA7JrY7njYPSmZoboc68bhGfrij57Fp+WaiSOxQdxTEzhWiFnQzYkqKx3+fr6vPjJ9cJ9uSIivGD9ixacvAC/W2YDOoSRAAiRAAiTQsgio6qYQSqpcQTRBk0FBgXfdW6PG6TtmXBNMxZAk0OwIoO4++szTv5i3NDEu+tE7JtuMpvqH1ux2zgWRAAmQQOsnwB3WnYAL5U0XapyulvCTQ3XfFj1JoJkQCPDznX7txc/cd9Ndf/pDRPUrsM1kbVwGCZAACZAACfxmBFDelEIVOZuq9NJYT9UFGwmQQIsjwAWTAAmQAAmQgCbgqpZNlXHyqbomTEECJEACJEACJEACvxOB339alDjVIlDmbKqMkzVOBZgHCZAACZAACZAACbRVAuYbnChzupruPc6mrnHaKyoLiqy/uNNWbyX3TQIk0CACHEwCJEACJNB0BFDaRJHTlE1V4yy4dWqNX41knhY+eF+9NnbgSPqHX8397/c/m38+xxyLdPPHZWvyCorSMrNNS71kcWnZ94tW/rhsrfsXcK7cuPVg6lEESc/MXrhqA5Q6djhjSF2c8wqLflq+1vSs+yjTv4Y8cjTz25+XuxePqwCybY/6g0DHi3xCBwSp0T0XbPIBt69/WvreF98n79prOucWFM5bunptsvqVmRu37f5h8arsvIKycvvStZszsnNNH0oSIAESIAESIIG2ScCsbiopTvCz6r8/HyRS54844w9jR4SHBqOoiZxm3ZadWTl5R7Nyl61P/uanZWuTt2/Ytqu0rBxZEWTyzpSMnFy4rd+6C2nW1t37lq7bnJNfsHj1xtWbtlXqP4GD4fExkb4+3u5fsV5RgSvqr+MEBwb06NoR20YcZKXIqJBdIbvF1PsOpa7csHXBivVw3bXvEK6mZmQldUzAEMwyf+kaSAzcue8gVjXn5+XFJaU4dXeX01VurzBPMSoiNGTvodRvFixDEChzl6w6mpWDIea8SH+Rty1YsW7F+i2Yzr0Gc3haRrZhGNl5+fBBion0bu+hI8vXJ2/esQeRS0rLDqdn2u0Vm7bvOZSWgdwUSzpwJM3t4J76cHqG54xmcLf0XLDJZ8X65EF9ekyedH6vJPUrM+GJ1fbtnuR0OpN37s0vKh7ar9eWnSleXrYAfz9EhgM7CZAACZAACZBAmyUghZQ4hJCi2WecfXskIeFbsmZjaWn5snXJyGmQJmbk5MVGhY8Y3L99XHT/nt0ysvOQtKUezUL2hsInMqoRQ/rn5BUg39p/OK1/z1ORD53SOTEkKBCnQghvby9YunXpkJGVk5WbB4u7I0ndfzh99/7DyJ/GDh+SmZ1XVlY+sHePLbv2op7aPj4Gie/+Q6l7Dhwe2r9XXFTkkfSswuKStIysYQP7IGVMy8iGPHNA75jI8MAAf9RlP5nzE9JHd3xTwSgYd+8/NH7EGYY0oJ8zuP+OvQdLy+xZufnDBvRG/H2H05C3De7bc9+hNPcaMBypJNbWqX38kaNZKzZs6ZQQd+bpvWMjI9rFRPXpnoTITqdr3+HUo9k59oqKuOiIEUNP37P/sJ+vj9sB0yH+uLOHbN6RglzcPSOCOxxOJOtL1m6CXqMjGuCEBQfBblT/6SApJKj6+vhk5uTGRIYjiS8pK6/Uv/wcbuwkQAKNR4CRSIAESKDlEfgt3uNsLCrIk66YMCosNHjH3gMulws5TWhQoMvjzz/abEZUeEhaRhZy0H2HUmOiwvGQd92WHU6Xs7KyAvlQoL9fXkFhyoEjxaWlSEaRVG3dtW9g7+5rNm/fuf9QSGBg7aWiPhfk74+8CvVCZJApBw8jnbIZNoRC6maz2ZDRrkvegVnMsUi5Avx8vb28AgP8cvMKUU2MjVR/HKi0vBxVUgQx3TwlMrMAPz+bYZSWlSE53rp7L6aAA4yIg45Sa1R4KIqgqGW61wAHpJJ4YL11196DqemIHBocZFTnf7iKDlbISvcfSe/QLhaV3Z0pB7BlhwMlbVxUHVP7+6rVBgX6w+6eEdeQzSNLzi8ogl6j+yBP97KZZVqHw7k95QB6MWLZK2w2w9/PF7vAHFCAqMZYnpIACZAACZAACbRBAhJ7lkJqaUA22+50uvC0+ueV65FeRYSFoL64eM0mPNGOjgzDmpEC4bEyTpFHZmTndWrfLj0zB0W40rLycntlQVEJfMzevWvH7Nz8SocTD3yRHvn5+RxJzywtt+PxOuqF8MFD9eXrt3w1fwnSO5wmtovdtf/QotUb4YzCHtIsL5sNdrOj2rd+2y5vb28sz7S4ZXl5BaZoFxMJHxiRLt9y9cWYHTo6MsXPfli4bN1m6DGRYXj6j63l5BcijUbqBiOmgzQ7ypBINBEtMT7Gcw1pGdnDB/a99PxzkWsin164asPiNRvL7HagQCEWY22GERMZjupseEgIVo4SLK76eHu5HeKiIpAHY+qyMjtSZAxxd7hdNXH0BSOHmRbPBcOS1LH994tX/rR87cpgGsyIAAAPO0lEQVSNW5AQo3dqH4f0+mDq0Z5JnZDWoxQdGRZyOD3zwOG0fYfTioqtW4Dh7CRAAiRAAiRAAm2KAEpRAnVO7NnVvJ+qG4Ycc9agc4eefsl5Z3dMiDuj/2ljhw+eOGpY+7gYPI9GCXDS2HNO7ZzYObHd1ReOSYiN+uPFY2OjIuA/8ozTLx03oluXjsMH9cU2u3XucNGY4UP69kSKhlNkbKOHDYQDqqfRESp5hf81F5138ZjhSKqG9OuJtOny8SPPGdyvd7euk8aeg8fWndrHw46UF2kWwp41oA+uwoIeFx1pzgJZXlERHx054LTuyNXwfB9zuTvGXv+H8Zh02IA+5qg/jB2BpQ44rduEc888e3A/RIbPkH49USZEKKwEnhece2ZCbLR7DYiG5+zt46KhoM4KN5A5e1A/TAqf007tguEIYi7bx8cLlM4Z0h97aRcb5XaIiQq/cNRZmPq84YOxfQwxZ0RMz4447gVjIkDuktjuqgtGjzpzIBZmeiKZxhToocFB2MXY4UMwNZY37pyhI88YEBQYYLpRkgAJtCEC3CoJkAAJVBMwq5uQUopmXeMULa3FRUfYKytRl20XGx0Wol55bGk74HpJgARIgARIgARIoHEI6BpnlWDG2ThMzSheNtvQfr1Ql+3cPt601Ja0kAAJkAAJkAAJkEBbIIDqppBSy4bVOG1dujYWr0YM1VhLYhwSIAESIIFWTIBbIwESaDoCpRV2L8Omy5vClA2qcfpdclljrbURQzXWkhiHBEiABEiABEiABEjgJAiUVVT4eHmZ1U1TNijj9BkxMvDuvzawPInhCIJQJ7EfDmlKAoxNAiRAAiRAAiRAAidDoKTC7uvlo6qbQrjwxeVqUMaJJSBTDHnh1fBv5p10x3AEQSh2EiABEiABEiCBWgRoIIEWRqDS6Si2l/t7++jqphT4T0rjs8xDLWwfXC4JkAAJkAAJkAAJkEBzJZBWkBfg4+tteGGBUlU4UecUxsP7NzvME5jZSaAlEuCaSYAESIAESIAEmgeB/LKS8srKsIAgoUubQv+oupDC2FdW/E76XnH8ZkhZ+4/rHN+dV0iABEiABEiABNokAW66tRNAQoi08Fd2WeFwpBfkxQSH2YQhXbqi6XJJ8z3OUWGxt+9Zd//ejQ7hPGYIH2+vUrv9mJdoJAESIAESIAESIAESaCMEkBAiLTzmZh0u55H83AO5WfFhUd42LxQ1hWFIuKLGiS+GYXxy6hkbThs1JTguJTMDrqiFVjp/kXr6+3rnlZRhCDsJkEADCXA4CZAACZAACbRcAkgIkRa61+9wOovKyzKLCpBo7svORFYZFxLubdgE6ppCoMYpRJWExTBstsTgsHZhUeGBwchPs4qL9mSl78hIdfeDBVmlFeUZ+UWCjQRIgARIgARIgARaPgHu4CQIIBVEQoi00J0i7s5KP1pUUFpZEeTr3y40MtQ/yIZ0UwgpJQ7VhVBSnxrC6VRJqMsV4O0X4R8cFxzePjSqfVhMgpJRWsaEBgeVV1Yezs4rLivHI3zBRgIkQAIkQAIkQAIk0AYIIPFD+ockEKlgeEhI+7DoxLBoLWOgxAeHRwaE+Hv7gQTySReqmy71R4Y8dRQ40Q1psxlSSsMwpEC9Ewos0pBatykz7IYREOhn8/bKLi7dn5mTkp7FTgIk0KoJ8N84CZAACZAACSgCSPxyiku9vL0CA/0kckKpk0YDmaJSJDJF6KqmiQvqTOWQUp0bkDhkVTOQdbqElEI3l3qDE2mpOlW1TyGFS0oIJX19bEFB/qEhAeFhweGhgeHhwWGhgWHVMhR6WJBpUXq4h15lD/Kwa93T7qGHhqqroWFBIaGBkBgFqXTaTSZNxyE8CLTDQs17p3WP+xLmqR/j/gaH4jPwi8+D9QnBp8X9mfH8/Ch7WNVnSdnDgmAJ05bwKj0oPCw4IiwoIjw4XEtLD6NdMyGHcHLw+LcQ5qGHe+i0h/Nzws+D/gy0sH8Xes1hHvfuN18/vvki1/T19UZyKCWEkEgcUctE5uhyQVTljdCEhI6vqtKJL2aHD7xcLpWXGlIYEimpNAwbFHwxhDDQpJDSkEJoVUhZrUsBHykE7FWeQthwAg9IYdr1ialLU8cgRFC6FFqXx9YNQ101pLQhmvTQq+0w46q6oK7iq/oCSw271CZzEkjtp4Wyuw0YW603dzvW515/tY61Y9vYBAyqG1qorcNV61WXq3TlqXxsMFfzVLppF3qQdv31e4SrylUIHUYPNvVqqWPqq9oihQoNCX8pBcZKoa5K2HEilG5IIaXEJ8pmWqp0aWi7FMLAZeUvpdC6FBJNaJ12cpD8POh/C+QgyMHj34L00Pm/k0J/NsihETkY+juzlkqTGq5QnPHJQ5emQX9RNU4kpMhGUewUTqeSqtKp6poCzeWCG4zVUtmrdVHLjkf3yGOVHQeuIoCSLtOuJOaCpSqCtiOispg6RrtcVT4uD38PO5w9uprH41QN0afK7hEH+bi6hKxcuEy9hnT7N2c7yLjXCR07MlerlOp9QUeHGyS6W4Fudc3BjFBDKn8PbtVXcSNdbl3FMX1gVrQxt8vlqUuBu4pPjvbEBRd0HbnKbuqWFNouVFNx8GmUSodVfVWHgC7QqnX1VR2C9hoEBJrUNglNaM2SyqYOyyLQpPaS0ITWLKls6rAsAk1qLwlNaM2SyqYOyyLQpPaS0ITWLKls6rAsAk1qLwlNaM2SyqYOyyKE0pVNHUoXaNW6+qoOwTg1CAg0qW0SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHZZFoEntJaEJrVlS2dRhWQSa1F4SmtCaJZVNHaYF399d6nu9TgSq8gSBb+/qsOxIFfCdXZW5hK47oVqEr0hRYZJSSJ2x1pCYFhYhZS0ptAVDEQAqpBlMeeJcdamvIiy6VHYBxTCEVHZhmNLDXmWpaUcohK6W+KqGmBapp1UmpSh7tcWtQ5FCRa4h62+XNSJgnbBItGPHl7hq+njK+s0r9b6EXn/1NlU0qe2/tACJugRf2HHillCE8jcErhk1JNZTw4IguEeWHZo0lFFJiTMdCxalKh12oT8/Qio3IaW2mFK4dYmmR8Eitb+EC4ZAGnpGCauQElKiKU26dXWmDssi0YS+LqBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHZZFogntJaBJrVlS2dRhWSSa0F4CmtSaJZVNHdUWKfV3bikMaeY5Vd/BpW5CjxZaN5B1qsxUCJyr+pIwq1OqeGZWwgSaSmCRspqVKqdpl8ruoTtN3ansHro0dad6Q1TXUIWSGOvU/pDH0F3wqY5TW9cWLEmNha5zaq1Lp9KlE3Mpu3Bqie2ZCqTWJRSn8oTDyesu17HH1teOlTjVOlXV0HmcmNV24bF+U8cWVNd2KFgS7p+SLsTEvYB06cimDuaw/IouhNPtY+paetwjqPDBhJC4R1L5/9LHhfjags+TqUMq/RefHxUHi9R26Fikx+cKaxb4v0SIr6T2oU4+/AzwM8DPAD8D/Aw0t8+AQMIohJBqXcgNkFWa98jUhfpm7jJ1A1UpKQWkEPCXAv8ZOj+FVCUqXJG6QqallNKwqWzWMKD+QrfZBGyGTeBCLd2w2WA2bDaMNWw1dQmLIZWEk9INty4wkc0dU9mF1LLKfmxdSsOw2aRNy+PphmHYqn0Mw7BZuqylK4uh5pI2LRtbN2w2LFNC2gzDpnTDZpO1dTi57cfSpQ2s1AqFoaXUsg46gmGsRLPZfqlLw2aD2bDBXlMXuGDT9x3Srf/yMwAzhmEhEsvAiaHjaN2ANGA2YMYXga+GoWJCUgcUcgCBVsSBn23+G+dngJ+BVvgZQN4opRDq3uIrvntDl1IK/AcphJTQlFQ1TiSk6CgmnbArN6fzeFINd+o3QWtKJ+ph1Vd/obtUNGWBAh+Xq6YOo+rarhREbpqOBbg7Jqqhm5amk+Z0ZnxTh8RpU3VPnlo3p6vNHxZ1CcVIYNeex7qPLrXOKp+aukvba0oY2UmABEiABEiABFoLAXyj//WOZFTWvcEbdSm3lDjB4GoJFZZjSVXIqrb/QsdQda6+6AOlsKqvBuyIhjMltR1KY3TpEcTS1XyYrKrDpUo7nl2vCD7Ks1rHae1VKyOOah/lf7yYNexwre7WOhGn2oi5TrJjmDuIqesVqvlhh666qWl5zHuHK9V2vTycIICW+ITgag2pYyqzssONnQRIgARIgARIoM0Q+H8AAAD//woQvPEAAAAGSURBVAMAybIyjQWES3AAAAAASUVORK5CYII=', "After Step 6", "Related Knowledge — clickable reference"),
        ("product_families_reference.png", 'iVBORw0KGgoAAAANSUhEUgAAAZoAAAHaCAIAAAC2PBhYAAAQAElEQVR4AeydB6AdRfXGz5mb3gNJaCGk0AKhd5AuvVcFRRT+KipIQlNIAiodQVFUFEQUBaQ36b2HkgAJJUAaAdJ7L+/t/H/nzN777qtJICDgnTf77TffnHNmdu59Z2f3RQzTFi3yutCxwj/vdZi6cOHURYumLlzg2CBfONVswAVGFi6cYnWB4wrwyQsXFuuCIkFZLj5p4cJine9kgeNy8AXYYExdOGnB/IkLF+Q18QXe/JScINSFCyYsmD9hAbhgQjk3pahXuK3AAl+lFVgT+8gWzJ+4YIHVheB8VxY4flo+aWH+3XDifCHfmYXeBJvmGFDNJn2lg6ioCocj5wpnBT6/ddDAj0gItv6ASj2uYjropkAwjofpxlXropiiSRd4hGJjfiqqGkQ4G2otrmV64irm65wAokGNqxouB2eWmCpFYrCGqIcwqjZ/BYkDKkUa1enERjgVvYo81Oh4S5CgKJrzYFwNxVEdG+QioaQXuWKfOFjGy/X/PlepmUO6XmngWoI2fu14ia+YOtZwZU1UY8lXsVSKoGgT3LqLNjnnJOk7wHfYfBlKRSmCbicOFVEzzPVlcEYw+2BeeEiQyI9wxOgojhX+Ga9DxjJHyaKw7HAYxJBxs9ggjxk/kSPj04LjnEWK0K7LCUtowThmHjVKhg1I6IjO8AIlEgaG6RMvQzOPxDZLODY2z1j0Yp6MnexruIXEwk704hqZggnua5yG8TTziGGU+hzFxqPbe43jB4cZ0s+0M7TMvUF6QOKBiTPpIreBM6zNt3EuXC1DGjLFZB+FgUzhSlFyXRihRi+3+e9wRi3NM9ps83nW5674KtVfB64PX/ueZJzNsrTOtEucC4fburMKmfiKEdHsyzgUJenOI36Cl/kmDkZzZ7Ix2nQipfHvkpnyKeBSY8O8xGaBGI0HVVG6VYyAFc4iUD/jdQhiCx4YSLmliARVlXQraozTa1bc2QQfpWlV63Ax0ezQjYe0W6GFWORmw3DKIRpEOauISjl3R+ZieqBLc67OVUSUAkgZ1zpcVfENoIiocVNEiAWxKtoAx6GkN8TRQr4O7l2LCxdkhxkRhfjCwIqqK5dLWcz/FteyOXxqLjXrw2qxhiyh1lpbRedIupq9WqG9TI6diGAJUWKIqAYVWkGEEzzQEIGr6yrSEJcynXjYGOISVILlT8t5lhgrPN0BPst18Jtf5veiLGMDAQMFyieQcXCXyRBTRacn1ZhRYqQTh4yTcZrYgKlKZqJ11hChyXWB3B9BuBlHASVGRLteu0VmcI+NnGWZZHS6jhhjRJGIJtGKSFaHRzqRYpaBHGA5x96VCMIjBcbYmJZzRK9mk0U6qZzAVMu5uEF0LHLBTLKIyOWAK4UTk1AgiwA6ly8Aj2kO8hlcLzGLVxrLecwieqqwREDhyHzZ02dazmOMNJMeI5wW9mA5TwFilqHbUbKPlNrft4iJfQ+5fInYChbwQK4jsYlSRFRUQVXDCldd2esQlV2FCHc8FYGqwG0U4+oc9O6gGtTeNwUUpV+CKbVRsBFVR3GEi6iYvWoIIooOqpZzRSnqids7MgYwPZgS1FDVsBEeynTjKjZx0BolroqiFDG5xBHqcEHCSwO6gKX3NfihO4qjatEGXbmuYPa1uGIjYrpIfa5iXqqO4tgw570hvVElBBFVx2VwNUtxSxzwWgnco+VxGuRR87GUt0dlXJafB8E3t088qCnGVepyFAkBVEfjpc9L1XS1XlF13jAG1VJvzjlJ8f2a2js7BD4tqKCrikU0RG+Ac/XC5oyDVGgoYogUOVe4r8DKWBOWl42OZJGdDreSxBFRWGxDTpkJziN2GZ9AFu3WA9JDG+csUnhPEI0TyT4n4zH5Cfco94uGDMOZyGA5T0oRPbZdKSHNy67XBoFjYmNwcqFBboGjTwHEN4MzHnsjkSxSDIlVi1ukop44FpHrMquMc85hGT2Z+DoIvSgMYOjztmGyYgSmkniUoi541+X02hWZjcXwJutWn+OLLpFxLc7ycBuraL+y+PKMi02af5pz4oa2IDZ/eJoPBJu63M1yPXGwtJ51OVdITEe+NJl9Ghx8OoTNJBoyDIyIoHEOs4+csTMdikLbkXOGn/EMHuHCGISP0XmkCLH56Aw5uVDkRBRLcpbslKznVRxTs8JX3joElSCitn9hhTVxJX4ARMFaXLnpBQyxpwtUNRvnUKWvLhcEVcFR/R0ZXE0StSKKC0OkWs5x8aollKILxjWc/pJeiwe3CRgrg4im2aoqelAroCj6cnC7isD8RRXEARSFK0GCICi4XFzMUtXQXARaxiVxG0vpEogyVoNcmVXRpsgx1sY5XVYZ1wOuFB5SKI+ZeFNzLrsWzLCn2pzdPZ+Pcc3Fhrn3qmNwbISbt41oNmVcbcrWRk9crRBKVEEaoGgdjmY9IuhBWXwNImrFsWEuKsV3Z54NSXOVuvJXgG0Fd5iEbCvgJeTuArfF535UrGaZxYTRSQkhxSpFErktu7GAhDLMch4zRnDLCAq93OZKaI4RE8EKHqPxGHM0s0jBqwwxpRWxiQxEK2HM+DE1ixGWMFKMRZRIqcNpZtZjESB5FW8aMlV4Cd1UVgQxpkYurVjLOUPUVC4WGx9LanNzyZVYi2Nc0utxTImzctFmGykZkXNeNmefW5Z0xzKOmdVouhG6arg5JjFm0S4kE8fElxfLP6kSh3hl3GgkxmhDOMZoPMYGESv0EibfHM2DGcZiESc5BlX57GoWs5mz51ZXV3/6Id5+/4OfX3Id+OlDNRGB+Ct/FNutiBZRWXBWPUdFT8rdj75w0Innjh4/8emXhv/qyn99PHnasDff3+tbP39x2Dtmg0sQJ465e84DvSoJFTPrFa2DNGvXYE1NaMZKwatJDPRatbNdkeKfeFE1pYxLPZ4UwVFL93NCwB21UcSFqsoFliqX0ARPvbYyog2hhSrqiddDLVOWnyfL2shktLbCzFEa01NvDYrPJEVomqfeRjFYTKmPGkSYTBAtR5rUpECoNVxxKdbE62LwTzNH4uLuilLghhzuxZmuRtB60vcNWxUJ2KUqfsrR/rIp9kBqT57cbUr80WeHrb39t0EeUcHEh48cu/5uJ8CpWx94yhMvvIF9/1/+mWaqR/3ownkLF6U4EybP2O87gzbb90d7HPOzjyZOI07SQfj7YydssvdJeD3y7DAU4oDoYIN8wuTpN93zJNiETfId/k7tSb74RtLBZfoSP43S4ByIUNLnzl90+A/P59rLY/b/Va2lmLtgkfdyt+Pu4SucWQzuPBanNl+4eMnceQujxLffH3/zvU/OmDl38ZKlc+ctQDF77kHcNh2JRRSj1sdZiJQUR/r5RA1Tk4GNoGEvGPOWw9DnZpxem49FMq8yjgmK9dNJg4HsdQoNMT3NCJ0ZEjuXMaKTXaihD17G2QhgCeILwlkdPMDECQLHFazNBUWQ8MSUeYCpusIgVst56o1+sRFXq9hwPY1Vn66ZYfCJOPMz93q+Hs9Wit5ym8Z1n3C9OLhb/TQ6Q1LL16EOT03Qau11rhnX9MgnQiiQOYGRUqNHU+jIWP8yLjF97gkjY0ip1xYIxS491y18pPB9y4iHKYoh42RJLyLfW7IeSIITUp5yiKqKmsYhsLpc5NB9dvz9L3/UoX2bgZddP2nqLOzhl5594p8vOvW07x/eqkVzFDzfen/clGmzbrt64JKlS98ZPZ44SQfhb743dvGSJWuutuozL41gqosWL31qyBsffjxlyGvvvPTau6M/mPDoM0Nnz503aeqsB558efLUmcxERCZNnXn3w88PffN9uyaJY8dPuvuR5x96auj8BQuJSWQRVVg+yR8zsYGXXk+SenHYyHdGjX9z5Fj2PlXV2dTps+579EU442KPxwcfT77t/ufGT5gqFGWgfFzm8Mgzw5gPNgsXLX7yxeH3PfbipMkznn/lzXfHfPjhxKlDhr1TnVWLj4sNI1569v/9+aKfshQtmhfeHDnu7kdeeGbIcC62ujob8vrId94f/+Z74+59dMjUGbM/+GjyHQ8+N+6jKRp02003GHjKMV1X6cR8RASlzzprnTfguD7rrAmft2Ax6/DI00PnL1xEH9+QV4e/d+eDz7757gdZzPgU1e51nG0eyjw0cZrEUlUjjnTwMiI1VdGtXwRBVe0kUsNdUM1ProfUnXOlmJBGUwz9HkpvaJjbjZrx2Yu5jSauqkESh8JFiSNKCSLL4PSbZSCmCqhWiCsaTFdElRJXuIp3AFY5kgh+aq6qxKgfX7VGV10xbgFx4QSm+mm4WiGAlNaEz6oBXlxPTL3X7DmcK7g8XPkcLI7Z5xy3oPBgiK5iNiiNczpF1b4RohSHYG6qzlXsBNjujOxGdihWWiQ/MoxQXnn9vfsffxmEp8QM6dtn7UP33mmfXbaaNWfe9JlzcMy7YuzZffVCKKAQpVWLFnPnL/z5Jde3ad1q4/V6ukhkemRpVfVjz72+w5Yb7brdJi8MfXvGrHmz58w/+5K/7f/dwUf/+ELy2iPPDD353D+O/3jqiJFjTjrn9yNGjmUIhj7nsut/et7Vh33/l3c+9Pzrb4/Z57hzzrrouu///Lc/OPt389gKxTw+lj7JHdMkJ06Z+Ztr7zj0+7848IRzb77nydffGr37N8865bw/Hdf/su+efvmcuQueeenNXY8+47Tz//yL3/4TX8YaURyXOZx87h+YDxvMfb8z8DsDLjv9gmtvuf+Z6255CMeXX3/3N9feSU7kqqKNLu3btN5jx833222b7bbo++gzww743uCfXXQdA/H8OG/Bwt9ee8f+xw884LuDTx78h+0O/uluNuhfjjn5InayDz8z9JRz//jBR5OJwxy4Rw1/Z8yPuPZ3xn40YeqB3x10xgXX9P/l1SdxpfMX/PXmh4466YJL/nTL907/9Tvvf5hlkTtXlnESHIkQcy7OIx0ZKRCrLOcxRjOJFGzAVIs882YW+cEsi9Exi/ZDw89FntpIVjkyEzjbybn5ZpSYZalyEnjkzImJZdiiZFwFomNT3JzMK7OLMoJvLOcxWpM4WRYpdMUsJo6IEjOD/01esw5ZTNzXx9aQRanh1otIxQxMNXODaEuXYSExa4pnWYYl4JhzmpmVmCFYzbKMIKlmMUskRlhtTsu1zDHiBXM0jn3KkOS2ulXIeCJX/+s/Jw38PcgvGArJEPKfJ14+8azf/Omf923Rb92ea3dDnDNvwc8vuQ5L9lw0qaTD0R9MrKqumjhl+m8G/2DuvAX99v7BP25/lC4qSXDYiPc3Xr/HDlv1ZYc1ZvwEROKv0qn9s7df8bMfHQ1noCRCrOnz+dOFpzxz2+Wrdupw76MvbtC7+3N3XHHDb8/cdvMNRowcN2nKDOxTxaV8kut074rSvHmzO/5y7h8vOPn62x5p3qzZU7f8mok9/+pbj7/w2n2PD2nXtvVjN15yzSWnYlkaDkLApNx6/zPjPpr8pwtOHvnkX0/93iHnaB93NAAAEABJREFUnvqt1q1aHLn/1267emC7Nq0wS/XjydO3OejkdXY87vHnh+35tS2G3PM7Ruyzzhovvf7u3HkLRbTvuj1evu+q/XffplOHdo/ddOnJ3z14xqw502baJldEuAWpX6kGoYoYkoI/njT9tO8f8YNvHfDG22PGfjhp9LgJzZoVBvzfEU/cclm/DdbBUlVyhDRQNZiYozpvCoNS0l0TH2sFzupHgyjWxZyVXilxQcGZ+7NjruccM2qyh3hlYuZSxpNShqISsGkI1XVHdUxxEi8iobRxnnqXhUKEok19npSVizagSoqZ+LLRVgMrd2qaW6/4iqljLV5/tbERNZuE6jxHQlGTAqlVFcdiTdww5N8HLX1bNClqvc5hzjnTVYbWKvtm2u5MyD08tNZGdijovz33pKH/+QMIT4oRka6rdDz31G//5aJT27Zuxf1/tS6dH7rhIix33mYTDFDeGfXhpVffcvBeO9D1w7N/f9bFf1296yoH7rld6h0xcuxHk6Zd9fd72WpVVVend3DE32qTdddZazVs4CBxwLz6DHmSZcRWLZsvWLj49XfG8FbuzAv/yr6JRzmeBJN9Qrxskj/99jUXn9q6VUuaG627Tt911160eCmPwEQgTqcO7dA/nDAVhV1V547tmhUKKGl0I8U1QflwwhS8enZfLa1eGqUOMscO7dpccvaJV194Sr8Nej3x/Os7HTbg4j/+e+bsefMXLOJKide+bWvSX+vWLVu2bN6+Xeu2bVozEPcuhoBIpPiZdwOEg0aZOGXGkqVLr7nxgVvue3q1rp0LIZz+g6P23nnLcy792x7fOIu3kPjaTHx7aB7uiMI5Q2TUWojMOBiCiddCGpnPAoBHGmZIPIFbMIJKrM2JhpIjo9k93Gw4TEfxazRukl9dmSLOxceQWtw8hQHpSghprNZaB3esp3BBgrvr9XlScuQUKXbB5sEyoCREboKn3pWL9cdNSsLSWP4BMDWuket39Pk3cr30YbMM5OKbrqXPxUZJnyxTYPzEazDmn6wpiRtiiB4pmQdwRGFaRbRwtXlSfF4sAReBuwjZjRUg3dVFURTp2KFNty6dQWuYYgdZiTdlJ35jX34bRcx34eLFb7wz+pXh77/21qiq6kxVFi5asnhp1Sod21941nf5TR725ij45GkzsWeWT780okP7Nrxj+t0vfsTOhS0Sj4pisVUd4eSdJ154/fb7nxMrisJ56Ij3H3r61YlTZ7DHef2tUVl1duGZx/N2iS4RVRXJUQ7cc1vin/jNfdq1beO6eFFSySYb9iLC0y+9+fAzQ0PQbTbbgBw6YcqM+598mfhm5nEgT7044vYHnuVaRGW7zftCbrz7ydffHj1m/KROHdoSecr0WZOnzeKKRFRV+GFN9tppywP23G6Nrqs8+tywtVbvcsEZ3+2ySkd6sUmVGw4ERfP1F1NUhQJTWbq0eva8BZIrut0WGzLPU753CLvLi88+YaP1e46fMPmk7xx05S9+NG3GHN76YWm2SpFyTjuoC6CdRQzVUaSGJyVHdb0Gg6rNUwU9gIE5CkqjXFSK/7Kbm7Nib4oGETXeEKrW9CZeg2pFDIKjOha5N4hM//Jgck6W9XlSciSiVR/ACIdo0GKprRdVreX82duk+dg4duSDO22M17/2pCwXqhUPX1x/hJB48Ah1MX3iqqYnXhfVvkuianrOTVExFEcV4yqNoKIzD9G0O/PMBpDyLM/x60nGFCsolgGNkjnpTAwNbmgn4S3Szy6+7qRzfveba+9cvHgJ+mZ9ex17yO5/u/XhY0+5hC3P4fvuRBZ4d8xHhOY38PlX3tp6k/WP2G+nw/bZaZdtNxk5+sMPJ0xJ8fHFBrF929ZXXHsHr+d4SGQQelu3ajFk2MiBv/77uj3X+sG39ttkg16Lliz9dv/L5vKbz7TK529NOyxartMkMA0ZcOKhu+2wGW/Kbrv/mTN+cCQz+cGx+63Xa61zr7jh7kde9OGk3/rrkCVvuOPREe+OI/NKlIO/vv03D9rtxrufOPjE83hq7talEw+Mz7785jGnXMxfIVkvnuGJzohwdi7wbTZdn1zz7f6XFoKii9hdSSRiKXTndjaxmG5w0Jhtt/mGzZsXrrr+7kWLF7sQD9hz+6MO2GXw5f/Y8dBTL/njLZOnzbjroecP+u7g/r+4eqtN1iM7s3ux+B4T7sEIGTm4dcWywui0fPRo3EzNBM6sYswyWh4CHmtzcV1Q8arHY0aJ0U6Y5Dxyp80yw2gozulGqc3xk6SbIV4IIE5gpFgnXjCrWSxxgXPx0cqKchvL/DhHjizGGDnb4VGdRyYdi6UWL4qckzmEmng5IlLLlcQRqYmXIyK1XEkckVqLlzXyuZmSxQzD6Jg5Fnn8pGvlayu4W81icf2F6LYTR+CTymI5x8a7xXztu0833DE6MmW8nMfanJaFpBfG9y1izy+MI+FKinMs6Y9ifxPjV0ZVRVQMOSuH7LXzlh8OuREUTfxf8E369nrvqb/9+LiDzVLFUOTK8076cMi/vN7Ii6S2bexFEm92Ljjj+LHP/ePtx655/aGr2YW999R1h+y9g/CgumrHJ2+57B+/OaNlixZE+NXp3yHmrttvOuSe3/32vJNQRLTfBj3feOjqd5+87pY/njPm2b8zNPW9p66/+6/nIT5248XdV++y87b9Rj7x11HPXH/f33713lN/Y27uK5umSX7nIOKYosKOjIlR09zat2v7jyvOIM7oZ/5+yncPIXd0X6PLI/+66J0nrhty95VpuDW6rfL4zZe888Rf77rm3LcevfbH3zmoTZuWvx74fwzHu7NfnnZcq5Yt2HYR5JF/XrTqKh2UwYKyFFxFty6dudOIyrGH7jHyyeuY28P/uhi9Z/c1br16ELVd29Zc6Yv3/K5b185EHvnU9Ztu1OtH3zlo5FPXbdK395b91n3jkb9c/5vTD9xzhw9e/Odeu2zZtk3LywZ+/92n//bWk9fe/ufBxL/wrO+Neu7vbz/x19uvGbxal06M5VeqavMQ0M5KgedUVKzNPBWWuPo8lSIicKVwu+deF7DxXlF0cV1Mx05NEVW73zpK4EeDcQwV3bmo3W+9K3kRAM25SlEXVTGupouKBlEw1ZyrqhRt6NUiV4pzFX6UQ1RBpYifayMyqnASCidaoAnFk6SGYw2oiDXsULhSitzOdmgdXSmfpW6x7SgOU+Tp7Oig6ieff+JQJHFdRFRhCZUitJRDVLW4tkpxLkKPapGrqiQOsVr2eaXP0e3dJogopcQFA/M1PX1P1BR19O8PDEWdS9INVcTtc64YSM7ZJnKTlkjKFopxGjBhP8LxKXWSWvu2bVIcRiwoc1/e+Pi28XdeZXOwebZp3VJFmSJ6yxbNWzZvnuKj+JyXNz5x2DbiRRyQubVr0zLxFIc1a9emlXMbFxs4I7Zp1Spxxm3dqmWhmf8ll0bm3gaRe4XZZJFRCsGumrtIlMidBkw82h0mwpH9PgYnSiRAy+bN2rbkpSTmYu1o2KJ587YMzWeEFKVZoWDrY102FHNDtpC0EC0SGidr08WJKGAZhzIacg1ag4NAeCc5ccPM5oluPMIZx9H0iDH3UvNl0Gg7Ua7N7slinO5ybleRdIzwSRUjI4xAsFSdmw5JFT0RTMu4tbg678KeAcSkmMAxsuTuEWvQOuwwxc65k5+sncchsjuZboeb09+QbkMuS082hCR+4gSDW2w7LH5juvWXxzc7szedECmohzOK4twAbpLTImf9zdl0P4jsa2hGiRjDxBwYgpahd+GQc+u3eVgTtdSE8GUB+Z7gCTFH/ya4Iv7dsO8Pvc6j6xHOAY8U/17Zvg9nwjFlIuQ88qEK3EJnwTMDqYar0ApXXeF1UFHlAEN+FucgdxFQg4JwdV2FvQyWZEsE44IFNlrigoKValBMFF1UjYv656WJ04elqhqKiloRznY4pwMO0gJFFVQ1LHGFqZbQWhxBxWcAMhoortThihrw5VrMR2x2xkVtznhxzrkmHQ9B0QAq9uYWAEVRKwKgg6pCAw7m3JXGuWIpal6qKl61fhHryuX6XKyHEJxBGuAKcrsevOr5luuNcZtb0bfcplxfXs4ErIrZG+GozRGofIBgqYpin6YAodbnpogaYisc5dxmrarIymenlKQEMckRXTR9H/x70ghXMRtVbDTnYlwEXUVqc0wkNPruzJIdCdNScqxwWwGWgdUorUmGFtk+IRuio8SMH+4VhqjcSWBZjBklxpjxEzPHaCjOpTa3u1BGEbOPEjMsJXGzZw4xY2wBjFuDw4aKdQtuSPl8MDK36Ocsw12yDIxZBkoWsxg5QOfiXBLHwgbI0GJGidFODlmMMbO7aOboXJx7BAKYXuTYCodfYzS0hkRMuANHL1mkGb2Icylx5t8kxxEPTAyzzDCiceQ8WqmRo3fWQhvSFj0yVhM8rUZ9m3J9ZXFG8XmylJyFA6WJudHrNnFZ6AvhK5Nlzt0hy8p4pAgHMV1dNncz84Ckz9RIrPdZZ3xnhKHYZ8WGeXQ9mk2M0WxiRonRTgYxrTCY0csvQ0ZqFrG0CXASARTgJAJUOCvQ4DrYyqkGFX5AUeeOosaDo9o9RJ1bm7uTcZaWO4xIiasYV9dVNIiCShFxTluVXs6gaokjKIWTqEAAO+VGtFSNBzXVuQS1AsKC62Bdbk5MQ0w3MM7BZIJy5gDpNNR0LaLBvVQk1OZapquoFcFYi9zsNdmocy0a2dkOM7WzGNBrJzgnoUVvOjlPoKmInRwUexpFzmXoJ+ehEd+QBlER1kRr4i+XXu5bxrU8Thkv11cGz0OrlzR1tWlw0BIVOuCGicMU5rpaEVrqBxgE4FBVca5WpPS5i0rSRSn1eBDTwSC4cBgGBTkMg+JerhPC/ntnYk+eYoUEV+EsRJPrwI2Kfsco3BmwN6zPTeHAJmKd2cEbAe5vjpmj9Qv3KY+IRXQOsc+Bk+n2+UjijGVKtF7kdDak4RpmZmNG5lKbR4pZuVHiNTNnPu5lCtw8bYYcmfsYZrTMLyN2BufkmLg499kQye66Rc4ZN0PzgHpEcUQ1Ave+ck6U1MQmcdC4nYpR6nCCYOGIL1dKCxPjHOjLwqzM5hNwWzemFvNP3ahzrq/EbVbWMJtyvRYvzvMTzIHrjWVXsTyc0WxWeOHsaF51uM/ZIOmYWfUrNkIA4xbKmxAqcUCrePoV1uI4JZ0ur0X79P0xFP9eZY4RtCNjBDtzpN8YdOxIfUxBRSu4vCugymoVD24RtBpG35dxBwmq3E0CfqaIc0dxJBS6cXrUDtEyFMWA3rpI23rsEHjxUEwFUVTqoVLEdA5JPKj4DJtEUbG9Bih8X5hbEcV5Q2i2ojWoqsIPaL4wEbg2jVpmk3iOWiwYIBkqZzuMi3FCC4LU5zaDYm99npRPhUqRAGgRhUaRQ5epYyPyqebAiCIqqpLH0SZ56q1BGPY5EkS1xNULvY1UhsU21cRzxE9M5pAS5xuUOJ6OOoUAABAASURBVBIzrY3574XrNVxp29qYknNTVIQ/unmmtbzIUeFNrgDP6pG7gu2zYn5nzfcyuJki3CmKSsQaS1fgZu88tzEuidNLq4T2OVhbUOCc6mBMhQ4ndsYC7khQlCKmu1hkLjY/5sNdrDaKKzkybnH+gg+WhmL+zrGFl1DMXhpCs4oUHx8rqM+OWYjzyAnFMfFGEfvGKu7RZxPzMXw8V2SFMF01WwW8ipzIXL1hpCNrgGdlepPcPn2OejZoFr+eno9VrpfPoTHdbYhZc+3Fa6lRmDLXuCxkVmVrG8s5gzRYsWm64oWBo5359OHpc2M6Jc7s6HbMGsSIT4ZJRm+J24rFjCwnIspPBZe5Asodl5UqotkHlfzOoAJX7hgqEuy/m+5cHaUGxbmj5mg2KqocKPVQyhTn1lZNiINAhZY2xPl86VbmxlFCkXyewpylxEVsnuK9kri6kjApooqHuF7kKiqi6rqKiijFQFSc5chJ1HpF1Ws5T4oiWZeoYCOGHKLaGAqj0vkpMUgpTixyJWYQQ1XD+jwptdHiFJWVxW109ZkUI5uSuGodXpq/+PwNMYF/clSp8TXOgVIHo9mgFasUCXrdKmJKjqpYIuQYNdeRJPHgSg3yjS0pxtW/jSAmIpYn6yFCRWcFfB3s7sH9gLsUit0bfTvm3FbP7oHR74rRiuvR7iH4cQ+JAoeC9i4ABacMK2wsEDomZoAOc2TcpJidKd4CvG1WdtAhScPeeixerkQ68ooS/V4WsbHZMjqWmc+5hIRI15b5jNDN2v0ymMBs/pGzYJE4YWwQEYbDCoxGXYYZxxcZpMEYht6dK+ZO2yJYF4dbA/inahrWSCVEco6b0RKHUInJLMsRsZaSESza9WMEpxs0DivpiRMfJec0LA4trrYMbR4ZPcIRc924OGetSnri9XVbA+uLWNo5w9N45GztiB6dl2Ot+dSdP+3INXKpJSSAxWEw4peh6dGC5b1FXqMnxZCD6RtaZM4+jnNvwIgMTVYQmonXQxMwYORaiANRfPDIjOjO0TqynLMq/j2MYMSW3Zmk3MjXTKEckg6p6GIrIbYOdkCUu4BK8BYnMS6GwSyD6ao8wKvAsQ1w7h6KKIoEF0ixlzMiHUHwV1DVUDhhp85FrU+tYK808BEV4VDVEkrOXQ9Jt4jCHGjRS8v1UOLMR0XAoIamq2KPIkkxhp/a4YoIPKgaCic7nBspvWFTpYmVKPYqakX8bG2F2VnUSo5+AqjKHBQbtQgSDFVpS1E3LsFQlUCiOVdRqdGZMzqoKsH1IFC42QdV5eroAoWG64kbSjCsrWNMHNcFIhrUvMRQnAdxXVUCqqE4D1KjB1F6pZ6u4npjcVxXEXy1bG5Sxou6zbBMV9eVk1gQFWagYmhNoQMucIWL0Jtzs5Sci6hozmEiIaiKKodozlHUdVVkq6pFNFO4cKaqGldVa4NqmqTvj3JWCtEkCB0CU7smRJrqXAw1cUM7An/ZTJlVyJKkQMtx+VHRyffk/SxjXbhzGxq3+xx3h4xNTgYX7gxwnt2zdNAwHU/JMjTJYiza4Of2yZKeLGesOTaRkvmo1o6uSLSSGWDLiZjW658XczDOIXTiiWBRmWeGkBU558jYuBpaEClyHOFubR4xy+BYRUaHGUrijgxrUROPFLfBChqNMw7nItqXKIoJoJmUuFhXhgSxU2RcWlRaxl3PxHQQaxAdzHn0WYohl50JfmW8pMeM2XHka4JXzHJuPs5z3yJHlyLHviHO7wtWaVxbd2zcUpiwZD7DLOk5X249j1nPvqhzjcuYm12vlNnUv97oM4xuEzOzj5wsMpsf5p3G8pnX12Oum19tjhJjlsVUOOccxuWgCq5WESLclTKeZTTQwZjxE5lZpBjPmB1UMsC40Ffk0XkW2Z0pH4yKVrDhFVDuDKxOYIGUu4TknHsFiqPWoDp3xBDdEUPTS1yFHxRDDvR6aBYiolEV9JZKztWVhjHaPBkXU0ORBpEA6OK9IswkENsULeOiinddVDGlNqqq8OMohtYQTZZ2kpIoKnCphVwjbVFNKIwKFRVDm5nkyrJ4unb3ikHwIjJo3hxBtB6m3k+KFq/oW8aZQJDPUPfrKMa3T21FudZbB1OYtr3zFYvmHKsyrkI76ctCZYZFm8QTSq7T0locY6qogDWVlgo/pogqvws40S5D1lmLinF2Z0IhU9ZDBO4uJOs6+L+ksy4sAMg9gdtuLeQmw8pwG0PNaHDHiJG7RKTAMS+ihciw5X6CHd2OpsB9OW0EMRUv52bNYbz8oNvsrceCmkuJWwNbj2JxI/MSv6MZ+jxzhTcOphMp8/k6ojAjaxMoo8WpFtqAkdFM58wQjkSBoidMvYbe4QQ7McLUsPOKMWpCuuDJAg6phW7vSlYLTa+leDSuk2HQ6QbLuMXNuKLoh2HZOqAJR1IScr11lLq6xSh6Rcar4cnS8LPWs7JxM+ZbnENWSy+/FqEnwzL6UYO2uBl9KPlKRlpcA87o9ZBOKgtscejF2tHiFDmB3MbbmNI2NBPXbS4NfE+QCYUVfnWRDqLgTQcjG7rkCrszEp9YhvPDc50zkTKumjTD/yU9iLIOhsFXIARFCc41FDlP/MHukpwV26AgYMjJekxRCk3Byg568QfFdDHuh4KiqjYuqEJE2o6i/IirqsY1ODoX5kBL3dJ5QKcVLEZwHkINV+PK7IqHjQRXCirDgCVukZVe4ikxNIB2Nj1xVeOqoKiqGKhDbRTMleKRmJD1wkVyrsJP4BCMRAV70DhW6KCqBNeDqBpX01XVZiwBFFXHMi7O8cAXU7A+D2rRvnq6as11SRlX9VXytRBfn1wp4xLMV9WxjJuuolrUE1ep0VWUOKDaikoIWuQKNa6qYiHUoYTWKurO1RxExdw056p5W4gOV0qoeXdWlgnJdLRSHm2Uc0+PWeYYHb9CvLq6OmZZdQZG57Gaq6uG25XCs8SxyPyoBrFxpJveDP+InMFjltHjR2ZWdmQWGz1imHNiWyOajRECxSz1VSdrLGJW7Zg5GkdgjMzm6ecIYg9iA2IDZhapOuc+GxqcY35k2MTMMFnGrBa3OWFZ1LLqzMY3NBWeUaLJjfKMiBwZ4/op4m08ZpmPDecqM0aKWbVj5liPV5uOE14YOlY7Zo7GiQ2PDaxJtJgeITIaR4M8w9tGoT9rmltvzDJbBfPxI7WyT8wj0ezIOGdEiZnh8s0n1rM3pehbj5t1fjS8bpmvYcZMIitlNnBbQ04RT4+ceJbV6FmWOBYxyxyr8Y+ZcQCeceLIjEYjfHOqo2UdgPyTcWLzlZAGmQiMGT+mZYAdkT0alCePmN6dkRVlOY5IvLJjOTyWJ+oX0sZ2UJb7yfvCHYE5GnInU7tquCZuaLYcpohqMPsa7go+ig89Fk+VNno9FFcc3U4cG7Mu0xkd2yIyb/HRaqOoMluhl/dKRVtXkp4j/aLqMwGZAUjL0aLSX8ZV6Uv2SjGmKHbU5WI68UQZyTiRRBXmSA98+RAruzbi2PzLD4vpvfJJ0LztwDfNhSg13HpUHXmnWdJ5p+PcZ2O9tkrMyS0/IZdacfi8VE1hrco5s1PN9UY48yBWskn8k6D5pEP9Krne0njpClV9jGUiV5DbWAyLplYk6YbCg6YfPEHSJRyKbmfh7NzGVEmoiiqqtA2DmH8xSyVehmTCjIO8x1uVMj3PbOVK4ilS4uX45dFL7wm4M8C5CJt7ZncNFOPROGC9GavDmdUB6Qe9x84cpvDBmJVw/0gKAnpt9Lh8FljQASZee51tXMHVVOM4uSWj4+SYoaPVQzTGj9zJnHFOnJb7xRxjMTbD5JzQRE1IL9zmxgg2FeQaS/RoOkeu07QRxKOXIx3EkmRZhqbb+5rIDK23yCMR7L1YxlUIfTknBgrjGkZviWPmmJRybEo3u0jsiG9kThljcTCNaKorSbd2VtSxruEZPjEptfBT62l05pEVx01Kjhk9+dxQbAZ2uMo8Ei/HDNWOaH5l3DQ7anTBwpSMs9inkFY7OmfIpEfvLWIUOnKFb4LPCIlxTDclgwvRrAuObCicrRd/GyXL7LcLNBnduk03J19n5kZXQnxiDCQ3qmU2jrLqwwkCvfXrV1gPwi3IqiqP4qpiSyTBEMWWQo2zwcFSAh0BnbsXWKo4OVfQuhXgwMOUUhAazhMoJ7MTNTtQEfyQhMHahBbOxhnULUMJJaAH1WDTLkOcUHzOQem2tuKeziho5ei6C95hPk4kaelUhokyLetPzrnEhNSHdkQUU0SDGWPpSom7Tlfdis5sU7Vo5hho4k0zr6pODOk3rloHl18PqsRXlkwANW6KUuDEsTU3BUN6kUF61NqmJyXhStUlj8mpNIc0H4ZmJLB01egljn2Jr4COqfg4EnAnfi1U+xRQalc+TzN2UctRWB7mLyGooltc50IcVQuuiq4U4SCOISehWEIT8hISnRZBtYRO7L935rnOkhxHyobkw8wauHKynGg2UM+Axk2zw/hXRrf7Q5aBABg5cQ/Issh9BoxZdMwoMEdxtFbMe2OWRe46ACyL/CCAxtAjJeOItGN5wQrJMEabg6GtLqPHGB3pFHzpzWJmCGQonIooWWbGoGTYJA6LfmSUxIqIldvE6EoJ0TMaGWrMHKOjcIqUjCOW8XI9i3jEhNhAapC4kd7M7J2ztvSWUOiUDCWTzEoRcSopiScUbDky1ieLGeuVJUQjSOIxixEG2ucYY+Ixy7LIT0JjMcto18ZoSkx6dO7IXDinSKZmMdLO+Hw5YnRejuX2K0u3kbIYfaw8fppnOabecgWOT836uAUKZz6FGh1GdENbW870xsy+dfURmRmUEEvJzMvCSoY9SuKlzx0VBS9D1s2M8KFGujJUznyejtaKFGTDLMsib9yMeneWZTFyZJG/bHquE0uAGiX6ObW0jCfF0FIkdlL/sMRZX/0S2afrtytTvxZDjSxDBLk71KDy9kTVketDd8TTdLc0zhHNhrPFQ4cVV8iUxBPDk94cYZHFtv681xSFK7r6fBKKc0d1jEIMxVPpcU4g45r0HLEVCnaGipfknLbwe8/MTVExRDML41CUIpoPJq4YNwssiReFt0pJp4VeQp8nAXx+UoOuY+WKgvnhM1aiK5rZa+Ie07hpduS8qPtZ3NZHoy0ZNpIj0Zg61+nIWYpKzO3FFOdJSYjw5dDrX5crNv+0DlwvjSJytWXXhcpa2eqhKy0x7ocmLOlKb9lhusVHTdF91Cj++SalFtbSldjCJ84RlLMqKKplGFRp4YVN6lEKjI7A4ESIZDZyJA1mnaPlRlrWS/LDwpFOlKQLDQ5GNUSjzXxolOOXQLc3NCIgR2T6doXCvJ2zDuJ6hGWuZrBoKyJ8bsY1gqwDiC+NqNyeREBi0CsUOiUWVwZZjHO2caMaSkJBNHs/arjNgi7/nDLJjJYfhI/CfEwjchSLyVTELDOfSWQmpmMqnAltikhCoVfQbOYRe4bP0eyxoi+ajbjcT8m9AAAQAElEQVSloHCIDYEB86cfRHP0eTIK7dyS4Jhga+3UU4bYE4oVMAObczSOt/EsvyKzj4l7NNOJlhQQXkdnfsQBWQeRzK8omiKgTUdstiXdFL9GRoJniZejhZIvg17ruljHetcbUYTLsdWIXC82XBcYWUnaIN0Cs+utpZtmh9lzJkIRRezziqbY+uEtHFAfAStauQJLemobN19J34Tqan7Dsgyk2/zEI4B2TgdfGLxMwobfykCiIyrpDVKnonMJYJmOvVURsmOqxCmrUsYZqFS/2DqfW6qiXBRzrYemc7DjUDGjIqKJcUSxAueUVONKL4tQQjqppvhCCYmBPqEwAZbaq4+ODZVOkK2DIzYamYMII0TlJHWwlk5wbBz9LFaMiRQxnaWmIFBpGzIorBxpUnPFTHwChPNZRYbnunIuZpaupYQQKiHAupVLE8HXdOcWSqI1haEg4ko5NqVjl9trtABKBDVCh0agjKs1Gzi+1DrXWJx//etNPcqaCN8QXwo4Kl6Sr5stfuKmcyyHHv0TFGLaB2fxoxELBWEcMJYp1hTJ9YCZSsFeooFKCSF1iQ1NW8VA00szUWSJYpKYFLKMm7ElOfKfnSzTxSxmlDIFF8lQXbczpnbHirEMMzf50mF1Vl0tWcKMM1dRwljtimF1rDYbV+DoCTNXHFkhbDJ0bilghhG9MbPeGqRp1W0wtooBcwB99aprMINXo1cLkZlnjlnSG8ToNoZZzRyyjCAZhfnURdpYZtUEjTlizPT4oL1WN4bYMESOglk13HydZ00g43BFseaKbG2xz0xhIlxvRnHF4uQ8a4DHajqTbpPBmbCGGR0o9DKrcp6UOmjXiBQ5suXg1WU2/y1ePs/yOZTrOWcFsixrCKuTzom1cswcqxPSm9aWj7WG2+fS0GdhevoEM/8EjadP1oLziRMbGz4yqY6SRc1qYbRmllVjbJhFSiYxi1kWYxZxzrKsOlrJDGKOqDRpBKV4dtMikvCKnF7LlyjkM8t+So+1REmWTmINqqiVLx0GuwPYn1aYeQh2FSXkmo1zt+DEX7kScpXG1XoTdwRQiqq1GjzEVUfuPDQcfVzcTS/jDChBNTADIKGKj16D6koRCUK/I46ci2hnpRCvhBD1K1YrqceYuzkBklxEcUXwYlbGi/PHIKAqvsytSbR+O9SvKMdaV0Fv0NqKBS4qZVwTz3vUrt0UvEscE/gnQfOxaOlsaHE5+8kh714ejp9VN3XIfU3kcMkh15fBcckN81Oy/6RofvVWj8+Bz9SQ0bAQtZUUdXRr0TKelLponzJWRLEIHIFPSyhszcR3aglVBeKoIah602yCRyikb5rlIoYUVRBwFBFoiHRKBKmZZTsSZGqiZqZInD571sNPP/+76/71s4t+e8qgi08eeNFXqZ4y8MKTB150ysALT2kSTzYDbKjJ+MKfmpL4RcYHGf9pEX868CJ4AzjI9UEXnTrISAkh9atFKJqd6gGx6T/oIjgEdH6x8xrsP9A42H/gRf0HXoyNIaRePXXQxaciDrrYDIp8QJnSAMfY64DaSHP562mDLsG4Dp426GKUxvGSst6cD7A4hErNGoIlXWB5PX1Qboa4vHxw7mL2RY77aZ+MJ6+EaTKJl2O53jgvn09dnrwGXWL6oLL5L4vbdQ1qdA1ZT2qyKWL6vNIQiTeMdT5rml4Z62IndZHlHXTxlX/8242PPvM8ySeKJaIs8qRiL+aiRNu+kazIXZFklQGIMcuCilpeAyNH4r4jiyZPmznr2pvuOPHi3/zi3WH/WLvZrbutc+tBG956cN+vVt3IL6cpvO3gjVK9tUhuOxj7XKSrXKe5surtB2+UKgFvP6QeP8QmkOuptw7SpBIEbKjecchGqRIkEXB5OGZfhHrnIRulymQSAT9rzhD/9fpZX2MT8en6zOvBfW/fveeNazc/f+SwH13y27/dfNe0WbOUXRhJKdomTCVwIpuJnYyr8KMhJbYs800ZDSETOpf49JBX+l9y5XWF2cMO2/Jre+80aJuv3bnF7i9utc+QLfeu1MoKVFagsgKf0QqQZO7aYvdB235t1313fv3wrf7RbPaZl/7u2SGvZLxrixkv9UhUGZsx5xCaMWYxRktsJDlyHAlPyHBRSIIicv9jz1zy6COvHLzZnttv/cgWe57Xa9Njuq2zV+c1tmu/6nYdulRqZQUqK1BZgc9qBdqvSqoh4ZB2SD5f336boQdvfsXjjz302HOiVkR440busrdpcNuYWdrSkEX75ylR7E0Zz5/8fRXl2Zde/eOQ517fq++5G251xjp9t++was9W7doXmkulVFbgC7wClal9lVaAhEPaIfmQgkhEb+zV9y8vPffcS6+yC+NPomSsjHdonHiejPyZNSOJBRKbWKoTEUt7vFibPnP2NXf+Z/TufX+97laHdem+fusOzTRIpVRWoLIClRX43FeA5EMKIhGRjsbsvtH1dz8wfeYsieQtVbG8pFKwRlRVKf5lkzSmkX0Zye6O+x99f8seR3bvs/+qa6zTqq1USmUFKitQWYH/6gqQiEhHJKVRW65994OPRXIVWzN/lZZl9k/YMsli5N0ZGS2KBrWEJjpj1pwHR4yY2HPV76zRu3erdv/VS6gMXlmBygp8yVdg5U2fdERSmtSzy0PDR8ycPVfYmkW1zKXs0digBYmWyzJGzNIbtJi98tqbU/p0/cZqPbq3bM02j65KraxAZQUqK/BfXwHSEUmJ1DS1T7dXXx8RLWVlkeL7Mh4ro5DhRNOP6aIjx4xb3KvbDp26rdq85X/9AioTqKxAZQUqK1BaAZISqWlpr27vjxkv/NWS3ZlQlCNGVWV3xlNmtCRHestiNmHylJntWvRq1bZ95e+YUimVFaiswBdoBUhKpKYZ7VtMnDRFYrUIT5ZZZJvGSzTfo9kDp7A9E4mRk86fv3B+kA7NWshylDfmzjjj3Ve3evG+To/f3ObRG60+cmObR27yenObh29u+/C/vd7S9iGvD97a9sFbOz98+9bPPnzm268PnzNrOQb5AphUplBZgcoKfDFWgNS0IMj8BQuFJ8uo0TZlwZCdWbS3aJFCNmO2EKvOgKbrz98fusNLD/zpw3ffmT9nKQlS2OtRg+YEriJWNce8a2kmI+fNvXrcqB2fe/zsd4ZLpVRWoLIClRVYzhUgQ5GtYnqiZHcWo23T2IsZ548CSpwYbXcmlndkeco3hz/9+w/ewYu/JlATAYntyIA1FcVt6LRR/JT3XjXm/WOHDlmeESs2lRWorEBlBUoroCSrmB4uUxIz5IhuAZLSyHDeahJ+/t7Qe6d8KFK+/yKONVXyLZiIUrWIEC12QaSoQ+6bNPHst0dIpVRWYAVWoGL6v74CWayObNFso5RFe3HGq7TIYyc7JnKZiqcYWVbhfdnvx79DkPKa9l8SCeL+EEJGItsujF6qoFinGoHbc6/10rxq9Kjhs2dbZ+WorEBlBb6KKzA3Wzph6YJRi+e+vWj2m4tmgXAU9E96uUFUlb9v4k/CEWsGzyjkHcs9lu0s8dDfaP3XhDEiKmK7MIiKRUxIk6pCTBc1KKVoIKJUzdG7rDc3/tdH/OVVKqWyApUV+IqtwOzqpWSuD5bMn1G9ZFGszti/CH+SjHAUdHqxWaGrjlayGDMOsYCkLfZoPHyKJRUBhZI4pNH6xPRJ7k4GtDiEoaIkLBGa1FITkmpJNMKV2TbN4jw1ZWqjQ3rHpKnTz/v9NZsceMy3zzj3ndHjXPusYOSYcT89//I3Rr7/WQ1QiVtZga/cCjR4Qey/Plw6n8zVYG8S6cUGy9RcLrRtkYqEICoSVFQNJZBLLLNYdrLUwiFNltEL5rkz/mrFAqkUMRH1JpiaILXUhKSKFYQu6uj5C6TxMn7CpD2P/8klf/nHlGkz7nj4ybMuu2r+wkWNm3/antHjP/7Lv+/6aNKUTxuo4l9Zgf/hFRi/1HZky7kA7NSwX05jzKLt8DJ2RBJtXxazanZq9u6stG/yvIZlU3VJhpeapb38UjMlBVKNkRutpl4Py5syKn1qJhxRvNfMhCdeUwiihMWowRpjvOJvN4364MOLTvvxpBcfmvziQ/++8sK2rVvNmD3n+LN+2XqTnXvufshtDz6O2RMvvtJnz8Mu+vP1mx54bLvNdv39Dbf8/PI/tur3tUN/fObsufNS76kX/Kb7zgd02mqPex9/huGuueVuXOiaOmPmHsf9iK3fA0+/cNwZ59F14tnnn3np75dWVV3595uxp0Jo0lWplRWorEDTK8Bua0710qZt6vRij1cdsbGmpRG11BF5fRZFVFXUd2dCQRB0Dmmy4JOqEEAUFEmbPVWxigIBqYmApWoiVqJGBFZTpZEyfdacp18euuZqXb954F6q2qFdW3IZaeXkX/761gcfO2TPXdq3bfOds37xxIuvLly8hC3VRVf/ffWuq1ZnGcnohWFvrNdz7QeffmHI62+m3pvue2iXbbZU0VMv/M24jybOmTcPF7qqq7PxEydPnjajc8cOG/TuwVy27LfhRuv2JlGSE4/e7+vUgb+5+sXXKn+EZW0qtbICTa0A78LYbTVl0UgfXvg20lku8xss0bZW7M/IXbZL4iCdkSKwUwUsv9ipicOSIn4eQYqIuMsq3d7b7cB5+x01b78j5+1PPWLe/kfM3f/weQccPtfqYXMPOOwPm25pLpkYehAcUzWlkVGrqqrmLVhYCKEQCiUTctBzQ9/YYfN+11448Pz+J2VZ9p+nnku9g39y4p1/vOxrW2229hqr3fq7i48/7AD0Ee+NAqk/+8F3bvrN+Yfvs/uEyVPfHfsBSnmFb7J+nyP22QPyk28dddyh+7GJK4SwYNGimXPmkiJfHv4WXZVaWYHKCjSxAlOrPvm7oOX0VeFHKeKHoabdGVlJyCpML2bpDG2kEkNEEyYCT1VE5lVVff+NV9rdf0e7+++ktr//rvb/sfrUNHvTr1iIiiqHcCqrNKWR0r5d23XXWZt906jxH5VMFi9ZsnTp0kKhoKrNmzUr6ZBCgRzNeSXUqqrqhYsXp0BdO3c64YiDNll/3dSsYGUFKivQ4ArMzZbydr/BruUR8SXCMi15uUTSypHdECymv2yaqyp7piigtZo4InbFbtth1X6PRg8GIBViBtiXRTWx7D0aZlQXOTdYebQ8fO/dyV7HnjaY92InnH3+Kb/6Nc+eW/XrywaN91l/vPE2Vd112y0bdK8jPjHk1T/889b/PPEsETbotU7H9vbfdPvb7fee+7u/fDhxcjJO+fGuR54c/u6onbfevKq6ukPbtt857IDtNu+3yzZbJJsKVlagsgINrsDcFXxlVj/IMiOQu5SkoqpkLNUS2kaGZELKcSTL1Q9eRzFn/GuqC1q06tO23aR9DuEZc96Bh8898LB5Bx32ym57FjvdVHIUJyBVRaXx8r0jDrz0zJN5z/XLq/56838e6di+feuWLS//+U83WrfX+X+87skhQ39xyvcP2G2n7Mon+AAAEABJREFUxgPU9Awf+f7pl/yOP4z+9pwBPbuvsc/Xtmfr958nnxvz4cfbbrpRsjtsr92I/M97Hrzw6r+dcMTBjH7d7ffu+q0f/uKqa0eN/zDZVPBzWoHKMF+2FViQVTcx5emLFx7/3ANPTWrq96jpCCl4NSmN3RIZC2JoB2/xxbKLkE34Q6OdknWjSIiySgyr6XWY+zT1tFp0xMXMyKDFak13bxDYLvX/7jHTX37s4+cemPPakxcMOAllw949h951w+QhD8989fGf//B4FDLa4reeP+PEb7dr0/rRv/9h1GN38jcBmklMkS3OK49Ne/nRQ7++K0qPNVd//d5/TXzhQeyfvflaEF/EV++8gbFuv+qSVTp1+POvzp497EkGGvvE3ZWHTRatUisr0MQKLLH/IEWj/aQh3mj9fOjTj09s4M11cms6QrIhc1G1WMhcUPv/CiCVZP7SLOWWZN0Yar2/Y4oQh2xIFYqKgvVqLqooVQzZGKqTEkrThZdiXTp3JG2Vm3Vq365Vy+X6zxmVvEhY5S4EXKVjB1WmUTKROmNhz0A13RVWWYHKCjSyAvytsX7P4qz6w/lzP5g3Z2FV1Tmbbt+nfaeBw56998P8D3R17BuMUNcms/+hZuaZC0zv0TzBmaGqCj9qvKmDlMfeStjJwYqVhNiUD31Yglb9n6HVuKdXbyI1Bmb0WRx77rDNB0/de9IxR3wWwSsxKyvwP7MCy7jQQCKpZ/LK1ImHP3n3gY/fQT36qXuHz5w6v2rphcOHPDfl43q20mCEOmbkKxTyl6gmJHexO0uZxDISKYWKURNVhR8VIVqO6hxEFJExCxas8eD9He67p/1993a41+q2Tz6FXqwqyV5VE3FElM+4sMPi2ZOt2Wc8TiV8ZQX+p1eghYb6179+x1VO23jrs/ptSz2z37YbdlylWQjf6LnhTt3Wqm/cYIQ6ZiQsKm+5yF9g4imviWUWDoFI0wU3/KmJgKWaO5IR/R2Z/Y8QEkHJ+3wX5iLP18mxFKpkUiGVFaiswJd3BdqEmn8fWrqKbq3aHNOr73F9Nj6614ZszUbPnfXD9TcbsPHWWrIoIw1GKOs3Ss5kQ0T+KiGhfHdGfrFHzzyxmG3jh1rGI/sSREVSNONqurRr1uyaLbace/AhVg85eK7XOYcctFvXLh7SLEUMleLu3rSAUimVFaiswJd/BdoXmjdxEbw+GzFz6g/W3+yHG2yujdg1HSE5+X4o2v9oMyaUmEX7zznyNkstxUjCZN0ostXyGnn/lYwgUZ6eOnWDRx5uf+897e+9t/091Pva33Nfh3v+0+HuVO/vcPf9Jw8bHt23JnFaUy0MxE4NH/MXLHp26Ij3xn1k7g2brKgqCxYtvueJF+fOX/j+Bx9XV7NdtQhLq6qHvT2KCrF22cHrxnfGjH9p+MiFi/J/WFvWWYvOX7gIYyRmO2vOPAaCU6fNnP3KiHcRP5gwmcuZNXc+YhN18ZKl1GQAoSZewkdfGHbVjfc8++oIJlwSVy5ZuHjJzfc/CS4z7IzZc2+6/8llmq2oAZdGTV5MYzknk+wr+F9ZgfaheSttYIOWJrN+h8737HFYE7kMXyIk4yaQ7ZBXSchzH/s1251ZcnG/JlOKW4ioaHk1wRXRWju1ZFPsrduFjgFIhVhVNmjSWPn7PY92XaXThCnTxxf/pWtjlsuvt27Z8us7bBFCePqV4YuXLsWRFHb1v+8jGS1avOSqG+8uTx90XXv7A9Nmzum2Sqe3Rjf6B2aCkGIu/MtNM2bPgd//9EuPDXnt+jsfHv3hRJrjJkxu167NbQ8//cbIMX16rDnivTEp69FVv479aBJxXh4xkq5yTrNUSYvH7L/btptuuPG665TElUtijCwIuMywWZYt+Az+MydcGjWNzjSWczLJvoL/rRXo2qxVE0O3KjTTxrub9q3xs2dK25GxP7AMRvKK9rgnovwUUZZRcE61FKJE6ut0UZPuyEZQIIilSjPVJgbmVwX7XbfZbJ01V5swdfoV19/Or/rzr73FXunXf7vt3w8+de1tD86cM49f+9sffnbkmA8vve7W86++8b1xH9/75JAr/n7H/U+/fNWN9/z5lvvZHJVGWbh4sd/qa7ZaYz+a2KVTh5236rfjFhutvXrX0eMnlIzpat+mNV29uq++9cbrk93+ftcjjHLdHQ9Nnjbr4mv/zSjn/uGGGbPn7rXjlr27r4EjO74Zs+fsveNWu2+72SgPhWXbVi2nz5p7wK7brdl11Z232kRV2SFecu0tv73hzmmz5vzmH3f87p93nf3bv3EhDEQo4lDLOc1U2VQOf2/sfU8NYR3Iemkpbn3o6TsffZ7J/PJP//rdP+9mhuBb739w9+MvfDR5GmvCZvHvdz9CemXE8/98I8vFnhH7P91831ujPigtbBoiIUn86n//J81qxPvjuNhzr/rHyLEflkZ88/1xv/jjP/9258NzFyxk10klX/O5zFuwEPHVt96r78JYzI358BkxCmt1w72PUc/49TVPvPT6X29/kC0tc2Y7zKVR+dAx/uNN906ZMYtvS/mi1f9kCVip/90V6Fhovkphxf75VJowXvgmvgwMnrVAFUmoYrszvh9kOkeShjRdWoaCilIlx5AIilIInOua9IRqou3RUhOkuqgJWwbiSGPlhMP2GfbOqF/+6Z9Tps964OmX116j22Yb9uFXYsnSpVv3W/+b++3Wp8capAz2Tb3XXv2BZ1/mfr5Br7XffH9sVVXVfl/beqctN15aVbX3jltu3W+DxoZAr86yNq3zu8rSquryh8HUxZPjhX+5id9DfpPnzF+wTb/1wdnz5nXt3PH07x6xZd91J0+fSZxU2fEtWVrVrFlBg/J8SrblUtu0atm8WWFpVRW/hOdcef0b745+d+yH2266QeuWLT+cNKVD2zbfP2r//XfZ5qPJ9j9xTXEaw/XWWWvT9XsdtNv2BMQmLcXR++4aghLh6H13Wavbqmd898gWzZu1bdOKnDV+whSujgQRVHuuudr3Dttnx803Gj9pysLFS9q1afWjbx704utvlxaWgKXapXOHH33zQGIyq37r9jz2gN1tbd8bl0Y8Yq+dn3l1xI+POeiEw/ch4/fqvsa4jydPmj6TbS8Jt33b1ltttH4dl0P22PHh51/deN11LM77Yxlo1tx5S5Ys/faBe262Qe++ve0/Z4JYXp986Y0j99n5J8cezO54/MTJpUXjNsNi7r2sT7Y8VIV/PiuwZvM2HZp8iVZ/GtjjVV9vSInp3ZnYf6Yscy5gUBF+31RUIV6lydK7TTthX5dsIHkt++djKMVeUqQZo5RXYSh71KUXORn0btsmOTWIrVu15Ffi6H12eW7Ym82bNevbe+3tN93wmP13x7hZwfLgBj27k7xmzp7be+01SA391l1nj+0223fnbQK/3EE7tW97yrcOJdnddP8TuDRWySZz5y+wWUXheXONrquULOmaMXtu+7ZtBhx/xKod28cY2b6RTY4/ZK/G/nkt02CqPLEuWry0Y7s2bI7W7LZqq5YtllZVZzGe8q1DendfnThdV+lInKP22bl39zVLw30ykpai5EsmLRQCyY5s1bZ1q7dGjSP/vvHumO6rd+Wx994nXthkvV5cF/aFEPj4mW35wqLXrw8++8qYjyYxYYLTy4hV1dWqii9NKpm9ujr7ePK0zTboNfSt9/usvWZ9F8xYnNJnRLNzh/bgVTfe3WONbh3atYXXqUurqlq1qLnblxZt8w3XXZ5Ptk60SvPzWYEezduy21rOsbDEfjmNRfjSRTUUFfv22lc4SKBFQrG0QoaRZZfdu3QjioiCeYV6U4hWJOokIcZUuGOAlFUmoOi7rbaqNF4uu+5WHl5uffiZfuv13GHzvjxCPvXK8LfL3mGt2a0Lj1GkG5ILv2y3P/LcEy+9MX7ilBQScvP9T8yeO59fGx5kHh/yWtITTpo2k6eYq268hykT4aJrbmIbSAro3LE9Dzg85mBG/M4d2l10zc2YLV5axTRGfziR92JsTKqzDINU+a276f4nSRk3P/AUj1rr9ljz3w889eTLr2+8bk9e/PGk3K5NazYgF1/zby5n6szZvdZao6qqmmfhx158bfHiJSlIwhdee5tEQH137EflHOWGex5NNsuP3VfvwnAb9enBw11nzx2Tp88iFA+SpSDlC0tXg6Pw1M98Hn1xGFeaHFu2aL5h77X/dsdDXDIvtmi2aMEftiKJ7K1RH3RbtVODLnU+I6ItWrK099pr8KWsrqqev3DxP+5+lM1dGgLcpt8GPEffcM9jPGx2X61radF4IVD6ZDGr1C/aCrDbWrt5W97uNzExerHBsgmb+l0xCr94oOWuKBCJSnIRUX6KKMso3167h3uaP8SqRfVcSOA8bt4rxWaJmL2LEESzc/dje9r7psbG/vn3v3HmCUf94ifHrd+zO3XgD485cu+dt9t0Q14/UfEio5563GGH7rkjnEfLM753JE9bG/ZaG6Vv7x7c879zyF7fPmhPnpi27rf+qp06YMZz3/8duV+XTh0GnXQswdku8XSG/cAfHstAR++7KzuXrTZeL20KiH/Y13ca/KNv/ezEo39yzEGdO7RjCB5ysWcTRxwCwjddv/exB+x+1cCfEI3nLGbCY9ppxx/Rrm1rNmKkWsx4MXf+T49nxLO//81OHdrydHn8oXt9Y79d+c0nDrPiiqiYXdT/BOoGvbqXc1ISl0McKvZdOnfEuFQRmQaXTIXQLNmc84NjML78zO9v0bdPn7XXYAKMy8KyzcQGy/KFLY3CfOgF0xAH7rbdSd84gCs6ap9dkoLj7ttudvr3juSS+QhosgInHL5vr+6r//bnJ3VfrUuDLqwMC3j0vrvwGeEyceqMHqt35cGZdP/Ey6+fetyhrBsz5CrSKDyZ/uz/vsHQ5/zgGBa2tGgb9l679MkSp1JX8gqsjHC8C1u3Zft1WthOjczF/oWoIJwdGTq92CCuUFX2YiQuAqmwj+I3FAykINKKYcZ52QE37djxlD7rqgQVIngFREU4JSx1hZJYIupm6u4iKhDVn6y/zqad7HFDPuNCDuW3d7MN+izPOGwZ+K1bpeNKmBjPYrtusykLtDzjNm3DL/wuW2/atM2n7/18RinNk1vC9Nlz2PkOe3vU7tttVtIr5CuzAu1Dc/ZfZK6NWnXs16oTCEdB/2TXyJsyXtqUkO0CuzVSj5BVgv+qqSxXuWjjjQ9afXVLfmTBVPHzPRdndLKGRPsjptBLdNSyXhPRrZrNgWt2u3Dz9TD5HCpXydYMXJ6xeG7q2K6B9zjL41vHhhdYrVvWvPqp07tCTebPE9kKuXwC489nlNLEWGr+NMH+7sQj9l1Za14KXiFfzRVQtkKBPZqKGvJbzX5NKGQW0GoNs1bjx43bbnNyH/Y4gVgiRKqptJKY9ITK0HQIPSpixgl/sn6Pf+60iVRKZQUqK1BZgRVaActVkd0Z77lAtlAgz4Ni6UVEyEhdZeAAABAASURBVDOyAuWifhs9t/vOJ/XutWG79s01EI5tV0JIqtH3aDFtzYqI8Qbt25603tpP773N57EvW4FrqphWVqCyAl+SFSB1qaptqEBRFTiaZTcyHU1DVeVyHDg3XTft2OHSTTd6+eu7TDtk3zmH7z/78P3mHL7f7MP3nX0EdZ/ZR1L3nn2k16P2mnXUXrOPAr8+9cg9huy7/SVbrP/5vC9r+hIqvZUVqKzAl2kFVElQSrHtEbuzaIX3/pyy6OlMRf2CwLZt27TNZE5VrX804J0VqKxAZQUqK/BfXgFSU5tMSFMiqqISOAxFVOu8O1PVNbp26Txv6dhF85f5/z4glVJZgU+4AhW3ygp8khUgKZGaVpm7ZI1uXXiotJoRJ/I3TXgD7842WLdXizGTX5w1ZfrSmv8xIx6VWlmBygpUVuC/uwIkJVJTszGT1u21jvDWTLSIIs6DUCIvzchuMNlqk76rjZ12y+TxHy1eWEW6M61yVFagsgKVFfgvrwDpiKREauo2dvpmm2woGVmLfRnbMzBxf3fGM6fwnEm2i7JKp/b7bLTRGh/MuGHimDGL5v2Xr6AyfGUFKivwuazAF38Q0hFJafVx07++Ud/O9v9VFFSDBpAz2zIV3p2Rx+wfVYjtzuC8XTt47917vzr29o9GPzB94geLlvGfGMSvUisrUFmBygp8pitAIiIdkZR6vjr2gL12UxX+khnJXGzLEqrwBi2IiqraVEA13m2Vzt8/5MDej7115qhhd0376L2Fc9jmmUHlqKxAZQUqK/D5rgDJhxREIiIdrfPYW8cftH/XVVcRsY2ZUoLCaVjmYq9mqS1/d+bJzp9Id9p68x9uu8NmD735q5FDL//gnSFzpo9bNI8/K0ilVFagsgKVFfhcVoCEQ9oh+ZCCSET9HnrzhK23237rzbKsmhxGtmKDlmUkrBgz+4dnbNgCEyPHKf1BKelRlGy3/x47n7nnXlvd89rjLw3d+7XHfzl2+M1TPnh05sSX5k5/ac60L2atzKqyApUV+CqswNzppBoSDmmH5PPYkFc3v3vYqbvtvvceX/McVSBr2XZMQ8GyVoALh3r2klTsKZQEJ+zX/CTs0X494CfHL2y9xV3Dnn34uQteef7w157cYejD2w97pFIrK1BZgcoKfEYrQJI57LUnL3j5uacffGbTO4ces6DVBaeetMM2W3he8r0YKcueKWPmSOpis5ZlGYnN0hq9ZDYh1XFSDSAPoiF06dzpxG8cdvWAnwzsvdmxoxce8diYw+968/A7R3wR6mE+jWUiBp+mHnrnCNwbx+HemyNmh945/NA7RxxajneUKY3wQ1w/xL2Wkx9y54hD7hx+CHiHYw2n6dV0J3Q1wg92fdl4x3CzcYSs3HqQh10pSJAVqW+UGSf+v4DD/arTlSa+DOTjxuVTIu7LWflKH8y3+tFRR41eeGavflee+uPvfOPQVVfhfZmEULDcBNoJrmQpEcOQiu3FEFQjec6yXFTlzwbC06eIKLrIqh077vW1bX/ynaPPP/0nvxl8xm/OzesVg0+vXy8fdNqnqb8edFrT9bJBp6WKGaQ+Il7qNglp1qnolw4a0DAORB9waRle5rwuIg4ccMlALE8rR3MceFoJLxlovZcMcsy5eeFiddCAiwfm9RLj/REvHti/FjeD/hefM+Cigf0vHuhY5BejnIO76Ref0/+ic6z3ImwgqdbibnBO/wvPGXDhQEfnFxnv72juF56TuOEFGAzsn+PA/nhd4AipXzGjtwE8p3+D+vnnnHr+Of0TXuB8ORCXU39lxo2iBzSzpglBvPZ3tGi/Oifx/wXMr7fpJSrvZZVoNoZ0Lcdnl3/WGPvn3v/CQXwPB9TBiwaddtHg0wwHnTZowI//77gjd//a9p06duD1GM+MKiqSKTnKchbpyt6XsSlDztilcbA7I2FRRQRUFbKeiHFc2KKpIth/4pHc5xyQAtlRajKiuk0JC4UCxtSSskI8qOLfhG/qTdhY5AJRLA5hmH8ZGtWCYUjohqEGAzzgwyFQDdIAKgaKBV14NogMgW6GWGBeQgaA51gIxgvBMNiMQqiDNkoB/wJ6zq1VGj0EVOUVAudCyLn1FmpzDdgYFAKTKjia4hJKzkPBeCiIIhgvBFCDYQiGidfD4J9aQOdYHgyFApY5BudNISELBeZWKMNQMCUUGkRzUHcwDNiYImqIkq6rHENgZej9KmG6aruigq9AOQZXaqGtJCuAF3JTWDDLkJATPqEQlokFtwExDgF787Dvaonz6dA25CMS5i2cSUoqIRQ0oBRE0AJywIsTqkgI6HwDC4G0RaJDFy9wzKEopDxVKILnRG+oIrGDQ49Kj1eCwHMUJKUEH6+EIqaLNIWWeX0/qGKlPqKSj9ET8iwNL0fmj8JGs4Q+K8kRf6NivWIoUhdtDkUbkbq9Iii+GpKjWCEepzJU46r1ETNRNV1sEyxQRgTpKMPoPNJtV2qWUUWZl6qjcQ64+UXA56NokUMRlHPiIKsaTbOVIVoUD1MPJaJHMb1mXCTXa0bPFVYgYuqWcKiKiPIDmg1MRWowCsUti7NFMYtYsmEQeC1kzkggIYto87TPHefYILcOVbtU5dqFARgXy0bRZqZ4iepXBu1603UVPyNTSpxLhSdUzdeKpYc2ubZ0EgfE1ZAhIgtsJ7G1K+MqarnHkPdacENE7MX+T5j4HjCuIaISAdVQRDUT5Sca5VPLOMeYqWqaIVxgQr9E/rIJxIyxxDOOqCr5h6oqzgBrcYRgNgFbDrIqDjmiY1YbiaNi/UZw0MQLoVAIoRCawmah0CyEZtgUCmZZKNTBZgUMgBxpFayg4GPYrFAbGRGDMmxGfFea4YFeD9GbYVNPt5kU7ZuV5hnU9RKGQo0SgnFAYb54cKOJFxRjUwpoGkKtWvBmwWyIraEABmwKSQ+UolII9AZKQYFiLeeumVkIhlpoHAsFLTTDDMSsiIjFGpzk2MwMAljQHAt4hby3hiclmA3Bzd55cRqFQqAGty84r4XNrLdQjoXQzM2WE0Oz0Ax3sMC4Aa+QuMcJDSETKOmJf5HRrqgQ0gwTbwLza08rUIa447UCyJIyaDnSrFP96xoc1bKIakFoSkE1hCIGLRQ0kDQCSAcMlGBFC4hmWVDkUAhB6VPnaiVQUEC67D+7GLNoyZJUZ5lOKN5WJ6ic4cU+MqkJ9gQaBV1UbQQR42UYxcoKIINyr0goNiF7NhZxjIbRJwrSm1BcdxQQJ9elBpkDao7RCp2mcDA3r1BGsTRv/VFsOUAMBT31gj5WLM4hWmFWIqaXMBImmo1Ec3F0CQsaJaQX04gvqqMwFpYgjo6xbCZwDM2m2Ct4RYrPFhHqKA2hdxYBC2iTKPTaZCwWvKzaUvpBFxY5JgV0S84sbEJIbhM9ZkLMIGBDlfswLrwwqUHMImtiEaxXnDeB0S2LyDzzaD4nIpgC9wicTanNmbbrdrZD8LCzHSXO/GwIDoZbBpqtHXUtTbOjrm7j1dJ9VBQ7MzXr55Tm499DFARbGU71r6hGiW5TByMT4PtWRDEbguPVKEa3KWFs2Bf3WhUz5oeUo7LzohLLtmPs+vwaRSL5xK7LmvbGDD/MYmYaB+/6+UwZMosZ32YO/rCZkdEsxakYWlqyM4lTVcRPCgao5yysgwqJUF0OfhJHxFAoJI6A1wpzCxFwpDKGhoLaeAXjBeOhUMPVuIIaDIUhCzANhVDimjhB6A3ogaIKGOcMByUwaKHEQ/m4NVy1UNCgoVAQRzgKXAu06Q3CKHDVUKjh9AqmDAGajjG9geJ6AXMcanGMQ0GwMKTfxjWbAA/qKGDBuI1V5DZWOVdFAUBRxbLEUWo4w9EIITqKI1xCgIP1uKJoMJSgbtMED8kGe9HEl4mMy6XlqFqfq4m5nrihAIiOcNygIFzL9BXjwWMAatPGV5zHpBPddT8ryMWqne2ozXGgH+c8DhZYSmhAx662r6KYp+a+QkNDPgfnuX3ijgbY24khOLGehhzaiM4oogomG/uUaVmjzDeUOIZwQ3eq4WLfJaZqChMThfM9qYch8D0UVfXvEq2QcktQ42YerDNQOJtZ0gsWX8xIJLDbA0NBRVTpkJQohVLDaNBPl9pJHFX8KVeDZU4aKgjOYXDDfOSQz0Mpy8dFmYuK2XOo+dvZDuceX8W4mJUKPKhxQ8Z1ytwUrpgwrjhX5QSAwikEMASYpmFKXNQEwdl6xQBuhnBajICvcRFRbMW4ChYBTONKCHQmrqGom0niGKoatU5kOIih0hNUVeDiHAyKJOroXJ3TDqJmBJrOCU04AqpyhqdAIYhxBZHVOVjDaZiTBEd1hKvrYB0uqm5Te/6IxRqKJJ3FT45CSG8BiddF2swzoapYcazN1QzUO+mAFDHyDeVur6KOnFDgoiKqwo+jOKrSXoaOJRFiWbSkaJlSi0seM3JWDotvvKj7Odcb4+X2xhsby/XoKI5wIXZx3BXl0X1B8WjqGBUVwZCTCJC4UOgERR0cnasaVzgHvGHEyTKOAkoeybNNhCueduakWAlnP9iGiSo7sRLGaG/TYqwmCDHSd0/sKyv5F9S4wq0rBFEKNNhvDSoVZoJyvzc1qLknFsip1p2UhBrqKg3oBbMhWWNMWCXbmldBixgKhVAIIRS0wKmg1iiUYRDXHYMEughomLhhIUgBMWjCELR2JaTiWAgU5YA7GiReKOAbQkEKGBU0GFrkUKjBQkGpIbhSC7HXELRQrHCvwZGuRAwLQAiGRA7B4odQRLcsSCEoSsJQgKfq4xZqsFBQajBjrY3BvEwP6IWCIxCCUfRAwQUUjpArjEsrRxex0cCIGNRCtd4ClkpvwXohXst5wRXDEMDgsy04BrqClQKQeDkGn2hxbmbDWCEEfMGkl6OiBj6RAEEvIYRariSOSCVmHdSCRyjHYApmGgrJtxzVesv1nC+3ns+5gXF9DjV6IbdMo3v8WkpJD6FGTxzEHsTGsAAErp01Vo4QxDHxhKGekuuFgoagIDXA84q9ljVreEEVe38HpwHqRygEK/hy8hVT5hCCZSOiKHMrkOqC8UgAdJB3ZzEVIbn5k7I9rNLIKy6e+MidAs9VybkFMY5JXrk/EqgcaVLLlcQRqcYjz8K458h0hCdpHotLaI/HWQRREtpjPrPl8dkx8lcS42q9UTGz+dKVc3SLFxklUsR6jTsxs4hjpNvmEN29hNE80W0CgpnwJG9jic/Bx8oMo+kYWLX4FhYbKfIoBInCKIxDNZ1TXjGjzzCaWSQ4lkJkDGphsrErohdjUBirXqXLawRtLObMJWNmGJkHQ6BzRYwAGjcWKXAzwNIUTRw9co0oRTSzPCA2IjZPw8gleJXaSNMGxcVq9Cb2wuipEhCCbsjoNmfMIro55gouKBEl8rnbtCyCOJeYc4kwZlVCKSoQ9BJCROyK6FdzyjnNXE+9wrXYBJKe0OzRGbRxZALEkXo2K6qnERPm4zJbxi1hTDNPSHiwvmGfAAAQAElEQVSqc+G3gJnnyEzoKK1A4mBJhxM/R4LzWUTWlu8/QewTEVbe5lHDk5IwYi94Y8+KFdFc7EN3m5zAqTFGsgkjWkh4qkQQSrQjYm9n28fRm9HOlFdlMWrGvoxRUAT3LHi+UxIVh6U+JQl6hlSxPjHMdVPoEg5VFaoIqOpcjQuFZjkuh67lJQ1mSmKGNpPAsHDhpCo1GJwL81QVUTNxVEcUVRVBVrFOTaU079Skh/igNc1WS9xaHGJz0BwVwcbFWoyn0CK+mI6iio1ojlosghBU1IphMIKWeEJGQSE+aKbY0JFQ1XScJI2lyNiA6CA9JRR1G1wUE2uImr2o2slQBBbotaunQUspKKLwpBgiitlgCi1Hxq2tiIqNa/dS69BQG2lS0cxI6RVRX1tHUU4Wnl5R4+qIUTItoRupiAUzDGo2WAtHzpOSUJvSRUtx8HaelOSVI12q+fU6b1RPvY7BsT6ssC5lMcp4A/OpNX/GETF7EUNNc068HBvX7XMXYgY1+yBCnGBx7DNybhA0oaqanqN9soIQlAKI0oAqnBNI9AIndNrs6QKlUNBQEMKogghqJTgUFA91XTXCJeeiyruzKEpiE45ISs05eRcZQRVFVMHoNhldnn4tvasUkT6zEVV1WoNIdRRV6y3qdEYXLD6zoF0PmZikQxgQwzKMzolgMhcooo6ckkI0rRnQukVE+alBKSuEo1WDEUtbCMIIlHmIx4+OiYubM0oUKXK7lajb+FjWI0KnUFyBuxtn4YBHFRtFOWHOqT5yC6IXsxz95KNERzGMQjABbLKcI4f/QVqLKMWikGiWmKq6Gy0U56qmAD4tB2+omp84Mi4xRJNiVJxGR/ERxS0TJ3YUK/aBWkjBsKjA7TateMWcS+JEQFEcIpRloNvnLBRXRJVu4SQCGFetQSSRBAk9joUQihsmPUcxVxHvUPVGPVStp0uZ4jwNQyChJPsV1ZNXGaqWjVLGVU13EB9EEi9DFlEoRSXZGyZTVeOqNShQnOoikn8IWBqt+dQkfV45sr4pQBQ+OoHb7wVna2Njon3HJFoEWuSWGDMr4oVxJeJmXXayfRi/jtBIJmIKmfCVjxQs7N+dKW0SIKgKl5xbOyimHpUWOumvEaRHlVSJxwoiObjgLqEQghEOhVu0Qs4DHQUgaKG8qjcdGZeuoDxS4ygFLccQVF1BpArNwOZBtYiIIQRQC+ihPoaCGXuQAEqhHFULVs2moDkGV3IMGBRCoCtYUYCxHIFQ4mZjQn6YXUHrYKFggodS0HyxQSxDs0EhDFgIjB5CA4g7enlVbzj6uCHHGtkU+tEV93z0UMMZB70Mg41eSPbBeY6YFQohFDRQwCJn8RHqobqiocCPYhCCwgytodoAoqHXRXcxsQli1xXct+C4AjyYb8G9CmElcY+2/HNIlgzeULVJlek2w6I9i8KaFJXgvAZhKPmaFzgrR8C3YD3+6YT6WPfzLWBjjuXfBFGUgvl6KA8bQFEV/ugpIqSlYAXGSS0oVCkBzolJqxtZWuOwhCdWjGckOj/zlJruoQQkEZqGjXKUqmkxtcr15eW4+2AWwTjDRcvWzKdYkakMX14tlzMsHY64Q0HTOSGK5358CEl+N4WIwmg2N6N0oOJHw5AuzBPWmwM92NNZgzaWBaALHSQI6JyzGeYH9xF7wPeWEeaDQbHigCd6xgkxcgeKFNe9zSh2tiPNlQAQ0KvFpS9xdPMtuqD7/RAbpHQPpN+s6LIh3A1Sqn7ttlDWgxGuMLxzNENfQyNozomZm9oJYx8BjgFoczCFWFaxTpUuosCNYIC1o61tLU6/ddgJB7psggwDM7Q5o9ASvqsMIan43IyaEY7e44GsYaIdJe6N1B35Hpruzg1xbEs29XjysulE+jwOhEoHiAbmnDnRTjaOuW7c549BI9wtc5t6HCEtF1jixWHtbIcFtk4OW0OfSM452QX4emLmPIXi8zJDTugWJf9eeW8tnhQzxIxTsZruCoTvRsz47kPpNrTfFxQ0BLIRvxvVWcwiNauullgNSbNziVaW0RttdyaiyoooRQROPBHOkZQXVEPIM2QQE2i5pFYws24TSjoCXia5b9M6ZgVzKB0FsrJyIBQKQHAsFAqhQF4vhEBnITAACvHBpBg32XI8PFDsxFEwyhEw9FPpXPCmYYHoBfRCsGLIMMGEQihAC8E4fWrxEy84V3rpTBgIY41CCBqs5AvmsrWTXqBX0TQE0DTGIHAh0GEKHQXG1cQLIQSGoc+E/LaGUjAFN6XbeIALDTjjljAESdwjSiiYTShgqKASuQCgg6YENR7QgzhqwgJ6yHkociIXjEsITFbTWCEwE6qNm3gIxhW9gI0WChbdURxNMRt0egylNLdA/IIofiGfIacQNKAHKB2goHACg+vBHBBMd04PlmCwYr125iiU8WI3cigEtw+p1PBCUS8ECjoVYyqEmkjqMl40C4WC9QZ3LwQtBIopjeqB3lQDxgVrBgo8FLk3seGcag0vBE2SY5lefr2JG2qot24snw3keuJaxgNeBQUCiQHduPKpBfXPlA6U4DxooaCYgZwcUYQh4aYHrGkppeDTTRiCCg7W6ScyEm1Lq0qHUkSAxIOYJymS/JVyLQkObhaWdO2w+wJ7BRxE0OnNqmOszjLLn9F8M0eAEKAkrzJsWseFSlyQcCCRQeMMArMBIhQdJFoJGcS59TIuhiXED0sQmxrkamg4sibQ4rgx4oxeB+nOMu4M1iv8aYZJ1aDdNExInQlreovxLQRHLBbTuRIfK3HrQcGoiDbzjFWXFNQMxAAnTjYfTvXmQ2/GfKI75ejziYZpLAvDKDGaUYwgK1YH0+hipkYjpbxRl/v6mA1hGN7dfETGYT5IJaQvc6sMNcYMCyyLiGYK8VFyzI2iKZGC7JxImBva/I3W8DRps/QljEWMZkKgGBtSopW812hM18VMadXCLGYxr/wq1PAsF7Nib06ylaNXe5zyEbNohTXjlHNYtOFijLXnX3bt9dchrVXSnbOE5pDWlkasvcJFPUYWzAzdhAHhTESEOZkNvdG5IzFMx8l0XGMyR+eqEuJdlWVcalVmHxJBrEYvGSZ4SFQVPxPNB851HsZitIxmeQoL4ax2YK9kw8ChaopYr7mKkgIVHbciFmiLkEfDJ8CgBUKqOGrBItRgIWjBukCvYhjKURWbAKqgh/oYxER0rwWLL7XQdMKKDSSg0lsgoHN8Pb7pxEcvYVCFB3cPcKuSK6JBpVDCIAUVUxJaZCmUo0rJAFIglPVqKEezUYIEepVoGix+bfQuc1cpWG9dLKCrOkoJg40iBcfEczRjBpKCMooRdBtd4ElxDBrUFcyo3luw0bUBxEC1oOZCb0GlPgbBQHLEIEhQV+AqQQSvOhhcr4uIDVTFt2hZixdEgiZF6/Ok5EjYNElVm8x/EX3O+azqcb9STb2JO4pjulIJXEuD1aKJW9ZFu2TxT0HLUU1PSlocx0JjqMrQBUMtCLswtnWFEAoFzZEeId+JCkWDBkrBTnBRTSj0Oqebs3Ii81muwwlGoivDaPmPmCQy8pl12Ai0Ij2ZHRhkpFY2DBmaxCbQ4sQ68WPMMlzsFCnk95jFIkKyzJqQVC3TZ7GEdKInxFsygN4ajJn/mJAZkPDtxGxjTG6OiYM2f0ZHBGnjlFHcuGzcaBzRKiY080p8b/tMsoR0ZRxl1UbBEiWNAqFmMRpyyjibDaPAMuSaOSCUqsXHxg0QzQhuLEbiO5eMEzPJMbPejN6yGm0s7OmyT4feGGOW2TmLWYwcjBIjZ4mxHLMsowmAkVNGL2PVx8x7M0oihoyYZZjWmWFkDqZHbOi2U4zwaMUEifXR+sTADuuOMcWI5Y0iN0tbeTNqmscYsyyL/IA5zyiRkpkaDaNjapbzpKxczONnGWEZFswoMSbuGCOnpq+rvDdSsgyIDglzISl8ItbN98RO4hYJo/Mck0/0k0PtT5YsUS2WK3KMzsEMkmp1dVZdHdma8elkNoIwATgjZ1hl1fRmMav9bbR2UWEmgTwmZDYRUlsIwbi9EhQrqqQ/V8iLBVV1bhjgHEElqAZx1KhiaQ8kgmE0JXGmptH2g4YZbwvhqRoP4uj2xsW6LKyIYYwlJA5cXdFMQsQygqYXuWYWzQaKPp8oNgeQ+KCPZb0iFgpEB9EZN8vHUpGQmQHBzQw9Sog2FvERiQAKNjGfg80KLmYj0edAQBQ4xOObXuREUzgGYGZrlZQA97GsN2PQYs0sMjaMm+aQczPmGs2M+Nbrc2A+FgEuvg5iNqZEGyv1sm41CuNiE90MbmHtKmrmQ1e5bgY2qI0YfW6MZWINt/kQU1hJ+yN5EL44xgUzeMJMVPJe1lOxj5ipZqoQBhU1PXEQexA9lutStHcvt/GJaTmqewUfItedl+l2RTSDeJyIO99w4wEuOU8GQTDGwK4o8Rody1jULVQND59Yj0wjxWFQr4xiwdHzpk2g7Bppqk1bDJvWvVezZOkBM0N1fdnrzzRye1URs0fBFySmKiK5gh6QBkgtBA2BHRlGAmdxEfPeIIFSAM0GPaioeFU1oqEQFHcyBrL1KpTsF4USjWAWMyN8D/nKCReXkfWySJoUSowxs0qf2+dmzvFViviIhmqKEVFKIoaKLl4gXqNhFEfn1suvUhQBmUQRlRU2A8ZVzCMHvYJsyy7OaUleTBfRWtXdkUyMIsmF+OhUVR/Rx8UIXYwL47plFBXjqAI3c1HxLjE955yEEjmgXgmOhVh8gVPpLKEFys0iPAXEHc04pooejSOp+FjeFCEkB9MUPBWR+Iai3odiDFEiRrmYulRxTooJZiBmo+qi+SnjStm4Qj8HCiagczFiGTImAvrSYRLFHAw5c/Jvl+nwpFgDIz5TsWEjbTpUFJHDZqjqxCSVGu4UK6vonJRDkGts8vkjlmoUzLQMJRW+xsLMhcK1GHIIShQBhVLSzT3aVJ3Q61yc1OgswkrRPSzBqdG4CCiU0nxqeISyfumU8zrXa02xCEZi2VpJzkVE+ZFiMdnaKqA1nIhSTBH7jADnjO1ymgCILYgczSRGmH3WMUaySsz4QQEjSTBSUGx75mIW2a+xt8Mp41oxiRleogTlXmxxMjNEhQfT6fKKC9UTHmmvAKdiMH3OnIefefF3f7vpZxf/7uTBl5w86OJPX08ZdDGVaIaDPGaKXMRTBl1yyuBLcoTUrj8dfEnT9dTBl1CxSQj56eBLfzoYL8NTB1966mAMHM8FLz21NvYffGn/cy+phedeakqOlww412r/wTla1+BLyvBSMxhcxMGXDIDneKlzw9MGX5rqgHOtmeuJO552rhuUEELFCyzW08+9tKxe5rwunnHuZdTTiwipXzFouJ5nvmeU4ZnnXXbmuZfVxl9b87y6eNZ5vz4L4yKeaQbY1KmEKimJN4iIVCzBvBJ82fUXl9XY/MLmc1ZSyvjPfvHrn513WcKznDeIFucXKcKvzyJIw5yuks3Kjxv8OgAAEABJREFU4ped1eBY55nOzOltArmuspn7auTREnckFJWLAhupZ9pHma+88/LPIvEmkK6GK3M7078hZ9nXg/i2aOdd+oc/X//vJ557aebcOeL/tiIqqShooRACWctvvL4jo5cOapTi7ixSyHkkTkiUmMWYsXWUqTNnXnPT7SdcfMV5I1/9+9rNbtmtx60HbXjLwX1XVr3VQ4GfUb3t4L5EBj9R3ci9GkREq7cfkuPtB28EbxQP8d7G8Y5DNvoc6p0+Cvjp6kYr6n7XIeYCNlI3LtMTbxARqRvddQj4Gda7PX6DiPgFrCwIs2oC6fpsavlnkXgTSFfDla/TXf4NgeT14L537tHzph7NL3x32E8uvfLvN90+fcZM24ORl+wNmrBBi1nMqtnEkami8EBJjRl5zpKekMjUCQlOhFRHznt6yKv9L7nyr4XZww7b8mv7fG3QNjvducXuL261z5At967UygpUVqCyAp/RCpBk7tpi98Hbfm2XfXd+4/Ct/tl83tmX/+GFl1+TgkoI+TaNrZnv2tQJ+YqsZe/OYsbTJ7lPYkaqg8cYs/seeeriRx9++eDN9tx+60e22PO8Xpse022dvTqvsV37Vbfr0KVSKytQWYHKCnxWK9B+VVINCecXvTYl+Xx9h22GHrzZFY898tAjT0dLVuSnLGM7ljZnvkezDVq0P0GR1NiXhSgCI8lBnnlp6B+HPPfaXn3P3XCrM9bpu32HVXu2ate+0FwqpbIClRX44q3AV3JGJBzSDsmHFEQiGrH3Rte88uILr7xmF1sIyh/DyFvszXyPRuJC8d1ZZF9miY4dGslv2oyZf7nzvlG79/31ulsd1qX7+q07NNNgISpHZQUqK1BZgc93BUg+pCASEelo3O59/37Xf6ZNnxl5g8b7sSzL+LunPVNGI+zObEemluVIb+zLaN7xwGPvb9njyO599l91jXVatf18J18ZrbIClRWorEDdFSARkY5ISqO37HHvw0+Iv0Gzf6rG1kxV2asFS2LsziSS5wAeSmOcPnPW/cOHT+y56nfW6N27Vbu6USvtygpUVqCyAv+NFSAdkZQm9ery8PARM/lDJ5sytmYqvDUjgcWMn/TuTOwfcZDj2J29+vpbU/p0+8ZqPbq3bM027/OcdmWsygpUVqCyAo2tAOmIpHT0aj2m9ek2dPg7SiJjR0YmC0Y1BN6IsTtD4OVZzHx39s7osYt7dduhU7dVm7dsLG5Fr6xAZQUqK/D5rwBJacdO3Zb06vb+6LExi8IbtCxGXptlQHWWxSCBLCeiVtidTZg8dWa75r1atW1f+TumVEplBSor8AVaAZISqWlm+xaTJk8T8lYIvDUTsKCGbNAiOY2HzcwKL9Dmz18wP0iHZi1kOcqbE5YOunfurldM63H25DXOnLSm17XOnLzWGZO6nzGZuvYZk9c+fXKPMyb3OH3yOsW63llT9r5s+i/vmvv2x1XLMUjF5L+3ApWRKyvwBVsBUtOCIPMXLIhZjL47E1JXdZTMhKCqZDFRKyL2P/Jkj2aKLKOc95+5X79yxnXPLXhvcnVVtahosYoT0CsNcSIYSBDFeNTk6n88u/DAK2ZceM88qZTKClRWoLICy7kCMZKg7FBV9mUqYkS0oNBAL32RDJcZleUrJ9ww6y/PzLesZ86xNmGzx4hgnYoZClhT//bU/B9fP3v5xqxYVVagsgKVFUgrEC1l8bqsOpMsi5k1YybsllRJa2olJstl4S/+M/fBNxcpWbFOJYYpHKoitaqKUkxUo7REfWeojw5ffFFljyaVUlmB5VmBik1agZRFAknEEkn+r88kxCwjtbG9slzGVisZN45vTlj6l6fZl9lzKY+pKk6iwMVQKBrZhaFAjaB7RUVJW7NEHEWuf3LBO5X3aCxGpVZW4Cu6AguWyNT58uEsGTtDRk83hKOgf6IrjmQttmRkrJhV07A9mojtkNTikY4s3UT6rdnoccuri5QiqnkVIxyiQeASRNVIjkFp00WzrCIi15jJnS8vki9Sqa7Ops6YOX3W7PIFmb9w0b/uefC9seO/SDOtzKWyAl/oFZi3xLLYxLkyZ5EsqZbM0owhHAWdvIbNCl0DWUqFDAKI+hu0hIFfV/ocI3s0MQNpojz73mIzMx92Xrhg68gs82pN/qbAjixh0d7+y5xwEiddEEeCWH3h3SUEaqxef8d/fnzepc+++joGb4x8H/7ws0Pg9evIMeN+ev7l2NTvWiGFXLbD0Sd+s/9AUljJ8blXXz/xnAsuvfaGklIhlRWorEATK8D+a/Jcy2JN2JDXsMGyCZu6XRqzvERYquzSgloRA1VOdd3qtcdNqzY7pYgRKSEkVe8yoKnY2H5N1UYSa6qAYqKAoiJBZfzUamm8PDf09etuv/ecK/44e+68jyZNgY94b1SD5qPHf/yXf9+FTYO9n1LcZZstHrn+qgsH/OhTxqm4V1bgf2EFJvmObDmvlJ0a9stpLKJqm7IACKnFTxokWI6LAkZKlsmyytKqyK6q/O1Y4rkfGzS6QbZvJmFMZf/lKFERI6/bOEnOoVGqCAtpsr48/O37n3q+3GTytOnHnja45cY7bXrgsc+88toTL75y3BnnYXDi2eefeenvz/nNnzY/+FvD3x31yoi3NznwmMuv+9e8BQuPGTDoqJ+eTVp84Onne+5+SOtNdv7W6YOJwy7s22ecS9fA31zddbu9h741kjjUxUuWnHbxlRvv/80Hn37hxdeGn3D2+fc+8Sx7tz2O+9H3fv6r4848r1W/rx364zMJuLSq6tzf/aXdZrtS++x5GAaYEaFSKyvwP7gC7LbmN/XE1cCSYI9XAx0NSGzIeHuW+ZHxF84YoRKUIuIgGsgwVGmiqJoVRkEEzKvqWh0Llx7d4cHTV/3PgFXPOqB9u5bEUhWlHLZNq5tPXuU/Z65y3Q86b79uC61xpFNUVE0BpImyWpdVt9ts4z/+67ZpM2clM9LHgIuuJMv86tQftmrZ4rSLrmzWrNkGvXvQu2W/DTdat/cm66/7zuhxb70/5oVhw3nhxbPqhxMnPz9sePfVu414b/Q3Th3Yvm2bQ/bc5c5Hnjrp3EsXLlw0edqM+5549q+33n3MgXv3WntN4lBve/CxP914+wlHHLTvLjssXLyEfd+cefN4szZ+4uSb7nv4g48nEo05DHn9zbseeerSa244+bijqR9PntqnR/fmzZsToVIrK/C/tgLzltibsk9w1ezR8F0OR/+bZkodyt5MHcTfnQkvsmz3FPMtVZPRsIk4lFXMYzz/iPbf2Lb1Rms169e92Q93b3Puoe0ISj12h1bnHtZ+2z7NN1yz2a59W5x/dPt1V2+Gh1gQGxSbnJva6LFg4cIDd9/59XfeG/bWu61b2f+elMzy3NA3VunUYdzHEwqFMHr8R+3atDlinz0I8ZNvHfW9Iw7ceL3eJCzSGS6ktrdHjR3x7qjpM2dtu+nGjzw3ZMnSpef3P+naCwfusHm/Ia+PGD9pEo6tW7a875rf/n7wGat06EDzpTfePP3i3+23y45kKNW6CXenLTe9/9orTzrmcCyrqqtnz5sH2X7zfhut24uVPHLfPTu1/9L8J0mYeaVWVmBlrcDMBZ880vL5RtuNxSxjT+aYZcaCUkQM1DCSraSpYlaCpVeRtEfbfO0WG67RbM6ieObNc+4Zyp8+ZYf1Wmzeg42Y7rdZqw6tdeSEqiN+O3Pc1OrVOoYd1muuRLCqTlSVaCpNlurqbKetNtv7a9tfd9s9ixbbLpbHwKVLlyanLfpu8L0jD1q1k+WgpIBrr7Ha+j17kKreen/M/rvtSP567IWXW7dqufF6vefOt/Vu3qyZqhYKBYxTXbVzxx5rrJY4uHjJUvaxT7009I2R79OsU3HEvSSSSbfZZKNvnHrO9wdeyDzJa6WuCqmswP/OCixYIrzd/8TXiy8RluWuPCTaYQnM3qAZ54jskthp+SnLsmVFwbSBOnch2UaaF6Rbh/DhdHjs3Db07hZIUSjEnDkvW61DaNtSGWfREtuU0UUgf+9mTdujYddkbd2y5cnfPor8knJu99VX67d+nxmz5my9yUYH77nLfrvssM5aa5ChiHHXI0/yoo3NEZmLHdmkadN22HyTzh3aP/rCyyQ40hzGZKLLrr2BB0meE4mzVrduONapvbqvef81v+3Qvu3Zl/+Bt2N1eus0X3/73TffH33dRYNev/fGf17+y46VrVmdBao0/zdWYH6+x/jkV7vMCDFmZIwSxrxIUM8rqnZSNWx6Fr4dMzsVVT+AcVOzl0cvbdlMzzyg3cl7t21eUDLX5d/q+P5vu6VHyx3Wb/GnEzt27RDGT6seNmZpwFcER1VLst5UWY6yy7ZbHLbXbsmwXZvWV517Ru8ea/34vEsP+dEZ19xyN6/z6eVZ75/3PHjh1X9jb7XtZhvPmjuPh9BNN1x3w949J0yeul7PHqS53bff+henfJ+UN/jKP6/fqwdx2rZpncKWI4lvu837nXLc0TzV/vXWe1i08t46fPWuXdq2bnXiORfw94fVtt/noB+evswMWCdCpVlZgZoV+NKyRU2ms7mLsl/cMnnE+Kb+nWnTEXxhbCNE8hAyiKUQdRB7d0Y3v6jRnj19q0Z7GTVaMDZamBva6Vd3zbnmifmjJlUtWEwwUqcsXBLnL47VvuFbUhXZtd39yqIfXTNr3ORq9mW8L7OtWnJP2PigbHlmDn18y403aN6s2U2/OX/xW8+fceK3MSdDDb3rhslDHp497Mnbfn8x2aTHmqu/eucNHz/3wO1XXdKqZYvvH30oxiMfvm2t1br9+8oL4f+47DxWgDg//+Hxc157cuILD7565z827N2T5Pjo3/8w6rE7V++6KpFBOAo6Y+F4+onfOnD3r0Fo1u89YLedeB/XsX37YXf/8/1H79hzx234M+gHE+x9HNEqtbIC/zsrUOW/8o1db5bF6igX3Tnl9XELG7NpOoJ7Rc9XmVjWyowTN0Z2Z/m2SFVUg5s2BSqigi2HWcM9hC5YLFfcP3/fS2dc//SCqiwuWBLPu23ulj+bOnaK/VeAho5Zuvf5M87+15zJsyLOeNWpBJFPWthtkblK3oVC6NK5IwmrpDRGsFmlYwdVbcxgRfUxH368zwmn7HzsD5548VX+PNq3T88VjVCxr6zAl30FMrY39a5haXWcNrdq6pwqdjY/2GuVPqu1uPTuqUPes1fY9WzJUfW1OoqmYo+Xqo6BlBTYJQmF/ZHalgq6jBrFdlXimHi0zRreZAXftUn9QpdFj2zL0psyTGgUm0zCW6hf3nrq8d8c8Z+b+avoVYPPYGd31blnki6/vJdTmXllBT7ZCgT7ba/r+t6ExT++dsKJf/qI+pNrPx7+waL5i7Ir75/29kcNPHU2GKF2RPJFZENGVomRlJJXUpqbkeGE7CbLLGpWWHtVL6YI+om7tbl9QOejd2jdLGibFvrLo9oPu7Rrr272zzI2Waf53T9b5fxjOvAq904AABAASURBVKzeucCQKpKjBTBfXebAX3gDLmW9nmsfvvfu/F2i++oN/GHhC38FlQlWVmAlrECz0ECQ7qs2/9Heq/x0/y7UU/brsuFaLZsV9LDtOvTt3qq+dYMRapup8PumCQM81fzdGUnOdkjkudpODbTMJs+F7pW4XPzN9mce1G6zdZrzvt9HkdYt7A8CZC+CtG2pG6zZ7MgdWl39g449u/qWkDiWYXEv7dEwrNTKClRW4Mu9Aq2aNzD/jm0Ku/dr9/VN2+26cduRExaNmbzke7t3PnrHTtqArTQYobZhjFmMtf++KVkWVD0gqCQ4re3TQEspGJZX1V7dCtv0aVFVHf/25IKrH5nPc/L8xfFn/5qz0alTR0+qIsqQ95ac+fc5U+dka3cpbNmnRVDhRyVHz65SKZUVqKzAV2AF2jaUzkrXNWV21dsfLj5+t84Hbt2BDFDSy0nTEdySNGRVDNTRskiIGXss3yJhxXYJbLLai7MovCnDKudR2rcK7A8XL5W3P6zqvkqhEHTW/GzsJPsLJikMy85tw4gPqmbPz1Q99UY0grAv4x1ckdu54YOd3MixHy5ctLjhbpEZs+fedP+TjfUuU1+waPE9T7wIzl+4yB7Iyxweeu7VUeMnlAkrgdYfpRQ0XcjCxUtuvv9JsKSvPPIJIz0+5LWxH9X6Q+3SqupPM0M+U9YBHPb2KOonnFbF7Yu3Am1aSIuaf5led35rrdL8dyes2UQuw5cIdd3qtUkhfHlIXjxWgjGzREZKc0MtQ6eNgUrxD5pCPlKlLfLm+KXvTaxq31ov/06Hg7ZuxTCvjl765odVdD7y+pK5C+MGazV76LxV1l2j2dTZ2SvvLVXzxT1VYkDQpLEyY/acWx546uUR7zZmkGXZgoUNvFNszL6O3rply6/vsAVIEmGs8t5Fi5csrbINZrn4KXn9UUoB04XEGP1XnY+s1PNfJjtusXH31buWT2LU+I8ffu7VcmWF+MLFi1kHcON116GukG/F+Au+Ap3bNDXBFs3sV74xi6Z9S14qSiay3xAVIZ5XXmOZQnpLSU6WWchVeCQU31tFw1/eOveeVxdNmpWNn1b9r6cXXHDrXN+7xVueXfDru+a9+cHS6XOzV0ct/eXNc0dPrPaxbGsmVnxvSEDjDR9jPpzIr9OYjyZWV2cvDR/567/ddutDT9/56PPvjBlPvfvxF3Cj9zf/uOP8q28cN2HyX29/8OJr//3nW+7Hnq5Up0yfdeUNdyG++f44DObMW0CcN0aOefH1tx9/aRi/Wm+OGjf8vbE3P/DU+x9MIM6l19369ujxVVXVdzzy3Hl/uOHeJ4csrar++12PoF93x0MTp85I0cr3bhgQn6HvePS5kWM+JAj8xTfeYYhLrr3l/D/feO1tD7w37uM0yriPJ1/x9zv+dPN9b4364KJrbiYsQ5QvA5w9I46/veHOabPmXHXjPdShb71ff9zZ8+ZjQ/zbH34Wr9sfefbCv9zEJMdPmMIciPzW6PF1QjHPV0a8i8gE/v3AU9Xc3ERwvPSvtzAZyJKl1cwW3z/cdO+8BQvJXOQvlppFYzXuevyFZ4e++dywN1947e06QdJqcy1c3aMvDLvrsecJiBfr/+zQEdRpM2ez/q+/M/qNd8fc8cizz7/21ssjRpYvUVVVdsO9j1HP+PU19KaAFfwSrUC7FtKhgVf8y74CvPBdth0Ji3u+vT7jRAKh5rszNWfLbWJJTpZRLLOq8KMCelVR1cmzsnP+NXeP86btd/6My+6av2CxmCwaVO8asujYy2ftfs70E38369X3l4bUUXIXVRHCSuOFtLL26l3JTZOmzViydOnW/dY/et9dA4HKXHp3X+O044/YbdtNySPfOnCPPbbbfNaceTPnzC2Z8CzJY/Che+7Qb72e7dq2fnfcRy2aNx8/acpHk6f1XHN1zPqsvcam6/c6Zv/d+O06cp+df3bi0Rv16dGsWeGIvb92yrcOmTJ95mvvjJozf8E2/dYHJ0+fmaKt2yP/b28QgUTZpVOHs7//zUP32PGJl17/4TcO+MkxB/N7u2DRkp223Oj07x6pqt1X75JGademFfVH3zzo9ZGjD/v6Tmd898iPJ0+bNdf+d+yEoo6fOPndsR9uu+kGbBs/nDSlRfNmJxy+z6qdOtQft2O7tt87bJ8dN9+Iyxn94YTps+ac84NjTzxi32eHvZkupE3LFqVQYz+auLSqau8dt9y63wbkwQ16dT9i750LITBis0Lh4D126P+dwydPm/nKm+927tCeRVh/nbWGvzuW3lQP3G07VmP6zNk7b9Xva1v223GLjeoESWZVVVX7fW3rvn16TJkx6+zvH/P1Hbask5g279tnsw16M3TzZgVcqqqrS0s0afqMJUuWfvvAPTHYeN2v/r/d4/K/erVrW2nbYsUuC3u8ls9HRUWU9BJEjRmqsjsT1fQqrLRdkiZKzy4F21sli+guhmRHJEe2BxIZgQTq/wzN/lUafa7YPk6ioXWJETMT6dHVvtPSSJk5Z96EqdObN2tGDsKkWSGADdaq6ozHtL/f/ShZpkvnDuU2Pdda7f+O2PfBZ18h0fRZe022OZuu35MtA0+T5IhySzImY5UrJU5W3XT9Xscfstem6/cuRSv1kiZIfzT55VTVxoJgkGohBFXBq1WLhj/5rqt0ZLij9tm5d/c1CwFjLb+KFAQc/eHEe594YZP1enVo24ZopXHLL6QUavMN1z3lW4e+NfqDm+5/4riDvt59ta6//ccd7L+Ikyr3O257xGnVsuFZJbMSNhgkMNugS5dWtWzRnGssGS8P6dS+HWZX3Xh3jzW6kanhlfplXIHV26/AHo19GfYrcpmWbUqHCkksBhKaEsMOQUlnabzstH5LbEgnKvYtVYoRTupKCVFVkTiLqFVakCJq+qdnNr6KbrdhU785m27Qi93Wfrts8/4HH2fFf3TcoV2be5948ZHnh4qXCVOm83gy5PV3Nlmv55x583l1zdPc3PkLedoiFWLyxsgxPCstWVrVvm2bddbsNu7jSev17L5oyVLejvErhwGV30Gej7bou+4tDz5903+e4HEMsVTZ1pE4Hhvy2jOvjmCnlqIVQigNgQHPnjjy/AW/4Z7Hrr/rEbZ4LZrXStYhKKOQdlNkdlU83P31jgfJOGyIJrO5mreADN6iefOqqur7n375sRdfW+z/ERHsy6+CRzYuE5GKF5l62sw5pKfq6urr73r4xv88seVG66ULYWtZCjV6/ISb739i9tz57dq0ueXBp0aO/bBdm9YpA7IU/3nqpb/e/lD31bpsu8kGrPa/7nt8xPvjWH+GqFPZM3J3eXnEu6Ug7Oku/MvN5ZmRfERG++e9j/Gsut1mG3bq0P7pV0bc8tAzZMxCCKzAUy8PX1pVXSdyVXUVn0vvtdcIQT/NK9E6YSvNz38F2G2t1r6pvwwwJd79Y4MlfEUqaUO4w3tOgatyC2WvFIlhWyrbahmn2Xg9fLtWuHh1WxzrVPpMIVqEsvly9KaFLRfhBDE8cHv7T5hZf0MHuQy52yqdfnj0Abtus+nOW21CE/Fn//eNU4877NA9d+zSueO5P/72dw7++tk/+Ob6Pbuj8zR6wanfXWfN1bbaeL2099lsw97HHLA7EXhaJNQlp53IL+1Pjjno/47cr02rlglPOHxfbDbfsM/AHx5z7IF7bLPJBgTv27sH8ZPBGd878pv77YZIWCyJxmTgaQjiYIDj0fvusvNW/U497lBqmvDOW21CbwqSRmFuNLkQJnzmCUcR6si9d+7SuQNDr96lMwoPv98/av/jD93rG/vt2m3VThgTofwqtu63ftpXYok9lj///jfat22NF4+lPHFvtkFvojGf7Tbri4gBoTbsvfZ3Dtnr2wftecge239z/90O2WOHHx9zUErobMd4lmTOPPwy1oDjD8fs9O8eQb7jkvv27pGwS+eOTIZ9Io+iZL1SkO6rdyWJp8yYLAuFwLjHHfx1JrZm11W50ww66VjWHJERTz3usP132YbPkcVJlUGJPGvO/B6rdz1ot+0Z99EXh7FElfrlXQHeha3dSdbwnRqZK6hdCghnR4ZOLzamrsCR8obwGEF64e6YahDlR4Q9EwYqyyx912p2wu5tMFRhYyWOqqoEMg4VhaiICoRDFWZNUUNakCKqqui39myzQXf7Hw/Iyi5Lq6p22nLjVTq2X9mBa+J9DkPUDFbGuGXw5LvZBn3KtE9Lu6/e5dM83LEB5JUceepTzmONrqtMnz3nqhvvGfb2qN232+xTRqu4fxFWoE0LYf9F5uq1ivRZVUA4Cvonmp6al4FaNhFVzmoZiewmcOtevuNnh7bba1MeOdlVifJbJf52LNYgT7FEsrdjEQP2ZcRHMGKbNaPmazzG3TdrMeCItqZ9Bge/Wp/m93N5ZvQ5DNHgNPj82JqBDfZ+MnHrjddfs9uqn8wXr7ZtWrX2/1Aw/NNUlpS/bPAHB/6a8Vl/fJ9mnhXf/+4KqKrlFMslMYrt1AITQiPfGJKYaC9H/f2JHb9rezTipSpqRVRURVQsTapKUFWxJhgEUVWKqFa+vWeby37QQSqlsgKVFaiswIqsAJkEc0MVcouKkG2CWG6TWijLVc46tN0dZ3X+9q6t1129YG+6yY9kSCreYF4tuu3grAm3E8a9Vy98c7dWN/68U//PbF/GLP7Xa+X6Kyvw1V0BsolnLTvbPzmLaXem7J6Eg80SSIZTEVEDWVbZsHuznx/R7u5zVhn2264jruo2/Kquw//Q9Y2rur7xh66ve33tD11f+2PXYV6H/rHr0D91BV/8fZdbB3c+/ah266/9mbwvW9asK/2VFaiswJd2BVTJTapKphLlRzhUBGSPFsR2ZnawcYK3bdumbSZzqpbAK7WyApUVqKzAF2oFSE1tMmnTpnXKV4aRXZrNkT1aILEpmc0Oe/hcs1vXzvOWjl00f271UjOpHJUVqKzASluBSqBPtQIkJVJT57lL1litqwppi6rpUCF9heCvvCy/kefgG/Tp2WLM5BdnTZm+dLFUSmUFKitQWYEvzAqQlF6YNaX56El9evZI+SplLpDclUX7752R3YRDAZWtN91otbHTbpk8/qPFC6tik/8fBlIplRWorEBlBT6nFSAdkZRunTy+67jpm2/aN+UrsJS7VDSQ5KytcPsLZOeO7ffdaKM1Pphxw8QxYxbV/M+hP6cpV4aprEBlBSor0NAKkI5ISquPm75n376dOraPnrZEKQJGf/8fgjVVUdCUvxeEQ/fZo8+rY2//aPQD0yd+sGi+1C2VdmUFKitQWYHPdQVIRKQjklLPV8ceuPfuZK0gxawloqJeiu/OePJMz59g186d/++QA3s/9taZo4bdNe2j9xbOYZsnlVJZgcoKVFbgc18Bkg8piEREOurx2FvfPmDfzp07Rt6TCcDfNGNN7iq9O1Pym5Di7H+TpBp32WaLH22342YPvfmrkUMv/+CdIXOmj1s0jz+wrkTGAAAQAElEQVQrSKVUVqCyApUV+FxWgIRD2iH5kIJIRP0eevOErbfbYdstVGIIoaAUURFVUUfbsvHMKZQkWI+K4/577HzWnnttfc9rj780dO/XHv/l2OE3T/ng0ZkTX5o7/aU50yr1C7UClclUVuArtQJzp5NqSDi/GDuc5PPYkFc3v3vYKbvu9vXddxLSExuuKPZnSstUJpDEqCLi7840qIgKnSJgFOUksvM2W15x2snfXdhmi7uGPfvwcxe88vzhrz25w9CHtx/2SKVWVqCyApUV+IxWgCRz2GtPnv/yc888+Mwmdw795oJWvzr1Rztsu6WKksokglLr3ZlnMHT7y2YUe/5Mz5nCTo5DyWfKedVOnU78xqF/GXDyoN6bfWv0oiMfG3P4XW8eceeI+vXwO0d8FvWwO0c0Wu8o61pxfuidI/J6x3AnjnU5olfTndw5/NDEG0TEO4Yf8pnVg+8Y/tnVg+4Yvux6+xtu41iXI3o13ckdbxxUxg+8/Y0D7yjWT89ThDpI8/Y3DrjjjQPASv0UK2CflC9m/qnB+ewaQ7qoqRdCXRF+0B21vid8ww+6c/jBj4w6YtSC03r2u+KUHx539CGrdO7EOzIyklDU/4oJCnu0/A2aeJ6zfVnQYJs0EQxUaKihSq6rdunUcZ+dtz/l+KMvPOPk35175pXnnXXleeCZvz33jN+em/CMK8+lGv/N4NO9nuGY8ytcBK8YfMYVg0+7YvDpVxieccUg45cbngFePvg0x9MNBxm/YpDxK5xf7vzyEnfjX1vz9F8PPg3itRa/bPBplw1K9fQyfhocY7oMB58OuWyQY12O6NV0J4NOv9Ri1uClruQ4+DTrHWy9EK+1+CWDzAC8ZPDpjqc5Oh/kfNDpFw8acPGg07yWOMrpFw0acIl1wekFU63FsSnW04pkwEWDTrtwIDjA0fiFg4xfaHp/xAsHguj9LxqIDk8Ioeb8ArMZcMGgAZALBoK1+PmmDzh/IPU0uJMB5w88rUgS74/j+ef0dx0c4HzA+QP7G0E/Z8CvBvb/1TmpLoubZdGmxCEDiYZ+qoUaCDbG+//yHHpPdTT+y4HGfzmwHj+nTC/yXzhx7A8Waxk/Z+XzX3pMJsxwJU6zgTnXuxZWA8tfDbRLhv/qHEh/bza4Pv1/dc4AwprlQONGzunviH3/X9XhNBv57M53/fxz+FD6nz/QP2vwnP4XDDqd79KFg07jK5H4BYNOu3Dw6RcOPuPCc0+/cNDpA0//8Q++c/QeO2/fpXNnEaWQo8hWlpTYiwVVU6VQ4vZMqWnLZolPKIqfRIiwYYNHDlXlMVVVnIOkOW/QZoQiZQMocBFPjaoqFNV0gubVNn+Sq8ZLlGFVVfgBRdQKKGpcVKmAYw2n6dVsVGt0Vc3nA9NaOvaa95pOk4pSRDuLWsB6iJiq1nSJuaKqihhXlkwSt5PLAiJRMRER9Z/i2VrGEaNCCKBqRJgmh3Mt7p0hqipS6ihy3MQ4wMlNMKJlmI6ERUOlD66q6FqMbzwpqoiqaopwsioCgxgKRZVuEQdx7iC5YJ9pzvlCSd7HyccVQTFLBC7V0BUjqnSoFbMxbl2+PuKojhJFxL5GOSKi+H94Dx8R+hShyM28xOmEq6gKRVWFH1BEFebo3FoqFCcRLHJshB4VVT9UREWVA1zp3MKKElbFh3B0LgKHOTbMhYKXf8xQLA1F7XMxpKXWmXPxNQSF5UNztKWs4SrGDVWUQhOsy0VNF1WqiiRU8gaBEUVQbA6q1huZAs+K4h9pFIqy/8JUkDBT74fQNge+NpaorIU1SkhzDEoJKhQNUAeJnERtcBWoUpgISEM4kQuDFgwlFESD0It38FKghAAUCo7wYKUAaCihhoCBIZ4aCmUY4IHy/+ydB4BdRdXHzzmzgUDoJQHpTQgdBAGRqqKIdBSBzy6KhZLQSQMLhKbY22cHFIUgKL0oSu+9SEnoJBASAoGEZO98v3PmvrdvN7tLQlH0e7Nz//Of/zlzZu68tyfz7ktA1dQ5LUoroif6vVkbupgyVrQ7JhVCWUERVbUaoXA1WlET5c4b1YJYiDVXVfFIKg00VRVzVO2BdJO6yRqo7mMFEZPrONQVo2FzkVCW6Dc4HvRaMZn7GBIXwxyRvGFQ8jbAdxg9cWkLN+dESAwNnYnNuTriDHekQamRxry4gpvT8AESHSWm1dyCF3TdwyfGF24Wd2d4u8F17cHpdlW/dxaLDxrvQEcG0yRvNGFRQpqpc625K4U3keXBcatfNPEhvXANvRuqjwqFd4RzcZ/C3zzsilnWyUTUwrthzN7Lyl1nBHECabWXfUjqhkC30nFuLXvrkyHjVxDSVft61TRRzIBkZqmjoOGuKfYomQcJtA7DYA4+l9JRcU/3oIOniHMFTERU/TLHxh8y/uwMOY5nGQfyHRwFLkrqUxSIMsgbUpxqcEcVt4pE0mQ4mTUr+TSDmdb/nUEmmQbHC3eQLAmKsiTNvhoRRxVzRQOdKy5cXtWLAYylL2pKWFWQmMr8rtCiKz5CcWOJrxIx+XpXRAUXxoI4FFQTpThCRAAFonFdhR93FXoigSjKJYwV2toqFFVAQNpuKCxTXISUqtEtWBSwdMGoHlh9zSq+VwQMxe/Fua85Q1TEUf3enRTdUF3nfhFFpUbXtV65iogqqMyi0NgrMVOKozAv1K282uq6BDbm8kG+BnEi3KR6fHyCq1AU1IZPWM1RHUVMfagjxNSLo6ijehHG0gpXWYOEZ9dcyiT4gKJw5vL4oo6+V6INru5joHJfohqzgwZX1oPCKFXFR1W8AhChoarCvWoUFCmu6hGaHI/Xz5XRMX23mIheNUrMS5fWkYaqDX/Bp9xLjfiU+xUhMO8HRxWs7GGgqAk8kOHBVSiq2qUzzPcWBR/11wKQwkVNvThK4f6fU5TMW1cEn5zhnhPokB2Ewge/TBSYhFJcJRoyh/hVMQzCQzINPyEtVYyi8YhugjJGKnMPTCrKAkRUSY8ijdWoioiZejETxaoqYKJTyymZ/5iRPmFK3oW6jwJQLJZ8kPkgel1cTYxhHhJfdLUmV/5gFVPFQUEuMNEQITDh3MUxhqc15zJKknBqQRXczALruRiqpoEGMTitN43Ja46KwthABpjfER3x0eg4NDACiWOZpczb4LWutLUPkyo+SvHXg9cYFig18kcPwzUL2I2LopRqIkEaiEVUYxYQI+gO6o7BIXVFF4GriKmoSM25I7iGYqoigIpYNKpqKsHV6NCNJoCOOxsOUBPscMfu3JSxEjoAD0RV57SWWhTnvLQo5rpRlMuSRBPonF5w5rLCCScs1RS9Vro4xlpXCrpFA9aVxn1QGRvoUDhRvRNNALL7vwbX3n0YTLQydk4eg1iJj3VOiyvIGK/BrL4XwaMHd2cTUFt9fJesxx6aF2Xz+FVkHTHELKlPbqAGV3R6IBLEMRojvruKITU4NLoW/iQM2ppjEkuiuJriBCRV8oniYKWlT2WYiCI70heuDGeQku1gIBroJn5nkLAggeKXOOc3jdZtqhAuyaRSWjOpP7mKF0SaSMV8Ds+iHjkwM1KyNDhMXJHQkfEMTgT/Pa65+7iCQ12zd7PgHrPwJ4hQnCszCv3g4gWniAMnJksJlF7mxU3EIxfSiE8AdU5kGlAYW+JjEuUndBGfV+DiRYGMZ2PG7ILgLK6Ik1DQI6aAGX+fS4kUHFEzS8JbGaK0Hl9FYKFjhYOuC8VtmTg+i2aRKu4XZU7uJtWC7qDaKwrR3CRBHDNTu1gvT5gzKnoLJ1o3Z2kMkRiOc24JC0eX2ofIjSoiERzkFkChEEFw4I/ognhgVN+3pgd/bhMNN1SQuXygAj4XTbGCaGDxITKejkQWaXBxH7pFdIKSG2IrR4yaA90/XlknlYet9RhS8xZPd8NUKnohYK8csVSp11nvucS9gOhgjruDs0dwsMmxinAbOQc2uR+LPAg6S8pCYfGBAL4uMQ4xKoowi3MVFWnwHBzMoRRdPGO5gSBZICKiDBXlp2wSc8OzepuF4neW3YkVheqAhSZ8DFTiiqjCRdUwOaqo0mkAVFTIWBDlEhKmqqiqqUV1klIyMmdgSvCORLHA1GEpWYLjAnfsAFKHIaeOZGY4mIGFg4mCYFAzS8kSWCrDknXAwZTQOywwUawDCB07umPChSvhReNolJRSIBBS8uIsJUtmHRaYAuERE1NDT+Y+gagJt7jgxpUMpC2YKIVZsuAgHqD3Y67k+5A8pjlywcGYFwlqZj4YNO85T7C4UkJO9CxRnCd6KaWgRkkJjhUKJtOUEGo0L+pgPVGNV9otGEr1TsulltQYltC0hVtw0Py9kgCswQ00CoE9OhZXvAdVeEuFuk/riJqbJmKA0QRoiiaAm3SzKwmulrzrV1JHC+zB3UfdM2lyBzVKgyeU1+RJLWlqVLPggbVeeGBxKz7F6mjqCkgQ0LTvec39k/tbcu5odIMbpTeOJ5ZkKdFYgPXBNXSNGSKUv0ZmIENBqk9Xv0C8eG5DhIHu56Z6JI3rllRFjDOQhJLV1DUiQqyDgeqcTnISl4rSRq4iWcEVbqbeiOCKKiKVoHjCU5HsXEiBhWN1BTs6qIouSBiEJTDAB2baxp9/FVw1kMj4tqA6F5AoYKYJRUBFzxoIE/9zlUmyhyUrR0X0yIpY+R8dWinxNeZqQSEauiORsDYxKyExeUMQdBAxq+YIlRnLH6Qg86pWiq5EyBooDcQ5FCJgFTgjRbIHFqc4MFUgIkpUrdF1RY/4UvmMRFaslWgmmgZ6BK1QBat3MvExOS2z+AkRO7pjLMCJxxfECkXYB0WsREMJJIj6vERCr9FF5V5wrpEh2ECVomR8hMiKD7WCC/FRuqNKN6v7qKhwv45c0vX6Epqw4pEVkn06zThAmqi+A1h9RrhKzYW7EJzZJUdp4f4OCWtuYmxXdp/sew6JiQoHMVXhw9RwFCdNH0jULj2c6dbVbyqXga3IqnDoC4tnsRY+J7IMKrpjyxro1gNZCXWOBRRrE7v88fRaVYEeOXZJuqOwRQ0Fzs4XbO5/UQjrEWJh2oXlNa1fXwZmfw9kcawKqiBXIvze+zr8cjuOwhQiBbOwBvYVM4ghZwb40zTn/tJnnz6XdOYmJTdlUXKciMKV4onPG5XABmjRAz2jEtLvQLIyNHtM9pXIiI4ipkLphvQF3Wf02xdmyCJCKBFxRXMTEeuqAuEzLSaiOeIs2XkL4sS8jioUrbHEF3QGikgTsRPWUX3SRvyMyLNMR5Ua1ecSyR4fVKFE/OyouCGIc2liVucZkRlrXhSRiC/dZlTWSfxAFeZiIO8WR8kugNHUivpEKuJVW9GdPb5GfOlCgQvxXYnI8MzafC7X4bgEaqBkRQCV+Jl52Q0VgbAPjioqRGvBYlUio3chA0XE51Li5K7ZJfvsiOI6bhpuNaqgdHuVsQq7JK2ziwjrcVTATTQMFOUnc/m8It1QskhRcnjmsGbnkmuu+DS4ehIEkgAAEABJREFUtPCii/hdoGtWkeANlFDmFeeM40qZV8oudaEUPYvkbquVHCvJDc9yjxL3JeEZqE0kcnABhcJuO3KpOAdFGA53bOURhBeCsF7FV4IPvDF7puvr0V7Xo3UmMYUlUMkqaJaMjl9qzjl+JRhTw5jFzF9uVUXxjCUOdJhExP/kdK6iqnRpalSljw+VPApSMdXoEyjOXKIi+LZg+PhNoMF7oioKQ7xRig9XJFUVfhrorVJEXXSUILFDwrapD2TPClfXpUbFJIzjcvQel6rrqj1QvC+hOnpkhWQUuCgj4QrnchRaUalRYFJzxRcOKtQbEe/TQWhFEdd9zQrpHj9cVTFgUqVVikTr2HopKqGaVYWCWJA3FDyQ1mcR3hFQFfXSE0VFVIUfpQiXwmuNRlSEq4ke2YWMAhcvWSkxCwrUTTSYVH1sA71VinApFr/gIs4DAXgX8l6ik1XE9w1QOMNaUYTxXMgtyBClhFK4UFTxVS0o3iqlkF5QVJpVnedACSw8u8yqFHHeeLkjR/VZSkwJ3kTE16qKs6qjqCqRVOdEkWIUEYxA+AjImmv0lYgG4oAu8c4HRRFU8JMGCh3nqnOghFJQfSAuokoOElDiPxabOXEJb8xO7+bOXFX+zkEQf8X9BCbCWMlZaTKia87Dh/SmqqCIYBd17hFqIhR0ryhUEfypdQMzVZIl4Gill+gqKVaTkVbjUzfEFUSzIAVxTBKKKB+krYHoJVQomLrVhFkJ26xEgBPckU4yinNtmcsUMaoim3dRoE10Cd2XzexJiVBX7xq6d/FIqq44oCBgIjL9GrGbh8WE0qioLjaGFO5GZ0k9CBiRCYWb0iSrreiNSovFGtEb3B1dS0ESdmP/zTxMvTlJCRscU+wDipkr4UYor1wM86rQZkWAE5MJArVGg7RWAlLr+Iwq8RvoMyGqUbrFZwDx0b0mgip7ghLO2hjuos+b1DCACQWuOCA4KhyjOTeKcmEivqPRxaOgW+iEDlf2hxmpoSikVlKTm4/EnCDqK0yBprWzd5WpfWDhmIKgMA4stRtvOGBq6BYBDcVnIQgGampMhFI4hJoUozsbDuYcgcZoUFqQnovuw1XXZD4uuRuLpxNoBMTBMQVPGlxZGz5gvTas9L2yzW51kzWJYaHjaEpkrDXiI25FpjEVhSX3McvqlbSDiKpiIqoGeuuXc3WChyc1IdXRj1TnZzRBJC9GCqx1VeleVRsKzpG5q8CsRGhUkapUFYJXKqUWH1eYCQctVv4okxyJ31EkS+NxiYtaOQqmiKk9MIuEwmds1t3AEl+kWEHyuVfh03BT9FBZAnnAkYnjj6sqVxBLxRkSJl+tVqA6iopzkUrwkaySIWBUN2ltcgf1e6zUB3ah4KCZ6dDhqllxKMgNqA90q6BnbaJknKWBShAJq1RwqggD66pEk0q1Qi9IQBx4EQM9lOLTqEwripjBGELkyh0aU4hkEQ+OFaISDlJpVJRS6/jFquGjMVCZIRNQMLEqhVd01bvFoWClUsUawkEqxVkqaVQVd1OpouYWJD5uxQpiClRer6xMJ2A4aMFida5aKQ4gPlRIqVKJVj61Y25yJoU30MPCVavu2Kr3xauWUbVPS+QqeNe8ZSU+xOcK/y7iXfHFZ5CBTVRuLXQFxa0Qtzqv2DWRCkVbsCiBONTVHZrTScUUKCBuohFWfQ0+HW7Oi8gMGQec3SQVKBlUkcwF4hGYAwEe7lU0UlVVZyfHN3fIiDkGYglrruhHkuMlFlEyor/7+EXHLKoEV8miYuql9nQanhoooqIw8iYOzkVNFOKowQNxRAkUEXQgUFSFEigS3JE+vMbQvSsIZT1CNJQm4gKvUYmMsUYVReeqURWu0oqi2EpVInvXpBWxC39KiATmVuQFUZEWzM7jBSxEVCBSY+ZPG389ja31OOhE40XpFVUEh15QResC6V4tusZAZQcETzg9sMHFFEFQRHAQCkILKqpf4q12Q1VRL6AQX+laN443egPdJFKjqAhXF3oPoVmJBlfVQGGF0G7c51LiKwVOgFZUURTNDjDBSSjqnA5UpMlR4TWG7t1uBCNr8CHGy6pw1YISPFACVVwp/K1AFY+vEqhNjLUJi0IRlT6qhh6Ii3JFBeA1qjpXdZQm0vPdxsf3WYMHYkDpBVX4UQZ0ryLofmESCQ4qxY9XomZK8fOaqiVMoqFwElQT8Qyj6Cr8rvjri7+hq6iKqKhy0RfMXGDm4EIC8w7+vHxUVcVCzcoIN3FhlUw6zAwgJwZKxlPFnZrov7YigRizujErxZ2VoFCaLlQ86DEE5lxo56iIVHWv3EBGwWv0yFJzd4C7MxRWIw0egYheRUUVP0dI1NxA9bvwToaRnVpQ4PVVj/aeNjgdEWGoiKPrKlo4v3fOsnrpBUVVtLiKwOAtmH1N0mfxP4xwwQ4yEmxw/szye0BR4knrLA2eVWpdlB86gd7m2F/pWZCRshCh39dXiIdPoLdKES51RaSJ0OCuOy8rykJR53H3GOFN9NFuRxBRVRFHGqUUrq4IRWHBafurWVSouPuMzplRfP+QfSU6By/Km4stM5Y1BIoj6xDV/qqEFRRVxV0pTmseiogrQlFk5zRKy6vZQN8BQcvh1boPqMI4qmBXCVRVUS9dKKEUVBUv5A/eUVVFOslkFTiY45MS8+QKL6pkTmTUTMZhlHtWFb6CIPgJMUlnxOTdHaI4VzUyIUYVlUZR5/77ITleO1BMhWyZTDuSZ9BA/wsjSa3DUAylwwJRSqWrPPXAzVIC+dTdDZO6kgwkgqbkGNG0RndQIidtRDbtgGsg3Cwig664W8uMKaxJCWuOdFVZZFLtSK7UsxgOCu/QFjRlXuZKEI9Z4gemgkSwZIE4UOHGOl2p4ytc6xnNfIrYByeNuZLHV0cLVE0lPl01dG4qIZqlxDoJQswG4qDBzZiR1Qb68uJeGrttxAyumDRZjUmd0/VauFrwGgnCDbIPSS3uooF0qUZMPK1DCmoSuLqzOCbQzJHg8DnQ4yv7oz4wrK5EEGLC/Y5Ma84sbrKI30DWgK5GBPSYC641x8S8ZQ0F3ZOAiifxe6IoIq9O0TtKN4J0rYEFoxTEocmLUvAN674MInucWBLEI8deBSkr7IncrHv6DsTrpSnu1zcHU+GOVutwqvGqucI9ojvW+6zO2X+1sh7fEEyqBETx2Q2uhSfXtYGGteZWeEHtEMEEmol/Fw+qGlXURJUKiKKoWkpOTAQdZ1qDqTLWDVKR9oT05jUyH+nO0xXcU2BYs0Qi9HHuKo1Ctqz4REtm7ezsDMxVZ3bq6BpXp3NEKr4Zv87Oip/OzlxVVTdkLSiucmVsFQoDGujOMbLymLDcidU5hMpsFWonCoMdc3A62XUP3tnp0ejmCsk5LbzoeMB7IjacIjoD3dpZFcyYCNPpc3V2Vm519G4xVYwKxa0VOnftA3LlSBDszqvQHSvEDFQ40DDEWe6ER/wKDBMjO2ueiR4cQsVScXUy1ptcOQERiYxUVVBm9NvPdHBwpFt1QvwiclVVoHdyZwNpGd8ZeqeHJRDWUrFULM115IgPryo62ZEL6rPAXKngnVVVo9uqCkTKnUDErxxzpyOx3Rockjt9GWDlig+s8Kg5JhQityJB6gkyYxjuA/CJpcIr5wzIWOl2dsKhRM3BG+huTY5PxrHy4LnTsVZaeK0XpWDlnrnwVqxa9OARjQlCd6XwGllb1ekcN3gOzvqrKhYZWHU6dwuunRWeuROFtoGdXRzn8jpWnT5jwdzCw9pZ5grdB0MIAdIBmYMumCuExiyVl1wR0gmW7G1nJslQhVxE5tHsOSZnlOxfceaq6YClqjLFQbI/L8tVp6jAyUVG6jJPaghe1S2CSIWrl6SaxI9+RHGTei7LQqPiiTGpmWriEvGk6Srj4AQWU6yCbiXnmqiI4axmhavgbaBAqOocCcVoUKIGtxAZqR5NHE0dRYjmc6Ejg8gguoq7qPBjXGqB4sgAV0RBZFBVWZUF0m1wWnRkRxFHdR8T4ihhTCW4GDK8VAwq7gjWXMxE4WCIanREBVGUgt7C6akKMR0FhwZXYgj7YIoSlTCFc9duEFU1xqODgrPUCi2KYhAVUaWqgjU3VXU7CGuiy8oYWg8funmfHjpzgckwiyvYgyuCc0IJFzVxMZtjMQbz8Urf3evx6IbMhQ6hJhh9ItcoiK6hENM8QgILp9eTaxRL0QT4yCBAkzMMH/PIfqE7J573/BJ1LJdoaR0FoOIt7KGpwBQURHhgg2uXzgTS5c8rKHhSfTR6g2k3ri3zqntqvWbXlZA4lNqNp6IFhkEYrIwxQArv2kMjEBc6WHNjmGnC1WJY6cNdwYYIowZXiij3youiZCvnqclNWbSaCFUx8UmRkxMowidb0gzHQCoefCxEEWUe91O8jUv59pMshyRKDBJXFomekh0j7QkKugjjSYGkScaZujsTaRYlr+FJ2nSs3KfTY/ogWs+aeGTOdJkPuOTNLCXRggTPOFdVzVl9lZkCHX+WEDwXHiiOPgSCnqWzQnG3ThF0UniGiM/lPAvIGsCMSC3xA5mXGUOv56JLjZgoRPaYKNQ6Pt4e060RszU+os9bCffOknLcSysPB5ZXZhd8iO/oK8zCvMwS68eTyL4PKOwbSCCWgWclUNZGfEc6KPg4CtHYUsfS7azgRPMd5k82j+P7XyuMRWEsyLyN+D5vT86GC+4+o5uknp09LxOBxAep7Ax76zGzz1tlR3bOB8a9E8h5RbTe5go97lQa8fHkvhwZyESxV74GnyuzV4RkFhRQcuauM2vAudu8gu73TnAGYOXeHf3uylxM4TXmxdMrc7mPz+I+PiMvIkNcIYzPTkB09pC1ERPEULGMyq3Omboicq557oO3+lSNsZl7YWqGutI1FzFzKC3rCav4mrMvnpXX98KCcUMMgs46QYIyxJcqvBaN/WnyypWcwVxl9hZ3R4b4/ZbZK+F+CYUSc+HM3SE64um1cs7+eM1MJBTnVXBXiM+KiO9Y5dypik9FquEEpaQwzzFM5L6CW/Y+hL4SBGdECJLUf68ajQSl6ORKgqmQ5vxzac6ZiJNfeOGSq6477ednHHH8aV8ZecJXRhz/5RHHfxUy8oQDR57w1ZFj4TWOCt7AA0eOPXDU2BohLfWgUWOpWHtFxIPDAXLQqBMPcl4j+sGjTjx41NiDR594sJPA0WOdjz7xEJTRYwMbHDd0cNTYYZBRJwaOPSR4jaNOPGTUScNGndioTQ4p9cTh4eA4OngXnjh89ImHjnKE9FZPOnQ01l7wsNGIJxU8NHgronerY9zzsAYePvokeODJhyOO6YZHjDnpcFcQS23tnnTEmJMbtXd++BjXDw+3w48NfqwPOfzYk49EPPYUH37syUd4PcWxIR557Ck4dMMxPuTIYx2PCt4XHj3m5KPGnIK1xmODd+EpRx97ytFj5kDEqEcde8oxQfpBTC311OC94IhjTxlx7KkjaoREPeZg0T8AABAASURBVK6ILfw458cU/bhTRzLkuFMC32Q+wiP7XCN8rlNHxLwjCsdE7eJlkWxFL/cV94upW2VX0ftBTF11jv33F+vYlldqTPAaT+Y15f3Aa9oPlvcGDl6PPeWI0Sfx/jnuxO//9FdnXfn366dOe1GUBKVkQSXJcSgUFXUmajBRdRRPbqStTJ4TIaWTaJ2r/00Cxmay46QpU350xh8/fcKpox64+ZcrdPx+2xXP2nmt3+8y9KxdhrYi3ddX/7CLh+oHMc1dXbvp9sddnP9xl6F/2GXoH1v5rqHvGvqua+Pwx11xGPpHOKSua7s4Bz87fPrAtUN3hMxZz9nVTa+JOMxrHbfr2gzpBzE16joNsva4XV+bnxs+5+6K89rn9sbP3dX1QMja5+62jvNeEZG66zrus9s6f4L3izi8vnpehH1NxOH11nVbBnbx83dzfv5u65y32zrnv2X8vIh/XsR/Le4rafi8Bmer8ewHMb2RyovO8H4Q07mN90bweC/tMvTc7Vf+3YoDTnjwtq+edNqvfnfOs89PFskVj8k4A5K+Mk2uOKhxLsvZM1igpzflRJeFDKd+KPMzmnPVv99wyyEnfufn6YVbd9/4vTtsOXKTLc/ecLtrN97h+nZt70B7B9o78JbtAEnmnA23G7Xpe7f+0FZ37PGuMwa8NOLUH1530+2WkvqJTESTCQ/mOI3B6ySm2EhsFZ9ahQznBzTPcUrmq8677G/fvPTiG3Zef/vNN7lso/eNWWX9fYas9MEllt18kaU2a9f2DrR3oL0Db9kOkGRINSScY1dZn+Tz/i02vXWXDb51+aUXX3ZVlf25cK78e9UqR4kzWs7e8282OYsJ57IsIiQ4JbH9/fpbvn/dP277wNDRa2182EpDN1tkyZUHLrRwGiBvWmkHau9AewfaO9DfDpBwSDskH1IQieiuHdb+2U3XcUYTUT+Z+TMxhYn6U35QufhESh7z3ManUCXFVc9OmfLjc//80HZDT179XbsvtcI7F1ikgwHSLu0daO9Aewf+1TtA8iEFkYhIRxO2G/qrc//y3JQpnM54si/e8B0pBzASGF+kZv/YyQJJbArNoqrnXHD5gxutsNfyq314yWVXGjgIa7u2d6C9A+0d+DfuAImIdERSemTjFf9y8ZXq2UrUeHymIlqKSPlnzzmT6qrK09tzz0/5y513PLXykp9adtVVBy6ER7v+P9mB9m22d+DtvAOkI5LSM6ssdcmddz0/ZUrlhedoJK+Kh/2ZD5Y5+8dOFfUPmoE33X7PpNUG7z1kxeXmX4Bj3tv59tpra+9Aewf+/+wA6Yik9LEhKz632uDb7riPrKVxRhP17zc1/j6a8c1m9mxGcqvg9z08YeYqg9+z2OAlB8z//2en2nfa3oH2Drz9d4CktOVig19dZfA/HxmfKVVnJVLlThJXVYFiqipCQlOKZHlq4qQpCw1Ycf5BC6f295hv/9e3vcL/6h1o31z3HSApkZqmLDzfMxOfExVVnp2Jqf/AyWDG586cM8hHUZLc9OkvTzdZYsB83eO0e+0daO9Aewf+/TtAanrZZPrLL2cSVtXJj+cuf4bmfSPHsUaQKnwjwFFNOL9lxNesz0yYfemvX/zJYZPHfmLS8R+fePw+E09o1LH7TKSeuM9E6kn7TKSevM9E6in7TPz2Jyb96vDJf/3Ni5MmzH7NKdoO7R1o70B7B5o7UOXsuckv40RmgHAsU3XmfxWNMxn5Czdx5n7Nsf2Ry3/70s+Pnnzzxa9MfnJ2ni3q6RAQWpUGqjrXZhFY1SnPP9l528WvnH7M81ed/pK0S3sH2jvQ3oF52wFyVZzPyFuZh/5SkVaymErXD/mM3tyEPftbU2+4YLr7kys9AxKdcc7E1QbPWZxiLRWjEy5hoORbL5z+52+/gEu7tnegvQPtHZjrHdB43m9K9uLZWSC5K56dVaQ3znAkGPJMSUn9hb38ty8+cNPMCCQ1ihO+L1WIqBNV56o1Kg6qDTTvifdEH7555t/bZzRpl/YOtHdg7ncgZ7JWrnJF6cydnfQ5Kxk5RUgz5DryGNnstQJOnDD7hgteVpzxLAhpqf63PuhmIhI/kPOae5bozjDgEp48rsu3XTj92Udno7RrewfaO/BfuQOvzpDpU2XqJHn+KZn8pCMcBf113G/2LKIimtW0UUXVPLN4nnG7t/Ia5c6rXlEvokIjjtJAFwRQ5UOsqttUQVHmpFUKdmohYKPed9Ur0nfp7KzOvexvW+2zP/WsCy+bNXv2Sy+/8o0f/uKX5/yl70FtS1870NbbO/Cv24FXX/Es9uJkmTFdOmfFQyaSTnaOgk5ew2feFqTkqpz9XwMQkA+XgVUmnZFnpPXqP+74u15lMZy3utAHeDYkZFBmwujo5y+3wGlaag4ONurjdxPWR/d6fe+3Z338kBH3PDj+tnv/edDXTrn57vtemv7yr8b95epbbu/Vvy22d6C9A2+HHeD89eLznrn6WQw5Dh88+/GZw6RxYuKYJOqHJVE1QePTJ7kJ5FNoyXbSb5k6sdNENPKfFfQuAqIqilIa3Fvp0vAU7epqV3lhYqf0UTiIXfC3axZZaNBVZ/74hVv/etcFv1t9xeX3/OqRjz898awLLtv+E1+aNPn5C6+6ZuXtdl1gva32O3TUxOcmE+mnZ/3pXbt/8ju//j36kSd//0tjxi60wTY4QKa/MgOHS/5xPab519lypW13We19u1953U0Tn5u87/BRKOt/ZN+/33QbPu3a3oH2DrzuHSBJcf6ay+F44j+Xzn5c6vTnZbnKOVdVZycXecwzEiclUkzJM479huycjbt7+MlLpCDnMhWfgog0NaIEC58Muo84wT/cyqkNmqtGWB/U/Rq0wMD111pj2kvT373nZ74w6vip014aMGDAemuuRi58x5Clt9ho/Vvv/efeB49YeNCCu75v63GX/u2A0Se+MmPmtJdeuvufDx91yg/WXWO17bfY9NGnntl7pw9sut7QX5z95wv+evWEJ54+YMzYFZYZ/L/Hj5z24vTFF1lk0UUWHnb8aRddde3XDv7iwPnnG378aZOntr9y7f5KtHtv2Q789wXmtDWvHyHxZ9TcbYVJ8mMV2UQ4p5lZOZ2RSyQTgczSREifVYUUpo6qNSrxVFVU1biES70HKVXVu6qgO6iXQhwFJ8UkfRRVJcUc9rn/mX++Aaefd9EGu+x39c23f2HvPRaYf/6tNtnwm8O/dO2td7w6a9bXDzngZ98cscWG615/+12PPvV0CXbcQV/4809O/eB7NzvtmOFLLLrIs89PRX/s6WdmvDpzxoyZq66w/OYbrDtowYGbrLcW1qtvuWOJxRaZ8ORTKdnDjz3x6JPP4Nyu7R1o78C87gCJidPWvI7Cn1GMhbxWzcITs3h2JpkTmnBG48GVZzhRxmrBHLmNfr81u5XRNCC9nFdce74Dvr/UoWcOLnX4mYOpw84cTD3kjMGHnDn44DMGv/8LC2f3Z5KorIEuMxYkWh914HzzfWPYAU9efeHRB3y6qvxrgVbHF6e/THdARweJL6UEb9Z11lgV8YprbyIJ3nH/gztu855iWnOVlQ7YZ8/fX3Dpujt9nCXtv/fuM18lJc4q1o2GrvmZvXZecrFFSreN7R1o78A87cDLL86TezfnuRurnKHEj09SUIOTzsgsnlG4pGQ76a+oiIoqqFqj0ldVobw6I1/8o2nf2nfit/ed9O19J56276TT9p34nf0mPX53/aTfj2PqMytFgoh6kRhPiDkq32Me+LWTN97tE78658//HP8Y9pWXW7ajg8Rlt95z/3mXX7X95psQ4aSf/eaHZ5x9/e13r/vO1ZZfZghuzXrbfQ+QBDdZb2hnZ/2EbvLUaede9tfP7bXLTeN+fcu5v91wrXcyhIHPT522yXpr7/K+rXfceouVllu2GaFN2jvQ3oG53IFXZ7zGs//+4/DNABH698lkqjidcTIja4EoHIpIZ0omKZcEk36LupXzGB9SG+it50Qs8w3UD31pkeFnDuFQNuzMIRzKDjlzCOeyFdZt/pt2PH2A5MzTNKIwKght73X27M7VV1ph/BNPDT/htHMuufKjO77voE/uPXS1lff5yA73PjT+gNFjV1n+HcceuP+Nd9476rQfv3OVFb83+rCFFlygNdYO7918qcUXPfXnZ9x+3z8XHDgQE8/jVll+uZ+fff6me3yKrwJW3OYjD4x/lIGrrrjcl8ecuOuXDuObhPKNAc7t2t6B9g7M/Q7M8m/a5t69F8+5iKCkD1Grz2imcFUhnZFfhLRChgP9jEav78oYVTUur+LAFS2DOJ1d8qNpfijbb9J39p3Euey7+02iPn6Pn85URAVvVYiqwoVQ4kSlr7LAwPmHf3bfKTdf8dQ1F75w619PP+Vriy68EB8tvzf68InXX/Lg5eM4VR31xU9Nu+2vT1970c3jfr3WqisTimdtM++5Zqdtt4Svv+bqE/52/qTrL7nqjJ9MueUKTCSvG++8+7cnH/fo384fe9hXnpvyAs/jGHjLub8hJrP88bsnkPIY267tHWjvwDztwCz/Xe9zxMxXqj///OmnJ/SX8/qPEKFz49lZFjJXLujpTNVTiZJUhEvlNQpHKx/MeCq+oEueCkU4nX3wS4sccuZgnpcdHM/LDjpjMHWFdVpOZ54ym0MY3gwC6bOq6pKLLcp3jq0eiy28UPMgRoJbYtFFcGt1aHKsJMFmF8Khj5Pdlh///KjTfrL26qvs/oFtEanE7DELYru2d6C9A3O5A1W//8DHk08l5/3s6Scf6fNvzvcfIZahYiqchbRGVb7cFD+dRXaqPMlISUo0fVYfrhTh8oDeCiHLgHI641z23TiUfW+/SaU+EaczfPDsqoSQuMTXhfVfVjcc+s57L/7Db0857uQjD/r7mT+9edxvVnzHMv+y2f91E7Vnau/Av3wHyCdzztnZmadPm/3SC7Nnz85b777UsisN/PPPn5lwn3+JN6dzrxG6u2VOZ7nqJDn66YzvNnPFwzTSmariqZ7ZRPmR/kvGJZP7Gg+8/DNs9yyYvcuKqITKOKDA6tocSJDwRO/mQP9fUDnr8VF0jx2223idNVNiH/4Fc7anaO/Af/8ORD7peZvPPjnzN8c/9r9jJlB/9Y1Hx987fcb0zot/O3HiYzN7uor0GkG6FRU/TFlBTkOi9elMsicTz1BCcU7TZ1UK0xUUNThVnYiWD5uLHnzmkIOiHnjG4APPHPLVM4YsHx82VURV+HFUinBp9FXapb0D7R34b9gB6+jlLhZbcsD79xm84yeXiTpkhTUWSB367h0WH7JCL/9Pkl4jdA+aOZ155iJfZQrnouynM7okFPHEAkSVfosnPwZ1r0XMwofNS3/0wnf3nfi9qN/fb+L39534g/0mNj9s+go4leHfVX0p/U7ZNrZ34L9mB/77b2RA8zl5y70OHJTW2GChNTdeaPUNBj09YebTj87Ybs+lN95mMentINNLgnU9AAAQAElEQVRrhJZgUFVOZ6bK8KB0TTI1+4dO0ownqEyqwbefyngVQGpUhRBWXPDT2Q5fWpSjGYcyKueyr5455CuN05mI4KmiFCc0XqOPrV3bO9Degf/8HRjgfxWqz9uYPnX24w++vO3uS6+7+SKivbv1HyHG5FyJkLk4kpHAquz/bDOTzoSMgkON0scEeLRUMh+9QA5ZpMCcn7hn5i+/+iyHsu/vN4kTGeeyH+zHuWwS+MP9Jv1wv4k/2m/S3342DV+RrAz1gaIFhT4B+6x43T/+8Vdm+MfsG+96oFlbB8ya3fnKzH6/Ig7vSc9Pfeixp6DEvPehR194aTrknxOe+Mctd01/2b88Js59jzwG4tOsM1+dde1t995w5/3Pv/Ai/k39iutvG//EG/23UFWVb7vvYYIzS6nN+P0QRrFORpVt6cezH1PZyVYHbnxutrF1SJP3v/hm5Ddl05qTtsnbbQfmGyj9/G/gFllywL6HrtBPLmMsEV7rppTjmKgJ+UqNHzgnJGNY+f0s+FqJRUyUUhDSrYbJ46PCqSr8aFxzoomoqqOoSn/l+RemnXXh3/jdw2nmq682K91mfeixJy+5+uZmty/ywPjHf/qHC196+RVinn3pP56aNPnWex8k8vJDlr729nvJEY8+NfGsi666/xH/5wfNIC9Of/m2+x9eZbllLr/utov+cVNTf89G6yy/zNLN7usjF1998zPPPb/IQguyjBvvup/6mnF4sc74yxUPTniSUd/69bjnX3id/6ik7GTrdHO5ja1DmpyVU5vdHqQZ+U3ZtB7B29231Q4suHB/y+GpmfT9C9//2GbcXHEGyhT/YMnvQ3aBjCYllzgyB7U5oneSJXsgonRVsiARmwgpNcfXmjmHJ8cySCCREUF3oyEg2Gd95PGn+R145ImnOzs5Yna5PfXs5FN/efY3f3LmVTff+Y9b7r761rs5Q53/1+tP/dU54Ld/M+7rPz7j7Ev+0TUg2OKLLMSRijOaGbnUpVmzZy+z1OIfeM/GZvrw409vudE6D4x/wg0t1/wDOgYvudi6q680Y+bMG+68/+Rf/PEPF1910d9v4rf0T1dcy4zH/fD07/z2Tyf+/A/g7NnVeVdeN/ZnZ7GG56ZO+94Z51Evu/bWE372+9Hf+zUnzZbATme+Omv1FZfbeOjq19x6L+ny/vFPnHv5Ncf/9HfMMmnyVO6F+Axvxnxy0nMvvTxjp202G7rqiu/ZaO07HnjkutvvZTru92d/vLCqhEzNtvz8nIsfe2pS2aJrbrvHZ4qLvceB+BfHHwB3PTihLOzuhyaUbfz7TXf979kXIf74rAvKnj835YVv/PjMb/36nK//6Ay2vbmk5jonPDmxLP6ehx791bmXsg/Mzj4wESthVVfddOfV8QJdcvXNbNr9jzxOKKa47o77WhfP2mKNbfgP3oH5FpCBg17P+hnF2LkZqaQuKmlLnXKZivHuoXqOIaWUKv0VMqvOWSSkJooanKqqgaaUYKUvwUHawI4Olb7LvQ8/tsIyS/N7xSmm1evCq25cYdnBG6y1Gtlnq3et+96N1+V3e/bs2Tu+d5Ndttv8M7t/8D0brv3YM5Nejk+pzYGrrrAs/o8//exK7xiMuPHaa2y67pqn/PJsstKMmbMefWriyssNeXbKC5zgsDbrY888+79nX8SvMbO8OmvWJuu+82Mf2ob0Vxw+vPWmH/vQ1ssNXvKwT+8134COux8czzHw3euvucD88z/+zCSUz+7xwfdvsfG+O2235ior3P3PCWVUwQ+9d5OlF1/0uB/8lqy05cZr77jVpgstOJAPxUfvvw9DSEPljlZf8R3NmE888+yA+FerROhINnHylNmdnYw99NN7sdGPPPHU5KnTjvnCvp/b80OXXntL2SKyDM6lPjnpORyIz9Qo666+clnY/Q8/XrZx603X2+8j22+/2YZTp700ZVp99CPjD//Untu+e/07HxhfljR0tRWb67ztvodYAIuf/sqMadNf3nTdd4L3j3+UiY75wr77f/TD22y6PlvHC8SMnVV15Q23f3Hvnb6yzy633/fwyzNeZWxZ/Csz/ZECPu36H70DgxaTuUxMzdvEn1HNbr/ED2Ocx6iSq6ryLsc1I4soF2nFUaSg9FkWXyYp+S9nHJ0IhCyIfy/oz8hwwJK56idlpE68GxhUZJFlrGa9NVOmvcSJYEBHxwMTnmi1owxddYXN119rnw9v19Q5c2kcss6/8tr11lhlkUELNk2FLLv0Es9OeWHW7NkLh6mqqnXXWHnkAftNnDzl7ofGz5rd+fgzz6oqJ7jiX3DFZZb+/F47HrX/3pzRUEgiYI9a/mE8iYZT6NJLLLr+O1f56Ae3WnX5dyQzAnLseuSJZxBxax2Yc95y43WO/sLH73pw/MxX/a9Uz5o1e/75BihbHH7ljqDNmOREupzpwMlTX1x1+W7/Wn7W7NnsDCYqZM4t6hF/zoXx6fVXf7qMBLrU4osQpLXOjgNyWVKPOE03/uzhNj+16wcWGDiQBTT1JmHPVbVXU9OnTf7Td2DhJebhjDZwkOA/17es/I6rqCUDnLsgdIigXJFweiQZ5J51xXXn01JEvAW8eiBV4acLlSJcGqrV/j08cfC63DrzS99l/TVX4aSw49abPvjokyrG2cosgZtvMJTj0t9uuvPehx/lHESyu/3+h6VRJk6eyi/qc1OmTXvpZT7XkBCLhfXs+YH3fmirTUuXD1zf/MmZJ/3iD6Scpyc9v/W71mWubTZdj+MMxzGeWxe3ecJ3DF5y9uzOC666kWdtMxtfUPA7zJovu+5W0g0L+815lxGzqvLvLvwrn/tO+824wUssxgfh6++4n2xIpvjt+ZfzuWyzDdbCjbrisoObMTtSWmeNlfmkxsqfmjR5o6Gr49CsPAfs7Oz85bmXnPGXKzcculpzi5qT8rzv5Vdm/mLcJZdcc4upVVXVXFjZxnsemjDtpencOx8hm2Efeuyp35x/+fW338eH4iKypNZ1LjhwIItfdKEF+cB++fW3/f3mu5ZbesnmShjSfIFIZOS735x3+S/PvXTt1Vacb0DC2q7/fTvAaYsklQb0d2dY8cGzP6c5bJzJOAd41vLGzXE681zjCYVfclHxKv2VoVsvIMSgCr90nLnqSsf1poiDV/KjW8IU3GOH4p50gue8xtb9fbtLfsGV3/Yvfmyn975rHU5JW260NrjmKsuP+OI+e+2w1Wbrr8UnxCM/97EN11ptt/e9hydKq62w7OGf/eindvvAUZynlljsXeusMXA+//swW71rPSonLBJH8dxpm3eP+OK+jOUD0c7bbb7eO1dhLiJ8/MPb8olyyfivni21+KJMh14qEajwEqEgQyCIeA5ZcnGiMfveO27DXCgLDpz/I9tudsDeO/F57aMf3JrfYXIBzpzb/mfn9x3zhX1Yw+7v35J1smxOWwz/xC7vZ/HvWHpJwhI8JUMsMRcetMAWGww99iufYNSX99mZ8yDroTILc2HFk4+3fGDcYM1Vm1vUnBR/RuHwjYM+zcGwdWFlGwl15Of3/tiHtvnGwZ/m3lknlcPaJ3d5P6dI7qh1Sc11lsWvucoKh31mr4/vuC0+AwfO11wJt8AmN18g5j34E7tRt9l0faajlsWDzNWu/zU7wEfIxQbLwkv6SY3M5XlGBIT7iWxJwYqPzGPxExkHEBWIqfB75Og5h5RChuMhe0lN0l9ZeuWOjXYapBSWpGqiBKQ60eiq0tW6SLSB0kBvJXykWNfdcdASK/X2V4nlzSmchvjlWWLRhecpHDvDh6YN1lxtnkbNpfNCCy6w9Sbrz6Xzm+X2Riadf775VlvxHW/WStpx/r/twHwDZdBiQuZa4h2y5HICwgctJuivaytyjudlgXy6yFXlfxGNjEZOkZJcPK469H+9d7+FVtu0fDYk/+GbY1D8bbI4c9X/KpN84MbiA6MWTx6iMSOHNScrbTL/u/dbCNtbV3kOtehCg+Y1vqpwNAPndeDc+BOZP0/mxvNN9Hkjk3Li2+7dG7yJi/nvD9W+w7dwB1RNvTjyvKSm5smGSzWSD+e0uVrCjocsuuGHF1QhitQIDWaBKmIoIoCCVBV+1C8paCrrfnjB7Q9eVNqlvQPtHWjvwLzsQOZRGf65tXjOEYkU44cFdS5zV7bcb6GPHb/EejsssMTyyXiSSzokLYIML8hJDd5ErPAsOC+2XFp7hwV2+cYSm+771p7LfP721d6B9g781+2AmlJFFXTKpY2/d0b+oZJtVFRFCshrlaVW6njvJxfe+8Qlv/DrwQecMeSLZwz+4hlDvnD64C8E7n/6kP1PH7z/6UM+f/rgz58+5HOnD/7c6UM+e/rgT/9q8O5jl9zsEwu/pc/LXmvtbXt7B9o78B+4A6okKFV/oMUBrVR/cMazM45K6F4lMpjKQoMWGFTJtNmvSiltbO9AewfaO/C22QFSEwlq0KAFOJTVVdU4mvkDtDidcSgrHwRF9B2Dl17ipVnjZ0x/sbP+v7RJu7R3oL0D7R14G+wASYnUtNhLs5ZZemmpMkczz12ZInA+X5LWRJQC0spaq60y3yMTr5s6afKs9r81kXZp70B7B94+O0BSunbqpAEPP7PGaisJqYtK6nIUP6CpGilNyG4cz/hWM8u71h86ZMLksyY+9sTMV2b73+R4+9xLeyVlB9rY3oH/jztAOiIp/WHiY4MnTN5gnbXK6Sz7Y7Ocq0ybcyaziZpnOlF+ZInFFvnQ0LXf8ejzv3n6kUdmvPT/cdva99zegfYOvP12gHT066cfWWbC5PcNXWvxxRZWM0umyczUCZlMxTyJcS4TcpkWvtsHt1/15vFnP/HwhZOffnTGdGmX9g60d6C9A//WHSARkY7OeeLhlW56ZKf3byukqypzQGs+Owsihi5KcQchn5kNXnLxL+y682pX3Hv4Q7ee+9wT/3xlGsc8aZf2DrR3oI8daMtv3Q6QfEhBJCLS0UqX3/OpXXZacskllGNYVDIXvIlGVmt9dhYJL2+96UZfevcWG15yz9fuv+WUR++7ftrkCTNe4muFt27R7cjtHWjvQHsHWneAhEPaIfmQgo67/5Z1L777M5tstvm7NpCqM3f6P6nMVc48M2ugSpzOSG/GKY2PoKZwqph+5H1bH7H9+zc577Yrb7hlh9uuOG78nb+b9OhlU56+4cXJN0x7rl3bO9DegfYOvFU78OJkUg0Jh7RD8rn8+ps3/NOtB2673Q7bv5eHZVQemJGj/JFZUkfzxCUqpkB8rRnHNP9n4VKBme7Wm278rWFf/fQrC2587q3/uPgf37jpmj1u++sWt1yy+a2Xtmt7B9o70N6Bt2gHSDKkmq/fePVVF/9jvXG3fPzlgV8/6IAtNtnIPztWlXRmPlAGhwgki2Z/lCacykTUcxqHsqgc0lQ5rZnj0ksuvv/H9/jp8ANHr7bRJx6Z8dHLx+/5p3v2PPfuN1bviuGOe4y7a49z7+4dx4U+cXvoVAAAEABJREFU77j7uLt2H3d3C9J9o3U3j3kX+JbVOyNyF+46zjm467i7dh13565ziee0eBY+B+5yzp27nHPXLl1I943WnT3aneCbXj9yjof9yDl3QN5c3Cligjudc2e/eEdY+8azW0yFvxV4TsxSIhfeO/Z3L+wh9wK+uTtJwJ3PubMg5E2o4+7a5bKH9nr4lUNXXve0g7/8iY/tWp6XaUqaTLtQKWKAWiKVqT87UxKaqCe8LAU10xeKRrvkIot8cKvNDvzk3t889CvfHXX4d0fX9bSRh5426rDAQ8Fvjxz+7RF1/daI4X3XQ08dMfzUEY7fGnmou7XgqSNdP3Xk8O51WHSHnTJy+Ckje8eTRw47eeRwcA4fhqD3WU8aObz3OsL1E0cMPxGHBroneqmhn1S447DgNfrAEcMaCOlWx44Y1r0Oj67jCccMPwFrA10/xvWxxwyjnnDMsN7qIYjHH3PI8SOGOR4TWHgDv3nMId8cMQx0nxHhWeMwlGLqFb9xzCH91K8fcwgVh4KQ1orYUg8OHnj0wV875pCvHRPo/OCvdcfjjjn4uKO9okO+dvQhkNfCCBKjGNK9HhTdGo89+uBjjz7ouKNqPC54Kx571EHHhtUR0ks9aMxRB485KvDowFZelJ540JgupS9+cItPd94Sv+9VHXxsWflRfnetd1R4866PPYq9Pei4o3sg3brGbpf9nKudJxRDQK+8rEfzysbwbtxDdb0Hjjn4GyOGtdZvjhx+fKkjho8Y9uX9/2ev7d+72WKLLCxkoewHMc5gWnESy5zIhGMaJ7LQM0/Q4shGSsM5shZAz0wNJl5UxNS7jq6rO5AL6buuySigmgUmx+TcknMLrimpqToHE9ySBhqoyblbzQmRLC431cQ8vvMEaVjh+GMCnYeO3XmMRTfW6typc2KE4tyJX66VfsRyKa66h9lZw4NIHpAuVVVAc4SI+d/fa6Iavj5UDexZy2rD1OoWPLkc22KBGginqiWj+nYldHNiNVrykZZMXYGrEwMTizFEt3ChlNrF3Zqi62gRVgONgFS4pQQ2aiIm3KOiW3FuIF2vhg8OYKzEuzVJZuHgTUI3QilviwY35wjoxEy+A6lw82jJ95whYirBJbH5VBXiGsQUUxDpwlS4GtakCbRAuLqYQOeGKSGYJm1UU2vlxgKNktAN7t7JufsTv3CGdPFiDUSk1j6h1FyJE0uyEsdXYtYaP7iGD1hqOCfn7p+0IE0JEpgCLQbCfYvEwHpPgtM1NVXXnYuxnyoJDipdhgfio3B3TqYJnkDm9coSEc3CAd00uAX6EEhdtdF1F7ToA0LSclP8iomKpWQmQpdqJaCKqZqvzRKKc1PxcxlI5XkZpzNIXeszGgE051yL4gMUzEokhkS+VA50pig4sxLhmIfiVoKUsU5w8ziScWNeUUFxjrNzUfHKFaKAElKTCCUUvtoI0UN18Zg3YqIzowSXeq4yb8yFFMMFIiLORSQiF8yFi5eah9LK3RN76M5bCXr/VcUn7RXFTao0GZ9yC+yqc25BBE5FB0W9i6nm4ncNzyohOjoXERQRR0jULi7SxRXuQSJ+cIBNEyFss3rMENnbmquUIcxbqqAwCFQaKSIoja5QmtyJzxtTiL+mbs34O2cuhbsDsyAyL3FqrkUX1z0OI51L4SrootJVeZ9IV8nM5r2MSxfPSK4wETqdIK40OPEyl+tKG1wzc7Eq+tQevNxa0fvijbERRyKmdOdCgCzi6FMLJatE12Xn3B9K414wCRYUKQXWvbJOd0KUaCEqiFGzOM/OCdPkWuvcCCZ/jYTiIpwhrRUHbEWBq+cKgJheSwQ+KWLyv9nvU4RvZlCGcZsVz/RjRpy5ryyVz5Jx8G82hS65i4wnRoZTEVUlx5mKCoo6NzSImCg+9AJTSohUesmUrOyIf0hmHZZSh3Xg1kELD7RAepawYXU0SkoBNdK4zfCCFsQD7pgM7MBIfCiRGmjMG7YOc6cuNFT8UkrFwxCgiSZoAvFOrM9dEl4ojCiY0K3DzELpicl1AA/HhL2hwJPHSmxMQlcu50YbXEFNoDmqBsKbO5pQDN205ccsVPOiDtaK0UlgilmsBVGS0k9u1S6kS+3ApAbCcTMc1cxS4egMB1FKpQsBqR2WzLh6QaSExebcn8SIZCnGsu+pdA3qvjS8C7BiL4gdjr3G8OgwQ2FsT0wIUTssuUsL+lp83o4UGFa8kuspJZ8d745gjomeX/hYSs7KRcdgRWmgmlnyPUuObm4oSbuUotc4h27NseYuycxQLFlBGmvwosTvQq3BE1a/C1rWnAxuwc13MjkmQjZ058kSit+t4UmlQZkD2Z3kXikRGatjxKMPd0wUZkwM7yBmsgR2GH2sjnSNoqaeZ1RNEu85Kl0qchgEYq5aMkua4rLkTM00qaYkpqYQ+o7KcNJtrjKoIiAAKWcxOEpkRFVUUVHl8owoWqyi/ocJCm6M5xSnudJMwIpkmYMX1ODiKChEBomiOYNITcyuEBK9J+Im4vFzYOGB5GxmxI5VJGYp6PGzEJMFg41ZSO0RP+bCzpiwIlb8WSDEZ4lSCUMlYvaG7kIExgfi6/vg3OPTLTGF9eAaqIG5IDHRHUUDBSSaMC8Gn1eFUL52jyxu49IMoDcRTzijaiS+uHOXwhqwFT17fEKEFTeB+1Zn4fad+KQtHJ1oYA8dESX0WCHgo2iI0w3n8MHBq+tMLrH+QA2UQFbLnWbn/vr6DrBvvkD8WTSvlNA0XtNW7k7e9/U4V6lRiebc790jB2cW1+uVsHKPL+LvKI/PUAkls+dSFB+C3qJkONEccfG7o9FaYWBPpVg9TsMnFKD2DBlez4iBHXD/HvM21oa/FO6IE2tmuGPomfsVd2LfQncOwdOxXr8497nYhgxnRAMzMtFQulD9Fck5ZszBJbgGFp7FpwBVCOVVnDABb3vuScgiXNTM64yc8cEXxDFLJVVVZoRn8VkESTJQ8UUnVljOnOw8vylZTgjZ4NEVI+uhgKqeB01oxNE0maIX7EhKpgwuiWSbwFopuqNKoOumBS0lMXQDlYFwc26OHl+sRghV6SaUJMnMHCVQHYmTNEHRnUuqZ5EUMQO1CxG9hmJqzi3QYxpKEkMkInKSpBKzaKBgSqpmmlSTSQOlyQ3RcDAU54khmhyFCEZxrsEDI76p8wSamOIkCd1AOqCYaYKi1FVTIUkSJrdpzFJQujhxsDI2icENVHMkppjPyHsB3sRs1uRF11C6dP4gQ1EfKxoxu1BDcVTtQhfVi6gFN2YWB1YCdXRTUzFXTMGI7y5d3NyqrEq66+o6g7yGD0O8Fl0ae05XgxuYiJ8kmSupcCInTarJJKkRM6kk55ocJdC5aeGWDD8UpU0mZnAQHSy8oPs3rEVp8dEmVyu8jh+eLTzV3Mwk+ToV7LoXZTC6uIJLfXdqpihmrhscUlecpKGoISpdMeI411DCp5tefAgktac3Liqjaq7Ke8DoCKI6V3UUkJH+LhLnKty74pdUGZ1gapbUNCVHIQZtSmIqgcmcW3JUH5RJiKQ6URGynTcheFfJjgpRVK7M+5fkGQxn9xa3hL+pKldwzbVOxOD80cRYR1ek5F0smVg0oAiqaqCr9ILThkjrfYmC4j4ZkZjia8qMzUJP/dLsWOZSd2MEzswbyAjBX5Qf50LR4i1IGe4xGzxi4Y+urme4t7gyo0jNFZ25sopg4X5FsHBR0cULNhYBYnFsGR+6hq64as1iLo/nfRVHZvC5goug1FVK4d6d4IXuWO6M+ILASGmJ2Z2LiJaf3OITi1ffXcZjZj8kYtLNCpNMX3JgzV0oV63iV/r4d+el56gEryuNhj/IHCqN9Yj04CKuiDhqoEjNaZk9RBSh9OBlLeFDfObBBSwcX+4Uzm1y945Iil26c/HXutYLx46P+JZrdy6CGv60wdmwHj6sx6Sbj2aiqIio1OsRcUUDc2DhfkcZn6yizqVwJkERFoQqAqmrSE2UeM7xR1ERRxqlDS4CUxFHGqVt4URvKKE6NDzEi8ZqnMXl/hqFeRVPbdixSHkczztKsWQULVaNcxmnMO5HglfCa4Sxkirj6/uGL0xFPR+JOjKRWSgCRbFyIZqRNxGTJdBUJamZCpeqdpglczSLT9DJ4rO2oXfAzRKoqSOlDkspWdLUEZgCO3AwF5KChtXMsCQwBRSr4ciVEtG4QEsUnhwYXDsY3GEI1qGpg5VoSpY6tMO5YTQiw8FEZENPrCg4xFLjJ8EIADLK0S9ione4TyImEmgxSwdoGA2fjpgxmbGiDjCl4InpkyUv1pGcBbpi9PxSFCwdwSHGOo3ImhJoHWaJGZ1rURKSjzFIB4R9wJDociWKR0JnJQnF47vilyJYxE/1nhvBjPUnplGAfkqW/Dmoz54SUSylwA4QC3tnzItnBzrROhgQPsTsgCiCoSdiGtDg2uQwuCbF7GgGplT7w2s9oUuTW5KUVAvyvoQbEVSDmwlWM2KjmCVQ6arVmMxLMtfhoWPDChpFk4CWBGtBTUQzTejahQnOKFDMuSRTS+Kzu3+Ty1zqmuo4Zs2YWvNQkhk+ycySr8fwNxR8GmhelF60KRBwD1N0a8FkyOqrNbGEtQWT62riVmbBtUY15zUmMwamZOoM7NLVIxjObqm5exHTjLngasS3KHijpYSUDNTknJ6JqSXmpgluSRRvuqKUlMhXfJtAsiP/iReeJ9BkLlEHkNatOCk9ictRUSRLRVLszIFVDp47Kz7Nghm9s8qdVdUJZkjuLCi5EzG7TlbtrHIX5kwPN5BolXelxgoTkauqyp25qmIWR/TgxEH32YXIeObguTNLJzPm3IkOzxVKlV2vcsGKsVUFxlxVjZ05Ex8d7ITGLM75g6GLV50EKZF9lpg3FyQ4VqJVrKSqsiPOFQo8V7kiGlfRnWckVouWO6HMwlwZJbMbnXX8qspu5S46ncRceFaZOJ2ZuXJgmUWq0KucsVYZa3kSIVXOPmPErxx9f3wW5Owxg+eC7E9n5WM7s8fphhVK7ozIBSvWWXl8nxGe4/XNzJuJVtGFd6GPRSe+I3oVswd2VrmzorhCtM5M13ejhVe+D7H/lc8V8+YmYmXSishceFY5+123YkUhfkG3Mgs+nVVwdqbyGavK9yfmrdwac9WvSOE5E7+b0qq3WgtvxeJZsOiFt2I3vfKVcNesKress/CCVegAxNH3hJVXdKsGxzPD0fDw9Vc5MOatAjtzxX1VzJV9H2Le7Ihnxb4VXlVV7swEkXrfCkdglKNf/iqgZ2bMLZz4vENIJlHJNhwiQck5o2RsTvwUloPkyp+j5U7HqurMVc5VJSwg83TEyGyeqRSAi5YiKmTLgmpqhjmrCoqqglgt8WOanHZhSs4DU6AlFLXUQAtu6laDczXQgpuJh3TU4F2YMBCnRjXngUo4epZQEMzzdjLlpwVTFzeKchkelizBC1pwx2QWPIFh95g1t+COKTkSwpUmN0MxMymsThsAABAASURBVLea0RC/RvNYSCge2Sgo3vPVYk2azGMmc8TskQtPriQDzVIrxpiUPEoLWnLFEugW4ifiG5yrJ2JJZo25jOLrcMUsJeMnJVeaaGapoTAObmopMYslGDzQKD7OdasVRXMeTcLd3MMKpuCBZV6Dq6UGGtx8ncmImQKtJxpFDRmfGo34XJbQrQUTnCvQdbcnj+8KvjX3XjL0ZI4EhfdE0y6lGzfXi2LWB+9jbGrRGzwlY3bQULh6QTPXmSqZOapR3M/QjVK4obsHgjmnZVcZ0MSU0LVgl87olFiDJcdk/mMGN+uGpsViXpiLpjtiN0abGTNShexiXBlRTVJKZsoQCSzeomqJk5qUUZbgSuF05hkqsiC5MIvyQdSrCh/YUTxHZvKfZBQRt2NRoaCEmuFogo0xIl08Z8+p9BVf54yDleyrRFfGZC4V6cLwwBqKgkLRuEBGQT1u9lZAfBkDQw2eEeHYmRsOCnEUXWmDi8C8p+IFJAbIzTsqPXRHIim+AqeNnhAR7kh496WHPbsigjf36JidExIuFOXyOFL8uBdXynjXNXQVwT9UYgav+3AUDa/CHSM+M8GhGQ+pIwdHY0TBWL2KeLz6auHE0KJ6jJiFPio+EVGdE0MiGlBzw9fnUiQlQuZSnJRWuJR4Te59rtBplTHO/YKLoPmlWbCAzKGMF29dwZ6Zz321cFC46j0XCV5Qxdcmobi3UDQuR0H3Cx73B2cm7ykWZqVX5q25WKyCnkhjDe7T4IzqsbZiLTrIKBAfsFfe6l94E1lTg9OyEpGYl9XQZ53EhIPOi00a9+WqSCBxxIsC6koW0JlfhRfMrucWLuIKQX0eib0IdIaf8/BwjmdwXrPSluhdXEXw8UgmFL8HfF2N91L2c5l0dnLqE+6Ci2Qk9LDS99Oye6jzLJGjyHiqop7/NIqQ8JwwWHOtkwLJnMWnyZ2YJVWuZKDb4TSWjLgpaUpmyZKCCvFqZkkgKWkyinNjONTgCk+GSUDnxDcXjRI8mXlFS1zRoW1wb5Ok2gNrxHRVLTnX4GoCUSNi1KTmJbjFfNiMrnMu77muSJoEb01auMsJuyo6ItW6OFZs5qJBvCaIV0zWVQhJBzsoXMk0oSU4jZqqFZ5qjoPBDb1U84BJUzKv9NzkTVz40GoCTAKJQ0cMNyrxXVe6Cd5tLhctWegRITHQ6CN6NYMnRGY3ioTIqMIhjAKZF1QjvolZkxerJddDZBw8wY1Q3iZ84EURRLyNwjoxUQ0TOlhqN56Si8l9nBiF+HWXhRCcaslKVHzMC60PNFedi1GSpuA4oweiKxyjxXrMPVDEyRvlZkkSS02a1IJbSjU3SlKM6FCv+HjfKRYzvy/vdL9f9KiWjAGa3NWJ8bq0cNeTJrNU9ISP+VwuFR7WZG7xC6sZcyWsVHQQH0dDj+nwCC4puqCqMId5SYaYkrpzUlNLoNFDEVFRxUtZj2IjPojFVWyW6yKlVFEyWY88mbO3kWix5gaBi6gUB29Ifd6BOsOPCssk07joMtbRu0IWzpSwEgedkY7+pytU8CJeMMZ561dQNzEWb2fuBa0HoLhJ+JMCEXfHiOmc6fx0Ju7l4SAxvOYu4+Y9mlLpFIIRd8cYwr13cZ+Enjt2zYV3VNZDEKYGcQIJ4K5NCTXn3JS8y5+44htfCcPVTag+rpXzp5Gbm7Pj2uBB6TCbI4OJ4eMRGCMRX7ygcy/O4vL4EFRfvDOnTe4dD+aXx6FP/LoioAf29priJT6vdC+EaArh4gJxaEBMYM25K18JXvH+aeHcGvPiHPtPywhfA8xN3rjiAZz75X1vI0rTKSYDsOAQFcCnRvaqtiprCN37cByorsQa4F7fDB4xuWlv/epagzCv+Py+/pr7rGiu+JbQbbkXqNuK6B3cPDK0aD2x+fuCoXkvzAgHYxgt++pRYEXBuYvT8TmxNBqn9aVuhXuj8RGQRRNLPaQIqUh4zyBw8vJHcoSlVlXn7KoTCZ698SvnyhvSnpDcRMVRMplPgvNFMWmSDIiuRBVvIxM6I5IIRHx8qMXT+wxD1cifwT2miTp3FHU0Fdedqwa6rhQpl9FE9X4QQHyu+oKrOgelEd9jMnNzLtfdjqP6jKqBohqKo8DVe8wozmlpGOnIpUpE8Zlcc46CBgo93APhGhwUhsFVWYioI0Kt09CxiBWcyMIOOBdVdFEfIa7Tc2cBpIUrNuJHlUAUqvtFZPW/y+O3iiLORZTxjrzzVZkPrnB30JqrqvDjr703Yo48XxV0UwpciW8qxAyUQOIoOu7oKlhFBUQXVeGHmIFwce4dCDW7g3dVCQthN3uimhK/oHAfxv2LGgxUZZwxRMxUaY1L1FRpjUspRgCtuUYRR2Qa4jjS8bhcMViVlYiq90FxL7ioFl2NVl0xHymA4EMT+KZwghGnEd/nYqFKfK5A2sZ6sGKs1+RMGR0toO6Np1Pt2ofioV5ELXRTpTUu8Xkbe6vGaFEjvqN6vCbHGd4LqjIeva6N11rE3wMAes3pCClP6QqzMFuyJCLclSWzlDSpWqnm5z11joAHFJNRyH3kplzFD0yyiv8UnWxZ8iOW8MAPjfOPMhO+mUh4KF0VrdeIXnhWVui/M5ot/B1VHDNPIUTUOU4o6hxAUYVnoOYqzh25CscaXESiJb4ylz9PYC5P9yjYHPnDpI6fcfZ5CeeKEF/gCmBqxFTGoaj4T3CYRxb8lbs1+kLkmnt8H4uCIVDxhNdzKcHFu/VcTZ5VhHiKLWIyqoVjVVdUArkvIQDxA3MgimYVYQ2OConKjMFFvMvVLX5mbIkZKIEaCMfa5KJzvKYScxVd2HMFlLFCsSyYTbw4KrriYKIKJ7KEAkfpwcOqoQfmQAmMd1wOlOyz8C0Wc0nscLyXVMQqNlOwqgTmQOfchFjFEsRycFBwdq5z8opbdKvEWMndOBG4a5BDBNjFK+KzNmIGCshAsOivizdjejRtzOV3UThrizWE0linVngyHfMGho6k3Gl3zlh0x669YgjR2HaxqoX7QO6FLQn0/Wco3FGFuXKgdMOsSl+A3l53rFJ03tUe2fc9vgGvKjINH1E873RmOvGkjDMYNywq4tbsZzT6VcZLfKl+GvNvA5QiXKoNJNPBhcJ9BDMrGTBZ5EnJKjj7HatWQtdRlNWgF9QqAlbuiYPrPiqc6cAJHujWjF5Xzb46qTFi4hazEFlzc0aVrFLiB8aMjGXGGt1BCFsU1odOt7Wq+ESqGVG1MYuUyIGa0XELJJrPJTGvI6E1FCUC1hZEUVFQCmaCiEfGJ6PX1RUiZPWYmJyLx8w1CqsSaaIKulDUdZXstInamEsJKFi7Vc2i6A30mDkiB2oWj9lAtxKhzFLBWLDP5T70iNOCMS9zibouAuYu1IxOt0aPnKWJbq2EIR6ZgXAigw2uVfO95G4sAudWVInIBYk8R/UpcvjUqDF7T9SsPlcgU8BBqbQbirjSgj57xQoF/8IxEsqxQpDXwRlWovWYq14JoYncfZ3u31CkuXtZW7lmIQJIFed0u1dhtULBgTUIw+n0RN8lkRo9vnSfpSIIViIXXdyH6SqRjKlGppASWThpUZNZStZhtKrmfw765RMZJSF5H+pWz12iavyoqqhVFeGFklsKIhUBPSp/EElVkSWripRYdWYSJwtBpkrFn5qsLp71VJiatcpVqU0lSM65ypmrIq1WtH7gc54pzonOQJJujbhlZnETbogsusq5Rjdl77qSseaKHzqB3mZKVS4aagzx+IhNXvTKJ/JZXG/wKtdzuZUx6E3ElDMOOYO5UaogjlX2JeWCPjCCZ0c3CVbnLdPFwsQ3Iecas2QvXcgohBphmYCZUmUn6B6zCu6KVC2IirVGj0zYllrBM/ebc3a3Hlgh5oifKzfNgbwT0IkhrF0IzxusxsydCZCl9a3d4DgJ0ej6QFbBwBbEFFUIgkNvSGisjj4vsxCyB7JmlBoFN6arMabzbk0ynBkDpSd6EOYKvTuPtcVYFsL6WWjgvOj+WxX+vcTvWkmZl9eI+IG8ylh7r+Ecd8qyCFuw5RZy8EBsrLqBUlbSE2NGF312ArobU6ME5noxXfNKzB7ooUW9H4mtKpjZ7aqzqvgyczZtZ2flwQhQFlO5BSNt5qqwVp6dGBY852xKTlOKRL4TZ5JBdR2rilBdsaSkTLUQyGVCydiwB2YfYqKm2kALHpjVeaDW6Lpz0SZqcEexQO2BRFZVw+oIhQdGTOO0iV6jGhZ0UU7TPhBdmMucI8IbGLNYoHrBRNOCREYgzpxINKyB2kRCwRvYbUYioAdaF7J7dJrI1Lw24ivH2eftwcVjsg/M0lrdWVWxqja4+jTNyKgRWWt0N/YfuYEGJ2ZvKD6jKlbrjoxFCfS3pjE4S0EVVwJb5pLg3dDdia/q8edEc90ceU3hgT6j+C6xJO6oBSMIPsIrrhro1liU1hg6plJV3U21gUxFkIJK8eGKAvUZafx1QYm9bXLceD+E1XrhwtiG3gvn2WPMGCtvxtQ5ZulSuBesgX5fGIRLG3ehXlC6qrkJ/y5FXYkZy372tZOsx5T1qyMRurhSuNkWxEU9rPSG7JioVCqeq2CR48TfJwwjTjKyTepIborxSt84hFHVC62ZpyM6DFFmElWWRkAhrxHYkbikwkD+hKiqil5+burUi6+65lv/e/ph3/zWl0Yc/6UR32zFL4cCfnnEN/vGE8J6wlfwGVkjylecH9+FI4OPPP6r6IFf7QdHnHAg1hrHuueIsSgHjnAOOh+JckJwnJ0fFAp40MgTGnjCQaNO8G4LHhz84FHHo4MHjxx70KjjwcK74wmHjBx78CjHQ0Y2cNRY591xGFYUPEeNHRZ4yKgTho0aG3giyrBRTcQB3oXDRzlvxWGjxg53f8fhwR1Hn9jEQ0efeOiosQ086dDRcMfDRoMnHtodD8N59EkFDx19Ij7wGkfRPfGwwMMbePjoE+GHjx5b4+iTnDfwCPTRJxU8fMyJh8MDC+8HjyDCmBPBI3rDIyPOEaNPPnI0PoFjTj4C0ZWTjhiDctKRY1wBce7CMa7XypiTjhpz8pGBR70WHt3wPDo8W/CUGHtKKN3wqDEnHzXGlaMdTw48CUTvB48eUzzrsUTAHzza5/VoTX5UT4UhfkfoR4/p4kf1zcu9g0fG3XXh6K5dYlddH13vKnuLcoTvc2Pn3XTikaNPPqK3V+qI0SfxKhfk1Yf3g0eM9ndI8Tli9FjeaeCYsd/70S/OuPKq656fOiWThfy0xjEwc17L5KvKW3TPVBUpLHueypVnNrIiiY185ihK0hMJXfXZqVN+fMbZnzrh1JH33/yLFTp+t+2Kv995rd/vMnTu61nhDJbKQEhByOurf9hlKAP/sGs3RGwqhc81rh2ejn/cde0/7jIURAlcu4l/3MX5ayDD56hnh9JECJWw4GvWc3Zdu1Q8Ia1I943UcREZLJVQkIKQbnW3dbzbO64zzvVueG4oTTx313V68KLuRalqAAAQAElEQVS8bvzTbh6wiRAqUzQR8ibW83Zbh0rAHkiXOq86Q3qt8xpnTn+UN73Ouautyut+BQnSOnZcvEOKUvi4XdYet/3KZ6443zcfvO0rJ33nV2ed+9zUF/g2M4uqcUhLjqKi6rnLzFFVxU9nmQxXVRX5jSoiVRVc8lXX33Tg2NN+lqbesvvGW+6w5chNtjx7w+2u3XiH69u1vQPtHWjvwFu2AySZczbcbtSm793mQ1vdsce7ftsx7eiTv3vN9TeLn8sqCueyXFVec8Gc6QrPT3i0kD3NqXqS48ymRpqT8y+76huXXnLDzutvv/kml230vjGrrL/PkJU+uMSymy+y1Gbt2t6B9g60d+At2wGSDKmGhEPaIfm8f4tNb9llw29fefklV14tRqLyBMVRTC2JwNWLGax8s5kz303mioancVVVcS773nV/v/UDa41ea+PDVhq62SJLrjxwoYXTAGmXN7oD7fHtHWjvwFztAAmHtEPyIQWRiO74wNCf3XANZ7QqnptxUuNbTgLlqsqZmnmmZiQ2Ec5l6Mq5DO25KVN/OO7PD2439OTV37X7Uiu8c4FFOrQ44NOu7R1o70B7B/51O0DyIQWRiEhHj2w39FfnXfj8lBdElap8Dcwny3JGE0Hyf7NJnsukMclVVcH/eMFlD260wl7Lr/bhJZddaeAgaZf2DrR3oL0D/9YdIBGRjkhKD2+0wnkXXk6q8soxjazV2Snkr4ovOP3ZGRlN1DSLqOrkKS/85c47nlp5yU8tu+qqAxeSdmnvwH/7DrTv7z9iB0hHJKVnVlnqkrvumjrtRbJVqZzRhPxl/i8EjCTHiayKcxn8htvumrTa4L2HrLjc/AtwzPuPuM/2Its70N6B//odIB2RlEhNz642+OZb7xKOZv7ErJKCPDmTbHWGi+834fc/MmHGKoPfs9jgJQfM/1+/Qe0bbO9Aewf+g3aApERqmrXK4AcnPCbCszIRvrskh4EwDmkczXL2/Oakqp58ZtKUhQasOP+ghVP7e0xpl/YOtHfg37ADfUxJUiI1Pb/wfE8/M4msJY3PlFk5m/njf+ORmZDnPJnR6vSXX5lussSA+aRd2jvQ3oH2DrzNdoDU9LLJyy+/IpzI1P/Rpz87I4sF5/lZpointhynNGec1mRuynMPyrXflT98Un62nfz4vfIj6pbyoy3lh++RH24hP6BuLt+nbibfo75bvkvdVL6/hZz+MbnqFHn2gbmZpO3T3oH2DrR3oOwAqYnclaNE1hLOaCgiDnzuJK05zZnPolrGzBVe9wM553Ny1x/l+fHSOctDxBVzcPYjhk9QpumGOE9+RG77nZyxr/z9W/i1a3sH2jvQ3oF53AE+YapnG1XSllJE4nRGGsqknpz9s+jcxbx0pNz5e+HLBQa2IhFQ/K+AkNGyFFIQfc5682/lz4fO3ZT/aV7t9bZ3oL0Db9kOkK2k8sxSsKqc+7Mzkoyoiqg4ymuX674v4696bTcyJE6OfolPwyktOAkUU8EHr5C/neK99tXegfYOtHdgbndAxYyclVX9nKZQ9dOZkGT8Iyn5zhu6/VWel91xlucm0lN/tQqfgjk42NotPPDm38ik+/ubdK5tM2a+etu9Dzzy+JOdndVcD/oXOT435YVfnvOXZ56d/KbMl3O+4tqbLr36+rfhnb4pN9gO8l+1A7NfkVcmy0tPybRH5YUJjnAU9Hm/T978XquKIxm5xY9I5CLhw6aqEk61oHhDv+/6wIVkPzfHeGI5L+esLvT4tVsRi3O4NvS6Uzd3n1eT3hp+Y8+97G9b7bM/9awLL5s1e3ZvXvLYU8+8a/dPbv7Rz35x1Am33nv/Jf+4/pUZM3v1fFNEVvXrcy/Y4mOfPXTsd1548aUeMclcw0847dEnn27qZ5x/0QGjTzj9/Iuayhshzz4/5SvHnfTlY0+CvJE4/5/Gtu/137EDs6Z7Fps+UV59UTpfrTMGCQGOgv7SU4LPPC1NI1EpxzFRCmPVn6P56cxzD7mOvMMcGPqvT9zkC6qfkcXZimdnJMkmEqSv6m5ZQCo+BSHUCdf1M+33fnvWxw8Zcc+D42+7958Hfe2Um+++r1fnS66+/qFHHx/+2X3/+L2xP/7dOH7b58wyvQ58feIV1934lWNPWmf1Vc+++Iqf/eFPPYJcfcvtvzz7z5OnTmvqn9lz57/89Ftf2Hv3pvJGyNJLLP6H754w7gcnDVlqiTcSpz22vQNv4Q5w/nr5Wc9i/cxBXsMHz3585jDlyFd8BPPUBY+Dmj87g5PRSCngHKPmEF54oqfUY1iPrnu3SoX7YupjGhOzgimPumNv10svv3LB365ZZKFBV5354xdu/etdF/xuiw3X4x4uvOqalbfbdYH1ttrv0FETn5t85XU3HXnS9wlAEnnXbp8464LLHn964mYf/cxFV13zqSOO22a/L/Ip70+XX7Xa+3YH4SgHf+PUpyY+96UxYxfaYBviQKa/MoPDzvaf+BKckxTOjzz25Gm/+t1i79qeCmk9GN738IQlFlt0350/uMD8c/UvKH5/wWUHjB578133ssi/33TbRrv+z8B137vDp7/KjA9OeHzrfb9AF4SjsIbPHPW1Txw+BnG3Lx9OXqaisM4hm3/we78565nnJh/yjVNP+tlv+Az7P4eN3v3LRxz09VOwbrbXZzilMgWHU/Zn/nW2XGnbXbgR9gexXds78K/bgZcn+YlsLufjpIb/XDrjRiJRP52JcC4DOJ7xYdPPWiQWjEigSv9ldhwXyUGvURsHty63opDISi3dglkI28e8gxYYuP5aa0x7afq79/zMF0YdP3Waf6y75tY79z54xMKDFtz1fVuPu/RvB4w+caEFF1xz1RWJsfG6a22/+SaLLjxovgEDtn33u94xZPBKyy1z8133Pfz4kxf//donnpl01Y23wlHWfedqAwakR596Zu+dPrDpekN/cfafL/jr1XyEfOzpiXxUvOK6mzlPXXXTbUed8oOP7fh+6ohv/ei62+5iilI3WW/tyVOm7vj5Q1ZZYblP7vbh626/iwzSz8fbaS+9xOyvzHz1+tvv3mn/YfC9P/yBfXf5kKl++sjjnpz47NcO/iJ4zLd+SFZlDWf++RI+qy6/zOCLrrqWIRdedS3KF/bebfhn93vnKiupKD4Tn3u+qirwwquu+cfNt6+16kq33/dPPphPeOLpA8aMXWGZwf97/MhpL05ffJFFllx8sbLsNrZ34F+xA69Mllkvv+ZE3RzwZ1Q3qa+OZpJYrnjYTwLL7kWiEeO3Qumogkqeg89bjVAcr7z1i9To1alfzj3gnNxVxnlTjM56uVSV3/PDPvc/88834PTzLtpgl/3+8terL736+ldnzfr6IQf87Jsjtthw3etvv2uRhQft+cHtGf+V/T767RHD1159VT6FnXj4VzdYa40tN95gdmfnXQ88+M/xj633ztXv/ufDVDzx4SPbaccMX2LRRZ59firKY08/A1JXXm7Zy3/1/aO++KlL/nFdMnt5xowp017srKob77wHK5Wz0td/8HPCsjwyHfv6xZEn8FxswIAOrP1XUg+L/+6ow3590phP7/GR+8c/Rg4i40x44qmUEl9lkICIsOXG61/ws9MO2GcPOBMNXmJxM+OguujCC2296UaIrfUdQ5bmk+w3hn0JkSPkjFdnzpgxc9UVlt98g3UHLThwk/XWYh8wtWt7B/4VOzBr+jycy1oX9OqLc/kczURMVBmr6ug8/l8Bnkz48OZZjsZ7+PRZ/RlZJWBdOWrRzbL8u+SzF8hBN8nBN8vBt8gh1FvlkFtl2G1eh98u1B3GCDm098qnYOmrDJxvvm8MO+DJqy88+oBPV1XF6ePF6S/jPKCjQ1VJAfB+6jtXWXHIUkv+9YZbnp0y9SPbv/efEx678vqbVlh2yBorLc+Xg+THO+5/cMdt3tMaASvHmdmzO1+ZWX+ZsPTii312z53JhsWN49Jfr7/52AP3/8CWm/Gx9N17fWb8E0998eN7dKRUHPrBsvgFBtYfUWfOnEmiLP7v32JTjmzzz+//yIz74u6KDm6/xSbn/ejkIUsuwWfkTx5x7IxXX0VsVnJusq6p11xlpQP22fP3F1y67k4f50Xd/016YNecrk3aO9DfDsx8oT9r/7a5G+tJh3c2oRzJQp5WTMpvjKO+ViZjaEvt1fulSfL0HdI5Sx75u+NTdwgK9anb65FM6qwMbmIhbpjz4qxx4NdO3ni3T/zqnD9zvMKBo9OOW2/Bknly9MMzziaz8LFx+WWGYCoVU0qJj2C/+8slT018liPYmqus+Lfrb8HKkYckCN9w6Dt58nXbfQ+QHzdZb2hnZyfWHnXg/PNttcmGnIwWGTTok7vvtNmG6/Y4Ft394MNf/Z+PLrP0khOfm0wiu/GOe1o/bJIKT/jJr7485sQeX3Fu8+6NWeHxP/rlb8+76Kdn/WnVFZdfdYXlpkyb9r73bEr96I7vW3ShQT1WQvfOBx7iad2Pv34U56wnn5k0Y2a3dIZDa508ddq5l/31c3vtctO4X99y7m83XOudrdY2b+/AW7gDs18Rnu6/7gkYS4R+h5PByFytuatw48Og5xLsPt6pt/1c5KO+ahlVxyhNwWII7Brr6VX8iOeZVdDDPidwRFp9pRU4+5AUzrnkSn7bD/rk3tttvgknoxvvvHfUaT/m8PW90YcttOACzbE8bjtgnz24Jx57/fHiKzCR76a++NJyQwavvfoqfGqDb7T2miSgHd67+VKLL3rqz8/g496CAwc2IzTJl/bd6zN7fuTnZ5+/zX5fPPZ7P3vosceLibz22b12PvviK3c54FA+eO7/sd040F11063kqeIAsoDzr/g7Y/l2gvyCUupO2245/LP7ciT8/DHf+OEZf+So/MNjj+CAtu/wUdQzzr+YBFo8m0jOveaWO8Z896eb7P6pex8af8QXPrnYwv39xzXZgVWWX46pN93jU3wVsOI2H7nlnvub0dqkvQNv4Q7wCOwNRp+LCJVnDL+EXzOfznn97Mx/CXNWUdf7vzxB+UUedMfcwr3fx9XNi/wVbkUM2g/woYxf/ik3X/HUNRfyzebpp3xt0YUX4oTFg61pt/316Wsvunncr9dadWUi8Hxt5j3XkCzgu71/m+duvIwh5D663z5mGKaLf/4dMtq1f/g5/MgvfBJ9/TVXn/C38yddf8lVZ/xkyi1XEIGj1kOXj7vsV98nCeJAXvjx145m3onXXzL+yj81P2yyqh8ddxT6E//4y+QbL//+mMPv/MuZZ512PAc6RlF/fvxIZimVyBuvsybB6bI8Fn/88C+/dMdVLP62805fabllt93sXQSni3jykQctv8zg5hrKqJ233+rL++3FRE9efSHRuLvmOocstQSrxR+F4EzBkAfGP3rjnXf/9uTjHv3b+WMP+wrffl59c+OAzOLatb0Db90OdNbPZ3qdIb/8/PTffXX2w9f0aq3FfiPgQ56Kqo6qoIiaf9zj/gAAEABJREFUNp6d8Yc/uSVzyWuVcp7q5e+ddR/cvdcISiLLUmUSao1MWbqEbTj12qrqkost2kwWxYe8sMSii2Aq3R6IM0P6sjadCUJ+bHZ7JYTq9TSEzifZlKyMolvI3CDzti6eddJF7GcsE3GW7N+nOZxT7QGjx2758c+POu0nnEl3/8C2TVObtHfgLdyBalZ/watOydUrF3x99oN/79Ot/whxkiLBcCxrouTMeY2MpuQ29SJAnxP0YyBkP9aepuLdQG+5SHA9/dr9N7IDPBm89+I//PaU4zjr/f3Mn9487jcrvmOZNxLwP2xse7n/xh3gjDLn7LNfraY+WU15PM+aMfB9h9iSq7xy0fGz7rl4TkdXeo3ghvrylBXUSUlbqiaczupU5wHIK+HTL+DYV63HlTCBATFD2Oh2jZ3bZ2cxsg3zvAOcTPnsuccO2/E5l2PdPI9vD2jvwOvbASXJ9Bw5+/Fbp//6My/9fD/q9N9+vvPpe/KrL8+4/Nuzx9/Q05V+bxGQm9UTSaSVIOQUt5BQjJmjegCIy/N2EZABBSFz1mIqWKytPJQ5hFDb0N6B9g78B+6A9fLf5U9Lrzb/1l8auN2BXrf9Shq8hljHgA1361jl3b3cYW8RWt3IVI1K4vJzGVYUyxTyXANRX6PykMsfnJEKS+VzIiSwjCy5qSAKhArxmrsemZFSu9X+/t6ZD21f7R1o78B/xA6k+i9Uti5WF1p6vo12n+9dH51vg107n763c/KE+bf45MBtDhAhC0nP0luEVh8yilfPWlykpMyzM5gJ4XJWCkx7Cy3dC2GKUEjBooALDZblNpQ0QFbf1hG+8BChLrcRxqgxIAd2/ZuA0g17b/Dgo/V/8OflGTPPu/I6sDevPrUrrr9t/BP1X/fv06lfw8VX3/zQY0/161Ibb733IWrdaTTTX5lRVZmvF/9y1Q2db95/uag17OzZFV329ca7HqA2Zp7nlgglzryOfM0tIvLvL/rbL8ZdMqPfvzE3r/O2/d92OzBgwX6W1Dnl8c5n7pt/80/Mv8WnhOQjvZV+IzCAPKWiXiQa8YIS/wSd01lkltdIKj4EV0+E5MKe9dEb5Efvk5PW9XriOnLi2jKWOlROoK4lx1PXlAuOaRnVONPxNi+1xO8Nr7rpzpmz/OuSBeaf//1bbAT25tWn9p6N1ll+maX7NM+Fgd/AWX38V4l6jF5n9ZWoPcTfXfDX51+YxheX2717g9T4GrSHz+votoZ9dfYsuq/MnDnz1VepryNaGUKEEqd05x5nzHy1/y0i8kvTX9lrh63+dOW18/oH0twvo+3579+BjgUkzdfXMvjUOejTv+kvlzGWCH2ND71kquz/bNNzBxoKicmfndEhsQUCc1MZiFsDnRLNGy4SXlRoD7EhR+rE3Gydz93FrwS/bPwynH3pP775kzN/fs7Fjz016dRfng2/5rZ7rrv93rE/O+vrPz7jZ3+8cPorM793xnk/PuuCm+564BI/Wz35pyuuPfkXfxzz/d+c/9frifDD3/3513+67LCTf/rA+K7/RggRvvPbP/3q3Esfe3rS/559EYsCOVXxq/rHi/9+/E9/x7wTn5t62m/OJfKDjz1Fl6lZxtPPPn/qr84hJifBG++6nzgs48Sf/wG3+x55/M5/jv/dhX8jJoufNbuT5WH6/pnnP/v81G/9+pzv/Pbco7/9C86PjCqzT39lBlOX+tSzk3u9wX9OeLI17O33PXzHA4+cc+k/iM/AHqO4cdZDHBbJgl96+ZX7H3mcNXz9R2cQp3VnmnHOuexqbuf3F/6tk2cLIn+94fYT//es437w20mTp7LOsY19fnVWJ/dIqJvv/ifzUtnb5s4zC1Oc8LPfX3fHfSXyD353/jW33sM6f3b2RZOen/qLcRdfes0t/5zwBGtonYLXiKVedu2tHMaZ69u/GUdYgrfrf8YOzL9oP+vUAQNFVPoq/Y4tg8pgzRIHNAf/mKn+zWbGg4+dNCU/0e2v+gdVHFvOVrzdSZGg/w2yqu+nY4yK6m65y63uzsOzs6efnTx56rRjvrDv5/b80KXX3rLCsoM3WGu1ex56dHZn55Ybr33op/fi/qZNnz5r9uwd3rPxJuuu2bydj2y72YH77Tpp8pQnnnl2yFKL7/XBrTZYc9VVl+/6GwxTX5y+3JAl995xmwUb/6CyjJ1/vgEf/dDWh3xyj4nPTXlu6gucsHZ73xYLzD/f5MYyBnSkhRYc+KWP77zwIP/HCaxk203XP/JzH1MVlPXfuco+H94WB6Ld/eCExRdZGNM7V1ru7gcfXWTQgvt/9MMf3nrTJyY+25x90AK83vh6vfCqG3u9weWXWao17IZDV+Ne9txhK1bCsB6jNt9grS/t/ZGFBi0w/FN7rrLcMo8+NfHCf9y4zuorrbnKCnc/OB7/5s4045BS11xleQIm460i22y6wS7bb7HM0ks8MOFx7q65z7fe9+BSiy3C7Wyybv3vqF6dNbvs/EZD17jyhtu/uPdOX9lnF3LZmquuwAr332tHkLCrLDfkvocf6+ysyLxPTHxupXcM6TbF7Nk7vneT1Vd8xwPjH3/3+msuMP/8rJl1tut/xg4MGCTzLfx6lsooxr7WyEhWHIdo+bzn2St7L3M607qIKD/yWoXnYq/l4nafyNu4WjuFN9BbVpL7OZ1GhG7Ab0vzr5JChq66wubrr7XPh7drdVp0oUEH7rfbPQ8/euYFV7bqha+wzNJktJ+fc/H671y19b+BseNWm26xwdDvnnHexOf9P7BRnJvItlU5M/bze37oon/cdO1t9zB705rMVJs9JyR5fl2dtVwsvq+/bducnd/w5gim6PUGmw69kl5HJfMVLrjAQA5cJIh1V19p+802+NBWm/Ya4RM7v3/5IUt/+9fncJTjRk7/yxWzZ3eutcoKPZxnzZrd43YWW7hr51WVlfQYUrqrrfAOzq3LDl6yI6UJTz5DOmudwliq+W4uvcSiZO2PfnCr1Vdcrgxs43/GDiywpLzWI7CeN4I/o3qqvfR5Z/BBU0W9qgYaPYPi7lml+VmQfj910RXIh3NRq/ApGAmrPtYVDhZTwSxLrNTPnM88N+UHZ57PR5inn30eN37NOjs7f3nuJWf85UpOE3ww+dtNd977cLf/ACR/4P/ugitfeHE6RyGG9KgzXp3Fr/Q7V15+xquvPvzYU9/8ye/4pUU5/8rrbrnnwfkHDODQ9OyUF37yhws4fzGWHHTZtbf+79kXLz9kKT4V8sno1Vmz+Q1sLmPOj0KXXXfbz8+5iLPSskstwe8mwznvEGrdNVbmm43T/3zFXQ9OWHeNrrsmUTZnJ0fwGe2pZ/3/KrDFhkN7vUFCtYZNZsT/2413ckrm4/PmG/Q5ioGkGHLE2ZdefeUNd/ARGKVZS5zLrr3tjL9ccf/4xxdacAGcsbIkPsz+/ea7jK/F6TcqCe6+Rx771Z8ubX7YJGDZ+aUW90z0m/Mu/+W5l6692ooLxH8mZL4BA2a8OuuSq29eeKEFucEVhizNVj/97PPgnFOsuOxgcugFV914+XW38WyuMWe7/Q/ZgQUHz8MZjXMZ/nN3Z818RR6Bg4xTEaOhKpdIQem/rLR5dzvREBroLcG98fToLd0WByjHQsfu16pbdu936408YN/DP/tRPieutsKyn99rRz6+8QHts3t8cL+PbM8nlxFf3IcHzJutv9ZW71qPuuDA+fHho9wnd/3A/+z8Pj7H7fa+9wxddcWC/I5hfeTxpzdbb60PvXeTaS+9TOIjxfBLy28yn6eoh3xyd45gR+//8S9+bKcjP783Qz76wa0P+p/dDv7Ebru/f8sN1lp1n522w8SMzWXwi0dYFs0CqJAPbLERCg4p2Wf3+BBDSH8oCw6cf9in9mBhh356z6WXWKwoDNlmk/WZmhqzD37XOmsMnM+fp5Jze71B4rSG5ePwwZ/YnZvdZtP1iMnnxB6juAt0Flb2YcuN1znsM3t97ENbk4+KUhxKnF2335wV7rr9Fl/eZ2cU/szjcz3KqC/tx0BWS2UBBOT0xEvz6d12+MbBn2aTic9WNHce54M/4fu2zabrF/9FFlrwq/vuwjbyEfVrX/0km8ktj/jivh0d1jpFWRJbxwZ+arcP8PGfF53gc9a28rbeAU5bCy4tPN3vZ5VY8cGzH5/uJjKViv+5qlqjCII/OxMKH6NIMiXx0O2vrrurcADgqFUjg8oJqxURs9QPxXIc03LPLh9gWuuGH+1v0jfbtsryy9x+/8Oc+Dh8rffOVXjExi/tmzjJ4CUWp77ugBwGSQRLLPq6nj687lnbA9s78FbsAM/CFnqHDBoinL/IXKo+CQhHQceKj6tze3FMykIOIm95FcDzlz878xA8qqBRMhxN/3XwWrKp/7co+vfqbmVqhAY6ZT2c3gJZx3s+L8uug/wvq2SKr+y7C/XjH94WvkD3B/9vfBkcjqivOw65lWd/r3t4e2B7B952O9CxgHD+InMtspIsurKAcBT0eV+riqcqnnuoKIWeikDrD5tSl5Jf6k6fzXZHyBrvJyP2W+c4rPlZrIjkNWrhlaz1IfnQ6D7nahvaO9Degf+sHfiXrLaRqkobmHOPdDbXC9n9O/LuT3Gu4ojlY0hVNAVdjehOXOWq3YJ5EiwEfM/+ss9Padu1vQPtHWjvwNzvAIeyciLzIQrVrtMZ6UdVMShFpP6+k34/dfsj5bPjZJNPyJKriXWIP0qLM1cFZilIgpuz4rz06rLZZ+RLF7XPZf1scNvU3oH2DvTYAVKTIkWaoq3PRiQZDk+5cTpzDzfKQgsuOKiS52e9Gr3XgiFD5QMj5AsXyFH3yMgHoz4ko6gPy2jqIzL6ERkz3uuxE4R63KNCHfOwfPUK+fBx/+LnZa91M217ewfaO/B23wFS04KVLDRoQZFIWtpEFfVvOaVZVPUdg5de/KVZj82c/mKn/xvJpun/HWnfcHsH2jvwNtsBkhKpaYkXX1128NL18ys/l0H5eJmleTprLnut1Vae75GJ106dNHlWf//B76Z/m7R3oL0D7R341+wASYnU1PHIxDVWXUlURVSaqF6MtCae3bzl2mT9tZcZ/9xZEx97cuYrs3OFqV3bO9DegfYO/Nt3gHREUiI1DRn/3IbrrsVZjPOYI2kq18VUPMGZmSe3LEsstsiO66zzjkef//XTjzwy4yVpl/YOvF12oL2O/9c7QDoiKS3z6PPvX3tt0pRK5CwylwUljakaWUw4nXGJwKm7fXD71W4ef/YTD184+elHZ0yXdmnvQHsH2jvwb90BEhHpiKS0yk2P7LzDdqSqOmlxLGNhOatyUJP6751xRkMsdcgSS3xx151Xu+K+wx+69dznHv/nK9M45hVTG9s70N6B9g78K3eA5EMKIhGRjla6/J5P77zTUkssrsKPiqiXgsF4dsYTMxJboP/dDXi11bs3/vK7N9/wknu/dv+tpzx63w3TJk+Y8RJfK0i7tHegvQPtHUT1PW4AABAASURBVGjswFvaknBIOyQfUhCJaN2L795/08232HQDiYdlJLMsFYeznGtEMRVlTfExlE+eRo6z6Hzkfdsctf37Nz3vtitvuOUDt11x3Pg7fzfx0Uuef/r6ac/d0K7tHWjvQHsH3rIdIMmQakg4pB2Sz+U33LLRn24dtu32H9x+K9OkkaBEyVfJjCaRtWhE1SRo5lxWfxiVHAc1gDPatw498DMzFtz43Fv/ccnV37j5mr1u/+t7br1083Zt70B7B9o78JbtAElmz9v/+o2brr7qor9vMO6W/V4eePwhX37PuzcmN2Wp/ExGvsokqrgkq6izSkw5m4mokNXogGoONS69+OL7773H/w4/cMyqG37ykZkfu3z8Xn+6Z69z737r6p4RHOy9jrvb9V4R8XXWu/Ycd/ce42rcI3gD70J/U+ru4zwU+BbV3cbdRWTwLa27jruL+DWecxdk1xrvnIMX5U53OKeJkDdad/Fod4Lt+u/agV3P8f1vIuSN1btieGB5g5179+6XPfyxR2Ycuer63z3kK5/+2O5LxvMy5ZNjqZpok6cqAUTFjJOZ+P9nU1QljmZkuAzzipKLDi656KIf3HrzAz/5seMPP/C7ow7/7pgjvjva8TujDvvO6MN74GmjDjtt5KFz4rdHHfbtEcMbeGjwwJGO3xp56LdGDAfRC357pOsFUYr1WyPdZ048deSwU0cd2gcOP2XksFNGzYEjQxk5/OSRw6mnNBDSUg+Fnzzy0JPd2jueNHJ4//XEEcNPxKeB7jwihgQWa3cchv+JI+YBx44cjn/BXseOPWYY1t5xxHDXRww/4ZhDTuiGw0LpicePcMVxBP7DTqgRQi3KMKxddWR0axzu+sjhxx8z7PiRw785Ytg3wWMOCWzlRTnkmyOGf7O2HvKNEcO+cYwjChz04cfgM6wfjmcZ9abg10cM+/oxh7yt8E25r7JLc7OfxQd/9rxgUb5ZXqmCXa9p43VEGTGMV7m87rwHThh16AkjhzdwePDhY0cdWteRh44a/uUv/M9e22/57sUXXaTiYZknpexZquJQximtoq2qComOJ65OBE9qRupSitBRg1BFVOBCzlPVzEU1jCjlw6uRDFNKSEZRBQpPZimlwluRD7huwNOtBjcGwc3gjGqiUbxvVjBaKNW8MIwmkGhO4Sk5TzFjD7RkKeFi6GY1WrMkQzIKBDRzV0dvlV7iavBaZ66G4pHVCuLYS7XYMfNQ9fAuzp8yWDX0JtKDd6F36v1nYoOitKJL5v499KaP3yFDk+HZk7OY5IMtpeS8RqNAQXOzNXhrx40WAoy5G5iCFzTmdT2kJmeQcu+WVC2lQHPPlLqhubUoySxRzLwxA42SkK3BzUtRnPkVZmtFV3u/NOQ5kBW6wfWkjTWbwc3mAf0uGv598bmMye3gWdC8+NpKTLNW7rZaKXQOTKGA1KABpVNjNKkbeielgj7AqVoTNXgLJriZI2NUU0r1alWN0kS1JOqZR8RUsFBFFIyqxg895UqJjiX1YhQkM43TmUQhbYknO1ERVVFPhX4VHqiqrquEDgoRvHGVSz1A+MQEqj4DYKoqXsURJ7h4kiSzFkVViRnTo0OlFFTxC+8igIpSLmEUDGQEWLjm0ImjyIU7CjrD0V0WjCIMEnrBBcz0aKDKJVwqAgpFacslDKtZDt4NVZWQAqjPRS87FwhRBC7S4N06IqL8CEXjchQpvjUXEeVH6qJ127Np1Vt5T7/cJeDW1fM/AjH5wnkuAeOOcIgXWGqFvgi8tL4oFV+rIjaIiDMFuFq4CJKgKUWywkRVhRKo2p2zHlcyTqzE/X0Qy8mMQAnvoOilBVVUxMc1UEpRUYhfAittYVr3S6tCYbyGqiKFq0Lh3ijFdUXyi+lV1BXQ11m4iGAtvBVFWvVW/yZnHD41isBFRZWrwUWFEorCVaTmEkXFfwQUSuEQqnIJvtFI9JxytXBfiP+eovrNYclIbL0oKSODvv/qenDAjTSiUvKAN2QKsoEq6UA0foSicbm/ovk0xCameI+4Ik7cloVuzrzfIOLPz+hUWHiaxukt+yVGnym8Ch1TZbjSMLsqhCWgQ6kcyDBjxBCodN1q6kgPbxMxgaqJQUAVJVpS6yAvqyXr6Ah0MXWgo8A70C2BSVNKHcmS646Q1NFKNBldRTeczQwrMTtqJWJqBzGJkBSfUIx5UyjJ47szPgmFgY156Xo14nu1ZHQjPlLw5AMNNDPGJrMOZbUoxHdszJuSRnxFT+7jAxFTsmQtnFC+fkPF1IVmXTwF7wh03hiOwPBkxGdsgnMvoMdXgyS1Vp40maIjUpNR1FBKTdowWYKbJlWsyQKdm/MUQUxSCr1YHSW5jzDELBCHBKHiCUblzWLeVXz8vSGWRMNNw19TWLtxS+6jjtxghyjoinHjljQl7UKCJ0tFcd5i7TD397Hh3xHonpYctb67Lm6pi2sLt1T0Htgav0O75nojemrMm1itJUd9nevsUN6rHqF1PUldAX2vLCVm0eRca85Lw34yNqGI7zz7n4Kj+6ugyXgFNQU352bOxTTzPItDE1mozlkoVBHRRjUmo6oa1byngL8X6CRzU4BYMmrScAsfphI0VUfm8NMZWS5zeaol+TELHfFSFEGRgGL1FC0EoCJrFDOhaPEsSB//npgx5qqgBBfJQmblFtGFhVT0BaXojjl8WJUTtyLWNeIDZGvfLsbSqTwy34AQ35FRKIyrcpklkJv3WQqX5rw449lEZoYTk9mdR6eiL0T2puZF9/joxGxF7oUvXnyFERkrweZAj8AqiIm1CyM+EVypwieQ+L4DWNFqJDqrallD9nvEjbka6FZmYWeYBd1RuHvfMbjPEo2/FuhEnhOZhzVk4Y4I657OBV6vim5X9Rl9glCYEbcafTKf1xXnvCsIPVcYAfFkUoY0ETJP9Y2MnaeJ/u3OzTstZJ7WU4Y0kFeLd0hU/43LvAqY+kGfi/cb47pQ4nV3VJKMeOGgxVtEGl0fxRDeUai8Y90lQ6uK9yR9r7wLkHI8NctVJ9wozGJqjlYjDdOYkgGRsJlJ4WLu6lwphqxIAk8JtfB+EEeSeDc0C8XRF5NSoBFNUzc08y7gJtPumEJwVEsevwstFNOEtQsZjxJoambJKJqct6AGV3eICcwxqaGkGpmMwTUyPGFTa8Hk3CwZA4A+0CMnM9N6xuTumhLuBc2SmRK5oNIwS0Gc3MqQxKXorri34tDC3ZoMg+LImnsghmTmmCyQoZpSFxIZQ0ruZN3QJzI8maFbpYOpIBPCuYdAWmsozr3jcbt4q7UbN+K5uwRpIuQ1qsbb1RyV4sN9OgneRA+i7oPSN9cWa+GBDKQyRQ+kSy06hDonL0qfWH7dYhabkxelC9lkVtgDleD1WBOI17hTJ02lJj7WfKvxhEfVQBeVgnUOtFBAXkzQ/UMBWrmlhAMKscDERb+udBTdqyvJfFoG8GZ0zXvmBB1XqCP3kcl/ZEQyG0g6BIP7H+wqXlRFRUByIw2ErmRmm9NHKcRmErAPLgw39ksFFHFUcXTOGLiqaFdp5YqF21BV5gj01i/uR0zFDFVBWtDd8RcUU3GrBQqKqoqaAo60hWujCASpBQ1fo68M8JZLiEzPUUWY0UJREK/ABhfcFbFRRbWpiCpjVQUUVW0YpIVrBLJAAB/mAuGmGtxUxIKDCmnU7lyYBQfQdRhupqLaUAjmPsUq3nCpmoZTQVMRVTXQVEUpFgihooFUtt9UDKZKfAiVXs2j0+DeCS7gnD5FEfE46tEsuKm4f3BVcStcKaKm2uTaZ3EvjKYClmqlCeydc2NhdSg8/Fi5t34p03vrl6/KW7/65K3+vXCfqRFzTu6WojpafS+iPhuKcfWs4kKA4s7Km3sVXEAV0EQiSiA+rqsroo6Ebiqm+CsFxcLaBxeKqnDECyCZKC+nkGU46inGruofbujx3SW5iXMZvAWhpaqaCiGzMXcWf01EnIOKSZBVXDdVdEdaUTxN+HSbwBClRgz6GjyZpfBxpAM3MzCZNWbR5JKaMRPcMTgjaekWhCjFhzFSvCVP+qAmxxEeaIyWpI6WzBhhgckKT6opaTKl75icwL06F8QU8Quac3NUPLFaMsERNAoiMtjgSTUZs7PqQFzVPQxMiFgVH2zJxIyuwC24Y3IlFd18LvMiDqb8uI+hE8TRjNGhEx8zGDUZDnU15wIapqTGq2CBCQcJTyFK6lqDmfPA5LMkA/EBuZcazYOYI4Sqridj/ZrU0U2qVqpBpJUnxU3cn/UED3+pMZmhmzomczTWKClZUknBmS+ZmqEo3II7JrM+eDIWhq+mFGiOljSZmrXyoiD2XbXF9GZxX0ZX2BTdwNa1OU+x5mSWTFNyNC8KM1OnYKLtyZMKPslcTyo+FjSDo5uKUUDGgqaupEDVFAqYgierdVOiNTgiFUXVHcwxwRUHZkgmydydCcwsReMISW7wC1+LoI5ijsgqngaE53QkNv8MWglpkkcifGYlA6p4qdHzJbmzRXGKP41acWIYvblCHrYQkQcvBX2Mz8sxMcdCWFB2K8tA74nuwoESp8jjuQU9Eu6YiOyIkQa5IMmd0YHER+5ChsVcPKRy3wq/zCw0LShwHJnRsRKIO8RkhaMzHHQdKVgrj3snMgaQHWtBtMa80HouX5X4Opmlwd3q0zR0+t2rG0PJXWuIuQiCwspZU9h8FiyhwKENDDN7xZAGMrTHSgiDwnTdMKZ2kxNpkAiIK23RCUfwGnFjcmxg4cJKuuzcO5ZeESf0GrPvcPDMLL7yzNp8De4zbzwiMNRHtfBQmMLjCx3uwJE9RCn6W4Gt8cv7JDNt5u5a1gZFiVW9zvstEUBhJ4U7Ik7ZeW6qyYvShb4HvhIU3EBWCzrPYWNJ1MIdYwLCO8dAMgK71eJBUKko+LlALJpYVVlepnGz23m1yWwkJJWI55nRGSmPUxeZqsZaxxc3R67mKC1WwSQlzmsg/sQVEiv+TCUG14LqkbkEXUN3JDQ8UNXU4xdUda4aKI4eWa0gfcHddXNuIHEcOY4KqdwsUHHLyIpVtQVFWRfWrqrqXLWJPky0RmNA1+z00LshjswhcYcg0btQ6WHvQlW4qTbQjcxrSKI1qrqiGrMUFOzms6gyUlR7InZX3dt9BTs+PRAnVV4jLA1Uxcd1sUAFmYeOz+73xACFKyNFw6rKGFGK63gHc866XRda5RLVnhgRVRVvYRIhkrlPD66h18gYzK7AGFVQlZmtP1QKPoGA+ytjoU0UdUW0F0QTXJWV0vSFrL1Xa196zzishFl6oEdUbo2rRnxUe9wLpobi2wMvSjf0yP5a07qTv5gex1ifhOAovre10sWVtQpjUVgQaOo+qn0hK1BVUJR5ktARYqhZ4jKmM0uGYuZKUjWFgCnBRI0iphSx8DSjZ2RIUQ5mZDwyXFXzLBqpEoQFFh/JFSm6ePIR1/0wiGrxCZTAVqXwLj2r4uPItMrJjzj+0A4PPiaZBngxAAAQAElEQVTXuvCdWBbBKhT/NkywC0V9bR5BpAVVVERUQY8gEjyL8gOqiiND3erzRnysGc+YJTj3JlyKo0gMdZRmwYC9BZ0KMRnQxNiTGKzawlHwzIzg8jtSV7hq7ivxb7WzshsiBWnR6QhFFXDODYhHFhTMLajZdRXFU1sRXeiriKi2oNQ8wrTwsie+kqyO4ogTOqg+C1IsJN5GGcUN4gVOE9HQoRLckX6Ti2rwVpQeSlYNpacuyqqyKD8iNS8Kexh6zj42w2uddxcvU2C5iwYKhfcDKOqA7k3wHCjmgluzFHS9cBFXRAq6LlL4G8Geccpczfdmcz0iZZbu/ng3dXYAXhS4cnd0uN8sLfuWax7WmrNvjATxFBURdJHuKEX39XSz1nqxtmDRW5DfCSF+xXsm++ejigKphJK50AV3TGB8j4mCg3Aew5A5x0vhJrECJbWJivMcXEQUGwoTEVK9j5NiJQsGkiWbCt5qWhCfJi9KwS49qTAS5NOzqQZKoHN08m0yQ0mYk1kX+syWsKi63dGthl640E0lPjcQQ7GZrw0LVkb7XEmNKKbaHb2HR3KdcVh7oLiVyAlPUQurI5zJahSo+iJMBa7qXNURX8YYPKmBXGBSVRNHBU1Fk5oK3Bt0K7xGTTGBo4iZYu1CE+fECx/AnHMxr1LoNjAc8VDXEihcphKza6D4/WqNrfMGx5Phjok28bopsyRTrI7mBcVQWKaj+9BNYUmuMI8qVm1FCcUxMbu5jznWCtbk60Qro4QZrVYKV2bCpxXxSWaBCkJb0PXuSvHppns0I3KMS/QKB5vcYKn2qXVfrZnvpGnwXvRibRmLPz08mQts8lpHNaIlsxptDkVD6akndT0xLjGnEDexww1OMB/VtZNqNVcYngXNeEUkUEH3MW2gexqzuOI8BU/mnskoTvA2nPzy6dWSsBTj3ZvwsJQwqpmIkGeNjiUV0ZTMVEwpYKIxxUlRRPyP1Sz8ZFF+nHEAIz8L2Y90yFFMXPcrk9izisCbyEg8W5W55Iwilv85qcwWsxNeyp8YRKXjHOY+RXcUvztUOCyQVpGFmYUCJzIkC0rmIpc3EEsmLjELMje8BYVoUuJn59qFeKE4oqE74irEhvZEj89Nqs/ImFhDF3crs2Bl6JwoZRWMgzUR1y7uHX9Fes7LSJ9LfC5moFswRuZA6UJsfv94w0Dce0PkHH8GxozcVhaXumNG56kOBsfMuwd7RkUJlMCiFO4YrxerFVYF5376QiGQ4iV4a82z+xNS0btzaVEEb5xcYQ3ccUFU+OtCj9gY2wsnvtQX87LeRk+avKcurf65l5ixzlYdJ95H84wMYOUxG2vwPSEoSr2TEsqc++mKsHo8QfHV4Ck+SqQ7hh1/ETzUPV3xGeHC/FxNZA1YOIrh45x3L33/FJgr5uJNpVWuOiupqkzpnJ2zcxFiZ5dCF5hkE9KbqKhiBRUf8TlFhXynZEI1EU6bqnSETIjkiCiqMLVaKTy0bsr/sXcegFYU1xs/Z/bROyjYQAELoIJiLxixaxRrYouamMTEroBJDMUkNrAbU0yMSdRYo2JXLNgVlGLvCBZEpCNVeDv/35nZe999/WFU0P/Om/32m2/OnJk9975zZ/dSatTJvgmjk8TsybiBO5dY9tXEJUnkiUNxrojOKZx2hjQdRU2v4AltZxYK4Nkl9MEdnyguScAEhXN1dGaWOH5YhXPYOBdQnQuKgyTOCqh0OZSENhxmSAMHzOISxpsSuAuoAU1PGMzZmZ440xMXuaMkLnG0YEVMaBTWYNy5hJP5d66IECpdYGLWCagY2LycErgr8sQ5a7sEJXFwG0HEEqcBM8U5l/DjNEmcqwldotQkSczKkCYcjNU5DJwrQWVG5/AllVGUz2en1RBdzZIRjDIbh02SoEuSmH1lFOcqdOeKK1fWkCTOJfSWogSFXvRSrFl3rkJ3rsLeRd05l6ijJJoEdGZT5I7exDnQ1ayLy/yoS+CGXF3gAkae2Njiakv0pMI+YXiCjQujCuhQsNEkcS7ygGooUVHj1aLqmJKxoDLSbJzFHxleGaPuNBGzZEmcDFUNnStiErnZq0tMT0BEGysOPVunU1XBm1BcaLhwCiBO+Ym7M/IXyZKsZ5sxxZrDkynRyZop+VJUTSNVlqcpiZBkST+EcaBlW9rY07bcyglbWMCadVwLidbG2ryCX7ihBG45ODWFpaTCLMxgNRVvWbmIaPBYI2de7821iXjwJguDjIe5SrnNiE90zyzMDdIGPT6Yl1HWab3myGcFTWiHCsA9ltYpUvAfrzHoZo4u7HKyXvMu2DPKOpkTN97mDYqNZXY0kBcFjJzgQkK1cXgzTrcILks4vlDwK4g2LycP9yWcMSjRxjOvrQfN5kDnVFxPNc6M3nrpiPMSycDNh6eBHxxbDQoABz2epbAGPoHTwNMwuwTEDTbYBURDoMX6jcPQLT4m0LITCiygsCop+rGF0BNswp8pl9ROAX1AARkIMo0htqmFFi5cGjwqRU6zhPsSHuw9RThMtxNcSrgv4Q3WQ7RL5zUhZYWeK+WCzGe2Zo8Za7b3DEHy2HAKmHHFPozKrhduY0IcitwUL4YS0Lj5McW4jWWiWuwZw6g09Kbh9Q3jGJSiR+5Zs0/5ERbOktK0vNz78vKU9EKTpCO2ZZMU1QtX5NmNismp50dSj7uAmAncMlrIUzY7dvCAqiHziapm+zUmtDGWA1VDr43lECR2eeLU+kEnDePOq0ribCyooonCDY07n6iiO3Q+KUAV4yoJ3Anp2TnjoBOaElFFnIbqREWcwydNryLBv6GqmOeAKoqeiOmMCFxDr+I5ERAPik/jKtjAnSq9KgJGRUWc9RqqSOKMgCpwPAieVQyTgBgzS6KxV4kbOhi8+STOG2YxhVihqBhX4YfhxjlUmYUzqCI4dAFVjCegE4zpTSIHFf/02qqYEd2ihHPnbbjLdPzQqzYviq1QVYOlITo+naCLcYaLMjzjqvRqwESMOxF6k4BVODGkC2Sd+MfGuAizG1dmN84oFJXAmY6YBIyjVISxZoMuiqVxe0+yqsyDSuRK5NXW5l3AwHEbdAHFGfKewV4SOBMFXB25vV6FdYZrJ/Il11V6jQVu1+LtGmuOjxBJFUMXULl8ZglovPIrYjYFJTF7omdjS15TYijWpFcFm8jBRDRx6sDEdmbsy9QBToVKolKx4ii8UNaB7kQ0jBHlR0xRJ6oqnDwZjtfNMmY8DE0L+Y9eH3rFB92ynrHUigd8uXiSJBmXfBcwZlzLx6W6ORIg82M+7DAPXkD0NNiDjAVDTfGWWsHGB+5J22kK96lPjXvO5pn5vZUiT70Xn0aEwek2DHOlnh8bX+G5MG9UCv12ltR7s/Xe1gB4FDS8MQHch7m8rQHZo8O9p9N4yihPX+X4pCm6T81v8JAa2hq8xz71PmBq4zwt6w3cMxeqcWTvU4oXAx/mTQ0Z61OPb59iQaUfPY16RObyFOzNo9CHMYieonsUb7Eq5SkmHl82wk6SpnQbFtYTeGroU59agcOYnaZhnL2wNmG8cY+NUes1bnrgnvXQwYwg6wGLPI7xZhdmwTbjHpuUNaGk9fBgmaapT1MO88N60jRyRBSPVxOy6zIl2KwOeuk6I09TzhysLqXYasM1+ojEME3r5mmwwYhXAe5TfjycwxA/YXw1ntmihxEYmeJTsgQ7LBAZjLzc9LQ8tRxSTmFvlmKZYuMl9RQhyGy6UiueXRuvJ5SONPXGgVRS2sFeUocxyU4ss5HxXOCKSA0csD5V8qgThatJ5M5YUSLh45Bcx50Dt6ZU40IitGo8ZZtolQs1G59x5qfJLS86yEAQe+f58BDHVALJKr2aBu5xa7tGzNRDvOmIlSpufNYrHv82HIPAbV4x/xgxI12g8bTCPzaIoPWisySPQ6sV83oJ3EQzw0ACh7AwwWVYKqJ4GlwsZlR48MxwM3IpZh5XQTcPcIxdigG6XSNKrOhU44KlqJiNitmYjli10oUZUwj+sWEBoPEwb+S2HrscZG/r8Z4piBuIARg5Zva60Cs2b8ZpUpmXC5cQQ7jYvMRQIdY03aaGp4F7lsQlRBSKpmGdHhp0DHx47dKsyRpYTKxw4qPeZjEuENYc1u/NnqmDXsJr0vEWL8quRexdAaGy1KCzVFsVCg7BoCNaNS5G6DLuS/jXqFdfD8tm6qJOk2jwAkHQswjYlYrFk1ezgkvWa7HK4hNiaPE0J+IlxNnMsGFgrMZTE2katxeCpjA4vI6M9cbxqeElBr0nURComDcSJy6kHae2S1MnsTrT4ZokEriqGgkDleJoFobgUHjVVER5EWxGFsy8tiQGkBFpe3Q0Vuo9uZBPJ0RrpT7kQ/qwL0VV4SegVENV+piwMjoRMcVbry9FsW8fvCcizMmiMkQWi5gpYkUNRAUHolKoGohhPIRuG2KtMAuWgUtAmwu3Fm01zrTGMbK5whpUrYV5lYpcRSk0FYKjCsSF4gTPca6Mc/IqavNC6WUQ7QzRmZ8OQ+WMDmIPZlxMFyuR2JUKolcNaFyCfwZ4pcksGWrQYy8o9HKooIuaTZhXizzouEHxdFsvfeoBeHDO1dHjIw+6eBEISrFqxSLpUnQJK9GASFERhSqIDooaL0XVglJ4LzFA0Iq6IiiKmOIDRoUVokclQx9sMpTQG5Buu7oSDjUls8dh0ZvwnvEalcgVS+y/lG5jC96M4wdvNayHDvsdUbFMZZY+rE0CRh5QohIw8kqoylgRKcGgqAalGqqW6IVRHi2+jpWQaAjm9AqvvmUPsoq3zRUrtybbcE6+nIwjdIn3bN1IPggidHrxIffwjaeqongJCuhD1nJiska0Ti/qnCopM/EeW9zo7PnzH3ryucuuvXHw+ZefOPTCX/72gl8OueArwgtP/O0FJw4JOLQKXshcJ/32wpOGXnjSkBrxgpOHXHjyUMOTChgsLzhpyAUoYOi9oIAXnjwEXglPGXLBKUMvPGXIhRkOC7yApw658NRhF4KngJgNC2YlaL1Dg01VHHGqKSNOG3rhacOq4IigVODpQ0ecPmwEiG44fMRpQ0ecVoKnD7deQ8zgNeEZw0ecMYw60nB4wGFVcOSZwwp1KF3Gzdj4iDNK8MyhdI2owOHwEWcOG3Hm0BFnRh5w4PARA4eOiHjm8JEDh9E70syGV8WBoTfgyIbgoOEjBw0b+SVw8LCRg4ePrIL4QcnwnIsqPEd+Tpgr46G3wBk1+JyLakUmKvZW4hfZqKgMD5xVYRl5rfrIMKqafTY29FbnUSlBu7rC+mvmw0uul8UMszhzjVl8htcQPfMTR60MDhzWoNc6vB9G8P4pfecgDh151V/+dfOjz4ybu2CBOMt+QpYS55JE1ZERyU6mi2ao2Ghxd0ZCTMmG5DlymE+tSefMuXOuvvG/x154ydC3xv+zc3Lzrl1uOWCTWwb0WD1qz7AMw1sHZHjrgB63Duj5ldfbgk/wa63/HdATXm90jQAAEABJREFU/+CXrwf2rBgbeY14YM/bg377gb0gDUDM6qp3mJ9e4GpY7zzQFlZESF6LEeD1gldBmqu4Duh5e//1/9O57Ny3xv9ixGXX3nT7zLlzQ2ry7NeEwi2ij9yzX7PElZLf0mx35slhpDcqe9XInT7x/PhTRlzx92TehIP77rTXTkO33un2Lfo/13evsXnNI5BHII/A1xYBkswdW/Qfus3Ou+zT76VDtrqu0ee/uuiPz74wwTZiPCnzYoRklXFRVRQO2515blOxIPuR4eAQSe9++PHzHn5w3AG9d9t+60e23P2crr2P7LT+3u3X3r71GtvlNY9AHoE8Al9bBEgypBoSDmmH5LPH9luPH9DnokcfeeCRJy1Z8aSNp2nsxixrIQTmwYpvNnmAKGQ5KjZPjp1w1fPPTNyz5/AefQev33O71h02aNqyVdJI6il5dx6BPAJ5BL6aCJBwSDskH1IQiejlPXtePe7ZZ1+Y6HlyxgM0vhERUUmErZlXK+LIYnajmfK9gKU379P0szlz/nLnPe/273HxhlsdvEbnjZu1LlMneckjkEcgj8A3HgGSDymIREQ6er9/z3+Mum/WnDkp33jywCwlXZVzP2lpyzJY6lSF9GboYB5+xwOPvrtl58PW675fh7XXb9pC8pJHII9AHoFVGgESEemIpPRe3y53PfhY3F/Zn4Ulc2m2NxO13Rl3onxPICQ672X23Ln3vvzyJxt0OG7tbt2atlyll5BP/hVEIHeRR+C7EQHSEUnp0w06PPTqK7Pnzvc+ZCy1m0u+vCR3eV8e7yKzfRnfAYyb9Ppn3Tse3qnLuk2asc37bgQiv4o8AnkEvu0RIB2RlEhNM7t1mvDS68JmzPKWqpLNVG3DZrszTxYjt/H4DHxr8pSlXTvu2LZjh0ZNJC95BPII5BFYbSJAUiI1fdF1zXcmT7XNGXeVYYtmt5bZszNHbhNVqpLXps2YObdloy5NWrRK8u8xV5uXMV/IdykC+bV82QiQlEhNc1o1nv7ZTLIVWUvUKcz+IpOIJi4lw1lys2TH7mzR4sWLnLRv1Fjykkcgj0AegdUsAqSmxU4WLVos3vvUe19O1vJpQJ6dKXeenq2ZhMNsvEiKiTSgfPixv/kOP/SC9IQz0p+cYvXHJ6c/Prn8uJOsHnti+bEnrjjml1Z/9IsV1KNPWHH0CcsRzzqn/Lpb/NQPGzBHbpJHII9AHoEsAqQmEpSQqIR9mW3MrIM9mrXs2Rk9Pk2tksSo1t2Q49ZR/vcXyaNPyifTdUW5jQjzGOGI3AsT8uUDgt3TckJZsUI+/iR96NEVv/lD+Q23ouU1j0AegTwCDY8AacqTsVKAfZn3HubTNPy5M7UiSuIRH7B+t/7P/0hHj8FBqHhLfeo9B5imnIUj9b4+LL939IpL/lT/fLlFHoE8AnkEChGIaQpU9mUqqkqPU2fPztKUxCMh4RnSUXf1t46SCS+bg8zOh7MPinHuX6NSRJRiLxvFqEclfWFC+XU3ByWHPAJ5BPII1B+BkKx8wNS2TJbBfCqF3VnBAcmIWmjVeOZ5GfsyXAVn5L9AfSQgtagUeXUldoHUFfc+5Kd8UONsX6vIqh577sWHnxlbXp5+rRPlzvMI/H+PwNJlMne+zJgl0z6Vj6cbwlHQv0xoSFO2HVLbnXE4Hp05SkhK/F6TVQLW59o/M87cVDarrgjTYeOFrirPzpCrK+WPP4NeWyXdjHrkiX5H/px66wOPLOfpW22mK6PPnDP35N9fdNLvLoK89f7U08695OW33l0ZB6uVbb6YPAKrZQSWLLUsNmuO8HXk8uUkGlsleQeOgk5ew8bUhh6eWzx2Zd6naeq9T/kBy8tJamI3nmoIodbt0r/+pg8ufIoDasr6PIo1U5+ihAqDe8/ZjhJu35t6X0VPX36tjnmvuuHWI84Y8vq7Uya98c5pf7hk/Gtv1mHc8K4127e77Y8X3vnnizqt0X7yh9P+dsuojz/9rOHDc8s8AnkE6okA+6/Zc4XMVYcdvdhgWYdNlS7LV0ohf0Wk3zZplojISKkPxJM06airzphJL9srHoeRIuHZRqxwQqRLrYPD0ignmoh0wUNFD2cDe+jmp39qtKZj4eIl9z/xbOuWLZ686er5Ex9/9f6bd9hicxb63wcfW6/f95tt3m/QiCsXhew+Y9bsowcNQ0F/8oWJf7/1ru67Hzzm+RfZfO12zIk/Gjx85px54A9OO3vIZX9dc7u9Hnp67BnnXXrRNdc/+NTzxww+h8l/eva5g0dcicPtDvvJu1M/Wrrsi5/99rz9Txg4Z/4CevOaRyCPwEpEgCTF/quBA7DEvoHGXtI0JQlkR+A8PCO7sS9T4YfqhFRXt0O/fDlevDc/7MuoVXhQPEXbtdHO60JCrW5PEqUnQ9zWNm+LZk1799howcJF2x76kxOGXTBvwUIsn534yvFnn7vVZj1PP+6Iq2+647YHHp3/+cLDTv3N7Q+N2bvf9kcP2HeLnhsvWLiQ3daSZV9wr/rh9BkzZs1J0xS8d8zT/7jtriP332v9ddeKers2rTfp1gW3fTfrselG3ftvt9VLb77zxAsTps+c9fCzL3Tvsl671q3ozWsegZWKwP9rY3ZbYZOxEkHAnlENGKB8iZmoimgoQhqDOHVsx0gqbK0CIRdJ3cVc1G5Bb9bZupU78tDk5J8lW29ZIdLnOQrVi3X50IwYaBVgnX84/ReDf/qjJo0b/efuB/sMOPq+x5/h+T1P0Fj/9M9m8QTwmQkvvfbu+xNff/vgPXe9/aoRI886pU2rWv85kGZNmtz798v/OGxw+9at41ybb9z90L13g5989A9+cuj+2/butUnX9Zll/KtvsrP7/q47sgZ685pHII9AgyJAYmK31SDTykaMYmxlrXqLlBU2VN6nYZ9kYJz9mKgllQyFxFZ9dIniKdlgmPeptxRYWUF0O+/gtuytXdbTPb4nnTqhZNV7iGBP9R7uKYGXTFKVNm3c+LwzfzntmQfO/uWP0zTla4HPuexgxd7tmAP33XmrLZYtW1aepk0aN06cCz21Qod2bbqs3anWbpGOHdr/YN/deU53++jHNunape+mPeswzrvyCOQRqBqBcAtVVWxguwFjlXRlaUspItYSGLsb0hf7MiGNWV4itUDrrLaN8tmzs8DxwOhSdL17ue/tKC1bMIfrtYnbbWc8hmdnnL0CwdpGBY5S6LV2lYNd2Kl/uLjvQcf8+45735lify9qg3XX/t62fTFLfXr0gftu22fTAbv327jr+t06r3vPY09ded0tV11/67QZn8UN2j9vv2f4lX/7aPoM7OuojcrK6B318OMvvPIG5JC9+i9fvvzuR5/qv/3Wa7Rrg5LXPAJ5BBoUgaXL6nn2X7cXvhnAQ502PuyDSCAhd5G5Ak3FiYqqiEoFSqVSpeErlbTYEhyHhrRv63b/nq67jn//g/TRJ0XU7bit22Gb0FkBRfuiJLWUFSvKN1y/85SPPxl44RV3jB7Dvum0Yw///q47/e7Unz/8zDie8Z9x/mUTXntzvbU6Xv7bM8vKkl9f/KezL/3LuJdf33vn7RnIPeP7H03j/rEW95nMXWqvDbvecPeD5//1n0uXfcHzMrIkfQf0t1wMyWsegTwCDYpAnXeLfvHiJX++pvzdyXW5qtODDVQSlpCz1BmSxuDcldk/D0QiYreU2m2jJRmpp/jQb6giVBKjWImKd7vu7Hpv6hcuLB/zVPn9o9M339YO7ZM9d9X11mULVsneM7ywWfPmosajWdMmA48/au74xz559gG+2fzPJX9g28Vm6je/OG7BpMenP/cgXXvtvD1j+RLgs7GjTZnwGNurLuus9dI9/6H5yL//9PTN14Cd1mgPvvfonWut2QF7EI7SsnkzjMffeT33szx6SxL3ztQP3/vg420278m3DVjmNY9AHoGGRuCL5XVZlqc+TZf8/d8r3nqnVrO6PTCMTZDlK+8D8vCKzJWmltaETOdJaZZpVBTbOiuWKeYeR54SOO6gKMqt5fbbSJPG6cuvlT/5TPrhtPLHn/bz5rtNNnL9dsAgVrMPA2PTELd1TquqHdq2adqk0r9cRFJr36Y1XcWhcBT0qEBoIsZm3UgW476SIdxv7vCD4z+ZMXP4KT8jddY9Ku/NI5BHoFIEavxT7itWkAfSufP8ihVNf3CQ67zu0mtvWPFKLX/atEYPleZQcaSqShh2Z2yuvFSUUl6h1sBwxp6ODp6jgZH76TP8h9Nouo26u8031TatXZ/NtW0brsQX/jkgG4hF5VqjWNnkm2tttWmPyY+NYlu3507bfXOzNnym3DKPwOocgZq2JuUffLTo3IsWDTvP6nkXl7/znl+yZOkNt5RPnlLDpdTkobKZF/sjXilIAgsoYXemoipSqMalrhJ2Unix3Vg44i7Ls9kTn6az5pSPfozMpWt1TPbqz0M0t+2W8sXy8iefLX9mbHX7oNhYSF2zfrN97AG5DwW/2Wnz2fIIfCciUFMScR3XbHL4IU2P/mGsSdf1JUka775r0m2DGq65Jg+VzDDgqRlfZYKB8xAt7M5IJOzISIcBOVcaVleDx17UaGGDYWRF/+bb5Y895ecv4Ala2Q8GaPPm5S+/ikISLTw7wxB7MGbRyCOamB95BPIIfIsjEP6QQJX1a6uWjbbdqtEO2zbapm/5+1PLP57W5OD9G++9u5CMqpjSrMkDckUlT/HULKWQV9hd2VYqLfeW3IQkxFGBFaOqM+8Z7L3PNmXhnAaMmvEVPDV7/kUpL5fGjdOPppU/NCb99FOzSStsfJr61HsOMPDqc+VKHoFvSQTyZZZEoHFd/81IOmdu+eQpTQ78fuPv7Sw15jI81emBfmGkPTtjeyYijh+17ZmEvxVAnkOM26OINGupWtSxpIrw7MzEwOMTNFm8dMXDY9I33/Vz56ePPpFOeiXTxUpmb5TMGk/1ILn4nakfPz3h1UWLl86Z//lN9z9ezwCRz+bMGzPupdff+2BxfX+GpYqrKR9/+snM2Yh4oELmLlj4/kfTIVQS76Q3J4975a1lX9T59Q2moS5Z9sXN9z8+e/7nd495fmVXEhx8Q7B8RTlL/R8nW7RkKfGp1wmho9Zrlht8WyPQrGkdK+eus8XZg+rKZQyu0wP9JA6fknXYDgn3fDw78ySx1NKaIAiF3ljhtVeGebILgzmVVEEsNOF+yofL/3vXihtuWf7w48goVEix0qTGJsRq7ZP+d/STb73/0Xqd1nztvalpmi6u94+liLw95aOly5a3atHs77c9MHVaPX+GtnTm519+85rbHuD3DQ9Uusa+/ObNDzy+cPESfuevuf0BSOuWzd/9wL7xoLfuygXyS960ceM9dtiyWZPV9//6e+/DaaOfGV/3tdTbS+Ke04C/q//Cq29R6/WWG3xbI9C0iTRqVNfiG5VJbfsyhjEWD5A6qiq7MU00bKaUNBarszzCs3i2T95SHmepv2CKkaGlRmg2zOIwddQAABAASURBVBS8mIDfV14vf+xJsdRToddhH0bVDLPnff79723Xdb21tuvdI1qQWf496uGR19527R0PzZn/+eXX33nu1TfePvrpWXPnX/rvO/5y872fL1rSqkXTLmt3XK/TGrPmzS8afDZ73hXXj7r61vvf+/ATdkxY3vLAE+V8KRL9ijRv2qRDuzaksCiQ1+bO/3zTDTdg1zbl4+mtmjfrt9VmPbt12WyjikeYbLsu/ud/r7xh1O/+fMMbkz9kc3fB329mbfc8zrcf5mbJsmX8qmN2+8NPn/+3m1jzI89NZAqc//PO0eRHuljhVTfePeyP1/3ppnvO+dP1Y8a9NH/hIsR4XRhcdt0dTHH25f9kJY+Pe2nkP279/Z9v4HLYt/71lvtYAPXdDz659F+3n/+3m56d9PrzL72BwnAUYoJITLjkEdfcittZ8xYUHb47ddrTE157ZuJrz02yvxFhKxbB+G+33T/8T9eDfJZwOef+9cZ3pk7jovDG+kc9+uwFf7+ZKVgDQ979YNor70y5+YEn+PDAgMA+8NQLbKh5Rf5x+4Osn6sj7E+Pf/XZiW88+PSLb0/5mFF5/W5GoHWtf2O6/uttwFh2CbYjK+fBF4nGUNhgpZbTyHASDiayM6c6qo/bs4ASsIoinpKGnqooqGmlXhTW4YNex6SNypIk4TFfhcmkN99bsGjxNpttDLJf+8nBe++4Ra8PP/2MO6aWzZueeMQBrVo041f0yhvuIln02aR70WDOgs9xddDuO2zYZR3yxSZd1zt0r36Jq+R8hz49uUudv3Ax83306Uwv0nmtNclTZL3mzWreSLdr3fLnP9jv8H2/x8IeH/fywXvsNPjHh02bwe9yxb8sNH3m7NnzFvz2hKN+eug+fXp0m/zR9I9nzGrdonnL5s2YiESJvmXP7rtt2wdX3N62admiuGzSAZbo++2yzcczZn5vmz4DdtthrTXbvz31I3LN97be/Pu7bNu314ZPvvhy57U79unRnfWvKC/fvk+PEw/fv2WLZgOPO7TrumtNeP0d9pvb9t6EfeJHn35WdPjp7Dnk6J37brbjlr1YSawE8JgBe3Rs3/bA3XZ84OkXNt1w/U26dn7t3SkrVqzYd+ete3bvwp342T8/co8d+pI6GbLR+uv23rjrkfvt2rJ5U2p8CdBj/WL5iuUrVuy1Y9+dt9p8p7699u23zSZd14tdOX4HI8CvSYvmX+a6GMXY+kZannKi6tiggSGNKehEhL64PxMKv7tg7VXLynj4Zf3R0jPcdnylSrZXM8m6FOI5sn0bw6so2OM2WNQMy1eU8/tcpY8Uw+/PcQfuOXfBonvGPLf5Rl35/cQmcU5tAuFX9PRjDvrxwXt9OP2zogH7tZ8dug+7gzHjXjrmgD3W67Tm5dfdQcpjYLE2b9Zk6802nvD6u2TRt6Z8hDs2gNNnzqEJKeczoWhamZCWvffLV6zg7rJyj7XQG5WVGRNZs11bsvCLr7296UbrRyWiY/FOm4ZHoeS74rJjbxH/c99jK1aU9+jaGaVHt87smLhHJovhv2e3ztv37nHkfv3pijVxXIGQiPkoWbN9G4L2g737dVtvndhbB45+ZvxOW25Kpib9bbbh+rtt12efftvEFS5fvqJJ40YxztU9JM5mLNXbtmpx6tEHvT75g5vuH1Oq5/w7G4F2baQBianS5WPPqEpSbQ0vbMfIKilZBU4Wshs/pyQ0DQ0NIyMGWjN06shvrMQnZfx+0CggYvx9hlAjp59KkwoprUUFouusVfN0Qd2iZ/eLrr2N+5p/jRrNzDNmzyOX8dv+6NhJT41/NU3LUchQs+ZWbIXCuAooGrz27tS7HnuOzUKTxo1vffCJt6Z81LJ5s8dfePn6ux+psBbZqtdGJEduBj+dOeeA/tvvuWPf9dZaY+myL9q3aTXiH7dwk3XfE+OYsTiKZHfLA4+Peuy5bTffhH0iN5X/uONBckf7tm3mLlg4a858nJM6y8vLuYQb7xuzeOnSLXpuyMZq/XXq+rc9istmeGklm7z89vtcO6/fwsVLWzRvyp34nHkLtu/Tk9T2xIuvvDG5hv97oWOHtiTB+5984dHnJy1b9kWpQxLW21M/fumtydwYPjZ2El1EBj9jxk56+LkJZMDbH35mzLiX+WCgi8qnAmu44Z5HSXnb9cmeAJDEHnluIs8KMaC2bd3qyRdfvfWhp3jRGXjz/WPmf76oXetWzZs2HfvyWx98shIPNPGW129fBDq0E3ZbDVw3ltg30JgP0pC42JFJ5Hy9yS8DbzU2Rz4cJDrOdTt0vTc1A/KinaSY/YokyAYoVGM1HVW6kr59arLKNO7+fnfyMWcd/wNuvsgRQ35x5Nprth/8k8OO2HfXg3bfcZOunek67qA9f/Pzw0lzPztsX4b122pzKoTavfPaRQPSzZHf7/+LH35/py17HbHfrgfutsNJRx7Aryu/n1hScdizWxduSM887pBdt+2DN/Id+mF79dt84670DvnFUb894cj9d92uV/cuxVGs54j9+rOwjTdYj8p0TMGQ5k0bw3t064wfbt+4Wzz+kL2P3n+3Fs2arlixgj1g8/DUE8QAxH/Pbl3WaNeGZvfOazM2XtcabVujYMBFUbkt/dEBuw878eid+m76zgcfH77P9763TW/SzRrtWrMG5t2udw/MqGsEVyKC58023IAF4JCbYlJbqcMN1u3065/+cIse3VlSh7atsWfrd+XZJw76yWEH7Lo9sxDtH+6zCyJ+WCHxwdUxA/Yg5uus2QF76vGH7ENsSdB4prn5RhsM/eVRJx95AJYE6tgD92TN++2yzVabbsR1YYZNXr/jEWC31aFdfd8MNBJssGx4LNjUpOyCvD1Bi5xNVSpOyUjxMF/G7Fz7kfTf2XsbCkqBwAvVZ2Ka2ZhunDPzp1LgtH3gQUkb7VVxc1T75F9XDwlrl617r6z34qjEue5d1gEb7oGtJbuh3bfbouFDarPctPv619/zKI/et+jRrX0by0S1Wdar88bg84BHjfVa5gZ5BBoagWZNpdMaskZ726k1aiRsphgJwtmRodOLDWLDK8MdnlRKkBtNR/qxHZnazswOL3UX7bp+2QH7mk20DMjjsKjY/WtQrMnhC8/ObA5zjyb80nDyHPTaqfGhB7oNu1l7FR3sRxyJfSVnL47iKVL/bfuADXfAhpGtZbOwNWv4qBot2TOydTr16AO50+RVrtGmgSLDuSiwgfa5WR6BhkaAtzr7LzLXumvJemsLCEdBb6iLEjtyCE/NEMhfgXsrwi+xSuGHXrhhnUfZT45y221NVjIP+IondmORoNhTOu+rK94KA6nGwpHsvH2TE4+vc8K8M4/AahOBfCGrQwT4yA37Ms5KcSQyTmq7M/JPyHFslDg3aLGNf3NGcuC+ccsVUSVkQnxIVjIla1U9xd7Ghx3Y7He/qdqXt/MI5BHII1BHBNgJhd1SmoatUeBotjtTJRMpXwuIKCcVS3VSX2l0/I+aXHFB2f57a+d1feLw5bP9WPRdMwo2SaJd1is7+PvNrr4s35fVF+a8P49AHoGKCDhVFVHnEmfUxX2ZUaXY7oztFbuysK/yPN5ukcqc5ZW+xZdaiuu2QaMTjmv6l0ua33Vj8/tvbX7/bS0eoP63xYPU21s8RL2j5WjqnS0fpo5q9ciolo/c1Qr92quanvzzZJU+L6vlmnI5j0AegdU3AqQmElTz5s14dOatiPecJB5OJftRzurW6bhmu4XLP1y26PPyBv0V69X3uvOV5RHII/DdigBJidREglq3U0endiepJUU0e3ZmT89Ib+zRNum+QeP3Zzw377PZy5d9t0KRX00egTwC3+4IkJRITWWTP92waxcf/piX9+GhVoE725OJiqiGunXvXmtNmXXrjA+nLVuywrOhk7zkEcgjkEdglUeAdERSIjWtNXV23969NBbJvuC0lrjw7CxkMlG7/2zfpvU+vXqt88Gc66a///7Shav8Gv7fLCC/0DwCeQTqigDpiKS09gdz9uzZs33xX91QhviYu7i5dHyhqUpTKcadHrT37t3GT7n948kPzJ7+wdJFmOc1j0AegTwCqzACJCLSEUmp64vvH7jP7o68pY6UVagICW3bnXmK7cyyJ2gd27c94cD9uz36+lnvTRw166N3lixgm7cKrySfOo9AHoH/txEg+ZCCSESko66Pvv6TA/br0LY1GcvHp2YVyPOz1Imobc6K6FVEdtmm74nb7djnodf+8NbESz54c9yC2VOXLuRrBbrymkfg/0kE8stctREg4ZB2SD6kIBJR74de+/m2O+y0bV8lWSnbs8TO4kAXMNHECfsy8ap2/0liE5VY99+t369333PruyeNGTdhz0mP/X7KKzfP+GD0nOljF8wal9c8AnkE8gh8bREgyZBqSDikHZLPY+Mm9L1r4sBdd9un/07kJyFZeW4lU/KVV8NU+NKSLy5TkpyjkMNUsVQKyY2MpqK7bNv3soGn/nhp876jJj49+pnzxj972EuP7zjx4e3zmkcgj0Aega8tAiQZUg0J56nRz2wxauKPFjcdOfCUnbbbStWpshNzojAnYukLIXEJGzSrfB1gmU482zRMRSSiqFDWaNfmZz886O9nnjKsW58fvb/0sEffP3TUa4fe+er/w3rIna9+8/XgO19dyfpKiX3k9SIG1eodJcodrxx05ysH1YsYUKMlZHWtB97xSl5jBOw1Xf1eJt7AB/G2f2TyEZMX/2aDzf502ok/+eFB7du2tttIS1McnuRE4hL7x8VS79OUnRndPjw7sy80LYdhw0mtkAHtRC5UUdehbdt9+m1/6rE/vGDwqVcO/9Ufz/nVleecdeU5v7pi+FlXDj/riuGDrzDM+OXDBlOvMBwEoa5yftmwwazhsmGDIFSWVJ1fagsefOmwgZBLg+WlQyvxy4YNunSoXdGlphsPysBLhwZuWMEvGTrwkqGDqPReUgu/OOgXB7OLhw68eOgg6iWGkUcPhhcPNc9grAUbzOg1LOgDLwrDLxpq+kVDB100tKjUzEeazaAiQqgXDR00ctggCDVy8KJhtryLhpnDiyrz0DsQ0erQQRdhMwyM1XS8XTRs0MigjwxdI4cNHGmK2Yws5UNL9MiHms2IwEdU5iOHDRoxdOAIQ5xX40NRCnoJ56KorHnk0IEjhw6iwvEDGTF0UDV+JuKIYDli6GrKw5qztQVu1zWihmsxfST6MK7aODZZDIfWHKvaYhtGDeI1HUFkzJvxkZEPq8ZLX98iHz74ouGDLjIcfLGhccglw6158bDBvxt08onHHL5Hv+3at2vrVZ0mHHZWl2gipCVxogp3ThUuDnTQkOdEhFFeFEJFg6OoqtBAVysoVFWhCo7EmKqhRI5IjbxUX7Xcs0IWaMh1sTqx9XBlpsDp915BZFBCNKpwzBjglSLGRUBrqMIquNcKLnXy0FtpbFBUw6hSXvCplIIODXZ2tkMMMsUHHiHqJTwKzIumiHS6AAAQAElEQVRWCWmrZgqEKipgrKLWFVA8HAbGGrkIZzsQRYrcWkorHqbD1M5CUcG9KExEKaBYgSIqh9hbFQnuBRuh2NNeMa6GyNW4ohT0StzbWFPMIHKneDHd3gOVuGKgmaKrJw9rztYWePFauEDjTr2oxTCs34uIKtWuVwQSa1GnWeDWWwMXwUbowI+IcRHQWipWlCKScVEIVdR+FN0EEVpixfyISSoqwhrtsKyjXsNzMhXhDJoo7MR4XkYfyBp4amb2wuMz9U68hkIKcgxSutS4Ki9xaEFClYJOt1q3qlimBJ1gjFxA5V42cS5JXGLIGaKOQa4UXTUl9n7VOpMmeHZJhvh3CaItM+iO70Qcq6S/EjpnCj3qKEk4EqcuMd0lSWJX5FwVTBLnsMEzCFWOEpugY0BNQlfiHB6Sgg6nyzDocCq9TvksAmytqlgbpytwAF82ILEuLJUXD66KvcLNWhil9nqpGjrlhUs0Q5RETEeBO3FVUR1iIhmqOjg2iTqlIdWwBh1DVCcKsqgMUYXDmcLZuJZw051TKwEjd+EUQCujGdmh5p5RBW5njlijHjhgHnBSudoMlZVohhZJQGZxNhHWOKKvGqJVsVx5e6ZggjiX8eAwI9bBHNUq3cGMUVxtQFubTW5H4NYB4WTWwcZ4kHAYuDUKHMF8Fk+mi7WqvF4mocelRUwcU7jEGpqIcwKqoToGJ5oobyZVWg7unNr7i1OiznHQwdvPaZK4JHLMnSYipGtnqI4TyVFEPFswVSELwkXFOC0RlJAwTfFKUyjWqyi+0OvNxlKjZ3hIoqLYSUBvqMo6gRKUEl7a+zXoosqligbEPwtTQdGgsFS4eFNEKlDFuJquKnBDVQ02dpVOVUWrokcRUVA5ORVVuHA2roGD9CGJFZtd8S9aihKUiCr2IyhhXsHQuEpEdSjZOo2LqFM1FOzhnKphWKeKqoqI9YqKK3ANvBKKOJGoiBrP0HRVU1RLsbou6hgTdMm4KRq4sh7RjEs9OmbCm04UP1oFg45Ytf5vOhNlNfjJuCizwA0LuvGvWmeK4vUaF42owm+ercEIvLSG9Xw5HecMzDD4ybjU87rw+pa+jsLvviclkCeEbZWSZazpPd9I0ife3p3q7V1nCpOgiMdOKMbFdM848d505ZRK9CWpl5Trz9CXOwYxGlTWKaLqVEVEQ7Fe5X2rGgSAXrhT8qA62gpzlEStWC+KUTvEAHM7rWLOypxTkCOiZszOXIS1uA5sgh0UJeoBpYiOiGBTQDPjGp1qJTTZYaNWnEG1w9Q4D7NiX4G2ptCrxJi5AndOsI6omS6FVTECrhrXYJbqmN2pZoiBWjGlMg8tZ5MzxLrtqNk0GGFlQ1gLLENlRqVXtAZEE406rs2SVYvNCI9KCWrgdWG4RhfGunB1ETXwGtG8hVH0rjRXijDKZrHLEMeFMlfdXCniCjZatFfKyuj1zaXF61KKKFMWlZXkXJeyzjCqyF2J4kLMsbFZ7B3I5Und3DmzqYI07bKsy7FbU01wgqMQlzLDoAQbcZKo0m82Dl3C5Tkn4lQNQ/qCCAqSkjxV1FAzJOHRSUut0Kd0CKDeQKGqAUVFjIk6p04jX23Rq63WUETgItXRPgEsFnyICGeOgFijBNQMvXgxLjUjVsSxBMU4xlQt4TSt2qrso0qsKJ5FQOwMKzgdYYV48KiZjVCUQ+J6JONYCGugI+o1oAi2mS4VazAxdlVHH4ZUoK2HWVgnfqogprgy9OYo4yKmiLD6oNi8BS5WlH4Jh6GKFLmIMItQTM14VFiDyaxAhEiKSKaIlWAuXwIZQsUFGGs93NssKrYo1qJYl3At4VLCa9NFRAtVpKFcxCxFVgZZq9l7EcniZoq3y0BgESISMCgZj0oVNCs7sPF21iIqlioSUMkXZAxRXnkUvKoVbDmx5xJ+7+ijl7aX1CmtNDRZow9r9IoVb1s7mWhZTbWIGrghR6HiJ1RxStVEVIUf8pfG4lSz6qwkjpRbtTpXVYlmzn2DeuKSWDlDXFJms1dCDMqSpCxxZfRCCoi5KTQZCyZWyhLkBD0xS0ejrOCQc2IHPY4jSQKilNRgzER0ZbUMM3NuIrzMZDhTFJCJWJu5c2UJxZkNolliluDeFJyg0F+o1rIjMQtIVl3iiorLHCYU1CRxQTF0xp0pgbgyxGK19VhfWSDWBaHXEHOMbWCZmVRwhwEKNgnFlQFwlzhIPAyTrO0S5wKvhhV6UmsJzhNDhmNFcAzDpIGbd7rghi4sNQnoEpckcamVMaldLyuxLHLsi7zMJoqzuAbqScFn4hq4ntL1J+HaXcCMJxSWUURIbdVmTCriHJpJATOdsSgZOi4wthJe04TibGq4S1zhQsrUkTSSAorC6FbbFSVa5iRxSsUItCpKmnKqiQhY4EpJVCjOCilQJBXh48yryR4qFQlPIrcuFdBqptFUKAakzKzSVjxUq6uPzsWyQm8rVOFDwIthGjj35N6IEJCs0sRAfKaHZuSGntHCWE4pEaCXDw3QLtdiiI0XTUUKKN4MNAXRQY3co2OWIooYmhMN69EUjhm6EbGmeGZIxU4ZqngV7AXUyA0RrXqzN2JDzKFxzGiCOAeZGsy4MCCFY4BOhSv+xRvBgzKvGcSmxScoWApmmpqxhvUo3EZJHBtQfGGdih+PsSIx0HNdNDllNeieTkis5kfMhqbNa9ymEDHUgFFHEZjYFBLRAsHYoBSmFrG5wkDzwCVUcDFju0aGK71p4AX7oGT2kddsb9OZz2gTMPj5qvTS9fhsPZKqeC5NjdhE6LEJxkpXIKFXAnqulxrHGmKj+Ml0uvAZ9cA9r2Mg2IRqxrjy9joWnauEIRreD0UuPgTB04vmVHgsZkMYy2+oeNoGxlPR1PtUAwc964Gj2Otb7iUVnqN5T5JjYsVCSX54IPOJqDokkNwYiIBBCqDkwUxwLoGpOBZjSKPIGVTkq4/OpbEqLkycqBNxapgEnqhzQvpXp8p1GUrkTkWxMeQocqcUh6JhJJhxZAm64LGkKmJWnRkUuIMkTp2TDHFl3JkOV9NxmhjnU8vhCMsKpE+DrmjWcFxJkSfGcEUfurOWYOQcAlVMcWCsarpjRnEuq4mKUxQ11Orcme7UObNR0LzDnaokSoEHBFAcutqhEqgad5p9OtOAR8TeOI1go1phb9xa9Fl1YohIZYih0O2SgPQixhomdVHJuDI2gcdRqgVeoSfY16cnrmb7OnXiYytPSvwnK+nHldi7oh9EalwzBB3kGsGsioNk8VHjNFUJhVXsKzidpturUbAJHMOgI0JBU3FlLI4xDH5UJXEYBXvjqqBaSZIyOhJ6nCbOcU4cRRNNRNSJUkxRTZSmOKdOvAtc7TdOjEM4qZADhULyM66iqjQ5ZahK2wMh65muaop1C3oFj8q3AEnv4sOFe8vvcB+aJchnhSfrm6V9qogvIJdXwum3kZxwV4pS4qvAGYpVAW1caFYiasam+GxVzEsTZBxYyrFAyVBstVIxGnOWgDkYuXXSthOCWXIq6aUDIeiYWbU1WNt46K1C+FQsKpGnvBuElUjJqsyprVOCbsgg+mtHm7fyka0ijmLNxbGRg8wCooORV0Nia34ECHbBMr5oSMirALkiqq2EC7ZVfVVrKL2uyAPivjgLVwwPyNnWUAuPvVXQlp3Zs/TQsjOSYhnacM4FjO+BaBM57xnySUHh9dKQTzQV9cJuTqzXko2kRQxvMMxEhOFemMAkVu8dmpLqOAkaRqYIZxutoVjK04JNjAguQpcUMLMJzW8BJ/WT4GNVGhpoNXT2WcFHQWI2zlCrojq2FM6Qa4cHHwkcqWZUZ1FzYK1V1cUqDpIUUdR4Kaom6jQgp8S4rVC1BnQ44yhFRqKAisqpAl1Q6ESqQPzSYYhskyXwQo08Q2fjExtJDAMHHKvlBDomS8xBLRxL5+hPXBGDL2txBJUzFubPsQRrNYBjhiMMwVhLeVQq0PzaXKY0kOMus2RxJWOjXoHm0tYdFVqMMg5jVBxLO3C0QM2+fo5FXTV6sjgwCY2VQhxjn2F47eCs0fEqO3oSxxqDrnUhVgnWDhtO6sTxfrBDhCxEwgVRElXEDOlSTWJqUnShYAPDRlWU4khu5CYRBQMn4QXuM8nbX4ny9GKj9ChFRAICKMa9qLB3Ewn4LeH2ocCHAGkdtDghkMiJQEDvlRty8cKHgA8opoRgBxvTBRuxgsKpAlU8fYbCucghVBF6xVvENGBVnopiZusR+4yC82mGZcpARfFemd1GoeMtxVRYrY88YFXuTWD1Zsr2UjijGCr29FrLvqbEhk6QNYAFbpY0eYrBxhWCHjgrpMsQe5RS3XsfFJARhpxs/YK9UOrmdo1iliyNgUKxqYQpoKAPsZWADeLEzYtE+8gDsky8xeox4DAME60cUZwQB1CCZ8mGR91QatDjRLG3ghf8RKWhyNolm5SYRZ+GPs5beu3VeVCy+NTAbQ34Nv8ivDpFHl9HFOTIwVo5S7H3bfZ+YJQX5WI1oARUwT9vH+/LU8/WLLzuPhXTuSze8FSPETa0rdt7F3pJS6JKyqQF9xww7wVUMV3VKQ2xEzxIBiKqymFc8YAVqBSydkAAZbXUE1ujy3ZXjiUaV/ucgSfo6pLEURKXBARg9mlBW+l1GKI4dQ5LQxNc4sIJQKrMabmEoaHDgRZRR0GjaYgJawhdwVNCr0NJ8GdcI7e+oCSYJnQ6lMCdM8UlygmuQVdHyRRY4oIQMcqYOjsqOlgbChoSNXDH7HCq6YnDBqU2njgXekFlbXaYYudwKJjwFrG+Gji2LvRi4yjKtIwwRKc36LaEwAV0apgErIGzEXCa6ZFXwuAZA+YqQQaY5wolMyvRWVkmYvkN6iyUeasjYqzCely4xiRgpWsPUaqkByWzybj5Sbg+py5gYugCj+8x0zmCzll5heAgNmCBu8ANHVYJLVd4bzgnCkcGnYOTVEQcRSVh3Z6ExhtBeGTmVQTK2EQdhGqCqKg6sULi4pRhSrrzqYgP+dKMrcMOVRG1YicMfDwHxSDrNlULXK1UUoKLAAWb0MhszPwb1MXmYmrJ5oWilCAt5RARPj1EAiJ4ZYR4wD73OAVOTFQVRFCRUlSKiCn0q4oaB9UKXDQoAY1DYhXTzQDCK5KJCGE9KEY1rEc9PK4wYhiqmY1QFEVUQeVUECIFVQxAO4lUoBpVEzgqKo5iFRWIVKCnSQ2Kj6hKt+m2tsBFKfQacmTrDFdR4MKYImdipS2qAUVFOEow6iC6yaJqJwMt4RW9MHQRQIVilpErpyBoAZUioqIKGsCg1hApcrESeiXrjCeT5Uvq5kE1ovlRiuBMkZSzCKBCMSVy5RQEBSMPqAGlGpquIkU9cpX4gw6LiMI0EVHgpRh0HxQvSvGRw0pf98h5ZdUGeCyFewuJbw9aXhkmYdcl9KJwSj37OPEpP7SElOQBDgxFhCExneFThCYGoiHjFXUvIk5F1XKmsIrUW67DnOjPzQAAEABJREFUs2cx+AkoKlY0nCqhyVJJiTaria6izpanKqKiFM6liIbiTLIIqBjVbJBx0SKKGhcFRRQ0j05FsBEKiqGiKN0iYOQKtSpiqGCsGppxhQEduomqElZSRDOUbIVOVa0XzYnygyCilRAVxdGrdIhaMZSgqIigiBRRKKZwKE45RVSzUEMIVRUu4kSUqoY2iYihFtAJZi6iBi4BUQUbOzkDwUSCUgOqOKdSGVWDYqgUegOqqRycRSlSRNEKriW8Xl2tCGBr4KSRG4uKM8rxv+vmQYN/zUrFmkUruJbwBupqRQBRfliywLXIrSVaOc70Su2vi1i3+cFGrQFHcpHzcpq3oMMFP8J+wKnPMOQW8amKqKMAqtZQwV5VnaI6OIlJhR/lELUCOhElJ9m/GETOw0Q8OS0kPCxokAY9vZ47WE6WHUUZIcoPY5Vs5sOjFFbAsPRbyUnOaXm42ALaZcNT78sJj6FPvdmkHiVNPcjHCYgOBq6BB0w12GTcdOKTel4nz0MD4wJJU0lTbAK36DGXcXsJbH8sePYUVuI9s+OgVIGL96D5L1lnQSlZQ7kGG5XU+6pcgmKYpoYsrFh5wevmNoRRqTkOlt7Qe9DWlhrUwOMVxd7yMLY8WLI2zwqJA6hQ1mYdma4ZT32FnnrjqQdZTIpl6iMHmRpEByOPaNdlloIeuJTo1bnUYlOqryqutaytuq4l15jxqBTiUxo347HXXh6LlSnCS1KFByXTjUspR/Dl5sCmoFFunbi1U5Fbg1xi2Yd0IhRNvU/TdEWalpf7cs/vA17RQTjtULGwoSgF3amalQvnyEE0kCEgnF6MAlcKnOE4Ar1PWYGKB2mCtXDLfoovEZAzKLK6cJI612irinFwUOGDQCnGNXKzCYoLKIYGyni4ZtyoirMmSkWVqIioCkXVxpmZSOSgiHH6M+5UEQIyuytwDQpo5nDVAheomp2C1lPgiiNn5ugFLkbQQ3VaumZVEzmwUeVsR+BiKEFRYQgdohq8chIO1YDBsBbuBBunCjKSFlidR8X8YCGKTca1lIvpmPKIRcVhWeBanRcUCTYaMHDVujh9onG5TsxeKYKyKnnpehrERbkOVwO6Eh0eYxhRHderYBxrDSdKcSKcnTCU3lIeFXrQC5z3CZaizquZCiX0itKWgJbNLKkoBXNlITBNIEpSCQPNUoUfVQ2cPYBaEUBUHGlQyEMioY2N4NZ7X9TFeul0XrJeFFEVplApRVzbx634mtA2QClpWkjTqylneSkfBaywZuQbFvvzx9jE6gWFOBVQKriXNPV2pRF94BnSJZUikEqwDGgDJcwi0SbySmj2tsLy1CyZnd4iMgpuGKYzjhnLsKZPvTelwG1e48zlK3hawqv3eilaMi/c0AuLMZ5arymp1IiYoYO2wmBD07gPA+PwDL3pcC9hzbWihdqGm33GWTYKmOLWszDTIzfR49mUleC+xB4uaTZ2FXJfsoYauZSsuYSXxiFy0C5HzCFxS/EmYWyGMf5Ribx2ZEh4pVIwcC8h2hmylbKXw0Qlo3hu7UIOgZN5vFi2ydDTwf4IpCfo3rMyT8sOTFWVx190pd7T9GRBToIaTgKz7w7sLBRn1hh6Cl0SjjASl7SEEscaqjqnVgJGXoohzWY2q6PObXnCjbldhoNXUNaMUMCEDpcYhm9hrOUwTyDqiqg0E2QUc4iOgkAjxME5i1JAIDGZDyJs0GMlkkbsZCMidwkmilNrOgUThmtYSQERXOROk9ANYooOIqiagyKaXlDoUNqJAzWgw5ezgmKctaBYwxiKmcYjMdVFDIr1OuZBD1WDywJab4EnZq9JEswCJo6iSdQdJ7oc58Q5QpK4CjTV/FQoiTNuutMkcSy0BJ1LgoKNc6WcCbB0UTdEwA/GVk03wZTIaSkHlpWR1jeqs5qwBps38gwRbOWcKq4rQYkxcdm1J0HBg9PEcXWYGyauKrrKCq3EueiZc0K30yRhOEhDXeQOE01K0EVObhFxKryaaiiBIYiI0qE2juRDqmFnRkNsw6ckLxHz4BglShFDZKgjOZHqAkpIvym9oiK2cUvTkPZmzZv/4JPPXv6PGwdfcPlJQy88ccj5Jw65oFAjL8XYVapE/q3RT/otSz0fPGlIKV5QTTn/5N9ecNKQC8CThxgvwQtPNv3Ck9CHBPxtwCHosV5wsvESHJrppzBw6AXgKRleWMJL9QtOHXLBKUPPB0+thKZnypCs95Qh5uRUQ5QLTzW9iNhX4VFBpF5w2tCAQwIOvfC0IReinDbE9MpYST/dLC8wHHphQOOnDy1F0ysrsbdCPy3YBxxxWoWfEWFUEbEvcjwUeT36GUOxzGwCHxHwwipIk3r6UNPPqIRmX1mJNt+0XlibzRt5Zaykn25XHaOUXXuJslI6L7fZl7wupoQmOs5jjbwUq+i8o6w3DKzgZ19wxR+vvfHhJ5+bPW8+m6k0FNKUhNyUtbyQnxANU3aD8V/dwETJbl7VbFW5jRTaZNFZ8+f97abbfzzi0uFvT/hXl7Jbd+1y6wE9bh3Q8zteDwwX+FXgbQN63nZgz5XEXsG+Xux124G9bhvQqy7EgDqg138P/PbV2w/slddvYwT+9zfb7QN6/ne3DW7o0uj3b0/4+YjL/nHLnbPnz3Ps8JS8JJoVuP25WVU2a4XKbo6akt9IYFbtSwUUr/LE8y+eduEV1yTzJx7cd+e9dhq69U63b9H/ub57jc1rHoE8AnkEvrYIkGRINSScfvv0m3TIVv8qWzBwxJVPjR1veUnEkhWP3Dx3oezI2JwJD+esy1tSE2QynErYlHmBI937yJPnPzJ63IDeu2+/9SNb7n5O195Hdlp/7/Zrb996je3ymkcgj0Aega8tAiQZUg0Jh7RD8iEFvTigz8hHH37g0Sd59M9dI6hC8uKJGukq26OpOmepTrg5DZsyEp56lCfHjr/q+acn7dlzeI++g9fvuV3rDhs0bdkqaSR5qTkCuZpHII/AVxwBEg5ph+RDCiIRvbRnz7+Mffbp5yeUk6G87cjKw59I83wvy81lShZL+bJAlFwn3FyKOLZlfua8eVePuu+9/j0v3nCrg9fovHGz1mVKx1e81txdHoE8AnkE6o0AyYcURCIiHU3u3/Oau+6bM2+eF++csDEjYXGCcPcJOqEHlyD3miQ5kTvuf+SdLTsftl73/TqsvX7TFnTmNY9AHoE8AqswAiQi0hFJ6b2+Xe68/1H1whaN9fDkzJO1vNhfjbBnZ+Q0NCfs1aBz5s6775VXpm/Q4bi1u3Vr2pIBec0j8G2MQL7m71gESEckJVLTA6++OmfBfJIVd5XqlHtK5UQG4/aSJEdq85zEEt64l17/rHvHwzt1WbdJM7Z5kpc8AnkE8gisBhEgHZGUSE0zu6354sTXyFjhqVkqKdmLQ/kGwIlKyHPcanpReXvylGVdO+7YtmOHRk1Wg0vIl5BHII9AHoEsAiQlUtOybp3enjKVrBWel6knh6nCnarDMOVm04sX8amfNmPm3JaNujRp0SrJv8ckNnnNI5BH4CuKwP/shqREaiJBffLpTHZn9ryMvJVyeHgq4rzYd5qgkNRUFi1avMhJ+0aN/+epcwd5BPII5BH4iiNAaiJBLV68RMOOTNmFOVERZW/meYzm2ZORzTzZzJiduRvFSuotH3wy8z/3Pnn25Tf8ZMhVx/z68h9Rf3X5j351+dG/uuzosy476qxLqUcOvvTIwZccMcjq4YMuPnzgxUefdemgkf/816jHpk77rN4pcoM8AnkE8ggUI8CtJLnJszcL/xSLNcvT1E4GZDZSm5gFI9Q454bUm+9/ethVNz387Esfz5i9ory8OARXRV6VhL7l5eUfzZj1wFMTzrrk39fdNaaqTd7OI5BHII9A3RFQZTsmGvIV+7JYVe3PndlA9mbh2ZnxBhx//M99Dz49wVtSZDtnNdKAaZpmPRCfUV+ppLQ40nsef+Hia+9swISrjUm+kDwCeQRWeQS8T+2fTWMf5jVll+Q9mzXx9uwsPDULeS5AvUu96f6nxr/2Xr1mRQNmK3IjldtjX3nn36MeMz0/8gjkEfiORmDZ8hULFi+bvWDxZ/MWzZi7EISjoH+ZK9ZQHNsx9TxEI6WwQQu7M9uY+ZR0R+VET13+eV724NMTGeNrLaTHQl8JLUjFofRRTb57zLgpH8+oa9bvbh/XP3ve/Jlz5paXp9/dq8yv7P9vBJZ+sYLMNW/h0iWkrvIsw/C2X1GeoqDTi81KBcinqcdVefhLm2zTSCooEv7MhjlSFRWJKHWVpye8EbvrSXvcumLHdwxgqNE+YhCAitaYca/Qrq2Wl6ejHnmi35E/p976wCPLV6yozfKr0t96f+pp517y8lvvljr03j/63Aus4Se/+cO0GZW+x1i67IvL/nnTSeeMLNaBF17xwbTppcNr5IuWLD3ijCE7/PCnZLQaDXKxvgjk/atvBNh/zV+0lMxVxxLpxQbLOmyqdjmVWCUQEpezf4mWZ2chp5DehAwUOKT2+tq7H/Bbnfo0YDh5T2b0FcW6rFU4wyON6FN+gpaCWX3prSm1zylX3XArv/Ovvztl0hvvnPaHS8a/9mYdxl9J1+QPp/3tllEff1opZ73x3pSjBw7rss5ak954+3d/vIalF+daUV7+4FPPXXv7PcX6r9vvnT1vQdEgJ3kE/r9FgJ0X+68GXjWW2DfQmNs6spWk5Csf0ad8G2nPzqzL5JjR6vP36ax5lU08j+OigpNIqmItHaXy9M/mVB1VaC9cvOT+J55t3bLFkzddPX/i46/ef/MOW2w+5vkXu+9+8N9vvYvdzY8GD9/tmBPZ3dDcYsDRJ54zou1Wu23Q/8AJr78Ve/f92RkHn/SrppvtfNBJZ83/fCGO2Tft9eNTULY88EdPvTgJhbFbHXzsldfdwsDjzz73mMHnIP707HPPGvlHSKxkt8VLlx09YJ8ObdtEpYgtmzd75N9/mv3io7tss2XntTt98MQ9cyc8tsF6ax/3q98327wfPv/74GOkP+oDTz5LE/HoQcNmzJpd9ABhbez76Oq0/d5XXX8rKRIxr3kEvo0RYLfFzeVKrRx7RjVoSPa8zAm7M7V9mSi7NWOWj/g1I6sZo1vqKtzoYUzqCwiQzcDSap46tm/zs8P2HHz8wVtvtiFuS7sL3MwK3OO2tllbNGvau8dGCxYu2vbQn5ww7IJ5CywfLVn2BcllwcKFeJgxa86H02eUl6c035w89ZFnX9hjx20Rzzjv0vmffw4h90G27d3rwSef++P1t/Ks6qiBw1589c0f7Ls7G6gjzxzCtouxr70z+TeX/HmzjbrvvuO2m3Trwnr6btaj14bdILFu3LVLuzatDz7pVzPnzPv1Cce+/u779455es78mrdgXNEpv7/4tgcfPXD3XVq1aH7sr3435vnxz0585fDTh9BEvPPhJ345fOTSpcuic/CBJ5+76d7RJxx+0MDjj9646/plSYKY1zS3eV0AABAASURBVDwC37oI8CxsybLlX2LZjGJs/QM9JmlaXu7TlCrekFTgVJUMpgoqBauVreZZSGo2rlmTxht3XbdNy+abbbx+70026GO16zprtt+wy9qdOrQ1i4JlHIVSJPAaK6v6w+m/GPzTHzVp3Og/dz/YZ8DR9z3+TI2WUbxy6KB/jzxn5636kNpIc4hsl2667Lwrhgxq3rQpCeWlN9996c13Buy+y/UX/e7kow+bNXf+sxNfxoz6+9NOuPdvlx59wN6H7r0bzZOP/sFPDt0fQiU9XXzNDXE/dcje/dds347cd9E11yeOjwj6q1ay7TMTXt5hi82uOX/IuWf8Mk3T+5545uFnxn6xfDlNRLrGvvTqh59+WhzZsX075xxb0TatWrLLK+o5ySPw7YrAoqVffOkFN2iskq/4XtOJqvILyNbMWSrLnp2R2CwhcWLfVedCzCQF7ODkLS96K6lv2qjRYXvv+OufHnLZb356wg/2XmuNdiSgfXbue9HgH59z0hGnHPX9jdZfxyzjkZoHOzwurNYxbdPGjc8785fTnnng7F/+OE3ta4E6jOvtWvbFF+VpWlaWEIokcaX2m27UDbFUKXKSI1ut4w874Oc/POjCq/+9+f5H8bXAiUcdRuop2pSSZV+QuJYnic3SqKwsdn2+aDGEJrPQBS+tu+2w9d1/vbhTh/ann3cpu7mly778e6LUbc7zCHyTEVjGJ///8DU93wzgoZ4Fk6as2h8G8HyzSfLy9iDNiSojVVUUDVp/rWM/taK8vCxxrVo0mzl3/tiX337yxdfe+3A6iaNZ0yapT+mtwXsd7oI1wTn1Dxf3PeiYf99x7ztTPkTbYN212Qaq6p2jHz/vL/9kE4RYrHc+/PifbrgNsWf3DdgUonO/edO9o9lJLV66dKe+vbferMfGG3S5b8zTV/z7ZrZ7bNm22rQnZqWVjENz1MOPv/DKG5Bi5bb00L37b9FzY7ZpLIDvBGq72Vx/nbW32qwny2CWP9/4X4y/t23ffXfZAcJK/nLj7WNfem2zjbtj1rxZU1aI5Stvv0fSvPrc3/TpsdG0Tz9bsqziPrS4gJzkEVjNI7BsOU/l/6c11uuBnMHvkbApE8CRvDRwnqJ50pi3VAdgJnUXM6p0SNYS36xp485rrcmO7INPPrv6lgcvuvbOK2+49/yrb3t6/Btpmnbs0JYHapmxnSoGWosF1DLxihXlG67fecrHnwy88Io7Ro/hgddpxx5Opth1u6142D/66ed5UlY6lGdn5/zx7+3btLpi6KDWLe1foCSNkkFGPfLEXjtvz9iOHdr/cdigZs2a/vriP02dNv1P55y1Za+NSz3AD95z114bdr3h7gfP/+s/4y6JJs/LyG77/PR08s6PDtx3m817skFL7VOBEVVrs6ZNLvnNaYw698/XPj52wu9O/fn3d92p//ZbQ3Ay7IqrN+7a5arhg/lWYcDuu3yxfPn1d93/xLgJrHzrg48jaf7qhGPbtW5V1WneziOw2kdg+Yq60hnfzl11/ah3pnxcx3XU7cEGep+m4cGZZ5uU0qKSx5wqqU1UFSNVQ0gdtZjwIomIfaf2bffYYYuu63X8YvmKF15999NZ83566J5nHX8wN5hPT3j909lz27ZqsXPfXlv0rHiyziirRRfWqOEgLww8/qi54x/75NkH+GbzP5f8gfs76v1/vxxl/J3X87TrvUfvXGvNDnHwX3//68/Gjp76xD1bbdojKut2WnPCXdd/+vxD91x9CQMRSYVTxtw1/bkHZ457mG8qVZVnc8tef5aMQy+1yzpr4Znb29uvGtG0if37ImVJctbPjlkw6fGPn74PvPaCoU/d9Pdxt/+rQ9vW2MfaMny/WVxMj24bTBh1/Yyxo1n8b35xHDs+KoThTD3+zuswYOBPDxvAV6I83Tv9uCNmv/Aok/Kt6EF7fI+uvOYR+NZFoLyWD/h4ISnF+6tvvvfNyXanFcUqWLcHM+Y3lmdm9rzMcVa2ZupEJezORJiC3Fb7DkmKhWzofYSIpEYja63Rdpete7EnWrZ8xey5CzbfeP1+W/XarvcmO/btOf/zRQsWLmncqGybzTfqs0lX782+FHwoxSlqJKy/Q9s2MbNEgyRxKGBsliI5i8RRqiQuYbODk6IIb9+mdRWzYi8Ez2u0a1PFgCZfAoAYUBs3aoQfSB21bauWpcvGkuFMXTqQPEilq8ZJ0fOaR+DbEgF+m6svlXusuQsWzpn/OTuvw/fbtcs6Ha+59f6X35xc3RKlRg/oJdVMvLdnZwJyhwmKhN2ZiP1qkd1U5cuW6TPnPjn+dZbbuCxp16blO1M/mfTm+69P/nDSG5NbNGvaslnTsGt75+W3phR3Y0Vic1ZqmPAljl8eeegHT9yz+w7bFMcy9S1XnP/8bdeSg4piTr6tEcjX/a2NwNRpn/7+j9cNufRa6u+vuv7t9z9asnTZdaMefu+DT77UNalarrJDwr5MQQm7M19a6vNdahu4BPSfzpo7ZuwrU6d91qRx4y17di8rc5dfd/ewK/8z7pW3t++zyVprtluwcPGzk96c8MZ70Z6dYIGEM5vD+qaut58NDrecpVshVWUHRy5j11Pv8Nwgj0Aegf89AvzSVXfSaY12Rx6w2zEH7Un90UF7duu8Nk9v9txpq+5d1qluXKOHSmaWM7gvpHpuLa0Vbvlsd8bgilppUIMacV8Fzpg9790PPln6xRfdu6x1zAH9d9iiR+e11jh2wG67bL1pWeI++nTWlI9mFD1iH3lGslPUcswjkEfg2xqBxGn1pbdq0Xy7Pj137Lvptr17TPnwk4+mzzx4r5332WVbrcFWavRQyadSHDsyTuIgnB0G4dlZSG4RkOqu0cxSoWXG1KdppnjfumXzDm1bJ86paO9NNvj1zw69augv9vve1i2aN8OoDb1tWnlYNsCYHd7jA1L3vHlvHoE8AqtfBGpYUaOyuv42y5z5C7jBPGhPvuXfUrWG4Uh1e8BAsnvClLN4koeniPekNGVjpIUSTOsB7Gu04JH/tXc8/Ie/3PKzYX+6+pYH2awt+2L56Gcnnnze1Wdfdv1F197x8tt1/T3zGn1G8d0PppWHP5gHweeUjz8dM+6l9z78BI4+4fV3aVI/m2N/n5TbWL4MBuPYGvGhZ8YzvMauhoj4JwWXWr7w6tvUUqU2HsfOmjv/vifHzZwz/6b7H1+8dNndY54HY9fEN96j1jb8K9d5NLtkWUP/vC6WN9//OMgyuF4qpI7KC0Stw6C2ruKqHhs7iZe7NrNcXw0j0KRRXemsY4d2Q04+uv/2teYyrqhuDxiIkLCcqhNVUTZQKsI5/J1NEltFlXpKhaUxnnjZqXgsXvLFq+9MnTVv/sQ33p/w+nt8G/Diq+9++Mlnb7z34cefzi6a1UjqmPjJF19Ztnw5BpDPFy2e9OZ7zZs2SdP0j/+5ixT2wqtvb7T+ur037tqmZQtslixbFn7l6vozqEv5Jfsf/pUh/M+ZX+mvai774gsqs9db41i+2ey/bR9VWbxkabMmTfbYYUswdm264frUev18VQbvfTht9DPjG+iNF46cC2LP9VIhddQXXn2LWodBbV3FVe245abrrbVmbWa5vhpGoEmjsrLE1bEw+nnn12bAWDzU1lvQeQ+yYyEHsEFL4/aMExlOKSQ50GrBvLZzcWsWCS4rW0ZZ+Gbgr7c8eP7fbhv78tvRIHZEjMqXRudcm1YtNlp/PW5vV5SXszXlJneNdm2aNLb/TO+lNye//Pb7dzz89NRpMy74+80jr73tnsfHfrG8/Jr/PgD/0033LFy8JE7NnugvN997ywNPlNvfk4ia3P7w0yP/cSsDIVVG4efSf99x471jXnlnys0PPPHCq29feM0tw6+67q0pH8XBn82ed8X1o66+9f7SrR8bsX/c/iAG4IuvvR3Hfjj9M5IXIjXm39femxq7nhz/6guvvvXJzNmX/uv28/9207OTXn/+pTeuvOGuf496mFSCfazFuV57dyqe2dzd9dhzb77/IcbnXn0jV8oy5i9cfPE//3vlDaN+9+cb3pj8IT4vKASEjSHX8qcb737kuYnPTHztuUn2lx+WfbEcV1wUYz9ftOSy6+5g7NmX/5P9EWPP/euNf77pHj4/4gLA0c9OiBMRKII54ppbL7/+zlnzFlx14914eHr8q89OfOPBp198e4r9mUnicN7VN+ETP3iLwQQZwoJvH/00OzJGMftto596esJrz4RVjX6GffQ0Fs8lsOCpn8wgDkx67R0PTZ85p3q0WVVeV3kEWjS1P6r55ZbRsLGWq1TCvizszmiTy+zZGYmO207QqtRTksSGFNJhPBduXm18OEy2IzQMQoPsaZwjsiDSsorbOiYmOfKLxLs5pokvli9/6Onxl/zrv2ut0Z5tzoJFi58a/+rzL7/JbyNOtujZvc8m3Q7dqx+J4OA9dhr848OmzZg19uU32rVu9euf/nDj9dd95e3snnf+wkWbdF3v0L36Ja7iw6QsSQbstsMZxx4yY9Zcsk/pqBUrVuy789ZHH7AbO8Ej99t1m802Oer7/Tfp2vm1d6YyL5WcwoUctPsOG9b0fQ0GXdddK45t2bwpzdLavfPasatZE0vKDzz5Que1O/bp0f319z6Y9/midTt1OHzf77VoVjGqOFe3zmuX+iG/77pNb65UVWbOmdeudcuf/2A/xrKlfXzcy8WAzPt8IWs4+agD99yx7859N9txy1444fPg6P132227LeYtWIhB6xbNGbvfLtt8PGMmYw/bu9/JRw3o2D77pwSw33unreJEE9989+0pH23bexP2mFM+nr58xYq9cLvV5jv17bVvv2026boextS11mg38LhDd922Ny9BDOaA/tv/5OC9d9yi14effjbhjXfXaNv67J8f8cO9d+m31WbFVfFh88QLL5905AGDfnwoSZyXe5vNNgZnzJ5bd7SZMa+rJAJNG5fFt/HKzs4oxjZklGUNDkx9Rf5x7MvY7CCqKlxhdda112hXW38NO68aJO5Pa3CwbscONagFiV8DfpF4N8c00bhRo336bX3W8T84cLcdWDW/dbtsvfkOfXry21gYYWd+qZo2zj4ljIc/3G8dheOYA/ZYr9Oal193R3G/Vuix9E62rTKKQKmriBD7jvc//pQcxN46Dtxg3U4/O3Qf9DHjXorKl8ZGZWU9u3XevnePI/frT0bg6v54492fzKz499GKcz01voZ/yJddc3l42hgXEF9wu5xCQNAT54gepFjnzP/833c9QpDXaFfxVx1ib5WxUQSLE63Zvg2h+MHe/bboseGpRx/0+uQPbrp/DAbV64qwsBjMyR9Nv2fMc5tv1JUXcfnyFcVIlo5K05T3UeJcFDuvtSYTHXfgnr037vZVRTt6zvErjEDr5k2aNMr+5YUGusWeUQ00dvwiKr+OnMRBHEDTfnN5t9gpHPV4673JBr6i2IiKViUWugIU5EqNgpidt6z+l5/qWUiRhVgzAAAQAElEQVRFd3HvFm9qEue4KXvihVe23nRjbhj/cceD/KZt36cXXyP8597HXn13au9Nuq4oL3/r/Y9vffAJ7hNbNm/GRuz8v90ckxqP1e57Ytw/bn9ovU5rbLv5JqWjilM6p9yjcZPIjI88P5Ff9cQl8Jfeep87vi+Wr0ici7dUDGnapPHMufP/dtv9s8M/TuvCWFZIl6rOmD2Pr1Dg1Ni1NPxDUTts0ZMbsSdefIV7yXvGPM/XHU0aNeIXvuj25cJc7Vq3WrRk2XV3PcJScUJ95PlJ197xIPfgnTq0mz5zzi0PPD7qsee4FjZBxYC0bWV/mxVj9lNvT/34pbeyP5+9YOEinr5zn05Xad1ms01ue+jJ6+9+tHizyTb2qfGvxYn69thwxYry+5984dHnJ3FTe/P9Y7goFta8adOxL7/1wSfZH9Bhc339PY+OfenNvj03LDqfMXseHwCz5i7o2b0LBjfdN+a2h57ik6m4KjL7Fpt0+9eohxnbvGkTMuCjYyc9Nf5V9psx2q1aNC96y8nqE4G2LZs2C7caDVkSltg3xDLasNuICYUmSQQOsZTGL5VwxIpWZ91tu95V+i0X8sVpQa3SLMhVz9EMNZK9du4Lr63+7LB9eR/TC1mjXZuDdt+xZzf75xVR0If+8ih2auzd4k0NvwmnH3Mwt0ibbbQB+i9++P3D9urXolmTM4875EcH7I4Z+Qvl4D12PGK/XdnfcRfDw2aM+bXBYdMmjfffdbvTjzno4D12wnnpqOK8xx+yz5Hf73/onjv/8vDvc/f0g7134WaNtW3Roxs6M/bbavOtNt0o7g2ZjhsoxF//7HAWH8euv04n7HnkN+QXR669Zns4c8Uu7v4YvvEG69HFOslr3PxSzzj24M5rdSy67VOYa7vePVjtiUfsz8XGsOy5w5Y45D6RezGcH7Fff1zhkIoNK8EtKR4brpddHjeMW/ToDm/fphWL/OE+3zvv9B93XmtNDFgVi6Hy7QRdXO9vTzgSEeMdtuiFW2yYqHHjMvC4g/bkrrZHt87HHrgnoeYlYLXMyMViT2Xfd+yAPc4+4YiOHdrGYHKLjQEDf/Pzw9do23rwTw47av/dfrjPLt3WWzuuKprt1HdTuhjLywQ5Yt9d0XEeo73NZlX/BQHmyuvqEIHWzZu0adG0LMl21jUuiV5ssKyxtzaRz35LXk4td7mQvFzhm01ymyfdWZarbXimd12v0wH9t8Uw1DjGZ2ML56CmiNwleO8Dppx9LNACifSQvXbcsEulB0DZZN/UiZ3FXjv2JQ8y4XprrRG/IYV/6cp+jd9AssOX9lDjwIa47di+HTUOT5zr3mUdMDZXLTZp3JjFrNo15LN/8xHgWViH1s3ZeTVrwg2DJR7WQPohi6Gg04sN4kpVH3MHzyFIJgXuJGQ3vFONN8Dl8YfsscMWPSoM4/6qop2xWuSst3jauW+vE4/Yr9hcJaRF86bNmjaJU3OLuk6dD/KiWd1IZvzfc2L1KRrili0qNY7Fvv+2fcDYXLXYqkUzFrNq15DPvqoiQCpj/0Xm6ti2Rad2LUE4CvqXWxL5ikdlVIYbksfYnfFknuyGZGgHtP76m58fNmC37TC3als7O1c+SlSjdlQ2sNahe+34u1OOqn++3CKPQB6BVRuB1Wx20oflLnZnPOkCvbBfY3dmac72ZUp+42joqn966J5X/vaEA/pv12XtNctq/2sNxT1aJCDG66/b8eA9drz696es8n1ZQ682t8sjkEdgdYpA2JGpupDBQLZmjhNZji/byW2sFaKqItzg0qq3duu81gk/3Psvw0+8609D7r/6HOoDf/sd9cG//4760DW/p46+5g9W//GHh/9x7sPXnvvItec++Pff/+Pc004+6vur9nlZvVeXG+QRyCOwukWA1KQiqsrWTMhdaZph6tM0dUKhr4AtWzRvkcqc5Q39S3yMy2segTwCeQS+mQiQmkhQLZo3Y89lMzqSm9hOTUB2Z8KtJ7eAhqq69pprtFu4/MNliz4vt78jSed3quYXk0cgj8C3NgIkJVITCWrdTh2F52VcSECfIbszDTs3OhQmPTbs2vj9Gc/N+2z28rr+Cjfmec0jkEcgj8A3GQGSEqmJBLVx1y7qeFomhuStwHl05rgF5Q7U1sSDMy9bbd6z05RZt874cNqyJSt8+Ke4rS8/8gjkEcgjsCojQDoiKZGaSFB9e/fyKU/LxJB7S56g8QwtZXcmZDi14iy1tW/bep9evdb+YM51099/f+lCyUsega83Arn3PAINigDpiKREatq7l/0XS84ljrSVgKJJIirinCnCvgyHoKqKHLhX/+7jp9z+8eQHZk//YOkievKaRyCPQB6BVRgBEhHpiKTU7cX3B+zVn7Ql7MhYUAGVfGa7MxWxJBYw8I7t2/3swP27Pfr6We9NHDXro3eWLGCbJ3nJI5BHII/ANx4Bkg8piEREOiIp/eyg/Tt2aC9kKufEKTsyIX0VuGNPxu6M20+Q52iSIvh+W2/xi2136PPQa394a+IlH7w5bsHsqUsX8rWC5CWPQB6B73oEVpPrI+GQdkg+pCASUe+HXiMp7bTVFpKWi1flkT9faPIEjYwFph6F/CailubUKQVmZ6f779bvrN332PruSY+Nm7DnpMd+P+WVm2d8MHrO9LELZo3Lax6BPAJ5BL62CJBkSDUkHNIOyYcUtNXdk87qv/t+u/dzpKfEsTkjU1nKShRFHKnLslg82/9+IrZDE5BkZzs1kX5bb3nxmSf/eEmzLUdNfHr0M+eNf/awlx7fceLD2+c1j0AegTwCX1sESDKkGhIOaYfkc9ziphefcdLO2/bl3tFSU7mlqMAFpGH5i1Nq/5OTCKkNcHZWch3NkOXga7Zr+9PDD/rbGScP7dbn6MlLDnv0/UNGvXbona/m9dsSgUPCi1VEyHe1Hnznq1xajYhIjb2QvBYjUBqTyKsgTSr24DdTDx31miWZRycf9f7SIV03++uZJ//0iEPW6NCetKTsy0hQfJvpnBhXcpTp5CsXDvZktnNT+ztQcEt+xlXshySnHO3bttlr521PPvaH5w06+fJhgy8fPviy4YaXDht02bBBVfDSoQMvGTpwlWO+houH2gtBHEprw18XRkUP/yMy/KKhA782PBPnI4cOBC8aajwQmpFneLEZVOJR+epxyJk1+/yq9KG1+P+yeghX5jPygASQahGrEttwdV/Bq3nJsEEXDxtUBSsyydBBfxh48sk/OmyPnbfv0Lo1z/N5LiYiht6Q52ZhX+YtX4l4LKymltMKmUucit2OipjCMIeiKCKmO4WjiKMIDXKlU/vRUnQJP6YGe3oq9dLO9W8mAgmBd5oEZEbniL1GhKHUjVgmX3Zsqeein+gtIrOjF7HUvlbOYnjTOfo5BTSOm3CNClZwesy5XS6Wpteo4Aj9q0Sbymas6vOr0vGtHFzcl0cGc9V1YGJTVIpnkkXSOhzXZmc7qnLHqgp6NU4MsAcTTJw6EeMlSC/zONEwt1h+ceqxpCMJScfBVABqokjqlBIyjhp6T3ITQSPrSQVHkVRE44/C+DJBraUqYmehKBlPhG8Z0AooKqqsQeNsBUQV1RxXVQR4mYh/QAlrCBjbdaJKsRi1QxgrFHtmIZU5byeU2tF67AjvNMkw8xP0WjnGEt+jlWcUK5VWJZmB1FeqjML8/4PCZdZba41DfI3C+OyVKuWxtwQtN6lGVFErYmdyi53IMKIa0HnFk+2/NL7K8X5RwywaGsqODJusskOz5MUIb7szuBenIkrOExEVp6ocIoaqYjoKuqhlJRS1ErhaxpICitMCx0pE1cVrCGicQQmiSwK6gDn/huJQpi5xLmBSpkngYKx0RYJs3DlrOucSfpxxSFkgJViGmFBMN84ZxSW1c+ewMQ9JWZnRJGAlzrsCl4n1llXhjEUx/2rdJdwFngSs4EniytQlCdcL1sMZa5MG+5Xj+FZHsVFMVuSlegl3JTwp4ZX0an4SuhXrxGZZed7wOARLxywuSQK6gKXcXheXJEGv4IQZBdmWmpgeOYtNLBe4gIxS5+DisCM1MMapK0uSRNW40yR0JA67yIOujpNi4yhqfRzOoagmzs7oouLIpJbbhBLyoJ2FTCeqUKtkTeNeTFCJ6dIxUL3pgmNVlqZmrKCKYbQXUbGPYEPJeYjAqoxDeNhgL28kvDReKCicK1c0QUmxCW8IRmQcxgciug823mc6itTF6ccho7GPHExN8tFRSnwCFY8fBekETedErwRdMOcdlXEJug9YL2f2zMb+nBKtOAuf8SLsCrx55mgwJ0TqhYIHxhV5qV7CuWBmwRxkcjDySno1PzZBwf+X4OFKfUDxFjdv83oJEfMBG85rfi18eC2898KVVOZI6KCgh25hSrMk9KIK0snlCSdheHbAUg72cXTgwWTP/gwa3jMGuEwtfaVB9GnIbCoxG1mWg6s6FRWJfUoRDnVKFYeBg2TVTtapZEeXuEADlNCijXXkukXBQhnPht9gTHhNHYX5QeYFWQEYODIvpL1ewS5wXuxCTTLiKIlxF8zsdYdVcNxYL5+bBVdalYdJAfQwM4PV/KgaAxPkStxcMSDhoDtg3TyYxA/1SAM3t9Z0+LOpaHOiZYQJOYE2g0PnbAKKnYLwZblNWhhbG7dJCjb/I7crMhe25sgNuR7mDjpnq8btbEdikB2Rx94GcBdsbDKmSGwYnBOvqaGjpbzW5hyKlCA5lyS8DKY7u2wHJs6BoSZG6eREW7FIHE1N6LCWw5tzFWhNL+Tr8JkStl1qeVtEOYuKiBEFOWhQlflFhLbj4AyHWLUD76E6zAJhETknCNTVLQ5qr5FKfJnUKcU4rMCjRdQrOG2zMSM7rMkhHFnN1HAKwMVbdwlnIox5M6KXcASNvkXDIGLmsvdV0HlWi04fWA/HLe4keKjOgzeNemXOoKJe5HFGCfZfjkefcWxt/Ev7j+sE7WrjImkwTQmnFSs2kaiz7gou9GS1ntiqmZXa8DZSx6bIdF4vmzyz4YodLzW9dIPGOQlWKlZC3jHirct2bdYncasW0lFGhb5AM8hODCWJsT1TYUpa6m1CERBPoGqRS+QYKiX2QZAVzdqOgXZlogG5KKc5l9U/DqwwUU2cJLzcjgUr3KmiOxe4NUx3oVcDOocBRhJJEpogH7RgJiZmU4lHJSC7JOuq4OIq84SJEltA8Gm9KEXuVOAu2FTh5of3XqJOxSUCaqKuZm69BZt6eBI8gEkiIKPAleaulrErq9e0hpJrrOdasKw9JtVjpTXGOUmUV8QlZl/kKEWejXJiczkxS6fVuEbdOeWHB2igo6hziaNwBuGJdWBjglJwpYGLmsorbpWO+OyM5GZ3qHb7KtxhC2nP+kVAsaIGtIRDaSjnjCtL5g5XrenBnPMs4NsXB89rxwecIR+CXAINMLyanvdHKUeM1du7hjeL1Yzb4xGGhoce1Xlmb72ZfYkiJVyrc7x53pyM5MIv5gAABFFJREFUFYFLTdxG8dYMNpLyKgQ/IjVxeov6N8KJra2t2lwrq2fXVc3PSuhEpHjtdXOpMc7Ev/haCNHmEgylTh57SzF7D3gvJBWQpQTkDSReJPCUk7XpMy2cQh+PzGjQ5WEowmM1cSHToZtD4yIk04w4FVGnVmGBiFMJYwxc6HLkWA3FqajTilLgLtdDUFbjODheSj4HDV14aRMUdUY1vr5QJOPhs5IXGm5iaBY4Ixx+XCgJXYmrKAWeoLtqhXmKupb0Kv5YF4MTx6bAZjKumtC2Tse6EueccZaV6dbGJuiRu4RXIbRVg1LgX07HCdUxI85AzfwjUr8Z3WYpzBu54Ze8riQhkI6L4ZSEc2LesniiJKoBeUGoir1TF0o4BXBICYcLJZMCdzbGlZYkNvgiVOhzxVE41sRpkljecKEEVmhraCnLFBVRJlErTsPuTATRUp8Hwp2s5UWIF6w4SqsXK8qISlXNhoOO6vXr1VVy/zHmX1EcvKjw6gef3tACjMJJKChVUBRBatG98q6SShisvYhV67VHtz48JZGAkVegR66w8ayNsaBnTtM5eZGAaugBpc3Ji4Achoy0RapxKcEvp+OKGv1AqNEPhPrN6HGWOG/kpbiSOr/1qhpjFePmg7eo+xjzIvJKIfEqoMDBIkePSiW0lxvndFq11118RFT2YBpazA+XWGw1EgD0YR22FfOCFckqVLP0zC5hd0YH9k4pZsSSMu4FQlU13VKnau0oTsWxr8vx2x+BxAmfgyWoRe6c8SI6x4uudWBi9gypWh16wkDTq3AX9DpRSnqr86jUjLyZGassO1HllQJr4jzlKdoY51cAy4g8XMm44KFqr9mU6l8VVzXPqtWxYj1aac01XVfN1x4sXa2ojCr01s1jbyVkh+WcZug08irvCqeSOEd1VrDhVIF2yVwZYxVDoS9xqiqOUIgGrnAnpjjLcBxCglMhWRpPhRI5GZGkqdamhxY2whEUzijopUrOv2MR4PX19jHLS82rXgntSuM7B/QiIABW51HBF5+9geMI++rIaPQ60frtMD92tqOCW6u4CvzTU8T4jqUXpQ4eZmehgo1xu+4Cl+DfFK2hV+nVEv2r4lLiszIXZpSsF8rVKieh1H2N1osh9gE5o4B2vTaYgw7QXjNOQa+bx95KiE+8VEEsUMCiT29P45jE5sceHQWbyHkl6DfO+9CX2x8yY5MWrjnFi0/B1KfYk9SUyxfJ0NnZDhHLdqrOFM+mTVScE9WggCoUpRee43c4As4lib3GhuoSeBHtql2SoXPoHOoSV51HhQ6XJJWRUSgFtD44J+Vw+Es4FzgUpQRDRwBWUVXHNA4uIteRoKpDiTygjUWvkTs6zT688wOPyuqDiWNZzsDWaTwBauFcfeLotghEXgVd6K1AWFGJvB5M6FcOZ44TVwM6l70HXOiNSHyFYWQUl5jCxkvIMU5FFUXtzBZL1DkBEpATSHeGqv8HAAD//2xCncMAAAAGSURBVAMAaQHD8Kzkz0EAAAAASUVORK5CYII=', "After Step 7", "HPE & Aruba Product Families — clickable reference"),
    ]
    target = SOP_IMAGE_DIR / "KB-USE-001"
    target.mkdir(parents=True, exist_ok=True)
    conn = db()
    now = datetime.now().isoformat(timespec="seconds")
    for filename, encoded, placement, caption in visuals:
        path = target / filename
        if not path.exists():
            try:
                path.write_bytes(__import__("base64").b64decode(encoded))
            except Exception:
                continue
        exists = conn.execute(
            "SELECT 1 FROM kb_images WHERE kb_id=? AND filename=? LIMIT 1",
            ("KB-USE-001", filename)
        ).fetchone()
        if not exists and path.exists():
            conn.execute("""
                INSERT INTO kb_images
                (kb_id, document_id, filename, image_path, page_num, placement, caption, created_at)
                VALUES (?, NULL, ?, ?, NULL, ?, ?, ?)
            """, ("KB-USE-001", filename, str(path), placement, caption, now))
    conn.commit()
    conn.close()


ensure_builtin_usage_visuals()


# ============================================================
# SEARCH
# ============================================================
@st.cache_data(ttl=30, show_spinner=False)
def search(query, family=None, limit=8):
    """Search only when the query has meaningful evidence in the KB.

    Important retrieval rule: a generic word such as ``onboarding`` must not
    make an unrelated article win simply because it appears somewhere in the
    answer corpus. Product/model/acronym anchors in the user's query are
    treated as required anchors when present.
    """
    records = load_records()
    documents = load_documents()

    if family:
        records = [r for r in records if r["family"] == family]

    q = query.strip().lower()
    if not q:
        return records[:limit], []

    raw_tokens = re.findall(r"[a-z0-9][a-z0-9\-]+", q)
    stopwords = {
        "what", "when", "where", "which", "who", "whom", "whose", "how",
        "why", "can", "could", "would", "should", "does", "do", "did",
        "is", "are", "was", "were", "the", "a", "an", "and", "or", "to",
        "for", "of", "in", "on", "with", "from", "my", "our", "your", "this",
        "that", "it", "be", "i", "me", "please", "guide", "issue", "problem",
        "check", "need", "help", "show", "tell", "about"
    }
    tokens = [t for t in raw_tokens if len(t) >= 3 and t not in stopwords]

    # Product/model/acronym anchors are especially important. If the user
    # types an acronym such as NSP, an article must actually contain NSP.
    original_words = re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]+", query)
    acronym_anchors = [
        w.lower() for w in original_words
        if 2 <= len(w) <= 8 and w.isupper() and w.lower() not in stopwords
    ]

    lexical = []
    for r in records:
        question = r["question"].lower()
        keywords = r["keywords"].lower()
        full = " ".join([
            r["family"], r["topic"], r["question"],
            r["answer"], r["steps"], r["keywords"]
        ]).lower()

        # Never allow an acronym/product anchor to disappear during ranking.
        if acronym_anchors and not all(a in full for a in acronym_anchors):
            continue

        matched_tokens = [t for t in tokens if t in full]
        if tokens and not matched_tokens:
            continue

        # If several meaningful query terms are present, require at least
        # half of them to occur in the candidate. This prevents a single
        # generic term such as "onboarding" from selecting CPPM for an NSP
        # question.
        if len(tokens) >= 2:
            required = max(1, (len(tokens) + 1) // 2)
            if len(set(matched_tokens)) < required:
                continue

        score = 0
        if q in question:
            score += 30
        if q in keywords:
            score += 20
        for t in tokens:
            if t in question:
                score += 6
            elif t in keywords:
                score += 4
            elif t in full:
                score += 1
        lexical.append((score, r))

    merged = []

    if TfidfVectorizer and cosine_similarity and records:
        try:
            corpus = [
                " ".join([
                    r["family"], r["topic"], r["question"],
                    r["answer"], r["steps"], r["keywords"]
                ])
                for r in records
            ]
            vec = TfidfVectorizer(stop_words="english", ngram_range=(1,2))
            matrix = vec.fit_transform(corpus)
            qv = vec.transform([query])
            sims = cosine_similarity(qv, matrix)[0]

            for i, sim in enumerate(sims):
                r = records[i]
                full = " ".join([
                    r["family"], r["topic"], r["question"],
                    r["answer"], r["steps"], r["keywords"]
                ]).lower()
                if acronym_anchors and not all(a in full for a in acronym_anchors):
                    continue
                matched_tokens = [t for t in tokens if t in full]
                if tokens and not matched_tokens:
                    continue
                if len(tokens) >= 2 and len(set(matched_tokens)) < max(1, (len(tokens) + 1) // 2):
                    continue
                # Ignore weak semantic similarity. The old implementation
                # accepted any sim > 0, which is why unrelated CPPM content
                # could answer an NSP question.
                if sim >= 0.18:
                    merged.append((float(sim) * 25, r))
        except Exception:
            pass

    merged.extend(lexical)

    unique = {}
    for score, r in merged:
        unique[r["kb_id"]] = max(score, unique.get(r["kb_id"], 0))

    record_map = {r["kb_id"]: r for r in records}
    ordered = sorted(unique.items(), key=lambda x: x[1], reverse=True)
    result_records = [record_map[k] for k, _ in ordered[:limit]]

    # PDF text search remains available for Related Knowledge, but it does
    # not create an AI exact answer by itself.
    doc_scored = []
    for d in documents:
        if family and d["family"] != family:
            continue
        body = (d["content"] or "").lower()
        hits = sum(1 for t in tokens if t in body)
        if acronym_anchors and not all(a in body for a in acronym_anchors):
            continue
        if q in body:
            hits += 20
        if hits:
            doc_scored.append((hits, d))

    doc_scored.sort(key=lambda x: x[0], reverse=True)
    return result_records, [d for _, d in doc_scored[:limit]]


# ============================================================
# STYLE
# ============================================================
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
  --navy:#061d2b;
  --navy2:#073247;
  --ink:#0b3154;
  --muted:#637b8e;
  --green:#00bfa5;
  --cyan:#00dfcf;
  --border:#d8e7ec;
}

* { box-sizing:border-box; }
html, body, [class*="css"] { font-family:Inter,Arial,sans-serif !important; }

.stApp {
  background:
    radial-gradient(circle at 10% 35%, rgba(0,207,195,.16), transparent 23%),
    radial-gradient(circle at 90% 45%, rgba(0,190,177,.12), transparent 22%),
    linear-gradient(180deg,#eaf8f7 0%,#f6fbfc 58%,#edf8f7 100%);
}

header[data-testid="stHeader"] { display:none; }

.block-container {
  position:relative !important;
  max-width:1536px !important;
  padding:0 22px 18px !important;
}

div[data-testid="stToolbar"] { display:none !important; }

.hero {
  height:245px;
  margin:0 -22px 10px;
  padding:16px 38px 12px;
  color:white;
  position:relative;
  overflow:hidden;
  background:
    radial-gradient(ellipse at 78% 47%,rgba(0,239,218,.23),transparent 23%),
    radial-gradient(ellipse at 31% 12%,rgba(0,119,158,.30),transparent 31%),
    linear-gradient(117deg,#041d2a 0%,#062a3a 52%,#062b38 100%);
}

.hero:before {
  content:"";
  position:absolute;
  left:165px; top:-115px;
  width:850px; height:270px;
  border:1px solid rgba(0,230,220,.25);
  border-radius:50%;
  transform:rotate(-8deg);
  box-shadow:
    0 0 0 40px rgba(0,230,220,.025),
    0 0 0 80px rgba(0,230,220,.018);
}

.hero:after {
  content:"";
  position:absolute;
  right:-30px; top:0;
  width:500px; height:245px;
  background:
    linear-gradient(112deg,transparent 0 43%,rgba(0,255,230,.12) 43.5% 44%,transparent 44.5%),
    linear-gradient(70deg,transparent 0 56%,rgba(0,255,230,.09) 56.5% 57%,transparent 57.5%);
}

.hero-grid {
  position:relative;
  z-index:2;
  display:grid;
  grid-template-columns:205px 1fr 155px;
  height:100%;
  align-items:start;
}

.brand { padding-top:2px; }

/* HPE logo supplied by the user, with Knowledge Base kept as live text. */
.hpe-kb-logo-image {
  display:block !important;
  width:154px !important;
  height:auto !important;
  max-width:154px !important;
  object-fit:contain !important;
  object-position:left center !important;
  margin:0 0 2px 0 !important;
  padding:0 !important;
  border:0 !important;
  background:transparent !important;
}
.kb-brand-text {
  color:#ffffff !important;
  font-size:17px !important;
  line-height:18px !important;
  font-weight:800 !important;
  letter-spacing:-.35px !important;
  margin:0 !important;
  padding:0 !important;
}

.hpe-logo {
  width:80px;
  height:25px;
  border:5px solid #00d9bd;
  border-bottom:0;
  margin-bottom:2px;
}

.hpe-text {
  font-size:30px;
  line-height:29px;
  font-weight:800;
  letter-spacing:-1.5px;
}

.kb-text {
  font-size:17px;
  line-height:18px;
  font-weight:800;
}

.brand-rule {
  width:174px;
  height:1px;
  background:rgba(255,255,255,.48);
  margin:15px 0 9px;
}

.brand-copy {
  font-size:10px;
  line-height:1.45;
  color:#d9edf0;
}

.hero-center { text-align:center; }

.hero-title {
  font-size:43px;
  line-height:47px;
  font-weight:800;
  letter-spacing:-1.8px;
  margin:2px 0 3px;
}

.hero-title span { color:#00e5c6; }

.hero-sub {
  font-size:15px;
  color:#d7eef1;
  margin-bottom:13px;
}

.hero-search { max-width:820px; margin:auto; }

.st-key-hero_search_area { position:relative; z-index:20; margin:-130px auto 8px !important; max-width:830px; }
.st-key-hero_search_area form { background:transparent !important; border:0 !important; padding:0 !important; }
.st-key-hero_search_area div[data-testid="stTextInput"] { margin:0 !important; }
.st-key-hero_search_area div[data-testid="stTextInput"] label { display:none !important; }
.st-key-hero_search_area div[data-testid="stTextInput"] input { height:54px !important; border-radius:29px !important; border:2px solid rgba(0,205,190,.40) !important; background:#fff !important; color:#36526a !important; font-size:14px !important; padding:0 20px 0 22px !important; box-shadow:0 7px 20px rgba(0,30,50,.20) !important; }
.st-key-hero_search_area div[data-testid="stFormSubmitButton"] button { width:52px !important; height:52px !important; min-height:52px !important; padding:0 !important; margin-top:0 !important; border-radius:50% !important; border:0 !important; background:#08bca3 !important; color:white !important; box-shadow:0 4px 12px rgba(0,160,140,.25) !important; font-size:25px !important; line-height:1 !important; }
.st-key-hero_search_area .try-label { text-align:left; display:inline-block; color:#e5f6f5; font-size:10px; font-weight:600; margin:6px 0 0; text-shadow:0 1px 3px rgba(0,0,0,.25); }
.st-key-hero_search_area div[data-testid="stButton"] button { height:34px !important; min-height:34px !important; padding:3px 8px !important; margin-top:3px !important; border-radius:18px !important; border:1px solid rgba(255,255,255,.34) !important; background:rgba(255,255,255,.11) !important; color:white !important; font-size:9px !important; font-weight:500 !important; box-shadow:none !important; white-space:nowrap !important; }
.st-key-hero_search_area div[data-testid="stButton"] button:hover { background:rgba(255,255,255,.20) !important; border-color:rgba(255,255,255,.60) !important; }

.hero-right {
  text-align:right;
  padding-top:60px;
  font-size:15px;
  line-height:1.05;
  font-weight:700;
}

.hero-right .next {
  color:#e2f8f5;
  font-size:10px;
  font-weight:500;
  display:block;
  margin-top:2px;
}

.hero-right .dash {
  color:#00e6c7;
  font-size:25px;
  line-height:20px;
}

.family-row {
  display:grid;
  grid-template-columns:repeat(6,1fr);
  gap:9px;
  margin:0 0 11px;
}

.family-card {
  height:137px;
  background:rgba(255,255,255,.92);
  border:1px solid rgba(214,228,234,.95);
  border-radius:15px;
  padding:14px 13px 10px;
  box-shadow:0 5px 18px rgba(36,76,92,.07);
  position:relative;
}

.family-icon {
  width:38px;
  height:38px;
  border-radius:10px;
  display:flex;
  align-items:center;
  justify-content:center;
  color:#fff;
  font-size:21px;
  font-weight:700;
  margin-bottom:7px;
}

.mint{background:linear-gradient(145deg,#08bda0,#10a7bf)}
.blue{background:linear-gradient(145deg,#1476e7,#3768eb)}
.violet{background:linear-gradient(145deg,#8157e8,#a15ce9)}
.orange{background:linear-gradient(145deg,#ff9b40,#ff7a19)}
.red{background:linear-gradient(145deg,#ff6570,#f13f4b)}
.slate{background:linear-gradient(145deg,#607a92,#405e78)}

.family-name {
  color:#0a3154;
  font-size:16px;
  line-height:18px;
  font-weight:800;
}

.family-desc {
  color:#536f82;
  font-size:10px;
  line-height:13px;
  margin-top:4px;
  padding-right:19px;
}

.family-arrow {
  position:absolute;
  right:12px;
  bottom:14px;
  width:29px;
  height:29px;
  border-radius:50%;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:17px;
  font-weight:700;
}

.arrow-mint{background:#d3f7ef;color:#00a88e}
.arrow-blue{background:#dcecff;color:#2a74df}
.arrow-violet{background:#ece4ff;color:#7c54db}
.arrow-orange{background:#fff0df;color:#f17c1d}
.arrow-red{background:#ffe3e6;color:#f24f5a}
.arrow-slate{background:#e8eef3;color:#516c83}



.family-stack-title {
  color:#0a3154;
  font-size:12px;
  font-weight:800;
  margin:0 0 1px 2px;
  line-height:14px;
}
.st-key-family_cards_home {
  margin-top:-2px !important;
  padding-top:0 !important;
}
.st-key-family_cards_home [data-testid="stHorizontalBlock"] {
  margin-top:0 !important;
  padding-top:0 !important;
}
.family-link {
  display:block !important;
  text-decoration:none !important;
  color:inherit !important;
  -webkit-tap-highlight-color:transparent;
}
.family-link:hover,
.family-link:focus,
.family-link:active {
  text-decoration:none !important;
  color:inherit !important;
}
.family-link .family-card {
  cursor:pointer;
  transition:transform .15s ease, box-shadow .15s ease, border-color .15s ease;
}
.family-link:hover .family-card {
  transform:translateX(-2px);
  border-color:#9edfd8;
  box-shadow:0 8px 24px rgba(0,150,135,.15);
}
.family-link:focus-visible .family-card {
  outline:3px solid rgba(0,191,165,.28);
  outline-offset:2px;
}
.family-card-compact {
  height:68px !important;
  min-height:68px !important;
  padding:8px 10px !important;
  margin-bottom:7px !important;
  display:flex !important;
  align-items:center !important;
  gap:9px !important;
  border-radius:12px !important;

  /* Reference-style HPE family tile: white surface with HPE teal outline. */
  background:#ffffff !important;
  border:2px solid #00bfa5 !important;
  box-shadow:0 4px 12px rgba(0,128,116,.10) !important;
  backdrop-filter:none !important;
  -webkit-backdrop-filter:none !important;
}

.family-card-compact .family-name,
.family-card-compact .family-desc {
  color:#0a3154 !important;
  text-shadow:none;
}
.family-card-compact .family-icon {
  width:44px !important;
  height:44px !important;
  min-width:44px !important;
  flex:0 0 44px !important;
  border-radius:11px !important;
  margin:0 !important;
  font-size:22px !important;
}
.family-card-compact .family-card-copy {
  min-width:0;
  flex:1;
}
.family-card-compact .family-name {
  font-size:13px !important;
  line-height:16px !important;
}
.family-card-compact .family-desc {
  font-size:9px !important;
  line-height:11px !important;
  margin-top:2px !important;
  padding-right:0 !important;
  white-space:nowrap;
  overflow:hidden;
  text-overflow:ellipsis;
}
.family-card-compact .family-arrow {
  position:static !important;
  flex:0 0 26px;
  width:26px !important;
  height:26px !important;
  font-size:15px !important;
}

.main-grid {
  display:grid;
  grid-template-columns:1fr;
  gap:12px;
}

.panel {
  background:rgba(255,255,255,.90);
  border:1px solid rgba(215,229,234,.96);
  border-radius:18px;
  box-shadow:0 6px 19px rgba(35,74,91,.07);
  padding:16px;
}

.ai-panel { min-height:0 !important; height:auto !important; }

.ai-panel div[data-testid="stTextInput"] {
  margin:7px 0 4px !important;
}
.ai-panel div[data-testid="stTextInput"] input {
  height:38px !important;
  min-height:38px !important;
  border-radius:10px !important;
  border:2px solid #00bfa5 !important;
  background:#00bfa5 !important;
  color:#ffffff !important;
  -webkit-text-fill-color:#ffffff !important;
  font-size:10px !important;
  padding:0 12px !important;
  box-shadow:0 4px 12px rgba(0,160,140,.14) !important;
}
.ai-panel div[data-testid="stTextInput"] input::placeholder {
  color:rgba(255,255,255,.92) !important;
  opacity:1 !important;
}

/* Suggested answers and exact-answer controls: white, HPE teal, light-teal lining. */
div[class*="_question_box"] [data-testid="stVerticalBlock"] {
  gap:0 !important;
}
div[class*="_question_box"] div[data-testid="stButton"] {
  margin:0 0 5px 0 !important;
}
div[class*="_question_box"] div[data-testid="stButton"] button {
  background:rgba(255,255,255,.06) !important;
  color:#ffffff !important;
  border:1px solid rgba(255,255,255,.30) !important;
  box-shadow:none !important;
  font-weight:700 !important;
}
div[class*="_question_box"] div[data-testid="stButton"] button:hover {
  background:rgba(255,255,255,.12) !important;
  border-color:rgba(255,255,255,.62) !important;
  color:#ffffff !important;
}
div[class*="_exact_answer_box"] div[data-testid="stButton"] button {
  background:#ffffff !important;
  color:#008f7b !important;
  border:1px solid #9edfd8 !important;
  box-shadow:none !important;
  font-weight:700 !important;
}
div[class*="_exact_answer_box"] div[data-testid="stButton"] button:hover {
  background:#f3fffd !important;
  border-color:#00bfa5 !important;
  color:#007f70 !important;
}
.ai-panel div[data-testid="stFormSubmitButton"] button,
.ai-panel button[kind="secondaryFormSubmit"] {
  height:34px !important;
  min-height:34px !important;
  width:34px !important;
  padding:0 !important;
  border-radius:50% !important;
  border:0 !important;
  background:#08bda4 !important;
  color:white !important;
  font-size:15px !important;
  font-weight:800 !important;
}
.ai-panel div[data-testid="stFormSubmitButton"] button:hover {
  background:#009f8d !important;
  color:white !important;
}
.ai-panel div[data-testid="stButton"] button {
  min-height:26px !important;
  padding:3px 7px !important;
  border-radius:16px !important;
  font-size:8px !important;
  line-height:1.1 !important;
  white-space:nowrap !important;
}

.ai-panel div[data-testid="stButton"] button p,
.ai-panel div[data-testid="stButton"] button span {
  font-size:8px !important;
  line-height:1.1 !important;
  margin:0 !important;
}
.ai-panel .stForm {
  border:0 !important;
  padding:0 !important;
}
.docs-panel { min-height:255px; }
.topics-panel { min-height:139px; }

.ai-head {
  display:flex;
  align-items:center;
  gap:10px;
}

.robot {
  width:48px; height:48px;
  border-radius:50%;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:24px;
  background:linear-gradient(145deg,#e2faf6,#c9f4ee);
  border:1px solid #c0ebe4;
}

.panel-title {
  color:#0a3154;
  font-size:17px;
  font-weight:800;
}

.panel-sub {
  color:#718596;
  font-size:10px;
  margin-top:1px;
}

.beta {
  display:inline-block;
  margin-left:4px;
  padding:3px 7px;
  border-radius:8px;
  background:#dff6ff;
  color:#087da1;
  font-size:8px;
  font-weight:800;
}

.powered {
  margin-left:auto;
  padding:8px 12px;
  border-radius:18px;
  background:#f0faf9;
  border:1px solid #d8eee9;
  color:#008e79;
  font-size:9px;
}


.ai-answer {
  margin-left:12px;
  padding:12px 13px;
  border-radius:0 10px 10px 10px;
  background:#f1f6f8;
  border-left:3px solid #00bfa5;
  color:#263f54;
  font-size:10px;
  line-height:1.45;
}

.ai-step {
  display:flex;
  gap:9px;
  margin:6px 0;
}

.ai-num {
  flex:0 0 20px;
  width:20px;height:20px;
  border-radius:50%;
  background:#09b89f;
  color:white;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:9px;
  font-weight:800;
}

.related-meta {
  margin:-2px 0 5px 10px;
  color:#8a9aa5;
  font-size:8px;
  line-height:10px;
}

.related-file-link {
  display:flex;
  align-items:center;
  gap:10px;
  width:100%;
  box-sizing:border-box;
  margin:7px 0;
  padding:10px 12px;
  border:1px solid #9adfd7;
  border-radius:8px;
  background:#ffffff;
  color:#008f7b !important;
  text-decoration:none !important;
  transition:all .15s ease;
}
.related-file-link:hover {
  border-color:#00a991;
  box-shadow:0 2px 8px rgba(0,169,145,.12);
  text-decoration:none !important;
}
.related-file-icon {
  width:30px;
  height:30px;
  border-radius:6px;
  background:#e9f8f6;
  color:#008f7b;
  display:flex;
  align-items:center;
  justify-content:center;
  font-weight:800;
  flex:0 0 30px;
}
.related-file-link b {
  display:block;
  color:#173a56;
  font-size:10px;
  line-height:1.35;
}
.related-file-link small {
  display:block;
  color:#8a9aa5;
  font-size:8px;
  margin-top:2px;
}
.related-file-arrow {
  margin-left:auto;
  color:#00a991;
  font-weight:900;
  font-size:14px;
}

.ai-doc {
  margin-top:9px;
  padding:9px;
  border-radius:7px;
  background:white;
  border:1px solid #dce7eb;
  display:flex;
  align-items:center;
  gap:8px;
}

.pdf-icon {
  width:28px;height:31px;
  border-radius:4px;
  background:#ee4c45;
  color:white;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:13px;
}

.ai-actions {
  display:flex;
  gap:7px;
  margin:9px 12px;
}

.ai-actions span {
  border:1px solid #d6e4e8;
  background:white;
  padding:6px 10px;
  border-radius:16px;
  color:#466274;
  font-size:8px;
}

.follow {
  margin:7px 0 0;
  border:1px solid #d5e6ea;
  background:white;
  border-radius:15px;
  height:36px;
  display:flex;
  align-items:center;
  padding:0 12px;
  color:#728696;
  font-size:9px;
}

.follow-arrow {
  margin-left:auto;
  width:25px;height:25px;
  background:#08bda4;
  color:white;
  border-radius:50%;
  display:flex;
  justify-content:center;
  align-items:center;
  font-size:15px;
}

.section-head {
  display:flex;
  align-items:flex-start;
  justify-content:space-between;
}

.section-head-title {
  display:flex;
  align-items:center;
  gap:9px;
}

.section-emoji { font-size:23px; }

.view-all {
  color:#0069db;
  font-size:10px;
  font-weight:700;
}

.section-caption {
  color:#6f8595;
  font-size:9px;
  margin:0 0 10px 32px;
}

.doc-grid {
  display:grid;
  grid-template-columns:repeat(4,1fr);
  gap:9px;
}

.doc-card {
  border:1px solid #dbe7ec;
  border-radius:9px;
  background:white;
  padding:7px;
  min-height:178px;
}

.doc-cover {
  height:72px;
  border-radius:5px;
  display:flex;
  align-items:flex-end;
  padding:7px;
  color:white;
  font-size:8px;
  line-height:9px;
  font-weight:800;
  background:
    radial-gradient(circle at 72% 20%,rgba(0,241,222,.45),transparent 24%),
    linear-gradient(140deg,#061d2b,#0a4352 65%,#00a58c);
}

.doc-cover:nth-child(2) { background:linear-gradient(140deg,#071f2f,#173e63,#0bbca1); }

.doc-title {
  color:#153753;
  font-size:9px;
  line-height:11px;
  font-weight:800;
  margin-top:7px;
  min-height:34px;
}

.doc-meta {
  color:#8999a4;
  font-size:8px;
  margin-top:4px;
}

.topic-grid {
  display:grid;
  grid-template-columns:repeat(3,1fr);
  gap:8px;
}

.topic-link {
  display:block;
  text-decoration:none !important;
  color:inherit !important;
  -webkit-tap-highlight-color:transparent;
  margin-bottom:8px;
}
.topic-link:hover,
.topic-link:focus,
.topic-link:active {
  text-decoration:none !important;
  color:inherit !important;
}
.topic-link .topic-card {
  cursor:pointer;
  transition:transform .15s ease, box-shadow .15s ease, border-color .15s ease;
}
.topic-link:hover .topic-card {
  transform:translateY(-1px);
  border-color:#9edfd8;
  box-shadow:0 6px 18px rgba(0,150,135,.12);
}
.topic-link:focus-visible .topic-card {
  outline:3px solid rgba(0,191,165,.28);
  outline-offset:2px;
}
.topic-card {
  min-height:58px;
  border:1px solid #dbe7ec;
  background:white;
  border-radius:9px;
  padding:8px 9px;
  display:flex;
  align-items:center;
  gap:8px;
}

.topic-icon {
  width:27px;height:27px;
  border-radius:50%;
  background:#e6f7f4;
  color:#00a78f;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:14px;
}

.topic-name {
  color:#173a56;
  font-size:12px;
  font-weight:700;
}

.topic-count {
  color:#7b8f9d;
  font-size:9px;
  margin-top:1px;
}

.bottom-strip {
  height:61px;
  margin-top:11px;
  border:1px solid #dce8ec;
  border-radius:16px;
  background:rgba(255,255,255,.92);
  display:grid;
  grid-template-columns:repeat(5,1fr);
  align-items:center;
}

.stat {
  display:flex;
  align-items:center;
  justify-content:center;
  gap:10px;
  border-right:1px solid #dce8ec;
}

.stat:last-child { border-right:0; }

.stat-icon {
  font-size:21px;
  color:#00a991;
}

.stat-value {
  color:#173852;
  font-size:10px;
  font-weight:800;
}

.stat-label {
  color:#788c9a;
  font-size:8px;
  margin-top:2px;
}

.search-result {
  background:white;
  border:1px solid #dbe7ec;
  border-radius:12px;
  padding:13px;
  margin:8px 0;
}

.result-id {
  color:#7c909d;
  font-size:8px;
}

.result-q {
  color:#0c3455;
  font-size:12px;
  font-weight:800;
  margin-top:4px;
}

.result-a {
  color:#536c7e;
  font-size:10px;
  line-height:1.45;
  margin-top:5px;
}

.family-banner {
  margin:4px 0 12px;
  padding:18px 22px;
  border-radius:17px;
  color:white;
  background:linear-gradient(120deg,#061d2b,#074356);
}

.family-banner h1 { margin:0; font-size:27px; }
.family-banner p { margin:4px 0 0; color:#cfe9eb; font-size:11px; }

.topic-directory {
  background:white;
  border:1px solid #dbe7ec;
  border-radius:12px;
  padding:12px;
  margin-bottom:9px;
  font-size:12px !important;
}

.source-card {
  background:white;
  border:1px solid #dbe7ec;
  border-radius:11px;
  padding:11px;
  margin-bottom:8px;
}

.source-card a {
  color:#008f7a;
  text-decoration:none;
  font-weight:700;
  font-size:10px;
}

.footer {
  text-align:center;
  color:#7d909c;
  font-size:8px;
  padding:10px 0 0;
}

div.stButton > button {
  border-radius:18px !important;
  border:1px solid #cfe4e8 !important;
  background:white !important;
  color:#008f7b !important;
  font-size:9px !important;
  font-weight:700 !important;
  min-height:31px !important;
}


.small-button div.stButton > button {
  min-height:26px !important;
  font-size:8px !important;
  padding:2px 8px !important;
}

@media (max-width:1100px){ .hero-grid{grid-template-columns:170px 1fr 120px}.hero-title{font-size:35px;line-height:40px}.hero-sub{font-size:13px}.family-card{height:145px}.main-grid{grid-template-columns:1fr}.st-key-hero_search_area{max-width:760px}.st-key-hero_search_area div[data-testid="stButton"] button{font-size:8px !important} }
@media (max-width:800px){ .block-container{padding:0 10px 14px !important}.hero{height:340px;min-height:340px;margin:0 -10px;padding:14px 16px 12px}.hero-grid{grid-template-columns:1fr;height:auto;text-align:center}.brand{text-align:center}.hpe-kb-logo-image{margin-left:auto !important;margin-right:auto !important}.hpe-logo{margin-left:auto;margin-right:auto}.brand-rule{margin-left:auto;margin-right:auto}.brand-copy{font-size:9px}.hero-center{margin-top:14px}.hero-title{font-size:29px;line-height:33px;letter-spacing:-1px;padding:0 8px}.hero-sub{font-size:11px;line-height:15px;padding:0 15px;margin-top:5px}.hero-right{display:none}.st-key-hero_search_area{width:calc(100% - 20px) !important;max-width:none !important;margin:-110px auto 10px !important}.st-key-hero_search_area div[data-testid="stTextInput"] input{height:50px !important;font-size:12px !important;padding:0 13px !important}.st-key-hero_search_area div[data-testid="stFormSubmitButton"] button{width:48px !important;height:48px !important;min-height:48px !important;font-size:22px !important}.st-key-hero_search_area .try-label{display:block;text-align:center;margin-top:7px;font-size:9px}.st-key-hero_search_area div[data-testid="stHorizontalBlock"]{gap:4px !important}.st-key-hero_search_area div[data-testid="stButton"] button{height:auto !important;min-height:31px !important;font-size:7px !important;padding:4px !important;white-space:normal !important;line-height:9px !important}.family-card{height:142px}.doc-grid{grid-template-columns:repeat(2,1fr)}.topic-grid{grid-template-columns:1fr}.main-grid{display:block}.panel{padding:12px;border-radius:14px}.bottom-strip{grid-template-columns:repeat(2,1fr);height:auto;padding:7px 0}.stat{min-height:48px;border-right:0}.stat:last-child{grid-column:1 / -1} }
@media (max-width:480px){ .hero{height:360px;min-height:360px}.hero-title{font-size:25px;line-height:29px}.hero-sub{font-size:10px}.st-key-hero_search_area{margin-top:-116px !important}.family-card{height:136px;padding:11px 10px}.family-name{font-size:14px}.family-desc{font-size:9px;line-height:12px}.family-icon{width:34px;height:34px;font-size:18px}.doc-grid{grid-template-columns:1fr}.topic-grid{grid-template-columns:1fr} }

/* Smaller answer-choice text — keeps long suggested questions compact. */
.ai-panel div[data-testid="stButton"] button,
.ai-panel div[data-testid="stButton"] button p,
.ai-panel div[data-testid="stButton"] button span {
  font-size:8px !important;
  line-height:1.1 !important;
}

@media (max-width:800px) {
  .ai-panel div[data-testid="stButton"] button,
  .ai-panel div[data-testid="stButton"] button p,
  .ai-panel div[data-testid="stButton"] button span {
    font-size:7px !important;
    line-height:1.05 !important;
  }
}

/* The AI Assistant uses form submission only.
   No standalone "Ask this question" button is rendered. */


/* Admin gear: visually anchored INSIDE the dark-teal hero, directly above
   the “Accelerating what's next together” message. The Streamlit popover
   remains the real accessible/password-protected control. */
.st-key-hero_admin_gear {
  position:absolute !important;
  top:18px !important;
  right:18px !important;
  width:48px !important;
  min-width:48px !important;
  max-width:48px !important;
  height:48px !important;
  margin:0 !important;
  padding:0 !important;
  z-index:1000 !important;
}
.st-key-hero_admin_gear > div,
.st-key-hero_admin_gear [data-testid="stVerticalBlock"] {
  width:48px !important;
  min-width:48px !important;
  padding:0 !important;
  margin:0 !important;
}
.st-key-hero_admin_gear div[data-testid="stPopover"] {
  width:48px !important;
}
.st-key-hero_admin_gear div[data-testid="stPopover"] > button {
  width:46px !important;
  height:46px !important;
  min-width:46px !important;
  min-height:46px !important;
  padding:0 !important;
  margin:0 !important;
  border-radius:12px !important;
  background:transparent !important;
  background-color:transparent !important;
  border:1px solid rgba(0,230,220,.30) !important;
  color:#ffffff !important;
  -webkit-text-fill-color:#ffffff !important;
  box-shadow:none !important;
  font-size:23px !important;
  line-height:1 !important;
  cursor:pointer !important;
  appearance:none !important;
}
.st-key-hero_admin_gear div[data-testid="stPopover"] > button,
.st-key-hero_admin_gear div[data-testid="stPopover"] [data-baseweb="button"] {
  background:transparent !important;
  background-color:transparent !important;
}
.st-key-hero_admin_gear div[data-testid="stPopover"] > button *,
.st-key-hero_admin_gear div[data-testid="stPopover"] > button svg,
.st-key-hero_admin_gear div[data-testid="stPopover"] [data-baseweb="button"] svg {
  color:#ffffff !important;
  fill:#ffffff !important;
  stroke:#ffffff !important;
  -webkit-text-fill-color:#ffffff !important;
}
.st-key-hero_admin_gear div[data-testid="stPopover"] > button:hover,
.st-key-hero_admin_gear div[data-testid="stPopover"] > button:focus-visible {
  background:transparent !important;
  background-color:transparent !important;
  border-color:rgba(0,230,220,.55) !important;
  color:#ffffff !important;
  -webkit-text-fill-color:#ffffff !important;
  outline:3px solid rgba(255,255,255,.12) !important;
  outline-offset:3px !important;
}
.st-key-hero_admin_gear [data-testid="stPopoverBody"] {
  right:0 !important;
  left:auto !important;
  top:54px !important;
  z-index:2000 !important;
}
.ai-panel {
  min-height:0 !important;
  height:auto !important;
}
.ai-panel .ai-head {
  margin-bottom:6px !important;
}


[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  margin:0 0 10px 0 !important;
  padding:10px 14px 9px !important;
  border:1px solid #cfe6e1 !important;
  border-radius:10px !important;
  background:rgba(255,255,255,.72) !important;
  box-shadow:0 2px 8px rgba(0,130,115,.035) !important;
}
[class*="st-key-home_ai_question_box"] .stForm,
[class*="st-key-answer_"][class*="_question_box"] .stForm {
  border:0 !important;
  padding:0 !important;
  margin:0 !important;
}
[class*="st-key-home_ai_question_box"] .chatbot-prompt,
[class*="st-key-answer_"][class*="_question_box"] .chatbot-prompt {
  margin:7px 0 5px !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  margin:0 !important;
}

.chatbot-prompt {
  margin:9px 0 6px 12px;
  color:#526c7d;
  font-size:9px;
  font-weight:600;
}
.chatbot-related-label {
  margin:8px 0 5px 12px;
  color:#173a56;
  font-size:9px;
  font-weight:800;
}
.exact-answer-label {
  display:inline-block;
  margin-bottom:7px;
  color:#008f7b;
  font-size:10px;
  font-weight:800;
  letter-spacing:.6px;
}


@media (max-width:800px){
  .st-key-hero_admin_gear {
    top:13px !important;
    right:14px !important;
  }
  .family-card-compact {
    height:62px !important;
    min-height:62px !important;
  }
  .family-card-compact .family-desc {
    white-space:normal !important;
    display:-webkit-box;
    -webkit-line-clamp:2;
    -webkit-box-orient:vertical;
  }
}



/* Functional family tiles: compact tiles with LARGE icons, not large tiles. */
.st-key-family_tile_0,
.st-key-family_tile_1,
.st-key-family_tile_2,
.st-key-family_tile_3,
.st-key-family_tile_4,
.st-key-family_tile_5 {
  position:relative !important;
  margin:0 !important;
  padding:0 !important;
}
.st-key-family_tile_0 > div,
.st-key-family_tile_1 > div,
.st-key-family_tile_2 > div,
.st-key-family_tile_3 > div,
.st-key-family_tile_4 > div,
.st-key-family_tile_5 > div { position:relative !important; }
.family-tile-visual {
  position:relative;
  height:60px;
  min-height:60px;
  width:100%;
  border:1px solid #d5e4ea;
  border-radius:12px;
  background:rgba(255,255,255,.95);
  box-shadow:0 4px 14px rgba(36,76,92,.06);
  display:flex;
  align-items:center;
  gap:12px;
  padding:7px 42px 7px 12px;
  pointer-events:none;
}
.family-tile-visual .family-icon {
  width:38px !important;
  height:38px !important;
  min-width:38px !important;
  border-radius:10px !important;
  margin:0 !important;
  font-size:24px !important;
  line-height:1 !important;
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
}
.family-tile-copy { min-width:0; flex:1; }
.family-tile-visual .family-name {
  font-size:13px !important;
  line-height:16px !important;
  font-weight:800 !important;
  color:#0a3154 !important;
}
.family-tile-visual .family-desc {
  margin-top:2px !important;
  padding:0 !important;
  font-size:8px !important;
  line-height:10px !important;
  color:#536f82 !important;
  white-space:nowrap;
  overflow:hidden;
  text-overflow:ellipsis;
}
.family-tile-visual .family-arrow {
  position:absolute !important;
  right:11px !important;
  top:17px !important;
  width:27px !important;
  height:27px !important;
  border-radius:50% !important;
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
  font-size:16px !important;
  font-weight:800 !important;
}
[class*="st-key-family_tile_btn_"] {
  position:absolute !important;
  inset:0 !important;
  z-index:5 !important;
  margin:0 !important;
  padding:0 !important;
}
[class*="st-key-family_tile_btn_"] button {
  position:absolute !important;
  inset:0 !important;
  width:100% !important;
  height:60px !important;
  min-height:60px !important;
  margin:0 !important;
  padding:0 !important;
  border-radius:12px !important;
  border:2px solid transparent !important;
  background:transparent !important;
  color:transparent !important;
  box-shadow:none !important;
  font-size:0 !important;
  line-height:0 !important;
}
[class*="st-key-family_tile_btn_"] button:hover {
  background:rgba(0,191,165,.025) !important;
  border-color:#8ed8cf !important;
  box-shadow:0 7px 20px rgba(0,150,135,.10) !important;
  transform:none !important;
}
[class*="st-key-family_tile_btn_"] button:focus-visible {
  outline:3px solid rgba(0,191,165,.35) !important;
  outline-offset:2px !important;
}
@media (max-width:800px) {
  .family-tile-visual { height:60px; min-height:60px; padding:7px 39px 7px 10px; gap:10px; }
  .family-tile-visual .family-icon { width:38px !important; height:38px !important; min-width:38px !important; font-size:24px !important; }
  .family-tile-visual .family-name { font-size:12px !important; line-height:14px !important; }
  .family-tile-visual .family-desc { font-size:7px !important; line-height:9px !important; }
  [class*="st-key-family_tile_btn_"] button { height:60px !important; min-height:60px !important; }
}

.st-key-home_ai_exact_answer_box,
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  margin:9px 0 8px 0 !important;
  padding:0 14px 12px !important;
  border:1px solid #cfe6e1 !important;
  border-left:4px solid #00bfa5 !important;
  border-radius:0 12px 12px 12px !important;
  background:#f7fcfb !important;
  box-shadow:0 2px 9px rgba(0,130,115,.04) !important;
}
.exact-answer-empty {
  min-height:52px;
  width:100%;
  background:transparent;
}
.exact-section-title {
  margin:12px 0 5px;
  color:#008f7b;
  font-size:8px;
  font-weight:800;
  letter-spacing:.65px;
}
/* More readable exact answer presentation. */
.exact-answer {
  margin:9px 0 8px 12px;
  padding:14px 16px 15px;
  border:1px solid #cfe6e1;
  border-left:4px solid #00bfa5;
  border-radius:0 12px 12px 12px;
  background:#f7fcfb;
}
.exact-answer-header {
  display:flex;
  align-items:center;
  gap:10px;
  padding-bottom:10px;
  margin-bottom:10px;
  border-bottom:1px solid #dfecea;
}
.exact-answer-icon {
  width:28px;
  height:28px;
  border-radius:50%;
  background:#08bda4;
  color:#fff;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:16px;
  font-weight:800;
}
.exact-answer-label {
  margin:0 !important;
  color:#008f7b !important;
  font-size:9px !important;
  font-weight:800 !important;
  letter-spacing:.7px !important;
}
.exact-answer-title {
  margin-top:2px;
  color:#173a56;
  font-size:10px;
  font-weight:700;
}
.exact-answer-body {
  color:#243f55;
  font-size:13px;
  line-height:1.65;
  white-space:normal;
}
.summary-paragraph {
  color:#243f55;
  font-size:12px;
  line-height:1.7;
  margin:0 0 8px;
  padding:0;
}
.summary-subtitle {
  color:#173a56;
  font-size:11px;
  font-weight:800;
  margin:8px 0 4px;
}
.summary-bullet {
  display:flex;
  gap:8px;
  color:#324e63;
  font-size:11px;
  line-height:1.5;
  margin:4px 0;
}
.summary-bullet > span {
  color:#00a991;
  font-weight:900;
  flex:0 0 auto;
}
.related-empty {
  color:#7a8d99;
  font-size:9px;
  padding:8px 0 2px;
}
.msg-label {
  display:inline-block;
  margin-right:8px;
  color:#7890a0;
  font-size:8px;
  font-weight:800;
  letter-spacing:.5px;
}

@media (max-width:800px) {
  .st-key-hero_admin_gear {
    top:12px !important;
    right:12px !important;
    width:48px !important;
    min-width:48px !important;
  }
  .st-key-hero_admin_gear div[data-testid="stPopover"] > button {
    width:44px !important;
    height:44px !important;
    min-width:44px !important;
    min-height:44px !important;
    font-size:23px !important;
  }
  .exact-answer { margin-left:0; }
  .exact-answer-body { font-size:11px; line-height:1.6; }
}

/* ============================================================
   EXACT ANSWER / AI RESPONSE — READABLE REFERENCE SCALE
   Keep the typography visually consistent across the entire
   answer section shown in the supplied reference image.
   ============================================================ */
/* HPE exact-answer shell: strong teal outline on every side. */
.st-key-home_ai_exact_answer_box,
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  border:3px solid #00bfa5 !important;
  border-radius:12px !important;
  background:#f8fffd !important;
  padding:10px 12px 12px !important;
  box-sizing:border-box !important;
}
.st-key-home_ai_exact_answer_box > div,
[class*="st-key-answer_"][class*="_exact_answer_box"] > div {
  border-radius:9px !important;
}

.st-key-home_ai_exact_answer_box .user-msg,
[class*="st-key-answer_"][class*="_exact_answer_box"] .user-msg {
  font-size:10px !important;
  line-height:1.5 !important;
}
.st-key-home_ai_exact_answer_box .msg-label,
[class*="st-key-answer_"][class*="_exact_answer_box"] .msg-label {
  font-size:8px !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-label,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-label {
  font-size:9px !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-title,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-title {
  font-size:10px !important;
  line-height:1.35 !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-body,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-body {
  font-size:11px !important;
  line-height:1.6 !important;
}
.st-key-home_ai_exact_answer_box .exact-section-title,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-section-title {
  font-size:9px !important;
  line-height:1.4 !important;
}
.st-key-home_ai_exact_answer_box .summary-paragraph,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-paragraph {
  font-size:10px !important;
  line-height:1.6 !important;
}
.st-key-home_ai_exact_answer_box .summary-subtitle,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-subtitle {
  font-size:10px !important;
}
.st-key-home_ai_exact_answer_box .summary-bullet,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-bullet {
  font-size:10px !important;
  line-height:1.5 !important;
}
.st-key-home_ai_exact_answer_box .ai-step,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-step {
  font-size:10px !important;
}
.st-key-home_ai_exact_answer_box .ai-step > div:last-child,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-step > div:last-child {
  font-size:10px !important;
  line-height:1.5 !important;
}
.st-key-home_ai_exact_answer_box .ai-num,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-num {
  font-size:9px !important;
}
.st-key-home_ai_exact_answer_box .ai-doc,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-doc {
  font-size:10px !important;
}
.st-key-home_ai_exact_answer_box .ai-doc div[style*="font-size:9px"],
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-doc div[style*="font-size:9px"] {
  font-size:10px !important;
}
.st-key-home_ai_exact_answer_box .ai-doc div[style*="font-size:8px"],
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-doc div[style*="font-size:8px"] {
  font-size:9px !important;
}

/* Same-page Topics & Solutions tiles. */
[class*="st-key-topic_tile_"] div[data-testid="stButton"] button {
  width:100% !important;
  min-height:58px !important;
  height:58px !important;
  text-align:left !important;
  justify-content:flex-start !important;
  border:1px solid #dbe7ec !important;
  background:#ffffff !important;
  color:#173a56 !important;
  border-radius:9px !important;
  padding:8px 10px !important;
  box-shadow:none !important;
  font-size:12px !important;
  font-weight:700 !important;
  line-height:1.25 !important;
  white-space:normal !important;
  transition:transform .15s ease, box-shadow .15s ease, border-color .15s ease, background .15s ease !important;
}
[class*="st-key-topic_tile_"] div[data-testid="stButton"] button:hover {
  transform:translateY(-1px) !important;
  border-color:#00bfa5 !important;
  background:#f4fffd !important;
  color:#007f70 !important;
  box-shadow:0 6px 18px rgba(0,150,135,.12) !important;
}
[class*="st-key-topic_tile_"] div[data-testid="stButton"] button:focus-visible {
  outline:3px solid rgba(0,191,165,.28) !important;
  outline-offset:2px !important;
}

/* ============================================================
   AI ANSWER TABS — Excel/browser-tab style
   The active tab stays light teal until another tab is selected.
   ============================================================ */
.st-key-home_ai_exact_answer_box div[data-testid="stHorizontalBlock"] > div,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"] > div {
  padding:0 2px !important;
}
.st-key-home_ai_exact_answer_box div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button {
  height:34px !important;
  min-height:34px !important;
  border:1px solid #8fd8cf !important;
  border-bottom:2px solid #8fd8cf !important;
  border-radius:12px 12px 0 0 !important;
  background:#ffffff !important;
  color:#008f7b !important;
  box-shadow:0 1px 0 rgba(0,0,0,.03) !important;
  font-weight:800 !important;
  transition:all .15s ease !important;
  position:relative !important;
  top:2px !important;
}
.st-key-home_ai_exact_answer_box div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button:hover {
  background:#e9faf7 !important;
  border-color:#00bfa5 !important;
  color:#007f70 !important;
}
/* Selected state is injected by render_ai_assistant based on session state. */
/* The selected tab and its content now form one continuous browser/Excel-style page. */
[class*="st-key-answer_"][class*="_tab_content"],
[class*="st-key-home_ai_tab_content"] {
  background:#d9f7f2 !important;
  border:1px solid #00bfa5 !important;
  border-top:none !important;
  border-radius:0 0 12px 12px !important;
  padding:16px 18px 20px !important;
  margin-top:-1px !important;
  box-shadow:0 3px 12px rgba(0,150,135,.08) !important;
  position:relative !important;
  z-index:1 !important;
}
[class*="st-key-answer_"][class*="_tab_content"] .stMarkdown,
[class*="st-key-home_ai_tab_content"] .stMarkdown {
  background:transparent !important;
}
/* Keep the three action controls readable without making them oversized. */
.st-key-home_ai_exact_answer_box ~ div div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] ~ div div[data-testid="stButton"] button {
  font-size:10px !important;
  min-height:30px !important;
}

/* FINAL AI ANSWER TYPOGRAPHY OVERRIDE */
.st-key-home_ai_exact_answer_box .user-msg,
[class*="st-key-answer_"][class*="_exact_answer_box"] .user-msg {
  font-size:11px !important;
  line-height:1.55 !important;
}
.st-key-home_ai_exact_answer_box .msg-label,
[class*="st-key-answer_"][class*="_exact_answer_box"] .msg-label {
  font-size:8px !important;
  line-height:1.2 !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-label,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-label {
  font-size:9px !important;
  line-height:1.2 !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-title,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-title {
  font-size:11px !important;
  line-height:1.35 !important;
}
.st-key-home_ai_exact_answer_box .exact-answer-body,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-body {
  font-size:14px !important;
  line-height:1.65 !important;
}
.st-key-home_ai_exact_answer_box .exact-section-title,
[class*="st-key-answer_"][class*="_exact_answer_box"] .exact-section-title {
  font-size:10px !important;
  line-height:1.35 !important;
}
.st-key-home_ai_exact_answer_box .summary-paragraph,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-paragraph {
  font-size:13px !important;
  line-height:1.65 !important;
}
.st-key-home_ai_exact_answer_box .summary-subtitle,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-subtitle {
  font-size:11px !important;
}
.st-key-home_ai_exact_answer_box .summary-bullet,
[class*="st-key-answer_"][class*="_exact_answer_box"] .summary-bullet {
  font-size:12px !important;
  line-height:1.55 !important;
}
.st-key-home_ai_exact_answer_box .ai-step > div:last-child,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-step > div:last-child {
  font-size:12px !important;
  line-height:1.55 !important;
}
.st-key-home_ai_exact_answer_box .ai-num,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-num {
  font-size:10px !important;
}
@media (max-width:800px) {
  .st-key-home_ai_exact_answer_box .exact-answer-body,
  [class*="st-key-answer_"][class*="_exact_answer_box"] .exact-answer-body {
    font-size:13px !important;
    line-height:1.6 !important;
  }
  .st-key-home_ai_exact_answer_box .summary-paragraph,
  [class*="st-key-answer_"][class*="_exact_answer_box"] .summary-paragraph {
    font-size:12px !important;
  }
}

/* ============================================================
   COMPACT AI QUESTION / SUGGESTED ANSWERS — FINAL OVERRIDE
   Force Streamlit's generated wrapper gaps to collapse.
   ============================================================ */
[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  margin:0 0 9px 0 !important;
  padding:7px 12px 6px !important;
}

/* The prompt stays close to the input field. */
[class*="st-key-home_ai_question_box"] .chatbot-prompt,
[class*="st-key-answer_"][class*="_question_box"] .chatbot-prompt {
  margin:3px 0 2px !important;
  padding:0 !important;
}

/* Collapse EVERY Streamlit vertical wrapper inside the suggestion area. */
[class*="st-key-home_ai_question_box"] div[data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stVerticalBlock"] {
  gap:0px !important;
  row-gap:0px !important;
}

/* Streamlit may place each button inside its own element container.
   Remove both the wrapper spacing and the button spacing. */
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"],
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"],
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > div,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > div {
  margin:0 !important;
  padding:0 !important;
}

/* Only a 1px visual gap remains between suggested-answer buttons. */
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] {
  margin-bottom:1px !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  height:26px !important;
  min-height:26px !important;
  margin:0 !important;
  padding:1px 8px !important;
  line-height:1.0 !important;
}

/* ============================================================
   SUGGESTED ANSWERS — HARD 1PX ROW SPACING
   Streamlit can add spacing through nested element containers.
   Collapse every wrapper and put the only visual separation on the
   consecutive suggestion rows themselves.
   ============================================================ */
[class*="st-key-home_ai_question_box"] [data-testid="stVerticalBlockBorderWrapper"],
[class*="st-key-answer_"][class*="_question_box"] [data-testid="stVerticalBlockBorderWrapper"],
[class*="st-key-home_ai_question_box"] [data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_question_box"] [data-testid="stVerticalBlock"] {
  gap:0 !important;
  row-gap:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"] {
  margin-top:0 !important;
  margin-bottom:0 !important;
  padding-top:0 !important;
  padding-bottom:0 !important;
  min-height:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]) {
  margin-bottom:1px !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type {
  margin-bottom:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"],
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > div,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > div {
  margin:0 !important;
  padding:0 !important;
  min-height:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  margin:0 !important;
}

/* ============================================================
   EXACT ANSWER BOX — BALANCED INNER TOP/BOTTOM SPACE
   ============================================================ */
.st-key-home_ai_exact_answer_box,
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  margin:9px 0 9px 0 !important;
  padding:13px 14px 13px !important;
  overflow:visible !important;
}

/* Prevent the question row from consuming an uneven top margin. */
.st-key-home_ai_exact_answer_box .user-msg,
[class*="st-key-answer_"][class*="_exact_answer_box"] .user-msg {
  margin:0 0 12px 0 !important;
}

/* Normalize the last visible content so the box always has the same
   13px breathing room below it as above. */
.st-key-home_ai_exact_answer_box > div:last-child,
[class*="st-key-answer_"][class*="_exact_answer_box"] > div:last-child {
  margin-bottom:0 !important;
  padding-bottom:0 !important;
}

.st-key-home_ai_exact_answer_box .ai-doc:last-child,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-doc:last-child {
  margin-bottom:0 !important;
}

/* The related-knowledge list must not visually touch the bottom border. */
.st-key-home_ai_exact_answer_box .ai-doc,
[class*="st-key-answer_"][class*="_exact_answer_box"] .ai-doc {
  margin-bottom:0 !important;
}

/* Streamlit's inner vertical block must not replace the container's
   controlled padding with its own spacing. */
.st-key-home_ai_exact_answer_box > div[data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_exact_answer_box"] > div[data-testid="stVerticalBlock"] {
  gap:0 !important;
  row-gap:0 !important;
  padding-bottom:0 !important;
}

@media (max-width:800px) {
  [class*="st-key-home_ai_question_box"],
  [class*="st-key-answer_"][class*="_question_box"] {
    padding:6px 10px 5px !important;
    margin-bottom:8px !important;
  }

  [class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
  [class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] {
    margin-bottom:1px !important;
  }

  [class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
  [class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
    height:26px !important;
    min-height:26px !important;
    margin:0 !important;
  }

  .st-key-home_ai_exact_answer_box,
  [class*="st-key-answer_"][class*="_exact_answer_box"] {
    margin:8px 0 8px 0 !important;
    padding:11px 10px 11px !important;
  }
}


/* ============================================================
   USER REQUESTED FINAL UI OVERRIDES
   - Exact answer box aligns with the question box.
   - Suggested answers use 1px spacing and HPE teal.
   - Summary / Troubleshooting / Related Knowledge use HPE teal.
   - Product-family tiles use HPE teal while retaining icon colors.
   ============================================================ */

/* Align the answer container exactly with the question container. */
[class*="st-key-home_ai_exact_answer_box"],
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  width:100% !important;
  max-width:100% !important;
  margin:0 0 9px 0 !important;
  padding:13px 14px 13px !important;
  box-sizing:border-box !important;
}

/* Suggested-answer buttons: exactly 1px between rows. */
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] {
  margin:0 0 1px 0 !important;
  padding:0 !important;
}

/* Remove Streamlit wrapper spacing only for the suggested-answer button rows. */
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(div[data-testid="stButton"]) {
  margin:0 0 1px 0 !important;
  padding:0 !important;
  min-height:0 !important;
}

/* The final suggested answer has no extra trailing gap. */
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(div[data-testid="stButton"]):last-child,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(div[data-testid="stButton"]):last-child {
  margin-bottom:0 !important;
}

/* HPE teal suggested answers with white text. */
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  height:26px !important;
  min-height:26px !important;
  margin:0 !important;
  padding:2px 10px !important;
  border:1px solid #00bfa5 !important;
  border-radius:15px !important;
  background:#00bfa5 !important;
  color:#ffffff !important;
  box-shadow:none !important;
  font-weight:600 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button span {
  color:#ffffff !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button:hover {
  background:#009f8d !important;
  border-color:#009f8d !important;
  color:#ffffff !important;
}

/* HPE teal action buttons: Summary / Troubleshooting / Related knowledge. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button {
  min-height:30px !important;
  height:30px !important;
  padding:4px 10px !important;
  border:1px solid #00bfa5 !important;
  border-radius:16px !important;
  background:#00bfa5 !important;
  color:#ffffff !important;
  font-size:10px !important;
  font-weight:700 !important;
  box-shadow:none !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button p,
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button span {
  color:#ffffff !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button:hover {
  background:#009f8d !important;
  border-color:#009f8d !important;
  color:#ffffff !important;
}

/* Product-family tiles: solid HPE teal, while the individual icon colors
   and arrow colors remain exactly as defined by their existing classes. */
.family-card-compact {
  background:#00bfa5 !important;
  border-color:#00bfa5 !important;
  box-shadow:0 5px 16px rgba(0,130,115,.13) !important;
}

.family-card-compact .family-name {
  color:#ffffff !important;
}

.family-card-compact .family-desc {
  color:#ffffff !important;
}

.family-card-compact .family-icon {
  /* Do not override .mint/.blue/.violet/.orange/.red/.slate. */
  color:#ffffff !important;
}

.family-card-compact .family-arrow {
  /* Keep the existing arrow-mint / arrow-blue / etc. colors. */
  box-shadow:none !important;
}

.family-link:hover .family-card-compact {
  background:#00bfa5 !important;
  border-color:#00a991 !important;
  box-shadow:0 8px 24px rgba(0,150,135,.18) !important;
}


/* ============================================================
   HPE REFERENCE VISUAL — ANSWER ACTIONS + PRODUCT FAMILY TILES
   Uses the same dark HPE teal/navy visual language as the supplied
   hero reference image. Icons keep their original family colors.
   ============================================================ */

/* Shared HPE hero surface. */
:root {
  --hpe-reference-bg:
    radial-gradient(ellipse at 78% 47%,rgba(0,239,218,.23),transparent 23%),
    radial-gradient(ellipse at 31% 12%,rgba(0,119,158,.30),transparent 31%),
    linear-gradient(117deg,#041d2a 0%,#062a3a 52%,#062b38 100%);
}

/* Suggested / possible-answer buttons:
   exactly 1px vertical separation and white text on the HPE surface. */
[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  /* Do not let Streamlit's default vertical spacing create a large gap. */
  gap:1px !important;
}

[class*="st-key-home_ai_question_box"] [data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_question_box"] [data-testid="stVerticalBlock"] {
  gap:1px !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]) {
  margin-top:0 !important;
  margin-bottom:0 !important;
  padding-top:0 !important;
  padding-bottom:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] {
  margin:0 !important;
  padding:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > button {
  width:100% !important;
  min-height:28px !important;
  height:28px !important;
  margin:0 !important;
  padding:3px 12px !important;
  border:1px solid rgba(0,229,210,.58) !important;
  border-radius:16px !important;
  background:var(--hpe-reference-bg) !important;
  color:#ffffff !important;
  font-size:10px !important;
  font-weight:700 !important;
  line-height:1.1 !important;
  box-shadow:
    inset 0 1px 0 rgba(255,255,255,.07),
    0 2px 7px rgba(0,28,42,.08) !important;
  transition:filter .15s ease, transform .15s ease, border-color .15s ease !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > button:hover,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > button:hover {
  background:
    radial-gradient(ellipse at 78% 47%,rgba(0,239,218,.29),transparent 23%),
    linear-gradient(117deg,#052433 0%,#073b4c 52%,#063640 100%) !important;
  color:#ffffff !important;
  border-color:#00d9c5 !important;
  transform:none !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > button p,
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > button span,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > button p,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > button span {
  color:#ffffff !important;
  font-size:10px !important;
  font-weight:700 !important;
  line-height:1.1 !important;
  margin:0 !important;
}

/* Remove the extra vertical whitespace around the prompt itself. */
[class*="st-key-home_ai_question_box"] .chatbot-prompt,
[class*="st-key-answer_"][class*="_question_box"] .chatbot-prompt {
  margin:5px 0 3px !important;
}

/* Summary / Troubleshooting Steps / Related Knowledge buttons:
   same HPE reference background and white text. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] > button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] > button {
  min-height:30px !important;
  height:30px !important;
  padding:4px 10px !important;
  border:1px solid rgba(0,229,210,.58) !important;
  border-radius:17px !important;
  background:var(--hpe-reference-bg) !important;
  color:#ffffff !important;
  font-size:9px !important;
  font-weight:700 !important;
  line-height:1.1 !important;
  box-shadow:
    inset 0 1px 0 rgba(255,255,255,.07),
    0 3px 9px rgba(0,28,42,.10) !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] > button:hover,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] > button:hover {
  background:
    radial-gradient(ellipse at 78% 47%,rgba(0,239,218,.29),transparent 23%),
    linear-gradient(117deg,#052433 0%,#073b4c 52%,#063640 100%) !important;
  color:#ffffff !important;
  border-color:#00d9c5 !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] > button p,
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] > button span,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] > button p,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] > button span {
  color:#ffffff !important;
  font-size:9px !important;
  font-weight:700 !important;
  margin:0 !important;
}

/* Product-family tiles: same HPE reference surface; original icon colors
   remain untouched because the icon itself still uses its family class. */
.family-card-compact {
  background:var(--hpe-reference-bg) !important;
  border:1px solid rgba(0,229,210,.30) !important;
  box-shadow:
    inset 0 1px 0 rgba(255,255,255,.10),
    inset 0 -1px 0 rgba(0,25,38,.18),
    0 7px 18px rgba(0,28,42,.13) !important;
}

.family-card-compact .family-name,
.family-card-compact .family-desc {
  color:#ffffff !important;
  text-shadow:0 1px 2px rgba(0,0,0,.20);
}

.family-card-compact .family-arrow {
  /* Keep the original arrow/icon color treatment. */
  filter:none !important;
}

.family-link:hover .family-card-compact {
  background:
    radial-gradient(ellipse at 78% 47%,rgba(0,239,218,.29),transparent 23%),
    linear-gradient(117deg,#052433 0%,#073b4c 52%,#063640 100%) !important;
  border-color:#00d9c5 !important;
}

/* No-match answer area: no fabricated content, only a compact message. */
.no-match-answer {
  min-height:52px;
  display:flex;
  align-items:center;
  justify-content:center;
  padding:18px 12px;
  color:#637b8e;
  font-size:10px;
  font-weight:600;
  text-align:center;
}


/* ============================================================
   FINAL SUGGESTED-ANSWER SPACING / PROMPT VISIBILITY FIX
   Keep the explanatory prompt fully readable, then place exactly
   2px between each of the three suggested-answer buttons.
   ============================================================ */
[class*="st-key-home_ai_question_box"] .chatbot-prompt,
[class*="st-key-answer_"][class*="_question_box"] .chatbot-prompt {
  display:block !important;
  position:relative !important;
  height:auto !important;
  min-height:18px !important;
  max-height:none !important;
  margin:8px 0 7px !important;
  padding:2px 0 !important;
  overflow:visible !important;
  line-height:1.45 !important;
  white-space:normal !important;
  transform:none !important;
}

[class*="st-key-home_ai_question_box"] [data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_question_box"] [data-testid="stVerticalBlock"] {
  gap:0 !important;
  row-gap:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]) {
  margin-top:0 !important;
  margin-bottom:2px !important;
  padding-top:0 !important;
  padding-bottom:0 !important;
  min-height:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type {
  margin-bottom:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"],
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > div,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > div {
  margin:0 !important;
  padding:0 !important;
  min-height:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] > button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] > button {
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
  width:100% !important;
  height:28px !important;
  min-height:28px !important;
  margin:0 !important;
  padding:3px 12px !important;
  line-height:1.2 !important;
}


/* ============================================================
   FINAL UI POLISH — PROMPT VISIBILITY + WHITE HPE TILES
   ============================================================ */

/* The explanatory text must occupy its own real layout row.
   Streamlit can collapse the markdown element wrapper when the
   surrounding vertical gap is aggressively reset. */
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(.chatbot-prompt),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(.chatbot-prompt) {
  display:block !important;
  height:auto !important;
  min-height:22px !important;
  max-height:none !important;
  margin:0 !important;
  padding:0 !important;
  overflow:visible !important;
  flex:0 0 auto !important;
}

[class*="st-key-home_ai_question_box"] .chatbot-prompt,
[class*="st-key-answer_"][class*="_question_box"] .chatbot-prompt {
  display:block !important;
  width:100% !important;
  height:auto !important;
  min-height:20px !important;
  max-height:none !important;
  margin:4px 0 4px !important;
  padding:0 2px !important;
  overflow:visible !important;
  color:#526c7d !important;
  font-size:9px !important;
  font-weight:600 !important;
  line-height:16px !important;
  white-space:normal !important;
  text-overflow:clip !important;
}

/* Suggested / Possible Answer buttons:
   white tile surface with HPE teal + navy lining. */
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  height:28px !important;
  min-height:28px !important;
  margin:0 !important;
  padding:3px 12px !important;
  border:1px solid #00bfa5 !important;
  border-radius:15px !important;
  background:#ffffff !important;
  color:#062a3a !important;
  box-shadow:
    0 0 0 1px rgba(6,42,58,.72),
    inset 0 0 0 1px rgba(0,191,165,.28),
    0 3px 8px rgba(6,42,58,.08) !important;
  font-weight:700 !important;
  line-height:1.2 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button span {
  color:#062a3a !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button:hover {
  background:#f7fffd !important;
  border-color:#00a991 !important;
  color:#062a3a !important;
}

/* Exactly 2px visual separation between the three suggested answers. */
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]) {
  margin:0 0 2px 0 !important;
  padding:0 !important;
  min-height:0 !important;
}

[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-of-type {
  margin-bottom:0 !important;
}

/* Summary / Troubleshooting Steps / Related Knowledge:
   white buttons with the same HPE teal/navy lining. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button {
  min-height:30px !important;
  height:30px !important;
  padding:4px 10px !important;
  border:1px solid #00bfa5 !important;
  border-radius:16px !important;
  background:#ffffff !important;
  color:#062a3a !important;
  box-shadow:
    0 0 0 1px rgba(6,42,58,.72),
    inset 0 0 0 1px rgba(0,191,165,.28),
    0 3px 9px rgba(6,42,58,.09) !important;
  font-size:10px !important;
  font-weight:700 !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button p,
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button span {
  color:#062a3a !important;
}

[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button:hover {
  background:#f7fffd !important;
  border-color:#00a991 !important;
  color:#062a3a !important;
}

/* Product-family tiles:
   white background, HPE teal/navy lining, original family icon colors retained. */
.family-card-compact {
  background:#ffffff !important;
  border:1px solid #00bfa5 !important;
  box-shadow:
    0 0 0 1px rgba(6,42,58,.72),
    inset 0 0 0 1px rgba(0,191,165,.24),
    0 7px 18px rgba(6,42,58,.10) !important;
}

.family-card-compact .family-name {
  color:#062a3a !important;
  text-shadow:none !important;
}

.family-card-compact .family-desc {
  color:#536f82 !important;
  text-shadow:none !important;
}

.family-card-compact .family-arrow {
  filter:none !important;
}

.family-link:hover .family-card-compact {
  background:#f7fffd !important;
  border-color:#00a991 !important;
  box-shadow:
    0 0 0 1px rgba(6,42,58,.82),
    inset 0 0 0 1px rgba(0,191,165,.34),
    0 9px 24px rgba(0,150,135,.16) !important;
}

/* Preserve the compact tile geometry while making the white surface clear. */
.family-card-compact .family-icon {
  flex:0 0 34px !important;
}


/* ============================================================
   FINAL REQUEST OVERRIDES
   - Possible Answers: white / HPE teal text / light teal lining
   - AI action buttons: same treatment
   - Suggested answers: exactly 5px separation
   - AI question bar: HPE teal surface + white text
   ============================================================ */

[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button {
  height:28px !important;
  min-height:28px !important;
  margin:0 !important;
  padding:2px 10px !important;
  border-radius:15px !important;
  border:1px solid #9fe4da !important;
  background:#ffffff !important;
  color:#00a991 !important;
  box-shadow:
    inset 0 0 0 1px rgba(0,191,165,.10),
    0 1px 3px rgba(6,42,58,.06) !important;
  font-size:9px !important;
  font-weight:700 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button span {
  color:#00a991 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stButton"] button:hover {
  background:#f6fffd !important;
  border-color:#00bfa5 !important;
  color:#008f7b !important;
}

/* Streamlit wrapper gap is collapsed; the element container itself supplies
   the ONLY 5px separation between consecutive suggested answers. */
[class*="st-key-home_ai_question_box"] [data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_question_box"] [data-testid="stVerticalBlock"] {
  gap:0 !important;
  row-gap:0 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"] {
  margin-top:0 !important;
  padding-top:0 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]),
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]) {
  margin-bottom:5px !important;
  padding-bottom:0 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-child,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]):last-child {
  margin-bottom:0 !important;
}

/* Summary / Troubleshooting / Related Knowledge buttons:
   white surface, HPE teal text, light teal lining with a navy hairline. */
.st-key-home_ai_exact_answer_box div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button {
  min-height:30px !important;
  height:30px !important;
  margin:0 !important;
  padding:4px 10px !important;
  border:1px solid #9fe4da !important;
  border-radius:16px !important;
  background:#ffffff !important;
  color:#00a991 !important;
  box-shadow:
    inset 0 0 0 1px rgba(0,191,165,.10),
    0 1px 3px rgba(6,42,58,.07) !important;
  font-size:10px !important;
  font-weight:700 !important;
}
.st-key-home_ai_exact_answer_box div[data-testid="stButton"] button p,
.st-key-home_ai_exact_answer_box div[data-testid="stButton"] button span,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button p,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button span {
  color:#00a991 !important;
}
.st-key-home_ai_exact_answer_box div[data-testid="stButton"] button:hover,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stButton"] button:hover {
  background:#f6fffd !important;
  border-color:#00bfa5 !important;
  color:#008f7b !important;
}

/* AI Assistant question input: HPE teal background and white text. */
.st-key-home_ai_question_box div[data-testid="stTextInput"] input,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stTextInput"] input {
  background:#00bfa5 !important;
  color:#ffffff !important;
  border:1px solid #8de4d8 !important;
  caret-color:#ffffff !important;
  box-shadow:inset 0 0 0 1px rgba(255,255,255,.14) !important;
}
.st-key-home_ai_question_box div[data-testid="stTextInput"] input::placeholder,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stTextInput"] input::placeholder {
  color:rgba(255,255,255,.78) !important;
}
.st-key-home_ai_question_box div[data-testid="stFormSubmitButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stFormSubmitButton"] button {
  background:#00bfa5 !important;
  color:#ffffff !important;
  border:1px solid #00bfa5 !important;
}
.st-key-home_ai_question_box div[data-testid="stFormSubmitButton"] button:hover,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stFormSubmitButton"] button:hover {
  background:#00a991 !important;
  color:#ffffff !important;
}

/* SOP image authoring UI */
.sop-image-help {
  margin:7px 0 8px;
  padding:8px 10px;
  border:1px solid #cfe8e3;
  border-radius:8px;
  background:#f5fffd;
  color:#4b6677;
  font-size:10px;
  line-height:1.45;
}
.sop-image-placement-title {
  margin:5px 0 4px;
  color:#008f7b;
  font-size:10px;
  font-weight:800;
}

/* Related visuals inside the exact answer remain compact and aligned. */
.st-key-home_ai_exact_answer_box .stImage,
[class*="st-key-answer_"][class*="_exact_answer_box"] .stImage {
  margin-top:3px !important;
}

</style>
""", unsafe_allow_html=True)


# ============================================================
# UI HELPERS
# ============================================================
def eh(value):
    return escape(str(value or ""))


def render_hero():
    st.markdown(f"""
    <div class="hero">
      <div class="hero-grid">
        <div class="brand">
          <img class="hpe-kb-logo-image" src="data:image/png;base64,{HPE_KB_LOGO_B64}" alt="HPE">
          <div class="kb-brand-text">Knowledge Base</div>
        </div>
        <div class="hero-center">
          <div class="hero-title">How can we help you <span>today?</span></div>
          <div class="hero-sub">Search HPE documentation, guides, and solutions with AI.</div>
        </div>
        <div class="hero-right">
          Accelerating<br>what's next<br><span class="next">together</span>
          <div class="dash">—</div>
        </div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    with st.container(key="hero_admin_gear"):
        admin_login()

    with st.container(key="hero_search_area"):
        with st.form("hero_search_form", clear_on_submit=False):
            c1, c2 = st.columns([0.93, 0.07], gap="small", vertical_alignment="center")
            with c1:
                st.text_input(
                    "Global search",
                    placeholder="Search documents, topics, products, error messages...",
                    key="global_search_input",
                    label_visibility="collapsed"
                )
            with c2:
                submitted = st.form_submit_button("→", use_container_width=True)
        if submitted:
            value = st.session_state.get("global_search_input", "").strip()
            if value:
                # The hero field is GLOBAL SEARCH only. Never populate or
                # submit the separate AI Assistant question field.
                st.session_state["global_search_query"] = value
                st.session_state["global_search_submitted"] = value
                st.session_state.view = "global_search"
                st.rerun()



def open_group(group):
    st.session_state.view = "group"
    st.session_state.selected_group = group
    st.session_state.selected_topic = None
    st.rerun()


def render_family_cards():
    """Render the six product-family tiles in one horizontal row below global search."""
    st.markdown(
        '<div class="family-stack-title">HPE &amp; Aruba Product Families</div>',
        unsafe_allow_html=True
    )

    with st.container(key="family_cards_home"):
        groups = list(PRODUCT_GROUPS.items())
        cols = st.columns(6, gap="small")
        for index, (group, data) in enumerate(groups):
            with cols[index]:
                from urllib.parse import quote
                family_param = quote(group, safe="")
                st.markdown(
                    f"""
                    <a class="family-link family-link-horizontal" href="?family={family_param}"
                       aria-label="Open {eh(group)} product family">
                      <div class="family-card family-card-compact">
                        <div class="family-icon {data['class']}">{data['icon']}</div>
                        <div class="family-card-copy">
                          <div class="family-name">{eh(group)}</div>
                          <div class="family-desc">{eh(data['description'])}</div>
                        </div>
                        <div class="family-arrow arrow-{data['class']}">→</div>
                      </div>
                    </a>
                    """,
                    unsafe_allow_html=True
                )


def answer_card(r):
    st.markdown(
        f"""
        <div class="search-result">
          <div class="result-id">{eh(r['kb_id'])} • {eh(r['family'])} • {eh(r['topic'])}</div>
          <div class="result-q">{eh(r['question'])}</div>
          <div class="result-a">{eh(r['answer'])}</div>
        </div>
        """,
        unsafe_allow_html=True
    )
    with st.expander("Show troubleshooting steps"):
        for line in (r["steps"] or "").splitlines():
            st.markdown(eh(line))
        st.caption("Search terms: " + r["keywords"])


def render_ai_answer(query, family=None):
    """Compatibility wrapper that renders the real functional AI Assistant."""
    safe = re.sub(r"[^A-Za-z0-9]+", "_", (family or "global")).strip("_").lower()
    prefix = f"answer_{safe or 'global'}"
    st.session_state[f"{prefix}_query"] = query
    st.session_state[f"{prefix}_submitted"] = query
    st.session_state[f"{prefix}_action"] = None
    render_ai_assistant(
        default_query=query,
        family=family,
        key_prefix=prefix,
    )


def render_ai_assistant(default_query="",
                         family=None, key_prefix="home_ai"):
    """Exact-answer assistant with Answer, Troubleshooting Steps and Related Files."""
    q_key = f"{key_prefix}_query"
    submitted_key = f"{key_prefix}_submitted"
    selected_key = f"{key_prefix}_selected"
    action_key = f"{key_prefix}_action"

    if q_key not in st.session_state:
        st.session_state[q_key] = default_query or ""
    if submitted_key not in st.session_state:
        st.session_state[submitted_key] = default_query or ""

    previous_demo = "How do I troubleshoot ClearPass licensing issues?"
    if not default_query:
        if st.session_state.get(q_key) == previous_demo:
            st.session_state[q_key] = ""
        if st.session_state.get(submitted_key) == previous_demo:
            st.session_state[submitted_key] = ""

    st.session_state.setdefault(selected_key, None)
    if st.session_state.get(action_key) not in {"answer", "steps", "files"}:
        st.session_state[action_key] = "answer"

    query = st.session_state[submitted_key].strip()
    records = []
    docs = []

    if query:
        records, docs = search(query, family=family, limit=10)

    with st.container(key=f"{key_prefix}_question_box"):
        with st.form(f"{key_prefix}_question_form", clear_on_submit=False):
            c1, c2 = st.columns([0.94, 0.06], gap="small", vertical_alignment="center")
            with c1:
                st.text_input(
                    "Ask HPE AI",
                    key=q_key,
                    placeholder="How do I use the HPE Knowledge Base?",
                    label_visibility="collapsed",
                    help="Type your question and press Enter, or select the arrow to submit."
                )
            with c2:
                submit = st.form_submit_button("→", use_container_width=True)

        if submit:
            value = st.session_state[q_key].strip()
            if value:
                st.session_state[submitted_key] = value
                st.session_state[selected_key] = None
                st.session_state[action_key] = "answer"
                st.rerun()

        if query and records:
            options = records
            selected_id = st.session_state.get(selected_key)
            selected = next((r for r in options if r["kb_id"] == selected_id), None)

            if len(options) > 1 and selected is None:
                st.markdown(
                    '<div class="chatbot-prompt">I found these possible answers. Choose the one that best matches your question:</div>',
                    unsafe_allow_html=True
                )
                for i, r in enumerate(options):
                    if st.button(
                        r["question"],
                        key=f"{key_prefix}_option_{i}_{r['kb_id']}",
                        use_container_width=True
                    ):
                        st.session_state[selected_key] = r["kb_id"]
                        st.session_state[action_key] = "answer"
                        st.rerun()

    if not query:
        return

    if not records:
        with st.container(key=f"{key_prefix}_exact_answer_box"):
            st.markdown(
                '<div class="no-match-answer" aria-label="No matching information found">'
                'No matching information found in the HPE Knowledge Base.'
                '</div>',
                unsafe_allow_html=True
            )
        return

    options = records
    selected_id = st.session_state.get(selected_key)
    selected = next((r for r in options if r["kb_id"] == selected_id), None)
    selected = selected or options[0]

    related_files = [
        d for d in docs
        if Path(d.get("filename") or "").suffix.lower() in SUPPORTED_KB_FILES
    ]

    sop_images = load_kb_images(kb_id=selected["kb_id"], limit=50)
    sop_videos = load_sop_videos(selected["kb_id"], limit=20)

    with st.container(key=f"{key_prefix}_exact_answer_box"):
        st.markdown(
            f"""
            <div class="exact-answer-header">
              <span class="exact-answer-icon">✓</span>
              <div>
                <div class="exact-answer-label">EXACT ANSWER</div>
                <div class="exact-answer-title">{eh(selected["topic"])}</div>
              </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        # Dynamic active-tab styling: only the selected tab is light teal.
        selected_tab_index = {"answer": 1, "steps": 2, "files": 3}.get(
            st.session_state.get(action_key), 1
        )
        st.markdown(
            f"""
            <style>
            [class*="st-key-{key_prefix}_exact_answer_box"] div[data-testid="stHorizontalBlock"]
              > div:nth-child({selected_tab_index}) div[data-testid="stButton"] button,
            .st-key-home_ai_exact_answer_box div[data-testid="stHorizontalBlock"]
              > div:nth-child({selected_tab_index}) div[data-testid="stButton"] button {{
                background:#d9f7f2 !important;
                color:#007f70 !important;
                border-color:#00bfa5 !important;
                border-bottom-color:#d9f7f2 !important;
                margin-bottom:-1px !important;
                box-shadow:0 -1px 5px rgba(0,191,165,.10), inset 0 1px 0 rgba(255,255,255,.75) !important;
                top:0 !important;
                z-index:2 !important;
              }}
            </style>
            """,
            unsafe_allow_html=True,
        )

        a1, a2, a3 = st.columns(3, gap="small")
        with a1:
            if st.button("✦ Answer", key=f"{key_prefix}_answer", use_container_width=True):
                st.session_state[action_key] = "answer"
                st.rerun()
        with a2:
            if st.button("⌕ Troubleshooting steps", key=f"{key_prefix}_steps", use_container_width=True):
                st.session_state[action_key] = "steps"
                st.rerun()
        with a3:
            if st.button("▤ Related knowledge", key=f"{key_prefix}_files", use_container_width=True):
                st.session_state[action_key] = "files"
                st.rerun()

        action = st.session_state[action_key] or "answer"

        # The active tab is the top edge of this content page; the body shares
        # the active-tab background so the selected tab visibly connects to
        # the information displayed beneath it.
        with st.container(key=f"{key_prefix}_tab_content"):
            if action == "answer":
                answer_text = (selected.get("answer") or "").strip()
                st.markdown('<div class="exact-section-title">ANSWER</div>', unsafe_allow_html=True)
                st.markdown(
                    f'<div class="summary-paragraph">{eh(answer_text)}</div>',
                    unsafe_allow_html=True
                )

                step_lines = []
                for line in (selected.get("steps") or "").splitlines():
                    clean = re.sub(r"^\s*\d+\.\s*", "", line).strip()
                    if clean and not clean.lower().startswith(
                        ("prerequisites:", "procedure:", "verification:", "escalation:")
                    ):
                        step_lines.append(clean)

                if step_lines:
                    st.markdown('<div class="summary-subtitle">Key points</div>', unsafe_allow_html=True)
                    for item in step_lines[:5]:
                        st.markdown(
                            f'<div class="summary-bullet"><span>•</span><div>{eh(item)}</div></div>',
                            unsafe_allow_html=True
                        )

                if sop_images or sop_videos:
                    st.markdown(
                        '<div class="exact-section-title">RELATED VISUALS & MEDIA</div>',
                        unsafe_allow_html=True
                    )
                    if sop_images:
                        image_cols = st.columns(min(3, len(sop_images)), gap="small")
                        for image_index, image_record in enumerate(sop_images):
                            with image_cols[image_index % len(image_cols)]:
                                image_path = image_record.get("image_path")
                                if image_path and Path(image_path).exists():
                                    render_clickable_image(
                                        image_path,
                                        image_record.get("caption") or image_record.get("placement") or "",
                                        max_height=180
                                    )
                    for video_record in sop_videos:
                        video_path = video_record.get("video_path")
                        if video_path and Path(video_path).exists():
                            st.video(video_path)
                            if video_record.get("caption"):
                                st.caption(video_record["caption"])

            elif action == "steps":
                st.markdown('<div class="exact-section-title">TROUBLESHOOTING STEPS</div>', unsafe_allow_html=True)
                step_images = {str(img.get("placement") or ""): img for img in sop_images}
                step_videos = {str(video.get("placement") or ""): video for video in sop_videos}

                for i, line in enumerate((selected.get("steps") or "").splitlines(), 1):
                    clean = re.sub(r"^\s*\d+\.\s*", "", line)
                    if not clean.strip():
                        continue

                    if i == 1 and "Before Step 1" in step_images:
                        img = step_images["Before Step 1"]
                        if Path(img["image_path"]).exists():
                            render_clickable_image(img["image_path"], img.get("caption") or "SOP image — Before Step 1", max_height=180)
                    if i == 1 and "Before Step 1" in step_videos:
                        vid = step_videos["Before Step 1"]
                        if Path(vid["video_path"]).exists():
                            st.video(vid["video_path"])

                    st.markdown(
                        f'<div class="ai-step"><div class="ai-num">{i}</div>'
                        f'<div style="font-size:9px;color:#324e63;padding-top:3px;line-height:1.45;">{eh(clean)}</div></div>',
                        unsafe_allow_html=True
                    )

                    placement_key = f"After Step {i}"
                    if placement_key in step_images:
                        img = step_images[placement_key]
                        if Path(img["image_path"]).exists():
                            render_clickable_image(img["image_path"], img.get("caption") or f"SOP image — {placement_key}", max_height=180)
                    if placement_key in step_videos:
                        vid = step_videos[placement_key]
                        if Path(vid["video_path"]).exists():
                            st.video(vid["video_path"])

                if "End of SOP" in step_images:
                    img = step_images["End of SOP"]
                    if Path(img["image_path"]).exists():
                        render_clickable_image(img["image_path"], img.get("caption") or "SOP image — End of SOP", max_height=180)
                if "End of SOP" in step_videos:
                    vid = step_videos["End of SOP"]
                    if Path(vid["video_path"]).exists():
                        st.video(vid["video_path"])

            elif action == "files":
                st.markdown('<div class="exact-section-title">RELATED KNOWLEDGE</div>', unsafe_allow_html=True)

                if not related_files:
                    st.markdown(
                        '<div class="related-empty">No related PDF, Excel, Word, PowerPoint, or video file was found.</div>',
                        unsafe_allow_html=True
                    )
                else:
                    for file_index, d in enumerate(related_files[:8]):
                        suffix = Path(d["filename"]).suffix.lower()
                        file_type = SUPPORTED_KB_FILES.get(suffix, d.get("doc_type") or "File")
                        icon = {
                            "PDF": "▤",
                            "Excel": "▦",
                            "Word": "▤",
                            "PowerPoint": "▥",
                            "Video": "▶",
                        }.get(file_type, "▤")

                        if suffix == ".pdf":
                            if st.button(
                                f"{icon}  {d['filename']}",
                                key=f"{key_prefix}_open_pdf_{file_index}_{d['id']}",
                                use_container_width=True
                            ):
                                st.session_state.selected_document = d["id"]
                                st.session_state.view = "document"
                                st.rerun()
                        else:
                            path = BASE_DIR / "files" / file_type.lower().replace(" ", "_") / d["filename"]
                            data_uri = file_data_uri(path)
                            if data_uri:
                                st.markdown(
                                    f"""
                                    <a class="related-file-link" href="{data_uri}" target="_blank" rel="noopener">
                                      <span class="related-file-icon">{icon}</span>
                                      <span>
                                        <b>{eh(d["filename"])}</b>
                                        <small>{eh(file_type)} • Click to open</small>
                                      </span>
                                      <span class="related-file-arrow">↗</span>
                                    </a>
                                    """,
                                    unsafe_allow_html=True
                                )
                            else:
                                st.markdown(
                                    f'<div class="related-empty">{eh(d["filename"])} could not be opened from the stored file.</div>',
                                    unsafe_allow_html=True
                                )

    # No second question field: the original AI field remains the single input.

def render_documents(docs):
    cols = st.columns(min(4, max(1, len(docs))))
    for i, d in enumerate(docs[:4]):
        with cols[i % len(cols)]:
            st.markdown(
                f"""
                <div class="doc-card">
                  <div class="doc-cover">HPE<br>KNOWLEDGE</div>
                  <div class="doc-title">{eh(d['title'])}</div>
                  <div class="doc-meta">⌁ PDF • Indexed</div>
                </div>
                """,
                unsafe_allow_html=True
            )
            if st.button("View", key=f"doc_{d['id']}", use_container_width=True):
                st.session_state.view = "document"
                st.session_state.selected_document = d["id"]
                st.rerun()


def render_featured():
    docs = load_documents()

    st.markdown("""
    <div class="panel docs-panel">
      <div class="section-head">
        <div>
          <div class="section-head-title">
            <span class="section-emoji">⭐</span>
            <span class="panel-title">Featured Documents</span>
          </div>
          <div class="section-caption">Most useful and commonly accessed resources.</div>
        </div>
        <div class="view-all">View All →</div>
      </div>
    """, unsafe_allow_html=True)

    if docs:
        render_documents(docs[:4])
    else:
        featured = [
            ("ARBSOP040 - HPE ClearPass Licensing Server Overview", "ClearPass"),
            ("HPE Gen11 Server Setup and Configuration Guide", "Compute"),
            ("Aruba Switch Configuration Guide", "Networking"),
            ("HPE iLO 6 User Guide", "iLO"),
        ]
        cols = st.columns(4)
        for i, (title, label) in enumerate(featured):
            with cols[i]:
                st.markdown(
                    f"""
                    <div class="doc-card">
                      <div class="doc-cover">{eh(label)}<br>GUIDE</div>
                      <div class="doc-title">{eh(title)}</div>
                      <div class="doc-meta">⌁ PDF • Knowledge source</div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )
                if st.button("View", key=f"feature_{i}", use_container_width=True):
                    st.session_state.global_search_query = title
                    st.session_state.global_search_submitted = title
                    st.session_state.view = "global_search"
                    st.rerun()

    st.markdown("</div>", unsafe_allow_html=True)


def render_popular_topics():
    topics = [
        ("License Activation", "Software & Licensing", "⌕", "125+ documents"),
        ("ClearPass Troubleshooting", "Security", "⌘", "98+ documents"),
        ("iLO Configuration", "Compute", "▣", "110+ documents"),
        ("Firmware Updates", "Compute", "⟳", "86+ documents"),
        ("Switch Setup", "Networking", "⌘", "102+ documents"),
        ("Alletra Storage", "Storage", "▤", "74+ documents"),
    ]

    st.markdown("""
    <div class="panel topics-panel">
      <div class="section-head">
        <div>
          <div class="section-head-title">
            <span class="section-emoji">🔥</span>
            <span class="panel-title">Popular Topics</span>
          </div>
          <div class="section-caption">Frequently searched by HPE support teams.</div>
        </div>
        <div class="view-all">View All →</div>
      </div>
    """, unsafe_allow_html=True)

    cols = st.columns(3, gap="small")
    for i, (topic, group, icon, count) in enumerate(topics):
        with cols[i % 3]:
            st.markdown(
                f"""
                <div class="topic-card">
                  <div class="topic-icon">{icon}</div>
                  <div style="flex:1">
                    <div class="topic-name">{eh(topic)}</div>
                    <div class="topic-count">{eh(count)}</div>
                  </div>
                  <div style="color:#6b8191;font-size:12px;">›</div>
                </div>
                """,
                unsafe_allow_html=True
            )
            if st.button("Open", key=f"pop_{i}", use_container_width=True):
                st.session_state.view = "group"
                st.session_state.selected_group = group
                st.session_state.selected_topic = topic
                st.rerun()

    st.markdown("</div>", unsafe_allow_html=True)


def render_bottom_strip():
    stats = [
        ("♧", "10,000+", "Technical Documents"),
        ("✦", "AI-Powered", "Instant Answers"),
        ("◷", "Up-to-Date", "Latest HPE Resources"),
        ("♧", "Easy to Navigate", "Find What You Need"),
        ("♢", "Trusted Content", "From HPE Experts"),
    ]

    html = '<div class="bottom-strip">'
    for icon, value, label in stats:
        html += (
            f'<div class="stat">'
            f'<div class="stat-icon">{icon}</div>'
            f'<div><div class="stat-value">{eh(value)}</div>'
            f'<div class="stat-label">{eh(label)}</div></div>'
            f'</div>'
        )
    html += "</div>"
    st.markdown(html, unsafe_allow_html=True)


# ============================================================
# HOME
# ============================================================
def render_home():
    render_hero()

    # Product families now sit immediately below the global search.
    # The AI question + answer area then uses the full page width.
    render_family_cards()

    render_ai_assistant(
        default_query="How do I use the HPE Knowledge Base?",
        key_prefix="home_ai",
    )

    render_bottom_strip()


# ============================================================
# GLOBAL SEARCH PAGE
# ============================================================
def render_global_search():
    query = st.session_state.get("global_search_submitted", "").strip()

    if st.button("← Back to Knowledge Base", key="back_global_search"):
        st.session_state.view = "home"
        st.session_state.global_search_query = ""
        st.session_state.global_search_submitted = ""
        st.rerun()

    st.markdown(
        """
        <div class="family-banner">
          <h1>⌕ Global Search</h1>
          <p>Search indexed HPE knowledge, PDF documents, products, topics, keywords and technical information.</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    with st.form("global_search_results_form", clear_on_submit=False):
        c1, c2 = st.columns([0.93, 0.07], gap="small", vertical_alignment="center")
        with c1:
            search_input = st.text_input(
                "Search all knowledge",
                value=query,
                placeholder="Search documents, topics, products, error messages...",
                label_visibility="collapsed"
            )
        with c2:
            submit = st.form_submit_button("→", use_container_width=True)

    if submit and search_input.strip():
        st.session_state.global_search_query = search_input.strip()
        st.session_state.global_search_submitted = search_input.strip()
        st.rerun()

    query = st.session_state.get("global_search_submitted", "").strip()
    if not query:
        st.info("Enter a search term to search the entire Knowledge Base.")
        return

    records, docs = search(query, family=None, limit=12)
    total = len(records) + len(docs)

    st.markdown(
        f'<div style="margin:12px 0 8px;color:#173a56;font-size:13px;font-weight:800;">'
        f'Global results for <span style="color:#00a991;">{eh(query)}</span> · {total} matches</div>',
        unsafe_allow_html=True
    )

    if not records and not docs:
        st.warning(
            "No matching knowledge or documents were found. Try a product name, "
            "topic, keyword, model, feature, or exact error message."
        )
        return

    if records:
        st.markdown(
            '<div class="panel" style="margin-bottom:12px;">'
            '<div class="panel-title">Knowledge & Information</div>'
            '<div class="panel-sub">Matches from searchable HPE knowledge records.</div>'
            '<div style="height:6px"></div>',
            unsafe_allow_html=True
        )
        for r in records:
            answer_card(r)
        st.markdown('</div>', unsafe_allow_html=True)

    if docs:
        st.markdown(
            '<div class="panel">'
            '<div class="panel-title">Documents</div>'
            '<div class="panel-sub">Indexed PDF documents matching your search.</div>'
            '<div style="height:6px"></div>',
            unsafe_allow_html=True
        )
        for d in docs:
            excerpt = (d["content"] or "")[:420].replace("\n", " ")
            if len(d["content"] or "") > 420:
                excerpt += "…"
            st.markdown(
                f"""
                <div class="search-result">
                  <div class="result-id">PDF · {eh(d['family'])} · {eh(d['topic'])}</div>
                  <div class="result-q">{eh(d['title'])}</div>
                  <div class="result-a">{eh(excerpt)}</div>
                </div>
                """,
                unsafe_allow_html=True
            )
            if st.button("↗ Open document", key=f"global_doc_{d['id']}", use_container_width=True):
                st.session_state.selected_document = d["id"]
                st.session_state.view = "document"
                st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)


# ============================================================
# PRODUCT GROUP PAGE
# ============================================================
def render_group():
    group = st.session_state.get("selected_group")
    if group not in PRODUCT_GROUPS:
        st.session_state.view = "home"
        st.rerun()

    data = PRODUCT_GROUPS[group]

    if st.button("← Back to Knowledge Base", key="back_group"):
        try:
            st.query_params.clear()
        except Exception:
            pass
        st.session_state.view = "home"
        st.rerun()

    st.markdown(
        f"""
        <div class="family-banner">
          <h1>{eh(data['icon'])} {eh(group)}</h1>
          <p>{eh(data['description'])}</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    # Product directory cards were intentionally removed from this page.
    # Topics are the primary drill-down navigation and each tile is clickable.
    st.markdown("""
    <div class="panel" style="margin-bottom:10px;">
      <div class="panel-title">Topics & Solutions</div>
      <div class="panel-sub">Select a topic to view its related AI answers and sources on this page.</div>
      <div style="height:8px"></div>
    """, unsafe_allow_html=True)

    # Same-page topic navigation: clicking a tile changes the Streamlit
    # view in-session instead of navigating to a URL/new browser context.
    topic_cols = st.columns(3, gap="small")
    for i, topic in enumerate(data["topics"]):
        with topic_cols[i % 3]:
            with st.container(key=f"topic_tile_{i}"):
                if st.button(
                    f"{data['icon']}  {topic}",
                    key=f"topic_open_{i}_{re.sub(r'[^A-Za-z0-9]+', '_', topic)}",
                    use_container_width=True,
                    help=f"Open {topic}",
                ):
                    st.session_state.selected_topic = topic
                    st.session_state.view = "topic"
                    try:
                        st.query_params.clear()
                    except Exception:
                        pass
                    st.rerun()

    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown("""
    <div class="panel">
      <div class="panel-title">Source Library</div>
      <div class="panel-sub">Official/public documentation sources associated with this product group.</div>
      <div style="height:8px"></div>
    """, unsafe_allow_html=True)

    source_cols = st.columns(4, gap="small")
    for i, (name, url) in enumerate(data["sources"]):
        with source_cols[i % 4]:
            st.markdown(
                f"""
                <div class="source-card">
                  <div style="font-size:10px;font-weight:800;color:#173a56;">{eh(name)}</div>
                  <div class="small-muted">Official/public source</div>
                  <div style="height:7px"></div>
                  <a href="{eh(url)}" target="_self">Open source →</a>
                </div>
                """,
                unsafe_allow_html=True
            )
    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown("### Ask AI about this product group")
    q = st.text_input(
        "Family search",
        placeholder=f"Ask a question about {group}...",
        key=f"family_question_{group}",
        label_visibility="collapsed"
    )
    if q:
        render_ai_answer(q, family=group)


# ============================================================
# TOPIC PAGE
# ============================================================
def render_topic():
    group = st.session_state.get("selected_group")
    topic = st.session_state.get("selected_topic")

    if group not in PRODUCT_GROUPS:
        st.session_state.view = "home"
        st.rerun()

    if st.button("← Back to product family", key="back_topic"):
        st.session_state.view = "group"
        st.rerun()

    st.markdown(
        f"""
        <div class="family-banner">
          <h1>{eh(topic)}</h1>
          <p>{eh(group)} • Related topics, AI answers and indexed documents</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    # Search the built-in atomic records using topic + family.
    records = [
        r for r in load_records()
        if r["family"] == group and (
            topic.lower() in r["topic"].lower()
            or any(word in (r["question"] + " " + r["keywords"]).lower()
                   for word in topic.lower().split() if len(word) > 3)
        )
    ]

    if records:
        for r in records:
            answer_card(r)
    else:
        st.info("There is currently no atomic AI record mapped directly to this topic.")

    st.markdown("### Search this topic")
    q = st.text_input(
        "Topic search",
        placeholder=f"Ask a specific {topic} question...",
        key=f"topic_question_{group}_{topic}",
        label_visibility="collapsed"
    )
    if q:
        render_ai_answer(q, family=group)


# ============================================================
# DOCUMENT PAGE
# ============================================================
def render_document():
    doc_id = st.session_state.get("selected_document")
    doc = next((x for x in load_documents() if x["id"] == doc_id), None)

    if not doc:
        st.error("Document not found.")
        if st.button("Return"):
            st.session_state.view = "home"
            st.rerun()
        return

    if st.button("← Back", key="back_document"):
        st.session_state.view = "home"
        st.rerun()

    st.markdown(
        f"""
        <div class="family-banner">
          <h1>{eh(doc['title'])}</h1>
          <p>{eh(doc['family'])} • {eh(doc['topic'])} • {eh(doc['filename'])}</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    find = st.text_input(
        "Find in document",
        placeholder="Search within this PDF...",
        label_visibility="collapsed"
    )

    content = doc["content"] or ""
    if find:
        pieces = re.split(r"(?<=[.!?])\s+", content)
        hits = [x for x in pieces if find.lower() in x.lower()]
        if hits:
            for h in hits[:50]:
                st.markdown(f'<div class="search-result">{eh(h)}</div>', unsafe_allow_html=True)
        else:
            st.info("No matching text found.")
    else:
        st.text_area("Document text", content, height=620, label_visibility="collapsed")



# ============================================================
# ADMIN ACCESS + SOP AUTHORING
# ============================================================
def get_admin_password():
    """Read the admin password only from Streamlit secrets."""
    try:
        if "ADMIN_PASSWORD" in st.secrets:
            return str(st.secrets["ADMIN_PASSWORD"])
        if "ADMIN_PIN" in st.secrets:
            return str(st.secrets["ADMIN_PIN"])
    except Exception:
        pass
    return ""


def admin_login():
    """Password gate for the admin area. Password never lives in source code."""
    with st.popover("⚙", help="Security and Knowledge Base Admin"):
        render_access_controls()
        st.markdown("#### Admin")
        st.caption("Enter the admin password configured in Streamlit Secrets.")

        with st.form("admin_login_form"):
            password = st.text_input(
                "Password",
                type="password",
                autocomplete="current-password",
                label_visibility="collapsed",
                placeholder="Admin password"
            )
            login = st.form_submit_button("Sign in", use_container_width=True)

        if login:
            expected = get_admin_password()
            if expected and password == expected:
                st.session_state.admin_authenticated = True
                st.session_state.view = "admin"
                st.rerun()
            elif not expected:
                st.error("ADMIN_PASSWORD is not configured in Streamlit Secrets.")
            else:
                st.error("Incorrect admin password.")


def create_ai_ready_sop(
    family, product, topic, question, direct_answer,
    prerequisites, steps, verification, escalation, keywords,
    source_title="", source_url="", uploaded_images=None, image_placements=None,
    uploaded_videos=None, video_placements=None
):
    """
    Store one atomic SOP record designed for deterministic AI retrieval.
    The record is intentionally self-contained:
    Question -> Direct Answer -> Steps -> Verification -> Escalation.
    """
    now = datetime.now().isoformat(timespec="seconds")
    kb_id = f"SOP-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"

    # Build a normalized searchable text field in the existing answer/steps
    # columns without changing the database schema.
    normalized_steps = steps.strip()
    if prerequisites.strip():
        normalized_steps = (
            "Prerequisites:\n" + prerequisites.strip()
            + "\n\nProcedure:\n" + normalized_steps
        )
    if verification.strip():
        normalized_steps += "\n\nVerification:\n" + verification.strip()
    if escalation.strip():
        normalized_steps += "\n\nEscalation:\n" + escalation.strip()

    source = source_title.strip() or "Admin-created SOP"
    if source_url.strip():
        source += f" | {source_url.strip()}"

    conn = db()
    conn.execute("""
        INSERT INTO kb_records
        (kb_id, family, topic, question, answer, steps, keywords, source, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        kb_id,
        family,
        f"{product} — {topic}",
        question.strip(),
        direct_answer.strip(),
        normalized_steps,
        ", ".join([x.strip() for x in keywords.split(",") if x.strip()]),
        source,
        now
    ))
    conn.commit()
    conn.close()

    image_count = save_sop_images(
        kb_id,
        uploaded_images or [],
        image_placements or []
    )
    video_count = save_sop_videos(
        kb_id,
        uploaded_videos or [],
        video_placements or []
    )

    search.clear()
    return kb_id, image_count, video_count


def render_admin():
    if st.button("← Back to Knowledge Base", key="admin_back"):
        st.session_state.view = "home"
        st.rerun()

    st.markdown("""
    <div class="family-banner">
      <h1>⚙ Knowledge Base Admin</h1>
      <p>Create, update or delete AI-ready SOPs, attach images/videos, and upload related source files.</p>
    </div>
    """, unsafe_allow_html=True)

    create_tab, manage_tab, upload_tab = st.tabs(
        ["Create AI-Ready SOP", "Manage Existing SOPs", "Upload Related Files"]
    )

    with create_tab:
        st.markdown("""
        <div class="panel" style="margin-bottom:12px;">
          <div class="panel-title">AI Extraction Format</div>
          <div class="panel-sub">
            Use one question per SOP. Put the exact answer first, then prerequisites,
            numbered steps, verification, escalation criteria, and search terms.
            Attach screenshots, diagrams and optional instructional video directly to the SOP.
          </div>
        </div>
        """, unsafe_allow_html=True)

        with st.form("create_sop_form", clear_on_submit=True):
            c1, c2 = st.columns(2)
            with c1:
                family = st.selectbox("Product Family", list(PRODUCT_GROUPS.keys()), key="create_family")
                product = st.text_input(
                    "Product / Platform",
                    placeholder="e.g. Aruba ClearPass Policy Manager",
                    key="create_product"
                )
                topic = st.text_input(
                    "Topic / Feature",
                    placeholder="e.g. RADIUS Authentication Failure",
                    key="create_topic"
                )
                question = st.text_input(
                    "Exact User Question",
                    placeholder="e.g. How do I troubleshoot a ClearPass RADIUS authentication failure?",
                    key="create_question"
                )
            with c2:
                source_title = st.text_input(
                    "Source Title (optional)",
                    placeholder="e.g. ClearPass RADIUS Troubleshooting SOP",
                    key="create_source_title"
                )
                source_url = st.text_input(
                    "Source URL (optional)",
                    placeholder="https://...",
                    key="create_source_url"
                )
                keywords = st.text_input(
                    "Search Keywords",
                    placeholder="ClearPass, RADIUS, authentication, timeout, Access Tracker",
                    key="create_keywords"
                )

            direct_answer = st.text_area(
                "Direct Answer — write the exact answer the AI should return",
                height=130,
                placeholder="State the answer directly and completely.",
                key="create_answer"
            )
            prerequisites = st.text_area(
                "Prerequisites / Required Information",
                height=90,
                placeholder="Identify required access, model/version, timestamps, tools, etc.",
                key="create_prerequisites"
            )
            steps = st.text_area(
                "Procedure — one step per line",
                height=170,
                placeholder="1. Open Access Tracker.\n2. Locate the affected request.\n3. Review the failure reason.",
                key="create_steps"
            )

            st.markdown(
                '<div class="sop-image-help"><b>SOP Images</b> — Upload screenshots or diagrams and place each one before/after a procedure step or at the end.</div>',
                unsafe_allow_html=True
            )
            sop_images = st.file_uploader(
                "Upload SOP images",
                type=["png", "jpg", "jpeg", "webp", "gif"],
                accept_multiple_files=True,
                key="create_sop_images"
            )

            st.markdown(
                '<div class="sop-image-help"><b>SOP Video</b> — Attach an instructional or troubleshooting video. It can be displayed in the Answer and at a selected procedure position.</div>',
                unsafe_allow_html=True
            )
            sop_videos = st.file_uploader(
                "Upload SOP video",
                type=["mp4", "webm", "mov", "m4v", "avi"],
                accept_multiple_files=True,
                key="create_sop_videos"
            )

            step_count = len([line for line in steps.splitlines() if line.strip()])
            placement_options = ["Before Step 1"]
            if step_count:
                placement_options += [f"After Step {i}" for i in range(1, step_count + 1)]
            placement_options.append("End of SOP")

            image_placements = []
            if sop_images:
                st.markdown('<div class="sop-image-placement-title">Image placement</div>', unsafe_allow_html=True)
                for image_index, image_file in enumerate(sop_images):
                    image_placements.append(
                        st.selectbox(
                            image_file.name,
                            placement_options,
                            key=f"create_image_placement_{image_index}"
                        )
                    )

            video_placements = []
            if sop_videos:
                st.markdown('<div class="sop-image-placement-title">Video placement</div>', unsafe_allow_html=True)
                for video_index, video_file in enumerate(sop_videos):
                    video_placements.append(
                        st.selectbox(
                            video_file.name,
                            placement_options,
                            key=f"create_video_placement_{video_index}"
                        )
                    )

            verification = st.text_area(
                "Verification / Expected Result",
                height=100,
                placeholder="Describe exactly how the agent confirms the issue is resolved.",
                key="create_verification"
            )
            escalation = st.text_area(
                "Escalation Criteria",
                height=100,
                placeholder="State when the case must be escalated and what evidence must accompany the escalation.",
                key="create_escalation"
            )

            submitted = st.form_submit_button(
                "Create SOP & Add to AI Knowledge",
                type="primary",
                use_container_width=True
            )

        if submitted:
            required = {
                "Product / Platform": product,
                "Topic / Feature": topic,
                "Exact User Question": question,
                "Direct Answer": direct_answer,
                "Procedure": steps,
            }
            missing = [k for k, v in required.items() if not v.strip()]
            if missing:
                st.error("Complete these required fields: " + ", ".join(missing))
            else:
                kb_id, image_count, video_count = create_ai_ready_sop(
                    family, product, topic, question, direct_answer,
                    prerequisites, steps, verification, escalation,
                    keywords, source_title, source_url,
                    uploaded_images=sop_images,
                    image_placements=image_placements,
                    uploaded_videos=sop_videos,
                    video_placements=video_placements
                )
                st.success(
                    f"SOP created successfully: {kb_id}"
                    + (f" • {image_count} image(s)" if image_count else "")
                    + (f" • {video_count} video(s)" if video_count else "")
                )
                st.info("The SOP is immediately searchable by the AI Assistant.")

    with manage_tab:
        conn = db()
        sop_rows = conn.execute("""
            SELECT kb_id, family, topic, question, answer, steps, keywords, source, created_at
            FROM kb_records
            ORDER BY id DESC
        """).fetchall()
        conn.close()
        sop_rows = [dict(r) for r in sop_rows]

        if not sop_rows:
            st.info("No AI-ready SOPs are currently stored.")
        else:
            sop_labels = {
                r["kb_id"]: f'{r["kb_id"]} — {r["question"]}'
                for r in sop_rows
            }
            selected_manage_id = st.selectbox(
                "Select an existing SOP",
                list(sop_labels.keys()),
                format_func=lambda x: sop_labels[x],
                key="manage_sop_id"
            )
            current = next(r for r in sop_rows if r["kb_id"] == selected_manage_id)

            current_topic = current["topic"]
            if " — " in current_topic:
                current_product, current_feature = current_topic.split(" — ", 1)
            else:
                current_product, current_feature = current_topic, ""

            current_steps = current["steps"] or ""
            current_prereq = ""
            current_procedure = current_steps
            current_verification = ""
            current_escalation = ""

            if "\n\nProcedure:\n" in current_procedure:
                current_prereq, current_procedure = current_procedure.split(
                    "\n\nProcedure:\n", 1
                )
                current_prereq = re.sub(r"^Prerequisites:\n", "", current_prereq)

            if "\n\nVerification:\n" in current_procedure:
                current_procedure, current_verification = current_procedure.split(
                    "\n\nVerification:\n", 1
                )
            if "\n\nEscalation:\n" in current_procedure:
                current_procedure, current_escalation = current_procedure.split(
                    "\n\nEscalation:\n", 1
                )

            st.markdown(
                f'<div class="panel"><div class="panel-title">Editing {eh(current["kb_id"])}</div>'
                f'<div class="panel-sub">Changes keep the existing KB ID and update the searchable record.</div></div>',
                unsafe_allow_html=True
            )

            with st.form(f"update_sop_form_{selected_manage_id}", clear_on_submit=False):
                u1, u2 = st.columns(2)
                with u1:
                    ufamily = st.selectbox(
                        "Product Family",
                        list(PRODUCT_GROUPS.keys()),
                        index=list(PRODUCT_GROUPS.keys()).index(current["family"]) if current["family"] in PRODUCT_GROUPS else 0,
                        key=f"update_family_{selected_manage_id}"
                    )
                    uproduct = st.text_input(
                        "Product / Platform",
                        value=current_product,
                        key=f"update_product_{selected_manage_id}"
                    )
                    utopic = st.text_input(
                        "Topic / Feature",
                        value=current_feature,
                        key=f"update_topic_{selected_manage_id}"
                    )
                    uquestion = st.text_input(
                        "Exact User Question",
                        value=current["question"],
                        key=f"update_question_{selected_manage_id}"
                    )
                with u2:
                    usource = st.text_input(
                        "Source / Source Title",
                        value=current["source"] or "",
                        key=f"update_source_{selected_manage_id}"
                    )
                    ukeywords = st.text_input(
                        "Search Keywords",
                        value=current["keywords"] or "",
                        key=f"update_keywords_{selected_manage_id}"
                    )

                uanswer = st.text_area(
                    "Direct Answer",
                    value=current["answer"] or "",
                    height=130,
                    key=f"update_answer_{selected_manage_id}"
                )
                uprereq = st.text_area(
                    "Prerequisites / Required Information",
                    value=current_prereq,
                    height=90,
                    key=f"update_prereq_{selected_manage_id}"
                )
                usteps = st.text_area(
                    "Procedure — one step per line",
                    value=current_procedure,
                    height=170,
                    key=f"update_steps_{selected_manage_id}"
                )
                uverification = st.text_area(
                    "Verification / Expected Result",
                    value=current_verification,
                    height=90,
                    key=f"update_verification_{selected_manage_id}"
                )
                uescalation = st.text_area(
                    "Escalation Criteria",
                    value=current_escalation,
                    height=90,
                    key=f"update_escalation_{selected_manage_id}"
                )

                replace_media = st.checkbox(
                    "Replace existing SOP images/videos with the new uploads",
                    value=False,
                    key=f"replace_media_{selected_manage_id}"
                )
                uimages = st.file_uploader(
                    "Add SOP images",
                    type=["png", "jpg", "jpeg", "webp", "gif"],
                    accept_multiple_files=True,
                    key=f"update_images_{selected_manage_id}"
                )
                uvideos = st.file_uploader(
                    "Add SOP videos",
                    type=["mp4", "webm", "mov", "m4v", "avi"],
                    accept_multiple_files=True,
                    key=f"update_videos_{selected_manage_id}"
                )

                u_step_count = len([line for line in usteps.splitlines() if line.strip()])
                u_placements = ["Before Step 1"] + (
                    [f"After Step {i}" for i in range(1, u_step_count + 1)]
                    if u_step_count else []
                ) + ["End of SOP"]

                ui_placements = []
                if uimages:
                    st.markdown('<div class="sop-image-placement-title">New image placement</div>', unsafe_allow_html=True)
                    for i, f in enumerate(uimages):
                        ui_placements.append(
                            st.selectbox(
                                f.name,
                                u_placements,
                                key=f"update_image_place_{selected_manage_id}_{i}"
                            )
                        )

                uv_placements = []
                if uvideos:
                    st.markdown('<div class="sop-image-placement-title">New video placement</div>', unsafe_allow_html=True)
                    for i, f in enumerate(uvideos):
                        uv_placements.append(
                            st.selectbox(
                                f.name,
                                u_placements,
                                key=f"update_video_place_{selected_manage_id}_{i}"
                            )
                        )

                update_submitted = st.form_submit_button(
                    "Update SOP",
                    type="primary",
                    use_container_width=True
                )

            if update_submitted:
                missing = [
                    label for label, value in {
                        "Product / Platform": uproduct,
                        "Topic / Feature": utopic,
                        "Exact User Question": uquestion,
                        "Direct Answer": uanswer,
                        "Procedure": usteps,
                    }.items() if not value.strip()
                ]
                if missing:
                    st.error("Complete these required fields: " + ", ".join(missing))
                else:
                    image_count, video_count = update_ai_ready_sop(
                        selected_manage_id,
                        ufamily, uproduct, utopic, uquestion, uanswer,
                        uprereq, usteps, uverification, uescalation, ukeywords,
                        usource, "",
                        uploaded_images=uimages,
                        image_placements=ui_placements,
                        uploaded_videos=uvideos,
                        video_placements=uv_placements,
                        replace_media=replace_media
                    )
                    st.success(
                        f"{selected_manage_id} updated successfully."
                        + (f" • {image_count} image(s) added" if image_count else "")
                        + (f" • {video_count} video(s) added" if video_count else "")
                    )
                    st.rerun()

            st.markdown("---")
            st.markdown("### Delete SOP")
            st.warning(
                "Deleting an SOP permanently removes its AI knowledge record and all attached SOP images/videos."
            )
            confirm_delete = st.checkbox(
                "I understand that this SOP and its attached media will be deleted.",
                key=f"confirm_delete_{selected_manage_id}"
            )
            if st.button(
                "Delete Selected SOP",
                key=f"delete_sop_{selected_manage_id}",
                type="secondary",
                disabled=not confirm_delete,
                use_container_width=True
            ):
                delete_ai_ready_sop(selected_manage_id)
                st.success(f"{selected_manage_id} deleted.")
                st.rerun()

            existing_images = load_kb_images(kb_id=selected_manage_id, limit=50)
            existing_videos = load_sop_videos(selected_manage_id, limit=20)
            if existing_images or existing_videos:
                st.markdown("### Attached SOP Media")
                for img in existing_images:
                    st.caption(f"Image • {img['filename']} • {img.get('placement') or 'End of SOP'}")
                for vid in existing_videos:
                    st.caption(f"Video • {vid['filename']} • {vid.get('placement') or 'End of SOP'}")

    with upload_tab:
        c1, c2 = st.columns(2)
        with c1:
            upload_family = st.selectbox(
                "Product Family",
                list(PRODUCT_GROUPS.keys()),
                key="admin_upload_family"
            )
        with c2:
            upload_topic = st.text_input(
                "Topic",
                placeholder="e.g. ClearPass Licensing",
                key="admin_upload_topic"
            )

        file = st.file_uploader(
            "Upload related knowledge file",
            type=[ext.lstrip(".") for ext in SUPPORTED_KB_FILES.keys()],
            help="Supported: PDF, Excel, Word, PowerPoint and video."
        )

        if file and st.button("Upload & Index File", type="primary", use_container_width=True):
            ok, msg = index_supported_file(
                file,
                upload_family,
                upload_topic or "General"
            )
            if ok:
                st.success(msg)
                st.rerun()
            else:
                st.error(msg)

        st.markdown(
            '<div class="panel-sub" style="margin-top:8px;">'
            'Related files are shown only as files — PDF, Excel, Word, PowerPoint or video. '
            'They are not converted into AI answer cards.'
            '</div>',
            unsafe_allow_html=True
        )



# ============================================================
# PERSISTENT ONE-TIME ACCESS TOKEN
# ============================================================
# Security model:
# 1. ACCESS_CODE and TOKEN_SECRET live ONLY in Streamlit Secrets.
# 2. The user enters ACCESS_CODE once.
# 3. The app deterministically derives a signed token from the two
#    server-side secrets and stores ONLY that derived token in a browser cookie.
# 4. The raw access code is never written to the GitHub source code,
#    SQLite database, query string, or browser cookie.
# 5. The gear menu provides "Clear access token", which removes the cookie.
#
# Required Streamlit Secrets:
#   ACCESS_CODE = "your-one-time-access-code"
#   TOKEN_SECRET = "a-long-random-server-side-secret"
#
# Add this package to requirements.txt:
#   streamlit-cookies-controller

ACCESS_COOKIE_NAME = "hpe_kb_access_token_v1"


def _get_access_secrets():
    """Read the access code and signing secret only from Streamlit Secrets."""
    access_code = ""
    token_secret = ""
    try:
        access_code = str(st.secrets.get("ACCESS_CODE", "")).strip()
        token_secret = str(st.secrets.get("TOKEN_SECRET", "")).strip()
    except Exception:
        pass

    # Environment fallback is useful for local development/containers and
    # still keeps both values out of the GitHub source code.
    if not access_code:
        access_code = os.environ.get("ACCESS_CODE", "").strip()
    if not token_secret:
        token_secret = os.environ.get("TOKEN_SECRET", "").strip()

    return access_code, token_secret


def _derive_access_token(access_code, token_secret):
    """Convert the one-time access code into a server-verifiable token."""
    if not access_code or not token_secret:
        return ""
    payload = ("HPE-KB-ACCESS-V1|" + access_code).encode("utf-8")
    secret = token_secret.encode("utf-8")
    return hmac.new(secret, payload, hashlib.sha256).hexdigest()


def _get_cookie_controller():
    """Return one cookie controller per browser session.

    CookieController creates a Streamlit widget during construction, so it must
    NOT be created from an @st.cache_data/@st.cache_resource function. The
    controller instance is kept in session_state instead; this preserves the
    one-controller-per-session behavior without triggering CachedWidgetWarning.
    """
    if CookieController is None:
        return None

    controller_key = "_hpe_kb_cookie_controller"
    controller = st.session_state.get(controller_key)
    if controller is None:
        controller = CookieController(key="hpe_kb_access_cookie_controller")
        st.session_state[controller_key] = controller

    return controller


def _read_access_cookie():
    controller = _get_cookie_controller()
    if controller is None:
        return None, False
    try:
        # CookieController exposes ready() while its browser component is
        # establishing communication with the page.
        if hasattr(controller, "ready") and not controller.ready():
            return None, True
        return controller.get(ACCESS_COOKIE_NAME), False
    except Exception:
        return None, False


def _write_access_cookie(token):
    controller = _get_cookie_controller()
    if controller is None:
        return False
    try:
        # Keep the access token until the user explicitly clears it.
        controller.set(
            ACCESS_COOKIE_NAME,
            token,
            max_age=60 * 60 * 24 * 365 * 10,
        )
        return True
    except Exception:
        return False


def _clear_access_cookie():
    controller = _get_cookie_controller()
    if controller is None:
        return False
    try:
        controller.remove(ACCESS_COOKIE_NAME)
        return True
    except Exception:
        return False


def has_valid_access_token():
    """Return True only when the browser cookie matches the current secrets."""
    access_code, token_secret = _get_access_secrets()
    expected = _derive_access_token(access_code, token_secret)
    if not expected:
        return False

    cookie_token, cookie_waiting = _read_access_cookie()
    if cookie_waiting:
        return None

    if not cookie_token:
        return False

    return hmac.compare_digest(str(cookie_token), expected)


def access_token_gate():
    """
    Gate the entire app behind a one-time access code.

    The first successful entry creates the persistent browser cookie.
    Subsequent refreshes open the app without asking for the code again.
    """
    if CookieController is None:
        st.error(
            "Persistent access-token support is not installed. "
            "Add `streamlit-cookies-controller` to requirements.txt and redeploy."
        )
        st.stop()

    access_code, token_secret = _get_access_secrets()

    if not access_code or not token_secret:
        st.error(
            "Access control is not configured. Add ACCESS_CODE and TOKEN_SECRET "
            "to Streamlit Secrets; do not place them in the GitHub source code."
        )
        st.stop()

    valid = has_valid_access_token()

    if valid is None:
        # Allow the cookie component one browser round-trip to initialize.
        st.info("Preparing secure access…")
        st.stop()

    if valid:
        st.session_state["access_granted"] = True
        return True

    st.session_state["access_granted"] = False

    # Dedicated full-page access screen. No knowledge-base content is rendered
    # until the one-time access code is accepted.

    st.markdown("""
    <style>
    .kb-access-shell {
      min-height:58vh;
      display:flex;
      align-items:center;
      justify-content:center;
      padding:36px 16px;
      box-sizing:border-box;
    }
    .kb-access-card {
      width:min(520px,100%);
      padding:30px 34px 28px;
      border:1px solid #8fd8cf;
      border-radius:18px;
      background:linear-gradient(145deg,#ffffff 0%,#f2fffc 100%);
      box-shadow:0 14px 40px rgba(0,95,90,.10);
      text-align:center;
    }
    .kb-access-logo { width:62px; margin:0 auto 13px; display:block; }
    .kb-access-kicker { color:#008f7b; font-size:9px; font-weight:800; letter-spacing:1.3px; }
    .kb-access-title { margin-top:7px; color:#08384b; font-size:26px; line-height:1.15; font-weight:800; }
    .kb-access-copy { max-width:410px; margin:10px auto 0; color:#5a7180; font-size:12px; line-height:1.6; }
    </style>
    """, unsafe_allow_html=True)

    st.markdown(
        f"""
        <div class="kb-access-shell">
          <div class="kb-access-card">
            <img src="data:image/png;base64,{HPE_KB_LOGO_B64}" alt="HPE" class="kb-access-logo">
            <div class="kb-access-kicker">HPE KNOWLEDGE BASE</div>
            <div class="kb-access-title">Secure access</div>
            <div class="kb-access-copy">
              Enter your one-time access token to open the Knowledge Base.
              This device will stay authorized until the token is cleared from the gear menu.
            </div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    with st.form("hpe_access_token_form", clear_on_submit=True):
        entered = st.text_input(
            "Access token",
            type="password",
            placeholder="Enter access token",
            label_visibility="collapsed",
            autocomplete="one-time-code",
        )
        submit = st.form_submit_button(
            "Open Knowledge Base →",
            use_container_width=True,
            type="primary",
        )

    if submit:
        if hmac.compare_digest(entered.strip(), access_code):
            derived = _derive_access_token(access_code, token_secret)
            if _write_access_cookie(derived):
                st.session_state["access_granted"] = True
                st.rerun()
            else:
                st.error(
                    "The access token could not be saved in the browser. "
                    "Confirm that streamlit-cookies-controller is installed."
                )
        else:
            st.error("Invalid access token.")

    st.stop()


def render_access_controls():
    """Render access-token status and the clear-token control in the gear menu."""
    with st.expander("Secure access", expanded=False):
        st.caption("This browser is authorized with a server-derived access token.")
        if st.button(
            "Clear access token",
            key="clear_access_token",
            use_container_width=True,
        ):
            _clear_access_cookie()
            st.session_state["access_granted"] = False
            st.session_state["admin_authenticated"] = False
            st.session_state["view"] = "home"
            st.success("Access token cleared. The next page load will require the token again.")
            st.rerun()



# ============================================================
# ROUTER
# ============================================================

# The access gate runs before router/page rendering so no knowledge-base
# content is exposed before authorization.
access_token_gate()

for key, default in [
    ("view","home"),
    ("selected_group",None),
    ("selected_topic",None),
    ("selected_document",None),
    ("global_search_input",""),
    ("global_search_query",""),
    ("global_search_submitted",""),
    ("admin_authenticated",False),
]:
    if key not in st.session_state:
        st.session_state[key] = default

# Query parameters are used for tile navigation and as a fallback admin route.
try:
    family_param = st.query_params.get("family")
    topic_param = st.query_params.get("topic")
    if family_param:
        from urllib.parse import unquote_plus
        decoded_family = unquote_plus(str(family_param))
        decoded_topic = unquote_plus(str(topic_param)) if topic_param else None
        if decoded_family in PRODUCT_GROUPS:
            st.session_state.selected_group = decoded_family
            if decoded_topic and decoded_topic in PRODUCT_GROUPS[decoded_family]["topics"]:
                st.session_state.selected_topic = decoded_topic
                st.session_state.view = "topic"
            else:
                st.session_state.selected_topic = None
                st.session_state.view = "group"

    if str(st.query_params.get("admin","")).lower() in {"1","true","yes"}:
        if st.session_state.get("admin_authenticated"):
            st.session_state.view = "admin"
except Exception:
    pass


if st.session_state.view == "home":
    render_home()
elif st.session_state.view == "global_search":
    render_global_search()
elif st.session_state.view == "group":
    render_group()
elif st.session_state.view == "topic":
    render_topic()
elif st.session_state.view == "document":
    render_document()
elif st.session_state.view == "admin":
    if st.session_state.get("admin_authenticated"):
        render_admin()
    else:
        st.session_state.view = "home"
        st.rerun()
else:
    st.session_state.view = "home"
    st.rerun()

st.markdown(
    '<div class="footer">HPE Knowledge Base • AI-assisted technical knowledge search • '
    'Validate production procedures against current authoritative HPE documentation.</div>',
    unsafe_allow_html=True
)

# ============================================================
# FINAL UI OVERRIDES — SUPPLIED REFERENCE LAYOUT
# ============================================================
st.markdown("""
<style>
/* Product families: six compact tiles side-by-side directly below global search. */
.family-stack-title {
  margin: 8px 0 7px !important;
  color:#123c55 !important;
  font-size:11px !important;
  font-weight:800 !important;
  letter-spacing:.15px !important;
}
.family-link-horizontal {
  display:block !important;
  width:100% !important;
  text-decoration:none !important;
}
.family-link-horizontal .family-card-compact {
  height:78px !important;
  min-height:78px !important;
  padding:10px 9px !important;
  border-radius:13px !important;
  box-sizing:border-box !important;
  display:flex !important;
  align-items:center !important;
  gap:7px !important;
  overflow:hidden !important;
}
.family-link-horizontal .family-icon {
  flex:0 0 31px !important;
  width:31px !important;
  height:31px !important;
  border-radius:9px !important;
  font-size:15px !important;
}
.family-link-horizontal .family-card-copy {
  min-width:0 !important;
  flex:1 1 auto !important;
}
.family-link-horizontal .family-name {
  font-size:10px !important;
  line-height:12px !important;
  white-space:normal !important;
}
.family-link-horizontal .family-desc {
  margin-top:2px !important;
  font-size:7px !important;
  line-height:9px !important;
  display:-webkit-box !important;
  -webkit-line-clamp:2 !important;
  -webkit-box-orient:vertical !important;
  overflow:hidden !important;
}
.family-link-horizontal .family-arrow {
  flex:0 0 20px !important;
  width:20px !important;
  height:20px !important;
  font-size:12px !important;
}
.family-link-horizontal:hover .family-card-compact {
  transform:translateY(-1px) !important;
}

/* AI question bar: translucent HPE teal. */
[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  width:100% !important;
  margin:12px 0 12px !important;
  padding:9px 12px !important;
  border:1px solid rgba(0,159,141,.38) !important;
  border-radius:13px !important;
  background:rgba(0,191,165,.15) !important;
  box-shadow:0 5px 18px rgba(0,139,123,.08) !important;
  backdrop-filter:blur(7px) !important;
}
[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  background:#ffffff !important;
  border:1px solid #d7e7eb !important;
  border-radius:13px !important;
  box-shadow:0 5px 18px rgba(0,139,123,.06) !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stTextInput"],
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stTextInput"] {
  width:100% !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stTextInput"] input,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stTextInput"] input {
  height:44px !important;
  border:1px solid #9edfd8 !important;
  border-radius:11px !important;
  background:#ffffff !important;
  color:#08384b !important;
  -webkit-text-fill-color:#08384b !important;
  box-shadow:inset 0 1px 3px rgba(0,88,80,.06) !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stTextInput"] input::placeholder,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stTextInput"] input::placeholder {
  color:#356f75 !important;
  opacity:.9 !important;
}
[class*="st-key-home_ai_question_box"] div[data-testid="stFormSubmitButton"] button,
[class*="st-key-answer_"][class*="_question_box"] div[data-testid="stFormSubmitButton"] button {
  height:44px !important;
  min-height:44px !important;
  border-radius:11px !important;
  border:1px solid rgba(0,145,130,.45) !important;
  background:rgba(0,159,141,.82) !important;
  color:#ffffff !important;
  font-size:18px !important;
}

/* AI answer area is full-width below the six family tiles. */
[class*="st-key-home_ai_exact_answer_box"],
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  width:100% !important;
  max-width:none !important;
}

/* User-requested HPE teal outer outlines: full perimeter, not left-edge only. */
[class*="st-key-home_ai_question_box"],
[class*="st-key-answer_"][class*="_question_box"] {
  border:3px solid #00bfa5 !important;
  border-radius:13px !important;
}
[class*="st-key-home_ai_exact_answer_box"],
[class*="st-key-answer_"][class*="_exact_answer_box"] {
  border:3px solid #00bfa5 !important;
  border-radius:13px !important;
}

/* Compact image previews; clicking opens the original image in an in-page modal. */
.kb-image-preview {
  width:100% !important;
  min-height:70px !important;
  margin:5px 0 9px !important;
  padding:6px !important;
  border:1px solid #d7e7eb !important;
  border-radius:10px !important;
  background:#f7fbfc !important;
  text-align:center !important;
  box-sizing:border-box !important;
}
.kb-image-open {
  display:block !important;
  cursor:zoom-in !important;
  text-decoration:none !important;
}
.kb-clickable-image {
  display:block !important;
  width:auto !important;
  max-width:100% !important;
  max-height:190px !important;
  height:auto !important;
  margin:0 auto !important;
  object-fit:contain !important;
  border-radius:6px !important;
}
.kb-image-caption {
  margin-top:4px !important;
  color:#6c808d !important;
  font-size:7px !important;
  line-height:10px !important;
}

/* CSS-only modal: stays on the same Streamlit page and expands to actual image size. */
.kb-image-modal {
  display:none !important;
  position:fixed !important;
  inset:0 !important;
  z-index:999999 !important;
}
.kb-image-modal:target {
  display:block !important;
}
.kb-image-modal-backdrop {
  position:fixed !important;
  inset:0 !important;
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
  padding:28px !important;
  box-sizing:border-box !important;
  background:rgba(3,24,34,.78) !important;
  backdrop-filter:blur(3px) !important;
}
.kb-image-modal-content {
  position:relative !important;
  max-width:96vw !important;
  max-height:94vh !important;
  display:flex !important;
  flex-direction:column !important;
  align-items:center !important;
  justify-content:center !important;
  padding:14px !important;
  border-radius:14px !important;
  background:#ffffff !important;
  box-shadow:0 18px 60px rgba(0,0,0,.35) !important;
}
.kb-image-full {
  display:block !important;
  width:auto !important;
  height:auto !important;
  max-width:92vw !important;
  max-height:86vh !important;
  object-fit:contain !important;
  border-radius:7px !important;
}
.kb-image-modal-caption {
  width:100% !important;
  margin-top:7px !important;
  color:#466271 !important;
  font-size:10px !important;
  line-height:14px !important;
  text-align:center !important;
}
.kb-image-modal-close {
  position:fixed !important;
  top:14px !important;
  right:18px !important;
  z-index:1000001 !important;
  width:38px !important;
  height:38px !important;
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
  border-radius:50% !important;
  background:#ffffff !important;
  color:#07566a !important;
  text-decoration:none !important;
  font-size:27px !important;
  line-height:1 !important;
  font-weight:700 !important;
  box-shadow:0 5px 18px rgba(0,0,0,.25) !important;
}
.kb-image-modal-close:hover {
  background:#e9fffb !important;
  color:#008f7b !important;
}

@media (max-width:1100px) {
  .family-link-horizontal .family-card-compact { height:72px !important; min-height:72px !important; }
  .family-link-horizontal .family-name { font-size:9px !important; }
  .family-link-horizontal .family-desc { font-size:6.5px !important; }
}
@media (max-width:800px) {
  .family-link-horizontal .family-card-compact { height:70px !important; min-height:70px !important; }
  .family-link-horizontal .family-icon { flex-basis:28px !important; width:28px !important; height:28px !important; }
}

/* ============================================================
   CONNECTED AI ANSWER TABS
   The active tab and the content panel must read as ONE surface.
   Streamlit inserts wrappers/gaps between the button row and the
   following container; these overrides intentionally collapse that
   seam.
   ============================================================ */
[class*="st-key-home_ai_exact_answer_box"] > div[data-testid="stVerticalBlock"],
[class*="st-key-answer_"][class*="_exact_answer_box"] > div[data-testid="stVerticalBlock"] {
  gap:0 !important;
  row-gap:0 !important;
}

/* The tab button row sits directly on top of the content page. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stHorizontalBlock"]:has(button),
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"]:has(button) {
  margin-bottom:-1px !important;
  padding-bottom:0 !important;
  position:relative !important;
  z-index:5 !important;
}

/* Remove Streamlit's default element spacing around the tab row. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stHorizontalBlock"]:has(button)
  > div[data-testid="column"],
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"]:has(button)
  > div[data-testid="column"] {
  margin-bottom:0 !important;
  padding-bottom:0 !important;
}

/* Pull the content shell up into the tab row so there is no white gap. */
[class*="st-key-home_ai_exact_answer_box"] [class*="_tab_content"],
[class*="st-key-answer_"][class*="_tab_content"] {
  margin-top:-1px !important;
  position:relative !important;
  z-index:2 !important;
  border-top:1px solid #00bfa5 !important;
  border-radius:0 0 12px 12px !important;
  background:#d9f7f2 !important;
}

/* The selected button covers the content shell's top border, creating
   the visual effect of a single connected browser/Excel-style tab. */
[class*="st-key-home_ai_exact_answer_box"] div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button,
[class*="st-key-answer_"][class*="_exact_answer_box"] div[data-testid="stHorizontalBlock"] div[data-testid="stButton"] button {
  position:relative !important;
  z-index:6 !important;
  margin-bottom:-1px !important;
}

/* Access-token screen. */
.kb-access-shell {
  min-height:58vh !important;
  display:flex !important;
  align-items:center !important;
  justify-content:center !important;
  padding:36px 16px !important;
  box-sizing:border-box !important;
}
.kb-access-card {
  width:min(520px, 100%) !important;
  padding:30px 34px 28px !important;
  border:1px solid #8fd8cf !important;
  border-radius:18px !important;
  background:linear-gradient(145deg,#ffffff 0%,#f2fffc 100%) !important;
  box-shadow:0 14px 40px rgba(0,95,90,.10) !important;
  text-align:center !important;
}
.kb-access-logo {
  width:62px !important;
  height:auto !important;
  margin:0 auto 13px !important;
  display:block !important;
}
.kb-access-kicker {
  color:#008f7b !important;
  font-size:9px !important;
  font-weight:800 !important;
  letter-spacing:1.3px !important;
}
.kb-access-title {
  margin-top:7px !important;
  color:#08384b !important;
  font-size:26px !important;
  line-height:1.15 !important;
  font-weight:800 !important;
}
.kb-access-copy {
  max-width:410px !important;
  margin:10px auto 0 !important;
  color:#5a7180 !important;
  font-size:12px !important;
  line-height:1.6 !important;
}

</style>
""", unsafe_allow_html=True)
